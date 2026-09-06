from __future__ import annotations
import csv
import hashlib
import html
import json
import logging
import tempfile
import threading
import uuid
from urllib.parse import urlencode
from logging.handlers import RotatingFileHandler
from pathlib import Path
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from quality_knowledge.adapters.base import clean
from quality_knowledge.excel_utils import repair_read_only_dimensions
from quality_knowledge.repositories import IssueKnowledgeRepository
from quality_knowledge.services import KnowledgeIssueService
from quality_knowledge.services.intake_session_service import IntakeSessionService, IntakeSessionError
from quality_knowledge.mapping.repository import MappingConfigurationRepository
from quality_knowledge.mapping.service import MappingConfigurationService
from quality_knowledge.mapping.models import MappingItem, TARGET_DOMAINS
from quality_knowledge.mapping.runtime import runtime_detection_score
from quality_knowledge.human_analysis import HumanAnalysisRepository, HumanAnalysisService
from quality_knowledge.analysis_profiles import DOMAIN_PROFILES, ISSUE_TYPES, LIFECYCLE_PHASES, DOMAIN_LABELS, ISSUE_TYPE_LABELS, LIFECYCLE_LABELS, normalize_analysis_profile
from quality_knowledge.product_config import ProductConfigRepository
from quality_knowledge.adapters import register_product_adapter
from quality_knowledge.capability_extension import QualityCapabilityExtension
from quality_knowledge.batch_analysis_jobs import BatchAnalysisJobManager
from quality_knowledge.model_config import list_quality_issue_agents
from .statistics_presenter import present_statistics, present_common_gaps, zh_value
from quality_knowledge.product_report.legacy_service import LegacyProductQualityReportService
from quality_knowledge.product_report.service import ProductReportError
from quality_knowledge.materials import MaterialRepository, MaterialImportService
from quality_knowledge.scenarios import ScenarioRepository
from quality_knowledge.scenario_generation import ScenarioGenerationService

BASE = Path(__file__).parent
ALLOWED = {'.xlsx', '.xlsm'}
DIAG_DIR = BASE.parent.parent / 'output' / 'diagnostics'
DIAG_FILE = DIAG_DIR / 'data_intake.log'
_intake_log = logging.getLogger('quality_knowledge.data_intake')
if not _intake_log.handlers:
    DIAG_DIR.mkdir(parents=True,exist_ok=True)
    _handler=RotatingFileHandler(DIAG_FILE,maxBytes=2_000_000,backupCount=3,encoding='utf-8')
    _handler.setFormatter(logging.Formatter('%(asctime)s %(message)s'))
    _intake_log.addHandler(_handler); _intake_log.setLevel(logging.INFO); _intake_log.propagate=False

def _intake_diag(event,diagnostic_id,**details):
    _intake_log.info(json.dumps({'event':event,'diagnostic_id':diagnostic_id,**details},ensure_ascii=False,default=str))


def _safe_json(value, default=None):
    if value is None:
        return {} if default is None else default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        return {} if default is None else default


def _ev(obj):
    """Human-readable value from EvidenceValue-compatible objects."""
    if obj is None:
        return ''
    if isinstance(obj, str):
        return obj
    if isinstance(obj, dict):
        return obj.get('value') or obj.get('description') or ''
    return str(obj)


def _confidence(obj):
    if isinstance(obj, dict):
        try:
            return float(obj.get('confidence') or 0)
        except Exception:
            return 0.0
    return 0.0



_FIELD_LABELS = {
    'description':'问题描述','title':'问题标题','business_issue_id':'问题编号','product':'产品','platform':'平台','module':'模块',
    'issue_type':'问题类型','severity':'严重度','impact':'影响结果','industry':'行业','customer':'客户','department':'部门','business_group':'业务组',
    'root_cause':'根因','occurrence_cause':'发生原因','escape_cause':'流出原因','solution':'解决措施','verification':'验证结果',
    'issue_fact':'问题事实','product_context':'产品上下文','occurrence':'发生信息','escape':'流出信息','recurrence':'再发信息','product_extension':'产品扩展信息'
}

def _display_value(value):
    if value is None or value == '': return ''
    if isinstance(value, bool): return '是' if value else '否'
    if isinstance(value, list): return ' / '.join(_display_value(x) for x in value if _display_value(x))
    if isinstance(value, dict): return '；'.join(f"{_FIELD_LABELS.get(str(k),str(k))}：{_display_value(v)}" for k,v in value.items() if _display_value(v))
    return str(value)

def _readable_rows(obj, prefix=''):
    rows=[]
    if not isinstance(obj, dict): return rows
    for key,value in obj.items():
        if key in {'raw_json','normalized_json'}: continue
        label=_FIELD_LABELS.get(str(key), str(key).replace('_',' '))
        if isinstance(value,dict):
            for child in _readable_rows(value, label):
                child['label']=f"{label} · {child['label']}"
                rows.append(child)
        else:
            text=_display_value(value)
            if text: rows.append({'label':label,'key':str(key),'value':text})
    return rows


def _build_issue_view(svc: KnowledgeIssueService, knowledge_id: str, capability_extension=None):
    d = svc.get_issue_detail(knowledge_id)
    if not d:
        return None
    issue = d['issue']
    normalized = _safe_json(issue.get('normalized_json'))
    analysis = {x: svc.get_latest_analysis(knowledge_id, x) for x in ['occurrence', 'escape', 'recurrence', 'capability_gap']}
    results = {k: (v.get('result') if v else None) for k, v in analysis.items()}
    gaps = svc.query_capability_gaps(knowledge_id)
    grouped = {'TECHNICAL': [], 'MANAGEMENT': [], 'GOVERNANCE': []}
    for g in gaps:
        grouped.setdefault(g.get('dimension') or 'OTHER', []).append(g)
    history = svc.get_analysis_history(knowledge_id)
    latest_run = history[0] if history else None
    debug = None
    if latest_run and hasattr(svc.repository, 'get_analysis_debug'):
        debug = svc.repository.get_analysis_debug(latest_run['analysis_run_id'])
    recurrence = results.get('recurrence') or {}
    capability_view = capability_extension.project_issue(issue, analysis) if capability_extension else {
        'projection': None, 'tags': [], 'mrc': [], 'evidence': []
    }
    capability_view.setdefault('effective_mrc', {})
    capability_view.setdefault('pending_tags', [])
    if not capability_view.get('effective_tags'):
        source_labels={'HUMAN_CONFIRMED':'人工确认','SOURCE_DATA':'原始数据','AI_STANDARDIZED':'AI 标准化','AI_INFERRED':'AI 推断'}
        tag_groups={}
        for raw in capability_view.get('tags') or []:
            tag=dict(raw);tag['source_label']=source_labels.get(tag.get('source_type'),tag.get('source_type') or '未知来源')
            tag_groups.setdefault(tag.get('axis') or 'OTHER',[]).append(tag)
        capability_view['effective_tags']={axis:{'primary':values[0],'related':values[1:4]} for axis,values in tag_groups.items() if values}
    return {
        'd': d,
        'issue': issue,
        'normalized': normalized,
        'analysis': analysis,
        'results': results,
        'gaps': gaps,
        'gaps_grouped': grouped,
        'analysis_history': history,
        'latest_run': latest_run,
        'debug': debug,
        'quality_capability': capability_view,
        'open_questions': svc.repository.list_open_questions(knowledge_id,5) if hasattr(svc.repository,'list_open_questions') else [],
        'recurrence_level': recurrence.get('recurrence_risk_level') or 'UNKNOWN',
        'horizontal_action_needed': bool(recurrence.get('horizontal_action_needed')),
        'is_common_issue': bool(recurrence.get('is_common_issue')),
        'original_rows': _readable_rows(_safe_json((d.get('raw') or {}).get('raw_json'))),
        'normalized_rows': _readable_rows(normalized),
        'ai_summary': [
            {'label':'为什么发生','value':_ev((results.get('occurrence') or {}).get('root_cause_summary')) or '未分析'},
            {'label':'为什么流出','value':_ev((results.get('escape') or {}).get('escape_cause_summary')) or '未分析'},
            {'label':'再发风险','value':recurrence.get('recurrence_risk_level') or '未分析'},
            {'label':'主要能力缺口','value':' / '.join(x for x in ['Technical' if grouped.get('TECHNICAL') else '', 'Management' if grouped.get('MANAGEMENT') else '', 'Governance' if grouped.get('GOVERNANCE') else ''] if x) or '未识别'},
            {'label':'建议防控','value':next((g.get('recommended_action') or g.get('recommended_control') for g in gaps if g.get('recommended_action') or g.get('recommended_control')), '未形成')}
        ],
    }


def create_app(db_path):
    app = FastAPI(title='Quality Issue Knowledge', version='1.0-RC4')
    svc = KnowledgeIssueService(IssueKnowledgeRepository(db_path))
    capability_extension = QualityCapabilityExtension(db_path)
    app.state.quality_capability_extension = capability_extension
    batch_jobs = BatchAnalysisJobManager(db_path)
    app.state.batch_analysis_jobs = batch_jobs

    def project_capability(knowledge_id):
        issue = svc.get_issue(knowledge_id)
        if not issue:
            return None
        analyses = {stage: svc.get_latest_analysis(knowledge_id, stage) for stage in ('occurrence','escape','recurrence','capability_gap')}
        return capability_extension.project_issue(issue, analyses)

    def run_batch_with_projection(selected_ids, **kwargs):
        result = svc.run_batch_analysis(selected_ids, BASE.parent.parent, **kwargs)
        for knowledge_id in selected_ids:
            project_capability(knowledge_id)
        return result
    product_repo = ProductConfigRepository(db_path)
    for product in product_repo.list(False):
        register_product_adapter(product['product_code'])
    mapping_svc = MappingConfigurationService(MappingConfigurationRepository(db_path))
    app.state.mapping_configuration_service = mapping_svc
    human_svc = HumanAnalysisService(HumanAnalysisRepository(db_path))
    report_svc = LegacyProductQualityReportService(svc, BASE.parent.parent)
    app.state.human_analysis_service = human_svc
    app.state.knowledge_issue_service = svc
    app.state.product_config_repository = product_repo
    intake_svc = IntakeSessionService()
    material_repo = MaterialRepository(db_path)
    material_svc = MaterialImportService(material_repo)
    scenario_repo = ScenarioRepository(db_path)
    scenario_generation_svc = ScenarioGenerationService(svc,scenario_repo,BASE.parent.parent)
    app.state.material_repository = material_repo
    app.state.scenario_repository = scenario_repo
    app.state.scenario_generation_service = scenario_generation_svc
    app.state.intake_session_service = intake_svc
    app.mount('/static', StaticFiles(directory=BASE / 'static'), name='static')
    tpl = Jinja2Templates(directory=BASE / 'templates')
    tpl.env.globals['ev'] = _ev
    tpl.env.globals['confidence'] = _confidence
    tpl.env.globals['zh_value'] = zh_value
    from .scenario_asset_pages import create_asset_router
    app.include_router(create_asset_router(scenario_repo,tpl,scenario_generation_svc))

    def filters(req):
        return {k: v for k in ['business_type', 'business_issue_id', 'product', 'platform', 'severity', 'issue_type', 'issue_domain', 'year', 'month'] if (v := req.query_params.get(k))}

    def analysis_profile_from_form(domain_profile='', issue_types=None, lifecycle_phase=''):
        return normalize_analysis_profile({'domain_profile': domain_profile, 'issue_types': issue_types or [], 'lifecycle_phase': lifecycle_phase})

    def save_and_import(f, bt, issue_domain='AUTO'):
        if Path(f.filename or '').suffix.lower() not in ALLOWED:
            raise HTTPException(400, '仅支持 .xlsx / .xlsm')
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / Path(f.filename).name
            p.write_bytes(f.file.read())
            return svc.import_file(p, bt or None, issue_domain=issue_domain)

    @app.get('/', include_in_schema=False)
    def root():
        return RedirectResponse('/issues')

    @app.get('/import', response_class=HTMLResponse, include_in_schema=False)
    def import_page(request: Request):
        return tpl.TemplateResponse(request, 'import.html', {'products': product_repo.list()})

    material_workbenches={
        'itr':{'group_code':'ITR','title':'ITR问题工作台','description':'现场问题、客户影响、发生阶段与处理过程','scope':'软件 / 硬件 / 机械 / 跨领域'},
        'cs':{'group_code':'ITR-CS','title':'ITR彻底解决工作台','description':'责任认定、技术根因、永久措施、MRC与标准化','scope':'软件 / 硬件 / 机械 / 跨领域'},
        'software-operations':{'group_code':'SW-OPS','title':'软件问题考核工作台','description':'软件问题KPI、审核状态、责任单位与运营考核','scope':'仅软件'},
    }

    @app.get('/materials', include_in_schema=False)
    def materials_root():
        return RedirectResponse('/materials/itr',303)

    @app.get('/materials/{workbench}', response_class=HTMLResponse, include_in_schema=False)
    def materials_page(request: Request, workbench: str, q: str = '', domain: str = '', month: str = '', year: str = '', page: int = 1):
        workspace=material_workbenches.get(workbench)
        if not workspace:raise HTTPException(404,'MATERIAL_WORKBENCH_NOT_FOUND')
        result=material_repo.search_materials(workspace['group_code'],q=q,domain=domain,month=month,year=year,page=page)
        return tpl.TemplateResponse(request, 'materials.html', {
            'groups': material_repo.groups(False), 'items': result['items'], 'listing':result,
            'filters':{'q':q,'domain':domain,'month':month,'year':year},'group_code': workspace['group_code'], 'result': None, 'workbench':workbench, 'workspace':workspace,
        })

    @app.get('/materials/{workbench}/{material_id}', response_class=HTMLResponse, include_in_schema=False)
    def material_detail(request: Request, workbench: str, material_id: str):
        workspace=material_workbenches.get(workbench)
        if not workspace:raise HTTPException(404,'MATERIAL_WORKBENCH_NOT_FOUND')
        item=material_repo.material(material_id)
        if not item or item['group_code']!=workspace['group_code']:raise HTTPException(404,'MATERIAL_NOT_FOUND')
        return tpl.TemplateResponse(request,'material_detail.html',{'item':item,'workspace':workspace,'workbench':workbench,'related':material_repo.related_materials(item['canonical_itr'],material_id)})

    @app.post('/materials/{workbench}/{material_id}/review', include_in_schema=False)
    def material_review_save(workbench: str, material_id: str, review_status: str = Form('DRAFT'), analysis_summary: str = Form(''), root_cause: str = Form(''), improvement_action: str = Form(''), reviewer: str = Form('')):
        workspace=material_workbenches.get(workbench)
        item=material_repo.material(material_id)
        if not workspace or not item or item['group_code']!=workspace['group_code']:raise HTTPException(404,'MATERIAL_NOT_FOUND')
        try:material_repo.save_review(material_id,review_status=review_status,analysis_summary=analysis_summary,root_cause=root_cause,improvement_action=improvement_action,reviewer=reviewer)
        except ValueError as error:raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/materials/{workbench}/{material_id}#independent-review',303)

    @app.post('/materials/import', response_class=HTMLResponse, include_in_schema=False)
    def materials_import(request: Request, file: UploadFile = File(...), group_code: str = Form(...), header_rows: int = Form(2), workbench: str = Form(...), reporting_year: str = Form('')):
        workspace=material_workbenches.get(workbench)
        if not workspace or workspace['group_code']!=group_code:raise HTTPException(400,'WORKBENCH_GROUP_MISMATCH')
        if Path(file.filename or '').suffix.lower() not in ALLOWED:
            raise HTTPException(400, '仅支持 .xlsx / .xlsm')
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/Path(file.filename or 'upload.xlsx').name;path.write_bytes(file.file.read())
            try: result=material_svc.import_file(path,group_code,header_rows,reporting_year=reporting_year)
            except ValueError as error: raise HTTPException(400,str(error)) from error
        listing=material_repo.search_materials(group_code)
        return tpl.TemplateResponse(request, 'materials.html', {
            'groups': material_repo.groups(False), 'items': listing['items'],
            'listing':listing,'filters':{'q':'','domain':'','month':'','year':''},
            'group_code': group_code, 'result': result, 'workbench':workbench, 'workspace':workspace,
        })

    @app.post('/materials/software-operations/batch-year', include_in_schema=False)
    def software_operation_batch_year(material_ids: list[str] = Form(default=[]), reporting_year: str = Form(...)):
        try:material_repo.set_reporting_year(material_ids,reporting_year)
        except ValueError as error:raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/materials/software-operations?year={reporting_year}',303)

    @app.get('/settings/associations', response_class=HTMLResponse, include_in_schema=False)
    def association_settings(request: Request):
        return tpl.TemplateResponse(request,'association_settings.html',{'rules':material_repo.rules(),'preview':material_repo.link_preview()})

    @app.post('/settings/associations/{rule_id}', include_in_schema=False)
    def association_rule_save(rule_id: str, source_field: str = Form(...), target_field: str = Form(...), transform: str = Form(...), status: str = Form(...)):
        try:material_repo.update_rule(rule_id,source_field=source_field,target_field=target_field,transform=transform,status=status)
        except KeyError:raise HTTPException(404,'RULE_NOT_FOUND')
        except ValueError as error:raise HTTPException(400,str(error)) from error
        material_repo.refresh_links()
        return RedirectResponse('/settings/associations',303)

    @app.get('/quality-scenarios', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenarios(request: Request, ipmt: str = '', spdt: str = '', product_model: str = '', industry: str = '', customer_name: str = '', q: str = '', status: str = '', generation_id: str = '', activity_code: str = '', experience_code: str = '', qiu_code: str = '', quality_code: str = '', typical_problem_code: str = '', environment_code: str = '', deleted: int = 0, protected: int = 0):
        taxonomy=scenario_repo.taxonomy()
        filters={'ipmt':ipmt,'spdt':spdt,'product_model':product_model,'industry':industry,'customer_name':customer_name,'q':q,'status':status,'generation_id':generation_id,'activity_code':activity_code,'experience_code':experience_code,'qiu_code':qiu_code,'quality_code':quality_code,'typical_problem_code':typical_problem_code,'environment_code':environment_code}
        return tpl.TemplateResponse(request,'quality_scenarios.html',{'items':scenario_repo.scenarios(**filters),'options':scenario_repo.scope_options(),'filters':filters,'taxonomy':taxonomy,'lifecycle_labels':{x['lifecycle_code']:x['label_zh'] for x in taxonomy['lifecycles']},'activity_labels':{x['activity_code']:x['label_zh'] for x in taxonomy['activities']},'generation':scenario_repo.generation(generation_id) if generation_id else None,'deleted':deleted,'protected':protected})

    @app.post('/quality-scenarios/delete-generated', include_in_schema=False)
    def quality_scenario_delete_generated(ipmt: str = Form(''), spdt: str = Form(''), product_model: str = Form(''), industry: str = Form(''), customer_name: str = Form(''), q: str = Form(''), status: str = Form(''), generation_id: str = Form(''), activity_code: str = Form(''), experience_code: str = Form(''), qiu_code: str = Form(''), quality_code: str = Form(''), typical_problem_code: str = Form(''), environment_code: str = Form('')):
        filters={'ipmt':ipmt,'spdt':spdt,'product_model':product_model,'industry':industry,'customer_name':customer_name,'q':q,'status':status,'generation_id':generation_id,'activity_code':activity_code,'experience_code':experience_code,'qiu_code':qiu_code,'quality_code':quality_code,'typical_problem_code':typical_problem_code,'environment_code':environment_code}
        result=scenario_repo.delete_generated_scenarios(**filters)
        kept={key:value for key,value in filters.items() if value}
        kept.update({'deleted':result['deleted'],'protected':result['protected']})
        return RedirectResponse('/quality-scenarios?'+urlencode(kept),303)

    @app.get('/quality-scenarios/standardize', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_standardize_page(request: Request, batch_id: str = ''):
        items=scenario_repo.standardization_items()
        counts={key:sum(1 for x in items if x['quality_classification_status']==key) for key in ('NOT_ANALYZED','RUNNING','PENDING_CONFIRMATION','CONFIRMED','FAILED')}
        batch=scenario_repo.standardization_batch(batch_id) if batch_id else None
        return tpl.TemplateResponse(request,'quality_scenario_standardize.html',{'items':items,'counts':counts,'running':bool(counts['RUNNING']) or bool(batch and batch['status'] in {'QUEUED','RUNNING'}),'batch':batch,'batches':scenario_repo.standardization_batches()})

    @app.post('/quality-scenarios/standardize', include_in_schema=False)
    def quality_scenario_standardize(selected_ids: list[str] = Form([])):
        if not selected_ids:raise HTTPException(400,'QUALITY_SCENARIO_SELECTION_REQUIRED')
        batch_id=f'QSB-{uuid.uuid4().hex}';scenario_repo.create_standardization_batch(batch_id,list(selected_ids))
        threading.Thread(target=scenario_generation_svc.standardize_existing,args=(list(selected_ids),batch_id),daemon=True,name=f'scenario-standardize-{batch_id[-8:]}').start()
        return RedirectResponse(f'/quality-scenarios/standardize?batch_id={batch_id}',303)

    @app.post('/quality-scenarios/standardize/{batch_id}/retry', include_in_schema=False)
    def quality_scenario_standardize_retry(batch_id: str):
        batch=scenario_repo.standardization_batch(batch_id)
        if not batch:raise HTTPException(404,'QUALITY_STANDARDIZATION_BATCH_NOT_FOUND')
        ids=[x['scenario_id'] for x in batch['items'] if x['status']=='FAILED']
        if not ids:raise HTTPException(400,'NO_FAILED_STANDARDIZATION_ITEMS')
        new_id=f'QSB-{uuid.uuid4().hex}';scenario_repo.create_standardization_batch(new_id,ids,'WEB_RETRY')
        threading.Thread(target=scenario_generation_svc.standardize_existing,args=(ids,new_id),daemon=True,name=f'scenario-standardize-{new_id[-8:]}').start()
        return RedirectResponse(f'/quality-scenarios/standardize?batch_id={new_id}',303)

    @app.get('/quality-scenarios/{scenario_id}/capabilities', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_capabilities(request: Request, scenario_id: str):
        item=scenario_repo.scenario(scenario_id)
        if not item:raise HTTPException(404,'QUALITY_SCENARIO_NOT_FOUND')
        return tpl.TemplateResponse(request,'quality_scenario_capabilities.html',{'item':item,'dictionary':scenario_repo.capability_dictionary()})

    @app.post('/quality-scenarios/{scenario_id}/capabilities', include_in_schema=False)
    def quality_scenario_capability_save(scenario_id: str, capability_axis: str = Form(...), capability_code: str = Form(...), gap_description: str = Form(...), source_basis: str = Form(''), improvement_action: str = Form(''), verification_metric: str = Form(''), priority: str = Form('P1'), status: str = Form('OPEN')):
        try:scenario_repo.save_scenario_capability_gap(scenario_id,locals())
        except KeyError:raise HTTPException(404,'QUALITY_SCENARIO_NOT_FOUND')
        except ValueError as error:raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/quality-scenarios/{scenario_id}/capabilities',303)

    @app.post('/quality-scenarios/{scenario_id}/capabilities/{gap_id}/delete', include_in_schema=False)
    def quality_scenario_capability_delete(scenario_id: str,gap_id: str):
        try:scenario_repo.delete_scenario_capability_gap(scenario_id,gap_id)
        except KeyError:raise HTTPException(404,'SCENARIO_CAPABILITY_GAP_NOT_FOUND')
        return RedirectResponse(f'/quality-scenarios/{scenario_id}/capabilities',303)

    @app.get('/quality-scenarios/generate', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_generate_page(request: Request, product_code: str = '', start_month: str = '', end_month: str = '', preview: int = 0, job_id: str = '', source: str = 'cs_itr', ipmt: str = '', spdt: str = '', product_model: str = '', year: str = '', problem_domain: str = 'SOFTWARE', source_product: str = ''):
        from quality_knowledge.scenario_sources import scene_source_records
        group_choices=[g for g in material_repo.groups() if g['material_type'] in {'ITR_CS','ITR_SOURCE'}]
        requested_groups=request.query_params.getlist('group_ids')
        default_groups=[g['group_id'] for g in group_choices if g['group_code'] in {'ITR','ITR-CS'}]
        selected_groups=requested_groups or default_groups
        problem_domain=(problem_domain or 'SOFTWARE').strip().upper()
        filters={'product_code':product_code,'start_month':start_month,'end_month':end_month,'source':source,'ipmt':ipmt,'spdt':spdt,'product_model':product_model,'year':year,'problem_domain':problem_domain,'source_product':source_product,'group_ids':selected_groups}
        effective_source=source if source in {'operations','cs_itr'} else 'operations'
        options=scene_source_records(scenario_generation_svc,effective_source,{'product_code':product_code,'group_ids':selected_groups},metadata_only=True)
        choices={key:sorted({x[key] for x in options if x.get(key)}) for key in ('ipmt','spdt','product_model','year','source_product')}
        scope=None
        if preview and source in {'operations','cs_itr'}:
            rows=scene_source_records(scenario_generation_svc,source,filters)
            analysed=sum(bool(x['leakage_analysis']) for x in rows)
            from collections import Counter
            scope={'items':rows,'issue_count':len(rows),'analysed_count':analysed,'coverage_rate':round(analysed*100/len(rows),1) if rows else 0,
                   'domain_counts':dict(Counter(x.get('problem_domain') or 'UNKNOWN' for x in rows)),
                   'source_counts':dict(Counter(x.get('source_workbench') or 'unknown' for x in rows)),
                   'missing_leakage_count':sum(bool(x.get('missing_leakage')) for x in rows)}
        elif preview and product_code and start_month and end_month:
            scope=scenario_generation_svc.precheck(product_code,start_month,end_month)
        return tpl.TemplateResponse(request,'quality_scenario_generate.html',{'products':product_repo.list(),'generations':scenario_repo.generations(),'result':None,'scope':scope,'choices':choices,'job':scenario_repo.generation(job_id) if job_id else None,'filters':filters,'group_choices':group_choices})

    @app.get('/quality-scenarios/insights', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_insights(request: Request, status: str = ''):
        from quality_knowledge.scenario_assets import ScenarioAssets
        from quality_knowledge.scenario_decisions import decision_digest
        decision=decision_digest(ScenarioAssets(scenario_repo).report({'status':status}),scenario_repo)
        return tpl.TemplateResponse(request,'quality_scenario_insights.html',{'insights':scenario_repo.insights(status=status),'status':status,'decision':decision})

    @app.post('/quality-scenarios/generate', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_generate(request: Request, product_code: str = Form(...), start_month: str = Form(''), end_month: str = Form(''), selected_ids: list[str] = Form([]), source: str = Form('cs_itr'), ipmt: str = Form(''), spdt: str = Form(''), product_model: str = Form(''), year: str = Form(''), problem_domain: str = Form('SOFTWARE'), source_product: str = Form(''), group_ids: list[str] = Form([])):
        if not selected_ids:raise HTTPException(400,'SCENARIO_SOURCE_SELECTION_REQUIRED')
        if not scenario_repo.taxonomy_active(product_code):raise HTTPException(400,f'产品 {product_code} 尚未配置并激活场景词典，请先到场景词典配置中创建并激活')
        generation_id=f'QSG-{uuid.uuid4().hex}'
        if source in {'operations','cs_itr'}:
            from quality_knowledge.scenario_sources import scene_source_records
            rows=scene_source_records(scenario_generation_svc,source,{'product_code':product_code,'ipmt':ipmt,'spdt':spdt,'product_model':product_model,'year':year,'problem_domain':(problem_domain or 'SOFTWARE').upper(),'source_product':source_product,'start_month':start_month,'end_month':end_month,'group_ids':group_ids},selected_ids)
            if {x['knowledge_id'] for x in rows}!=set(selected_ids):raise HTTPException(409,'选中数据已变化或不属于当前筛选范围，请重新预览')
            scenario_generation_svc.save_source_snapshot(generation_id,rows)
        scenario_repo.create_generation(generation_id,product_code,start_month,end_month,len(selected_ids),'WEB_USER')
        threading.Thread(target=scenario_generation_svc.run_job,args=(generation_id,product_code,start_month,end_month,list(selected_ids)),daemon=True,name=f'scenario-{generation_id[-8:]}').start()
        return RedirectResponse(f'/quality-scenarios/generate?job_id={generation_id}',303)

    @app.get('/api/quality-scenario-generations/{generation_id}')
    def quality_scenario_generation_status(generation_id: str):
        job=scenario_repo.generation(generation_id)
        if not job:raise HTTPException(404,'SCENARIO_GENERATION_NOT_FOUND')
        return job

    @app.get('/quality-scenarios/generations/{generation_id}/issues', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_generation_issues(request: Request, generation_id: str):
        job=scenario_repo.generation(generation_id)
        if not job:raise HTTPException(404,'SCENARIO_GENERATION_NOT_FOUND')
        snapshot={row['knowledge_id']:row for row in scenario_generation_svc.source_snapshot(generation_id)}
        items=[]
        for row in scenario_repo.issue_classifications(generation_id):
            source=snapshot.get(row['knowledge_id'],{})
            items.append({**row,'source_workbench':source.get('source_workbench') or '',
                          'source_material_id':source.get('source_material_id') or ''})
        return tpl.TemplateResponse(request,'quality_scenario_generation_issues.html',{'job':job,'items':items})

    @app.post('/quality-scenarios/generations/delete-finished', include_in_schema=False)
    def quality_scenario_generation_delete_finished():
        deleted=scenario_repo.delete_finished_generations()
        return RedirectResponse(f'/quality-scenarios/generate?records_deleted={deleted}',303)

    @app.post('/quality-scenarios/generations/{generation_id}/delete', include_in_schema=False)
    def quality_scenario_generation_delete(generation_id: str):
        try:scenario_repo.delete_generation(generation_id)
        except KeyError:raise HTTPException(404,'SCENARIO_GENERATION_NOT_FOUND')
        except ValueError as error:raise HTTPException(409,str(error)) from error
        return RedirectResponse('/quality-scenarios/generate?records_deleted=1',303)

    @app.post('/quality-scenarios/generations/{generation_id}/retry', include_in_schema=False)
    def quality_scenario_generation_retry(generation_id: str, selected_ids: list[str] = Form([])):
        job=scenario_repo.generation(generation_id)
        if not job:raise HTTPException(404,'SCENARIO_GENERATION_NOT_FOUND')
        if not selected_ids:selected_ids=[x['knowledge_id'] for x in scenario_repo.issue_classifications(generation_id) if x['status'] in {'FAILED','REVIEW_REQUIRED','PENDING'}]
        if not selected_ids:raise HTTPException(400,'NO_RETRYABLE_SCENARIO_ISSUES')
        new_id=f'QSG-{uuid.uuid4().hex}';scenario_repo.create_generation(new_id,job['product_code'],job['start_month'],job['end_month'],len(selected_ids),'WEB_RETRY')
        snapshot=scenario_generation_svc.source_snapshot(generation_id)
        if snapshot:
            retry_rows=[r for r in snapshot if r['knowledge_id'] in selected_ids]
            if len(retry_rows)!=len(set(selected_ids)):raise HTTPException(400,'重试问题不属于原任务')
            scenario_generation_svc.save_source_snapshot(new_id,retry_rows)
        threading.Thread(target=scenario_generation_svc.run_job,args=(new_id,job['product_code'],job['start_month'],job['end_month'],selected_ids),daemon=True,name=f'scenario-{new_id[-8:]}').start()
        return RedirectResponse(f'/quality-scenarios/generate?job_id={new_id}',303)

    @app.get('/quality-scenarios/{scenario_id}', response_class=HTMLResponse, include_in_schema=False)
    @app.get('/quality-scenarios/new', response_class=HTMLResponse, include_in_schema=False)
    def quality_scenario_edit(request: Request, scenario_id: str = ''):
        item=scenario_repo.scenario(scenario_id) if scenario_id else None
        if scenario_id and not item:raise HTTPException(404,'QUALITY_SCENARIO_NOT_FOUND')
        product_code=(item or {}).get('product_code') or 'PLC';taxonomy=scenario_repo.taxonomy(product_code=product_code) or scenario_repo.taxonomy(product_code='PLC')
        return tpl.TemplateResponse(request,'quality_scenario_edit.html',{'item':item or {'scenario_id':'','scenario_code':'','name':'','product_code':product_code,'taxonomy_version_id':taxonomy['version_id'] if taxonomy else '','lifecycle_code':'','activity_code':'','scenario_chain':'','experience_requirement':'','concern_points':'','customer_perception':'','primary_typical_problem_code':'','secondary_typical_problem_codes':[],'primary_quality_concern_code':'','secondary_quality_concern_codes':[],'primary_customer_experience_statement':'','secondary_customer_experience_statements':[],'operating_environment':'','operating_condition':'','duration_frequency':'','disturbances':'','extreme_conditions':'','environment_condition_codes':[],'primary_experience_code':'','secondary_experience_codes':[],'quality_in_use_codes':[],'primary_quality_characteristic_code':'','secondary_quality_characteristic_codes':[],'quality_subcharacteristic_codes':[],'quality_classification_status':'PENDING_CONFIRMATION','quality_attribute':'','quality_subcharacteristic':'','failure_mode':'','failure_mechanism':'','trigger_conditions':'','preconditions':'','participating_systems':'','system_scale':'','user_type':'','affected_object':'','business_impact':'','recovery_method':'','applicable_boundary':'','validation_direction':'','measurement_suggestion':'','status':'DRAFT','scopes':{},'industry_variants':[]},'taxonomy':taxonomy,'options':scenario_repo.scope_options(),'quality_models':scenario_repo.quality_models(),'semantics':scenario_repo.semantic_dictionary(product_code)})

    @app.post('/quality-scenarios/save', include_in_schema=False)
    def quality_scenario_save(scenario_id: str = Form(''), scenario_code: str = Form(...), name: str = Form(...), product_code: str = Form('PLC'), taxonomy_version_id: str = Form(''), lifecycle_code: str = Form(''), activity_code: str = Form(''), customer_perception: str = Form(''), primary_typical_problem_code: str = Form(''), secondary_typical_problem_codes: list[str] = Form([]), primary_quality_concern_code: str = Form(''), secondary_quality_concern_codes: list[str] = Form([]), primary_customer_experience_statement: str = Form(''), secondary_customer_experience_statements: str = Form(''), operating_environment: str = Form(''), operating_condition: str = Form(''), duration_frequency: str = Form(''), disturbances: str = Form(''), extreme_conditions: str = Form(''), environment_condition_codes: list[str] = Form([]), primary_experience_code: str = Form(''), secondary_experience_codes: list[str] = Form([]), quality_in_use_codes: list[str] = Form([]), primary_quality_characteristic_code: str = Form(''), secondary_quality_characteristic_codes: list[str] = Form([]), quality_subcharacteristic_codes: list[str] = Form([]), quality_classification_status: str = Form('PENDING_CONFIRMATION'), confirmed_by: str = Form('WEB_USER'), experience_requirement: str = Form(''), concern_points: str = Form(''), quality_attribute: str = Form(''), quality_subcharacteristic: str = Form(''), failure_mode: str = Form(''), failure_mechanism: str = Form(''), trigger_conditions: str = Form(''), preconditions: str = Form(''), participating_systems: str = Form(''), system_scale: str = Form(''), user_type: str = Form(''), affected_object: str = Form(''), business_impact: str = Form(''), recovery_method: str = Form(''), applicable_boundary: str = Form(''), validation_direction: str = Form(''), measurement_suggestion: str = Form(''), status: str = Form('DRAFT'), ipmt: list[str] = Form([]), spdt: list[str] = Form([]), product_model: list[str] = Form([]), industry: list[str] = Form([]), customer_name: list[str] = Form([]), customer_level: list[str] = Form([]), customer_status: list[str] = Form([]), occurrence_phase: list[str] = Form([])):
        saved=scenario_repo.save_scenario(scenario_id,locals(),{'IPMT':ipmt,'SPDT':spdt,'PRODUCT_MODEL':product_model,'INDUSTRY':industry,'CUSTOMER_NAME':customer_name,'CUSTOMER_LEVEL':customer_level,'CUSTOMER_STATUS':customer_status,'OCCURRENCE_PHASE':occurrence_phase})
        return RedirectResponse(f'/quality-scenarios/{saved}',303)

    @app.get('/settings/scenario-semantics', response_class=HTMLResponse, include_in_schema=False)
    def scenario_semantics_settings(request: Request, product_code: str = 'PLC'):
        return tpl.TemplateResponse(request,'scenario_semantics.html',{'dictionary':scenario_repo.semantic_dictionary(product_code),'product_code':product_code,'products':product_repo.list()})

    @app.post('/settings/scenario-semantics/save', include_in_schema=False)
    def scenario_semantics_save(term_id: str = Form(''), term_type: str = Form(...), term_code: str = Form(...), label_zh: str = Form(...), definition: str = Form(...), inclusion_criteria: str = Form(''), exclusion_criteria: str = Form(''), aliases: str = Form(''), product_code: str = Form(''), status: str = Form('ACTIVE')):
        try:scenario_repo.save_semantic_term(locals())
        except ValueError as error:raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/settings/scenario-semantics?product_code={product_code}',303)

    @app.post('/settings/scenario-semantics/{term_id}/review', include_in_schema=False)
    def scenario_semantics_review(term_id: str, action: str = Form(...), target_code: str = Form(''), product_code: str = Form('PLC')):
        try:scenario_repo.review_semantic_term(term_id,action,target_code)
        except KeyError:raise HTTPException(404,'SEMANTIC_TERM_NOT_FOUND')
        except ValueError as error:raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/settings/scenario-semantics?product_code={product_code}',303)

    @app.post('/quality-scenarios/{scenario_id}/delete', include_in_schema=False)
    def quality_scenario_delete(scenario_id: str):
        try:scenario_repo.delete_scenario(scenario_id)
        except KeyError:raise HTTPException(404,'QUALITY_SCENARIO_NOT_FOUND')
        return RedirectResponse('/quality-scenarios',303)

    @app.get('/settings/scenario-taxonomy', response_class=HTMLResponse, include_in_schema=False)
    def scenario_taxonomy_settings(request: Request, product_code: str = 'PLC'):
        taxonomy=scenario_repo.taxonomy(product_code=product_code)
        return tpl.TemplateResponse(request,'scenario_taxonomy.html',{'taxonomy':taxonomy,'versions':scenario_repo.versions(product_code),'products':product_repo.list(),'product_code':product_code})

    @app.post('/settings/scenario-taxonomy/draft', include_in_schema=False)
    def scenario_taxonomy_draft(product_code: str = Form(...), source_product_code: str = Form('')):
        scenario_repo.create_draft(product_code,source_product_code);return RedirectResponse(f'/settings/scenario-taxonomy?product_code={product_code}',303)

    @app.post('/settings/scenario-taxonomy/lifecycle', include_in_schema=False)
    def scenario_lifecycle_save(version_id: str = Form(...), lifecycle_code: str = Form(...), label_zh: str = Form(...), description: str = Form(''), value_statement: str = Form(''), objective: str = Form(''), enabled: str = Form('')):
        scenario_repo.save_lifecycle(version_id,lifecycle_code,label_zh,description,enabled=='on',value_statement,objective);product=scenario_repo.taxonomy(version_id).get('product_code','PLC');return RedirectResponse(f'/settings/scenario-taxonomy?product_code={product}#lifecycles',303)

    @app.post('/settings/scenario-taxonomy/activity', include_in_schema=False)
    def scenario_activity_save(version_id: str = Form(...), lifecycle_code: str = Form(...), activity_code: str = Form(...), label_zh: str = Form(...), chain_text: str = Form(''), description: str = Form(''), participating_systems: str = Form(''), objective: str = Form(''), enabled: str = Form('')):
        scenario_repo.save_activity(version_id,lifecycle_code,activity_code,label_zh,chain_text,description,enabled=='on',participating_systems,objective);product=scenario_repo.taxonomy(version_id).get('product_code','PLC');return RedirectResponse(f'/settings/scenario-taxonomy?product_code={product}#business-activities',303)

    @app.post('/settings/scenario-taxonomy/activities/bulk', include_in_schema=False)
    async def scenario_activities_bulk(request: Request):
        form=await request.form();version_id=str(form.get('version_id') or '')
        codes=[str(x) for x in form.getlist('activity_code')];enabled=set(str(x) for x in form.getlist('enabled_code'))
        fields={name:[str(x) for x in form.getlist(name)] for name in ('lifecycle_code','label_zh','chain_text','description','participating_systems','objective')}
        items=[]
        for index,code in enumerate(codes):
            items.append({'activity_code':code,'enabled':code in enabled,**{name:(values[index] if index<len(values) else '') for name,values in fields.items()}})
        scenario_repo.save_activities(version_id,items);product=scenario_repo.taxonomy(version_id).get('product_code','PLC')
        return RedirectResponse(f'/settings/scenario-taxonomy?product_code={product}&saved={len(items)}#business-activities',303)

    @app.post('/settings/scenario-taxonomy/import', include_in_schema=False)
    def scenario_taxonomy_import(file: UploadFile = File(...), product_code: str = Form(...)):
        if Path(file.filename or '').suffix.lower() not in ALLOWED:raise HTTPException(400,'仅支持 .xlsx 或 .xlsm 模板')
        version_id=scenario_repo.create_draft(product_code)
        try:result=scenario_repo.import_taxonomy_workbook(version_id,file.file)
        except (KeyError,ValueError) as error:raise HTTPException(400,f'场景词典模板校验失败：{error}') from error
        finally:file.file.close()
        return RedirectResponse(f"/settings/scenario-taxonomy?product_code={product_code}&imported={result['activity_count']}#business-activities",303)

    @app.post('/settings/scenario-taxonomy/{version_id}/activate', include_in_schema=False)
    def scenario_taxonomy_activate(version_id: str):
        product=scenario_repo.taxonomy(version_id).get('product_code','PLC');scenario_repo.activate(version_id);return RedirectResponse(f'/settings/scenario-taxonomy?product_code={product}',303)

    def _build_intake_preview(file: UploadFile, business_type: str = '', issue_domain: str = 'AUTO', mapping_config_id: str = ''):
        if Path(file.filename or '').suffix.lower() not in ALLOWED:
            raise HTTPException(400, '仅支持 .xlsx / .xlsm')
        content=file.file.read(); meta=intake_svc.create(file.filename or 'upload.xlsx',content); diagnostic_id=meta['intake_session_id']
        _intake_diag('UPLOAD_SAVED',diagnostic_id,file_name=meta['source_file'],size_bytes=len(content),sha256=hashlib.sha256(content).hexdigest(),requested_business_type=(business_type or 'AUTO'))
        # Header/structure detection must see original cell content.  Formula
        # caches are frequently absent in exported workbooks and previously
        # collapsed a 51-column sheet into one visible header.
        path=Path(meta['path']); wb=load_workbook(path,read_only=True,data_only=False)
        requested=(business_type or '').upper() or None
        selected_mapping=mapping_svc.get_config_version(mapping_config_id) if mapping_config_id else None
        if selected_mapping and (not requested or selected_mapping['business_type'] != requested):
            raise HTTPException(400,'MAPPING_PRODUCT_MISMATCH')
        _intake_diag('WORKBOOK_OPENED',diagnostic_id,sheets=wb.sheetnames)
        candidates=[]
        for sn in wb.sheetnames:
            sheet=wb[sn]; declared_dimension,actual_dimension=repair_read_only_dimensions(sheet)
            if selected_mapping:
                scanned=list(sheet.iter_rows(min_row=1,max_row=min(sheet.max_row or 50,50),values_only=True))
                choices=[]
                for idx,row in enumerate(scanned,1):
                    cells=[clean(value) for value in row if clean(value)]
                    if cells: choices.append(((runtime_detection_score(selected_mapping,row),len(cells),-idx),idx))
                if choices:
                    score,header=max(choices)
                    det={'header_row':header,'business_type':requested,'score':score[0],'recognition_mode':'DRAFT_MAPPING_SCORE' if score[0] else 'DRAFT_STRUCTURAL'}
                else: det=None
            else:
                det=svc._detect_header(sheet,requested)
            if not det:
                _intake_diag('SHEET_REJECTED',diagnostic_id,sheet=sn,declared_dimension=declared_dimension,actual_dimension=actual_dimension,max_row=sheet.max_row,max_column=sheet.max_column,reason='HEADER_NOT_DETECTED')
                continue
            header=det['header_row']; values=next(sheet.iter_rows(min_row=header,max_row=header,values_only=True),())
            width=sum(1 for value in values if clean(value)); data_rows=sum(1 for row in sheet.iter_rows(min_row=header+1,values_only=True) if any(clean(value) for value in row))
            headers=[clean(value) for value in values if clean(value)]
            _intake_diag('SHEET_CANDIDATE',diagnostic_id,sheet=sn,declared_dimension=declared_dimension,actual_dimension=actual_dimension,max_row=sheet.max_row,max_column=sheet.max_column,header_row=header,header_width=width,data_rows=data_rows,detection_score=det.get('score',0),recognition_mode=det.get('recognition_mode'),headers=headers)
            candidates.append((det.get('score',0),width,data_rows,{'sheet':sn,**det,'header_width':width,'data_row_count':data_rows}))
        detected=max(candidates,key=lambda x:(x[0],x[1],x[2]))[3] if candidates else None
        if not detected:
            _intake_diag('PREVIEW_REJECTED',diagnostic_id,reason='NO_HEADER_CANDIDATE')
            intake_svc.discard(meta['intake_session_id']); raise HTTPException(400,f'无法识别业务类型或表头；诊断ID：{diagnostic_id}；日志：{DIAG_FILE}')
        bt=detected['business_type']; preview=mapping_svc.preview_file(path,bt,detected['sheet'],detected['header_row'],mapping_config_id or None); effective=selected_mapping or mapping_svc.get_effective_config(bt)
        ws=wb[detected['sheet']]; row_count=detected.get('data_row_count',max(0,ws.max_row-detected['header_row']))
        preview.update({'intake_session_id':meta['intake_session_id'],'detected_business_type':bt,'header_row':detected['header_row'],'detection_score':detected['score'],'recognition_mode':detected.get('recognition_mode','MAPPING_SCORE'),'data_row_count':row_count,'mapping_config_id':effective['config_id'],'mapping_config_version':effective['version']})
        preview['diagnostic_id']=diagnostic_id; preview['diagnostic_log']=str(DIAG_FILE)
        _intake_diag('PREVIEW_COMPLETED',diagnostic_id,selected_sheet=detected['sheet'],header_row=detected['header_row'],header_width=preview.get('total_source_fields'),data_rows=row_count,matched_structured=preview.get('matched_structured'),matched_extension=preview.get('matched_extension'),unmatched=preview.get('unmatched'),conflict=preview.get('conflict'),required_missing=preview.get('required_missing'))
        preview['mapping_config_status']=effective['status']
        meta.update({'business_type':bt,'sheet':detected['sheet'],'available_sheets':wb.sheetnames,'mapping_config_id':effective['config_id'],'mapping_config_version':effective['version'],'issue_domain':str(issue_domain or 'AUTO').upper(),'preview':preview}); intake_svc.save(meta); return preview

    @app.post('/import/preview', response_class=HTMLResponse, include_in_schema=False)
    def import_preview(request: Request, file: UploadFile = File(...), business_type: str = Form(''), issue_domain: str = Form('AUTO')):
        try: result=_build_intake_preview(file,business_type,issue_domain)
        except HTTPException as e: return tpl.TemplateResponse(request,'import.html',{'error':e.detail,'products':product_repo.list()},status_code=e.status_code)
        return tpl.TemplateResponse(request,'import_preview.html',{'result':result})

    @app.post('/import/confirm', include_in_schema=False)
    def import_confirm(intake_session_id: str = Form(...)):
        try: meta=intake_svc.get(intake_session_id)
        except IntakeSessionError as e: raise HTTPException(404,str(e))
        if meta.get('status')=='COMMITTED': return RedirectResponse('/imports/'+meta['batch_id'],303)
        effective=mapping_svc.get_effective_config(meta['business_type'])
        if not effective or effective['config_id']!=meta.get('mapping_config_id'): raise HTTPException(409,'MAPPING_CHANGED_AFTER_PREVIEW: 请重新预检后再导入')
        pv=meta.get('preview') or {}
        if pv.get('conflict') or pv.get('required_missing'): raise HTTPException(409,'PREVIEW_HAS_BLOCKING_ERRORS: 请先解决冲突或必填字段缺失')
        r=svc.import_file(Path(meta['path']),meta['business_type'],sheet=meta.get('sheet'),issue_domain=meta.get('issue_domain','AUTO')); intake_svc.mark_committed(intake_session_id,r['batch_id']); return RedirectResponse('/imports/'+r['batch_id'],303)

    @app.post('/import', include_in_schema=False)
    def import_submit(request: Request, file: UploadFile = File(...), business_type: str = Form(''), issue_domain: str = Form('AUTO')):
        try: result=_build_intake_preview(file,business_type,issue_domain)
        except HTTPException as e: return tpl.TemplateResponse(request,'import.html',{'error':e.detail,'products':product_repo.list()},status_code=e.status_code)
        return tpl.TemplateResponse(request,'import_preview.html',{'result':result})

    @app.get('/imports/{batch_id}', response_class=HTMLResponse, include_in_schema=False)
    def import_result(request: Request, batch_id: str):
        return tpl.TemplateResponse(request, 'import_result.html', {'result': svc.get_import_batch(batch_id)})

    @app.get('/issues', response_class=HTMLResponse, include_in_schema=False)
    def issues(request: Request, page: int = 1, page_size: int = 100, limit: int | None = None):
        f = filters(request)
        scoped_issue_ids = []
        for value in request.query_params.getlist('business_issue_ids'):
            scoped_issue_ids.extend(x.strip() for x in value.split(',') if x.strip())
        scoped_knowledge_ids = []
        for value in request.query_params.getlist('knowledge_ids'):
            scoped_knowledge_ids.extend(x.strip() for x in value.split(',') if x.strip())
        if scoped_issue_ids:
            f['business_issue_ids'] = scoped_issue_ids
        if scoped_knowledge_ids:
            f['knowledge_ids'] = scoped_knowledge_ids
        gap_dimension = (request.query_params.get('gap_dimension') or '').upper()
        gap_category = (request.query_params.get('gap_category') or '').strip()
        if gap_dimension:
            f['gap_dimension'] = gap_dimension
        if gap_category:
            f['gap_category'] = gap_category
        matrix_axis=(request.query_params.get('matrix_axis') or '').upper()
        matrix_code=(request.query_params.get('matrix_code') or '').strip()
        if matrix_axis and matrix_code and gap_dimension and gap_category:
            try:
                matrix_ids=capability_extension.matrix_issue_ids(
                    matrix_axis=matrix_axis,axis_code=matrix_code,capability_axis=gap_dimension,
                    capability_code=gap_category,business_type=f.get('business_type',''))
                f['knowledge_ids']=matrix_ids or ['__NO_MATRIX_MATCH__']
            except ValueError as error:
                raise HTTPException(400,str(error)) from error
        # Keep the old limit parameter compatible, but make paging explicit.  The
        # candidate set is complete before UI-only filters are applied, so totals
        # and semantic drill-downs are no longer silently capped at 100 rows.
        if limit is not None:
            page_size = limit
        page_size = max(10, min(int(page_size or 100), 200))
        page = max(1, int(page or 1))
        candidate_count = svc.count_issues(f)
        items = svc.query_issues(f, max(candidate_count, 1))
        human_field_id = request.query_params.get('human_field_id') or ''
        human_value = request.query_params.get('human_value') or ''
        if human_field_id or human_value:
            allowed=set(human_svc.query_issue_ids(human_field_id or None,human_value or None))
            items=[x for x in items if x['knowledge_id'] in allowed]
        for x in items:
            kid = x['knowledge_id']
            hist = svc.get_analysis_history(kid)
            x['ai_status'] = hist[0]['status'] if hist else 'NOT_ANALYZED'
            failed = next((r for r in hist if str(r.get('status') or '').upper() == 'FAILED'), None)
            x['failed_stage'] = failed.get('analysis_type') if failed else None
            rec = svc.get_latest_analysis(kid, 'recurrence')
            rr = rec.get('result') if rec else {}
            x['recurrence_risk_level'] = rr.get('recurrence_risk_level') or 'UNKNOWN'
            issue = svc.get_issue(kid)
            x['human_done'] = bool(human_svc.get_analysis(kid, issue['issue_version_id'])) if issue else False
            gs = svc.query_capability_gaps(kid)
            x['gap_summary'] = {
                'TECHNICAL': sum(1 for g in gs if g.get('dimension') == 'TECHNICAL'),
                'MANAGEMENT': sum(1 for g in gs if g.get('dimension') == 'MANAGEMENT'),
                'GOVERNANCE': sum(1 for g in gs if g.get('dimension') == 'GOVERNANCE'),
            }
        q=(request.query_params.get('q') or '').strip().lower()
        recurrence_risk=(request.query_params.get('recurrence_risk') or '').upper()
        analysis_status=(request.query_params.get('analysis_status') or '').upper()
        human_status=(request.query_params.get('human_status') or '').upper()
        all_items=list(items)
        if q: items=[x for x in items if q in ' '.join(str(x.get(k) or '') for k in ('business_issue_id','title','description','product','platform')).lower()]
        if recurrence_risk: items=[x for x in items if str(x.get('recurrence_risk_level') or '').upper()==recurrence_risk]
        if analysis_status: items=[x for x in items if str(x.get('ai_status') or '').upper()==analysis_status]
        if human_status=='DONE': items=[x for x in items if x.get('human_done')]
        elif human_status=='PENDING': items=[x for x in items if not x.get('human_done')]
        metrics={'total':len(all_items),'pending_ai':sum(1 for x in all_items if x['ai_status']=='NOT_ANALYZED'),'failed_ai':sum(1 for x in all_items if x['ai_status']=='FAILED'),'high_risk':sum(1 for x in all_items if x['recurrence_risk_level']=='HIGH'),'pending_human':sum(1 for x in all_items if not x['human_done'])}
        filtered_total = len(items)
        total_pages = max(1, (filtered_total + page_size - 1) // page_size)
        page = min(page, total_pages)
        page_items = items[(page - 1) * page_size:page * page_size]
        for row in page_items:
            if row.get('year'):
                row['month'] = f"{row['year']}年 {row.get('month') or '-'}"
        base_params = [(k, v) for k, v in request.query_params.multi_items() if k not in {'page','limit','page_size'}]
        base_params.append(('page_size', str(page_size)))
        def page_url(target):
            return '/issues?' + urlencode(base_params + [('page', str(target))])
        pagination = {
            'page': page, 'page_size': page_size, 'total_pages': total_pages, 'total': filtered_total,
            'previous_url': page_url(page - 1) if page > 1 else '',
            'next_url': page_url(page + 1) if page < total_pages else '',
            'start': (page - 1) * page_size + 1 if filtered_total else 0,
            'end': min(page * page_size, filtered_total),
        }
        all_issue_count = svc.count_issues({})
        month_options = sorted({str(x.get('month')).strip() for x in svc.query_issues({}, max(all_issue_count, 1)) if str(x.get('month') or '').strip()})
        year_options = sorted({str(x.get('year')).strip() for x in svc.query_issues({}, max(all_issue_count, 1)) if str(x.get('year') or '').strip()}, reverse=True)
        return tpl.TemplateResponse(request, 'issues.html', {'items': page_items, 'filtered_total': filtered_total, 'pagination': pagination, 'metrics':metrics, 'filters': f, 'q':q, 'recurrence_risk':recurrence_risk, 'analysis_status':analysis_status, 'human_status':human_status, 'human_fields': human_svc.list_field_definitions(True), 'human_field_id': human_field_id, 'human_value': human_value, 'products': product_repo.list(), 'year_options': year_options, 'month_options': month_options})

    @app.get('/issues/{knowledge_id}', response_class=HTMLResponse, include_in_schema=False)
    def issue_detail(request: Request, knowledge_id: str):
        vm = _build_issue_view(svc, knowledge_id, capability_extension)
        if not vm:
            raise HTTPException(404, 'NOT_FOUND')
        issue = svc.get_issue(knowledge_id)
        sequence = svc.query_issues({}, 100000)
        sequence_index = next((idx for idx, row in enumerate(sequence) if row.get('knowledge_id') == knowledge_id), -1)
        vm['previous_issue'] = sequence[sequence_index - 1] if sequence_index > 0 else None
        vm['next_issue'] = sequence[sequence_index + 1] if 0 <= sequence_index < len(sequence) - 1 else None
        vm['issue_position'] = sequence_index + 1 if sequence_index >= 0 else None
        vm['issue_total'] = len(sequence)
        vm['human_fields'] = human_svc.list_field_definitions(True)
        vm['human_analysis'] = human_svc.get_analysis(knowledge_id, issue['issue_version_id']) if issue else None
        material_repo.refresh_links()
        vm['linked_materials'] = material_repo.materials_for_issue(knowledge_id)
        vm.update({'analysis_agents':list_quality_issue_agents(BASE.parent.parent),'domain_profiles': DOMAIN_PROFILES, 'domain_labels': DOMAIN_LABELS, 'issue_types': ISSUE_TYPES, 'issue_type_labels': ISSUE_TYPE_LABELS, 'lifecycle_phases': LIFECYCLE_PHASES, 'lifecycle_labels': LIFECYCLE_LABELS})
        return tpl.TemplateResponse(request, 'issue_detail.html', vm)

    @app.post('/issues/{knowledge_id}/period', include_in_schema=False)
    def issue_period_save(knowledge_id: str, year: str = Form(...), month: str = Form(...)):
        try:
            svc.repository.update_issue_period(knowledge_id, year, month, 'web')
        except KeyError:
            raise HTTPException(404, 'issue not found')
        except ValueError as error:
            raise HTTPException(400, str(error)) from error
        return RedirectResponse(f'/issues/{knowledge_id}#issue-period', 303)

    @app.post('/api/import/preview')
    def api_import_preview(file: UploadFile = File(...), business_type: str = Form(''), issue_domain: str = Form('AUTO')):
        return _build_intake_preview(file,business_type,issue_domain)

    @app.post('/api/import/confirm')
    def api_import_confirm(intake_session_id: str = Form(...)):
        try: meta=intake_svc.get(intake_session_id)
        except IntakeSessionError as e: raise HTTPException(404,str(e))
        if meta.get('status')=='COMMITTED': return svc.get_import_batch(meta['batch_id'])
        effective=mapping_svc.get_effective_config(meta['business_type'])
        if not effective or effective['config_id']!=meta.get('mapping_config_id'): raise HTTPException(409,'MAPPING_CHANGED_AFTER_PREVIEW')
        pv=meta.get('preview') or {}
        if pv.get('conflict') or pv.get('required_missing'): raise HTTPException(409,'PREVIEW_HAS_BLOCKING_ERRORS')
        r=svc.import_file(Path(meta['path']),meta['business_type'],sheet=meta.get('sheet'),issue_domain=meta.get('issue_domain','AUTO')); intake_svc.mark_committed(intake_session_id,r['batch_id']); return r

    @app.post('/import/mapping-differences/apply', include_in_schema=False)
    async def apply_mapping_differences(request: Request):
        form=await request.form(); token=str(form.get('intake_session_id') or '')
        try: meta=intake_svc.get(token)
        except IntakeSessionError as e: raise HTTPException(404,str(e))
        # Difference workbench may intentionally preview a new product's DRAFT.
        # Do not compare that Draft with the unrelated ACTIVE version.  For an
        # ACTIVE preview, retain the original protection against an intervening
        # version switch.
        effective=mapping_svc.get_config_version(meta.get('mapping_config_id'))
        if not effective or effective.get('business_type') != meta.get('business_type'):
            raise HTTPException(409,'MAPPING_CHANGED_AFTER_PREVIEW')
        if effective.get('status') == 'ACTIVE':
            current_active=mapping_svc.get_effective_config(meta['business_type'])
            if not current_active or current_active['config_id'] != effective['config_id']:
                raise HTTPException(409,'MAPPING_CHANGED_AFTER_PREVIEW')
        actions=[]
        for i in range(int(form.get('difference_count') or 0)):
            actions.append({'source_header':form.get(f'source_{i}'),'action':form.get(f'action_{i}'),'target_mapping_id':form.get(f'target_{i}'),'target_field':form.get(f'extension_key_{i}')})
        disabled=[str(x) for x in form.getlist('disable_mapping_id')]
        try: draft=mapping_svc.apply_difference_actions(effective['config_id'],actions,disabled)
        except ValueError as e:
            detail=html.escape(str(e))
            return HTMLResponse(f'<h2>字段差异处理失败</h2><p>{detail}</p><p>请返回页面，选择一个已有标准字段后重试。</p>',status_code=400)
        meta['difference_draft_id']=draft['config_id']; intake_svc.save(meta)
        return RedirectResponse(f"/settings/mapping?business_type={meta['business_type']}&config_id={draft['config_id']}",303)

    @app.post('/api/issues/import')
    def api_import(file: UploadFile = File(...), business_type: str = Form(''), issue_domain: str = Form('AUTO')):
        return save_and_import(file, business_type, issue_domain)

    @app.get('/api/imports/{batch_id}')
    def api_import_batch(batch_id: str):
        return svc.get_import_batch(batch_id)

    @app.get('/api/issues')
    def api_issues(request: Request, limit: int = 100):
        rows = svc.query_issues(filters(request), min(limit, 1000))
        return {'count': len(rows), 'items': rows}

    @app.post('/api/issues/batch-domain')
    async def api_batch_domain(request: Request):
        payload = await request.json()
        ids = [str(x) for x in (payload.get('knowledge_ids') or []) if str(x).strip()]
        domain = str(payload.get('issue_domain') or 'AUTO').upper()
        if domain not in {'AUTO','SOFTWARE','HARDWARE','MECHANICAL','EMBEDDED'}:
            raise HTTPException(400, 'INVALID_ISSUE_DOMAIN')
        return {'updated': svc.repository.set_issue_domain(ids, domain, 'USER', str(payload.get('changed_by') or 'web')), 'issue_domain': domain}

    @app.get('/api/issues/{knowledge_id}')
    def api_issue(knowledge_id: str):
        d = svc.get_issue_detail(knowledge_id)
        if not d:
            raise HTTPException(404, 'NOT_FOUND')
        return d

    @app.get('/analysis', response_class=HTMLResponse, include_in_schema=False)
    def analysis_page(request: Request, job_id: str = ''):
        items = svc.query_issues({}, 1000)
        statuses = {}
        for x in items:
            h = svc.get_analysis_history(x['knowledge_id'])
            statuses[x['knowledge_id']] = h[0]['status'] if h else '未分析'
        return tpl.TemplateResponse(request, 'analysis.html', {'items': items, 'statuses': statuses, 'products': product_repo.list(), 'analysis_agents': list_quality_issue_agents(BASE.parent.parent), 'domain_profiles': DOMAIN_PROFILES, 'domain_labels': DOMAIN_LABELS, 'issue_types': ISSUE_TYPES, 'issue_type_labels': ISSUE_TYPE_LABELS, 'lifecycle_phases': LIFECYCLE_PHASES, 'lifecycle_labels': LIFECYCLE_LABELS, 'batch_job_id': job_id, 'batch_job': batch_jobs.get(job_id) if job_id else None})

    @app.get('/api/analysis-agents')
    def api_analysis_agents():
        """Expose effective routing without secrets for deployment diagnostics."""
        return list_quality_issue_agents(BASE.parent.parent)

    @app.post('/analysis/{knowledge_id}', include_in_schema=False)
    def analysis_one_page(knowledge_id: str, force: str = Form(''), analysis_agent: str = Form(''), domain_profile: str = Form('AUTO'), issue_types: list[str] = Form([]), lifecycle_phase: str = Form('AUTO')):
        svc.run_issue_analysis(knowledge_id, BASE.parent.parent, force=force.lower() in {'1','true','on','yes'}, analysis_profile=analysis_profile_from_form(domain_profile, issue_types, lifecycle_phase),agent_id='' if analysis_agent in {'','DYNAMIC'} else analysis_agent)
        project_capability(knowledge_id)
        return RedirectResponse('/issues/' + knowledge_id, 303)

    @app.post('/analysis-batch', include_in_schema=False)
    def analysis_batch_page(business_type: str = Form(''), only_missing: str = Form(''), force: str = Form(''), concurrency: int = Form(2), analysis_agent: str = Form(''), domain_profile: str = Form('AUTO'), issue_types: list[str] = Form([]), lifecycle_phase: str = Form('AUTO')):
        ids = [x['knowledge_id'] for x in svc.query_issues({'business_type': business_type} if business_type else {}, 100000)]
        only_missing_flag=only_missing.lower() in {'1','true','on','yes'}
        force_flag=force.lower() in {'1','true','on','yes'}
        if only_missing_flag and not force_flag:
            ids=svc.incomplete_analysis_ids(ids)
        job = batch_jobs.start(
            ids, run_batch_with_projection,
            only_missing=only_missing_flag,
            force=force_flag,
            analysis_profile=analysis_profile_from_form(domain_profile, issue_types, lifecycle_phase),
            agent_id='' if analysis_agent in {'','DYNAMIC'} else analysis_agent,
            concurrency=concurrency,
        )
        return RedirectResponse('/analysis?job_id=' + job['job_id'], 303)

    @app.post('/analysis-batch/{job_id}/retry-failed', include_in_schema=False)
    def analysis_batch_retry_failed(job_id: str, concurrency: int = Form(2)):
        prior=batch_jobs.get(job_id)
        if not prior: raise HTTPException(404,'BATCH_ANALYSIS_JOB_NOT_FOUND')
        ids=[x.get('knowledge_id') for x in prior.get('items',[]) if x.get('status')=='FAILED' and x.get('knowledge_id')]
        if not ids: raise HTTPException(409,'BATCH_ANALYSIS_HAS_NO_FAILED_ITEMS')
        job=batch_jobs.start(ids,run_batch_with_projection,only_missing=False,force=True,analysis_profile=None,concurrency=concurrency)
        return RedirectResponse('/analysis?job_id='+job['job_id'],303)

    @app.get('/api/analysis-batch-jobs/{job_id}')
    def analysis_batch_job(job_id: str):
        job = batch_jobs.get(job_id)
        if not job:
            raise HTTPException(404, 'BATCH_ANALYSIS_JOB_NOT_FOUND')
        return job

    @app.get('/api/issues/{knowledge_id}/versions')
    def api_versions(knowledge_id: str):
        return {'items': svc.get_issue_history(knowledge_id)}

    @app.post('/api/issues/{knowledge_id}/analyze')
    def api_analyze(knowledge_id: str, force: bool = False, payload: dict | None = None):
        result = svc.run_issue_analysis(knowledge_id, BASE.parent.parent, force=force, analysis_profile=(payload or {}).get('analysis_profile'),agent_id=(payload or {}).get('analysis_agent') or '')
        project_capability(knowledge_id)
        return result

    @app.post('/api/issues/analyze-batch')
    def api_analyze_batch(payload: dict):
        ids = payload.get('knowledge_ids') or [x['knowledge_id'] for x in svc.query_issues({k: v for k, v in {'business_type': payload.get('business_type')}.items() if v}, 100000)]
        return svc.run_batch_analysis(ids, BASE.parent.parent, only_missing=bool(payload.get('only_missing',True)), force=bool(payload.get('force')), analysis_profile=payload.get('analysis_profile'), concurrency=payload.get('concurrency',2),agent_id=payload.get('analysis_agent') or '')

    @app.get('/api/analysis/{run_id}')
    def api_analysis(run_id: str):
        r = svc.get_analysis_status(run_id)
        if not r:
            raise HTTPException(404, 'NOT_FOUND')
        return r

    @app.get('/api/issues/{knowledge_id}/analysis-history')
    def api_analysis_history(knowledge_id: str):
        return {'items': svc.get_analysis_history(knowledge_id)}

    @app.get('/statistics', response_class=HTMLResponse, include_in_schema=False)
    def statistics_page(request: Request, business_type: str = '', gap_dimension: str = '', common_scope: str = 'all'):
        # Deterministic, idempotent backfill for analysis completed before the
        # extension was installed. It does not alter any V1 source row.
        month=(request.query_params.get('month') or '').strip()
        issue_domain=(request.query_params.get('issue_domain') or '').strip().upper()
        lifecycle=(request.query_params.get('lifecycle') or '').strip().upper()
        scope_filters={k:v for k,v in {'business_type':business_type,'month':month,'issue_domain':issue_domain}.items() if v}
        scoped = svc.query_issues(scope_filters, 100000)
        for row in scoped:
            project_capability(row['knowledge_id'])
        scope_ids=[row['knowledge_id'] for row in scoped]
        if lifecycle:
            scope_ids=capability_extension.filter_issue_ids_by_lifecycle(scope_ids,lifecycle)
        stats = svc.get_statistics(business_type or None, 50, scope_ids)
        runtime = svc.get_workspace_metrics(business_type or None, scope_ids)
        common_rows = svc.get_common_capability_gaps(
            business_type=business_type or None,
            dimension=gap_dimension or None,
            min_issues=2,
            limit=50,
            knowledge_ids=scope_ids if (month or issue_domain or lifecycle) else None,
        )
        common_gap_view = present_common_gaps(
            common_rows, business_type=business_type, dimension=gap_dimension, scope=common_scope
        )
        def common_url(*, scope=None, dimension=None, clear_dimension=False):
            params=[]
            if business_type:
                params.append(('business_type', business_type))
            for key,value in (('month',month),('issue_domain',issue_domain),('lifecycle',lifecycle)):
                if value: params.append((key,value))
            selected_scope = common_scope if scope is None else scope
            selected_dimension = '' if clear_dimension else (gap_dimension if dimension is None else dimension)
            if selected_dimension:
                params.append(('gap_dimension', selected_dimension))
            params.append(('common_scope', selected_scope or 'all'))
            return '/statistics?' + urlencode(params) + '#common'
        common_urls = {
            'all': common_url(scope='all', clear_dimension=True),
            'cross': common_url(scope='cross'),
            'technical': common_url(dimension='TECHNICAL'),
            'management': common_url(dimension='MANAGEMENT'),
            'governance': common_url(dimension='GOVERNANCE'),
        }
        return tpl.TemplateResponse(request, 'statistics.html', {
            'stats': stats, 'statistics_view': present_statistics(stats, runtime), 'business_type': business_type,
            'gap_dimension': gap_dimension, 'common_scope': common_scope, 'common_gap_view': common_gap_view,
            'common_urls': common_urls, 'products': product_repo.list(),
            'capability_matrices': capability_extension.matrices(business_type,scope_ids),
            'quality_insight_summary': capability_extension.insight_summary(business_type,scope_ids),
            'classification_comparison': capability_extension.classification_comparison(business_type,scope_ids),
            'month':month,'issue_domain':issue_domain,'lifecycle':lifecycle,
            'available_months':sorted({str(x.get('month') or '') for x in svc.query_issues({},100000) if x.get('month')}),
            'lifecycle_phases':sorted(set(LIFECYCLE_PHASES)|{x.get('lifecycle_code') for x in capability_extension.matrices(business_type).get('lifecycle_x_capability',[]) if x.get('lifecycle_code')}),
            'lifecycle_labels':LIFECYCLE_LABELS,
        })

    @app.get('/product-reports', response_class=HTMLResponse, include_in_schema=False)
    def product_reports_page(request: Request):
        return tpl.TemplateResponse(request,'p1_product_reports.html',{'api_prefix':'/api','report_base':'base.html','asset_prefix':'/static'})

    @app.get('/api/product-reports/precheck')
    def product_report_precheck(product_code: str,start_month: str,end_month: str):
        return report_svc.precheck(product_code,start_month,end_month)

    @app.get('/api/product-reports')
    def product_reports_api():
        items=report_svc.list();return {'items':items,'total':len(items)}

    @app.post('/api/product-reports',status_code=201)
    async def product_report_create(request: Request):
        payload=await request.json()
        try:return report_svc.create(str(payload.get('product_code') or ''),str(payload.get('start_month') or ''),str(payload.get('end_month') or ''),str(payload.get('created_by') or ''))
        except ProductReportError as error:raise HTTPException(400,str(error))

    @app.get('/api/product-reports/{report_id}')
    def product_report_get(report_id: str):
        try:return report_svc.get(report_id)
        except ProductReportError as error:raise HTTPException(404,str(error))

    @app.post('/api/product-reports/{report_id}/publish')
    def product_report_publish(report_id: str):
        try:return report_svc.publish(report_id)
        except ProductReportError as error:raise HTTPException(404,str(error))

    @app.delete('/api/product-reports/{report_id}')
    def product_report_delete(report_id: str):
        try:return report_svc.delete(report_id)
        except ProductReportError as error:raise HTTPException(409 if str(error)=='PUBLISHED_REPORT_CANNOT_BE_DELETED' else 404,str(error))

    @app.get('/export/common-gaps', include_in_schema=False)
    def export_common_gaps(business_type: str = '', gap_dimension: str = '', common_scope: str = 'all'):
        rows = svc.get_common_capability_gaps(
            business_type=business_type or None, dimension=gap_dimension or None, min_issues=2, limit=10000
        )
        view = present_common_gaps(rows, business_type=business_type, dimension=gap_dimension, scope=common_scope)
        output = Path(tempfile.mkdtemp(prefix='quality-common-gap-')) / 'common_capability_gap_governance.csv'
        columns = ['优先级','能力缺口','能力维度','范围判定','关联问题数','覆盖业务数','涉及业务','涉及产品','涉及平台','建议治理措施','预期预防效果']
        with output.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            for row in view['items']:
                writer.writerow({
                    '优先级': row['rank'], '能力缺口': row['category'], '能力维度': row['dimension_label'],
                    '范围判定': row['scope_label'], '关联问题数': row['related_issue_count'],
                    '覆盖业务数': row['coverage_count'], '涉及业务': '、'.join(row['businesses']),
                    '涉及产品': '、'.join(row['products']), '涉及平台': '、'.join(row['platforms']),
                    '建议治理措施': row['recommended_governance'], '预期预防效果': row['expected_prevention_effect'],
                })
        return FileResponse(output, filename=output.name, media_type='text/csv; charset=utf-8')

    @app.get('/insights/capability-gaps', response_class=HTMLResponse, include_in_schema=False)
    def capability_gaps_page(request: Request, dimension: str = '', category: str = '', business_type: str = ''):
        filters = {k: v for k, v in {
            'dimension': dimension.upper() if dimension else '',
            'category': category,
            'business_type': business_type.upper() if business_type else '',
        }.items() if v}
        rows = svc.query_capability_gaps(filters=filters, limit=1000)
        return tpl.TemplateResponse(request, 'capability_gaps.html', {
            'rows': rows, 'dimension': dimension.upper(), 'category': category,
            'business_type': business_type.upper(), 'total': len(rows),
        })



    @app.post('/settings/mapping/preview', response_class=HTMLResponse, include_in_schema=False)
    def mapping_preview(request: Request, file: UploadFile = File(...), business_type: str = Form(...), config_id: str = Form('')):
        try: result=_build_intake_preview(file,business_type,mapping_config_id=config_id)
        except HTTPException as e:
            return tpl.TemplateResponse(request,'mapping_preview.html',{'error':e.detail,'result':None,'business_type':business_type.upper()},status_code=e.status_code)
        return tpl.TemplateResponse(request,'import_preview.html',{'result':result,'mapping_preview_only':True})

    @app.post('/settings/mapping/preview/add', include_in_schema=False)
    def mapping_preview_add(business_type: str = Form(...), source_header: str = Form(...), target_domain: str = Form('PRODUCT_EXTENSION'), target_field: str = Form('')):
        effective=mapping_svc.get_effective_config(business_type)
        if not effective: raise HTTPException(400,'MAPPING_NOT_INITIALIZED')
        try: draft=mapping_svc.add_unmatched_to_draft(effective['config_id'],source_header,target_domain,target_field or None,'mapping')
        except ValueError as e: raise HTTPException(400,str(e))
        return RedirectResponse(f"/settings/mapping?business_type={business_type}&config_id={draft['config_id']}",303)

    @app.get('/settings/mapping', response_class=HTMLResponse, include_in_schema=False)
    def mapping_settings(request: Request, business_type: str = 'PLC', config_id: str = ''):
        bt=(business_type or 'PLC').upper()
        products=product_repo.list(False)
        product=next((x for x in products if x['product_code']==bt),None)
        if not product:
            raise HTTPException(404,'PRODUCT_NOT_FOUND: 请先在产品配置中创建该产品')
        configs=mapping_svc.list_configs(bt)
        effective=mapping_svc.get_effective_config(bt)
        selected=mapping_svc.get_config_version(config_id) if config_id else effective
        if config_id and selected and selected['business_type'] != bt:
            raise HTTPException(400,'MAPPING_PRODUCT_MISMATCH')
        if not selected:
            selected=next((x for x in configs if x['status']=='DRAFT'),None)
        validation=mapping_svc.validation_results(selected['config_id']) if selected else []
        migration=mapping_svc.migration_info(selected['config_id']) if selected else None
        return tpl.TemplateResponse(request,'mapping_settings.html',{'business_type':bt,'product':product,'products':products,'configs':configs,'effective':effective,'selected':selected,'validation':validation,'migration':migration,'target_domains':sorted(TARGET_DOMAINS)})

    @app.post('/settings/mapping/{business_type}/bootstrap', include_in_schema=False)
    def mapping_bootstrap_product(business_type: str):
        bt=business_type.upper()
        if not product_repo.get(bt):
            raise HTTPException(404,'PRODUCT_NOT_FOUND')
        try:
            draft=mapping_svc.create_product_starter_draft(bt)
        except ValueError as e:
            if str(e) == 'MAPPING_ALREADY_INITIALIZED':
                active=mapping_svc.get_effective_config(bt)
                return RedirectResponse(f"/settings/mapping?business_type={bt}&config_id={active['config_id']}",303)
            raise HTTPException(400,str(e))
        return RedirectResponse(f"/settings/mapping?business_type={bt}&config_id={draft['config_id']}",303)

    @app.get('/settings/products', response_class=HTMLResponse, include_in_schema=False)
    def product_settings(request: Request):
        products=product_repo.list(False)
        mapping_states={p['product_code']: bool(mapping_svc.get_effective_config(p['product_code'])) for p in products}
        return tpl.TemplateResponse(request, 'product_settings.html', {'products': products, 'mapping_states': mapping_states})

    @app.post('/settings/products', include_in_schema=False)
    async def product_settings_save(request: Request):
        form = await request.form()
        product = product_repo.upsert(form.get('product_code'), form.get('product_name'), form.get('product_kind') or 'PRODUCT', form.get('default_issue_domain') or 'AUTO', True, int(form.get('sort_order') or 0))
        register_product_adapter(product['product_code'])
        return RedirectResponse('/settings/products', 303)

    @app.get('/api/settings/products')
    def api_products(enabled_only: bool = True):
        return {'items': product_repo.list(enabled_only), 'count': len(product_repo.list(enabled_only))}

    @app.post('/api/settings/products')
    async def api_product_create(request: Request):
        form = await request.json()
        product = product_repo.upsert(form.get('product_code'), form.get('product_name'), form.get('product_kind') or 'PRODUCT', form.get('default_issue_domain') or 'AUTO', form.get('enabled', True), int(form.get('sort_order') or 0))
        register_product_adapter(product['product_code'])
        return product

    @app.patch('/api/settings/products/{product_code}')
    async def api_product_update(product_code: str, request: Request):
        old = product_repo.get(product_code)
        if not old: raise HTTPException(404, 'PRODUCT_NOT_FOUND')
        form = await request.json()
        product = product_repo.upsert(product_code, form.get('product_name', old['product_name']), form.get('product_kind', old['product_kind']), form.get('default_issue_domain', old['default_issue_domain']), form.get('enabled', bool(old['enabled'])), int(form.get('sort_order', old['sort_order'])))
        register_product_adapter(product['product_code'])
        return product

    @app.post('/api/settings/products/{product_code}/enable')
    def api_product_enable(product_code: str):
        product = product_repo.set_enabled(product_code, True)
        if not product: raise HTTPException(404, 'PRODUCT_NOT_FOUND')
        register_product_adapter(product_code); return product

    @app.post('/api/settings/products/{product_code}/disable')
    def api_product_disable(product_code: str):
        product = product_repo.set_enabled(product_code, False)
        if not product: raise HTTPException(404, 'PRODUCT_NOT_FOUND')
        return product

    @app.get('/api/settings/products/{product_code}/mapping-status')
    def api_product_mapping_status(product_code: str):
        product = product_repo.get(product_code)
        if not product: raise HTTPException(404, 'PRODUCT_NOT_FOUND')
        cfg = mapping_svc.get_effective_config(product_code.upper())
        return {'product_code': product_code.upper(), 'initialized': bool(cfg), 'active_mapping': cfg}

    @app.post('/settings/mapping/{config_id}/draft', include_in_schema=False)
    def mapping_create_draft(config_id: str):
        cfg=mapping_svc.create_draft_from(config_id)
        return RedirectResponse(f"/settings/mapping?business_type={cfg['business_type']}&config_id={cfg['config_id']}",303)

    @app.post('/settings/mapping/{config_id}/save', include_in_schema=False)
    async def mapping_save(request: Request, config_id: str):
        cfg=mapping_svc.get_config_version(config_id)
        if not cfg or cfg['status']!='DRAFT': raise HTTPException(400,'仅 DRAFT 可编辑')
        form=await request.form(); items=[]
        for i,old in enumerate(cfg['mappings']):
            mid=old['mapping_id']
            headers=[x.strip() for x in str(form.get(f'headers_{mid}','')).splitlines() if x.strip()]
            aliases=[x.strip() for x in str(form.get(f'aliases_{mid}','')).splitlines() if x.strip()]
            items.append(MappingItem(mid,str(form.get(f'canonical_{mid}',old['canonical_field'])).strip(),headers,aliases,str(form.get(f'domain_{mid}',old['target_domain'])),str(form.get(f'target_{mid}',old['target_field'])).strip(),bool(form.get(f'required_{mid}')),bool(form.get(f'enabled_{mid}')),str(form.get(f'description_{mid}',old.get('description') or '')),i))
        mapping_svc.update_draft(config_id,items)
        return RedirectResponse(f"/settings/mapping?business_type={cfg['business_type']}&config_id={config_id}",303)

    @app.post('/settings/mapping/{config_id}/validate', include_in_schema=False)
    def mapping_validate(config_id: str):
        cfg=mapping_svc.get_config_version(config_id); mapping_svc.validate_config(config_id)
        return RedirectResponse(f"/settings/mapping?business_type={cfg['business_type']}&config_id={config_id}",303)

    @app.post('/settings/mapping/{config_id}/activate', include_in_schema=False)
    def mapping_activate(config_id: str):
        cfg=mapping_svc.get_config_version(config_id)
        try: mapping_svc.activate_config(config_id)
        except ValueError as e: raise HTTPException(400,str(e))
        return RedirectResponse(f"/settings/mapping?business_type={cfg['business_type']}&config_id={config_id}",303)

    @app.get('/settings/mapping/{config_id}/export', include_in_schema=False)
    def mapping_export(config_id: str, format: str = 'yaml'):
        fmt='json' if format.lower()=='json' else 'yaml'
        td=Path(tempfile.gettempdir())/'quality_mapping_exports'; td.mkdir(exist_ok=True)
        cfg=mapping_svc.get_config_version(config_id)
        out=td/f"{cfg['business_type'].lower()}_mapping_v{cfg['version']}.{fmt}"
        mapping_svc.export_config(config_id,out,fmt)
        return FileResponse(out,filename=out.name,media_type='application/json' if fmt=='json' else 'application/x-yaml')


    @app.get('/settings/human-analysis', response_class=HTMLResponse, include_in_schema=False)
    def human_analysis_settings(request: Request):
        return tpl.TemplateResponse(request,'human_analysis_settings.html',{'fields':human_svc.list_field_definitions()})

    @app.post('/settings/human-analysis/fields', include_in_schema=False)
    async def human_analysis_create_field(request: Request):
        form=await request.form()
        options=[x.strip() for x in str(form.get('options','')).splitlines() if x.strip()]
        try:
            human_svc.create_field(field_key=str(form.get('field_key','')).strip(),field_name=str(form.get('field_name','')).strip(),field_type=str(form.get('field_type','TEXT')),description=str(form.get('description','')),required=bool(form.get('required')),enabled=True,display_order=int(form.get('display_order') or 0),options=options)
        except Exception as e: raise HTTPException(400,str(e))
        return RedirectResponse('/settings/human-analysis',303)


    @app.get('/settings/human-analysis/fields/{field_id}', response_class=HTMLResponse, include_in_schema=False)
    def human_analysis_edit_field(request: Request, field_id: str):
        field=next((x for x in human_svc.list_field_definitions() if x['field_id']==field_id),None)
        if not field: raise HTTPException(404,'field not found')
        return tpl.TemplateResponse(request,'human_analysis_field_edit.html',{'field':field})

    @app.post('/settings/human-analysis/fields/{field_id}', include_in_schema=False)
    async def human_analysis_update_field(request: Request, field_id: str):
        field=next((x for x in human_svc.list_field_definitions() if x['field_id']==field_id),None)
        if not field: raise HTTPException(404,'field not found')
        form=await request.form()
        try:
            human_svc.update_field(
                field_id,
                field_name=str(form.get('field_name','')).strip(),
                description=str(form.get('description','')),
                required=bool(form.get('required')),
                enabled=bool(form.get('enabled')),
                display_order=int(form.get('display_order') or 0),
                default_value=str(form.get('default_value','')) or None,
                validation_rule=str(form.get('validation_rule','')) or None,
            )
        except Exception as e: raise HTTPException(400,str(e))
        return RedirectResponse(f'/settings/human-analysis/fields/{field_id}',303)

    @app.post('/settings/human-analysis/fields/{field_id}/options', include_in_schema=False)
    async def human_analysis_create_option(request: Request, field_id: str):
        form=await request.form()
        try:
            human_svc.create_option(
                field_id,
                str(form.get('option_value','')).strip(),
                str(form.get('option_label','')).strip() or None,
                int(form.get('display_order') or 0),
                True,
            )
        except Exception as e: raise HTTPException(400,str(e))
        return RedirectResponse(f'/settings/human-analysis/fields/{field_id}',303)

    @app.post('/settings/human-analysis/options/{option_id}', include_in_schema=False)
    async def human_analysis_update_option(request: Request, option_id: str):
        form=await request.form()
        option=human_svc.repository.get_option(option_id)
        if not option: raise HTTPException(404,'option not found')
        try:
            human_svc.update_option(
                option_id,
                option_label=str(form.get('option_label','')).strip(),
                display_order=int(form.get('display_order') or 0),
                enabled=bool(form.get('enabled')),
            )
        except Exception as e: raise HTTPException(400,str(e))
        return RedirectResponse(f"/settings/human-analysis/fields/{option['field_id']}",303)


    @app.post('/settings/human-analysis/fields/{field_id}/delete', include_in_schema=False)
    def human_analysis_delete_field(field_id: str):
        try:
            human_svc.delete_field(field_id)
        except KeyError:
            raise HTTPException(404,'field not found')
        except ValueError as e:
            raise HTTPException(400,str(e))
        return RedirectResponse('/settings/human-analysis',303)

    @app.post('/settings/human-analysis/fields/{field_id}/toggle', include_in_schema=False)
    def human_analysis_toggle(field_id: str):
        f=next((x for x in human_svc.list_field_definitions() if x['field_id']==field_id),None)
        if not f: raise HTTPException(404,'field not found')
        human_svc.update_field(field_id,enabled=not bool(f['enabled']))
        return RedirectResponse('/settings/human-analysis',303)

    @app.post('/issues/{knowledge_id}/human-analysis', include_in_schema=False)
    async def human_analysis_save(request: Request, knowledge_id: str):
        issue=svc.get_issue(knowledge_id)
        if not issue: raise HTTPException(404,'issue not found')
        form=await request.form(); vals={}
        for f in human_svc.list_field_definitions(True):
            key='human_'+f['field_id']
            vals[f['field_id']]=form.getlist(key) if f['field_type']=='MULTI_SELECT' else form.get(key)
        try: human_svc.save_analysis(knowledge_id,issue['issue_version_id'],vals,'web')
        except ValueError as e: raise HTTPException(400,str(e))
        return RedirectResponse(f'/issues/{knowledge_id}',303)

    @app.post('/issues/{knowledge_id}/analysis-confirmations', include_in_schema=False)
    async def analysis_confirmations_save(request: Request, knowledge_id: str):
        if not svc.get_issue(knowledge_id): raise HTTPException(404,'issue not found')
        form=await request.form(); answers={}
        for q in svc.repository.list_open_questions(knowledge_id,50):
            qid=q['question_id']
            answers[qid]={'status':form.get('status_'+qid) or 'UNRESOLVED','answer':form.get('answer_'+qid) or '','evidence':form.get('evidence_'+qid) or ''}
        svc.repository.save_question_confirmations(knowledge_id,answers,'web')
        if str(form.get('reanalyze') or '').lower() in {'1','true','on','yes'}:
            svc.run_issue_analysis(knowledge_id,BASE.parent.parent,force=True)
        return RedirectResponse(f'/issues/{knowledge_id}#human-analysis',303)

    @app.post('/issues/{knowledge_id}/quality-confirmations', include_in_schema=False)
    def quality_confirmations_save(
        knowledge_id: str, occurrence_mrc: str = Form(''), escape_mrc: str = Form(''),
        reason: str = Form(''), evidence: str = Form(''), confirmed_by: str = Form('web')
    ):
        issue=svc.get_issue(knowledge_id)
        if not issue: raise HTTPException(404,'issue not found')
        current=project_capability(knowledge_id) or {}
        effective=current.get('effective_mrc') or {}
        try:
            for side,code in (('OCCURRENCE',occurrence_mrc),('ESCAPE',escape_mrc)):
                code=code.strip().upper()
                if not code: continue
                original=(effective.get(side) or {}).get('mrc_code')
                capability_extension.save_human_revision(
                    knowledge_id=knowledge_id,issue_version_id=issue['issue_version_id'],
                    target_path=f"{side.lower()}.mrc.primary",original_value=original,
                    confirmed_value={'code':code},status='CONFIRMED' if original==code else 'CORRECTED',
                    reason=reason,evidence=evidence,confirmed_by=confirmed_by,
                )
        except ValueError as error:
            raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/issues/{knowledge_id}#quality-capability-evidence',303)

    @app.post('/issues/{knowledge_id}/hardware-components', include_in_schema=False)
    async def hardware_component_save(request: Request, knowledge_id: str):
        issue=svc.get_issue(knowledge_id)
        if not issue: raise HTTPException(404,'issue not found')
        form=await request.form()
        values={key:str(form.get(key) or '').strip() for key in (
            'relevance','component_category','component_name','manufacturer','model_part_number','lot_batch',
            'serial_number','board_module','reference_designator','installation_location','hardware_version',
            'failure_mode','failure_mechanism','failure_cause','customer_impact','reproduction_condition',
            'detection_method','disposition','evidence')}
        if not values['component_name'] and not values['model_part_number'] and not values['component_category']:
            raise HTTPException(400,'HARDWARE_COMPONENT_IDENTITY_REQUIRED')
        try:
            capability_extension.add_hardware_component(
                knowledge_id=knowledge_id,issue_version_id=issue['issue_version_id'],values=values,
                actor=str(form.get('confirmed_by') or 'web'))
        except ValueError as error:
            raise HTTPException(400,str(error)) from error
        return RedirectResponse(f'/issues/{knowledge_id}#hardware-components',303)

    @app.get('/api/issues/{knowledge_id}/analysis-confirmations')
    def api_analysis_confirmations(knowledge_id: str):
        if not svc.get_issue(knowledge_id): raise HTTPException(404,'issue not found')
        rows=svc.repository.list_open_questions(knowledge_id,50)
        return {'count':len(rows),'items':rows}

    @app.post('/api/issues/{knowledge_id}/analysis-confirmations')
    async def api_analysis_confirmations_save(request: Request, knowledge_id: str):
        if not svc.get_issue(knowledge_id): raise HTTPException(404,'issue not found')
        payload=await request.json(); updated=svc.repository.save_question_confirmations(knowledge_id,payload.get('answers') or {},str(payload.get('confirmed_by') or 'api'))
        result={'updated':updated}
        if payload.get('reanalyze'):
            result['analysis']=svc.run_issue_analysis(knowledge_id,BASE.parent.parent,force=True)
        return result

    @app.get('/export', response_class=HTMLResponse, include_in_schema=False)
    def export_page(request: Request):
        fs = filters(request)
        return tpl.TemplateResponse(request, 'export.html', {'filters': fs})

    @app.post('/export/download', include_in_schema=False)
    def export_download(
        format: str = Form('xlsx'), dataset: str = Form('issues'), scope: str = Form('all'),
        business_type: str = Form(''), business_issue_id: str = Form(''), product: str = Form(''),
        platform: str = Form(''), severity: str = Form(''), issue_type: str = Form('')
    ):
        fmt = format.lower()
        if fmt not in {'xlsx', 'csv'}:
            raise HTTPException(400, 'format must be xlsx or csv')
        if dataset not in {'issues', 'capability_gaps', 'ai_analysis', 'human_analysis'}:
            raise HTTPException(400, 'unsupported dataset')
        fs = {} if scope == 'all' else {k: v for k, v in {
            'business_type': business_type, 'business_issue_id': business_issue_id, 'product': product,
            'platform': platform, 'severity': severity, 'issue_type': issue_type
        }.items() if v}
        suffix = '.csv' if fmt == 'csv' else '.xlsx'
        stamp = __import__('datetime').datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        name = f'quality_issue_{dataset}_{stamp}{suffix}'
        td = Path(tempfile.gettempdir()) / 'quality_issue_exports_v1'
        td.mkdir(exist_ok=True)
        out = td / name
        svc.export_issues(out, format=fmt, filters=fs, dataset=dataset)
        media = 'text/csv; charset=utf-8' if fmt == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return FileResponse(path=out, filename=name, media_type=media)

    @app.get('/api/statistics')
    def api_statistics(business_type: str | None = None, limit: int = 20):
        return svc.get_statistics(business_type, min(limit, 100))

    @app.post('/api/export')
    def api_export(payload: dict):
        fmt = (payload.get('format') or 'xlsx').lower()
        dataset = payload.get('dataset') or 'issues'
        fs = payload.get('filters') or {}
        suffix = '.csv' if fmt == 'csv' else '.xlsx'
        td = Path(tempfile.gettempdir()) / 'quality_issue_exports_v1'
        td.mkdir(exist_ok=True)
        out = td / (payload.get('filename') or ('quality_knowledge' + suffix))
        return svc.export_issues(out, format=fmt, filters=fs, dataset=dataset)


    @app.get('/api/common-capability-gaps')
    def api_common_gaps(business_type: str | None = None, dimension: str | None = None, min_issues: int = 2, limit: int = 50):
        return {'items': svc.get_common_capability_gaps(business_type=business_type, dimension=dimension, min_issues=max(1,min_issues), limit=min(limit,200))}

    @app.get('/api/capability-gaps')
    def api_gaps(knowledge_id: str | None = None):
        return {'items': svc.query_capability_gaps(knowledge_id)}

    return app
