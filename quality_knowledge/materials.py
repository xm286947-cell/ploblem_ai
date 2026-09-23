from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from quality_knowledge.sqlite_tuning import configure_connection


MATERIAL_TYPES = {"ESCAPE_ANALYSIS", "ITR_SOURCE", "ITR_CS", "SOFTWARE_OPERATION", "BATCH_ISSUE", "TEST_ISSUE"}
DEFAULT_GROUPS = (
    ("DG-ESCAPE", "ESCAPE_ANALYSIS", "漏测分析", 1, 1, 1),
    ("DG-ITR", "ITR_SOURCE", "ITR原始问题", 0, 0, 0),
    ("DG-ITR-CS", "ITR_CS", "ITR彻底解决单", 0, 0, 0),
    ("DG-SW-OPS", "SOFTWARE_OPERATION", "软件问题运营数据", 0, 0, 0),
    ("DG-BATCH", "BATCH_ISSUE", "批量问题", 0, 0, 0),
    ("DG-TEST", "TEST_ISSUE", "测试问题", 0, 0, 0),
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS data_group(
 group_id TEXT PRIMARY KEY,group_code TEXT NOT NULL UNIQUE,group_name TEXT NOT NULL,
 material_type TEXT NOT NULL,enabled INTEGER NOT NULL DEFAULT 1,
 include_ai INTEGER NOT NULL DEFAULT 0,include_insight INTEGER NOT NULL DEFAULT 0,
 include_report INTEGER NOT NULL DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS source_material(
 material_id TEXT PRIMARY KEY,group_id TEXT NOT NULL REFERENCES data_group(group_id),
 material_type TEXT NOT NULL,business_key TEXT NOT NULL,canonical_itr TEXT,
 version_no INTEGER NOT NULL,source_hash TEXT NOT NULL,source_file TEXT,sheet_name TEXT,row_number INTEGER,
 raw_json TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,
 UNIQUE(group_id,business_key,source_hash));
CREATE INDEX IF NOT EXISTS idx_source_material_itr ON source_material(canonical_itr,material_type);
CREATE INDEX IF NOT EXISTS idx_source_material_group_created ON source_material(group_id,created_at DESC,material_id);
CREATE INDEX IF NOT EXISTS idx_source_material_group_business ON source_material(group_id,business_key);
CREATE INDEX IF NOT EXISTS idx_source_material_group_itr_version ON source_material(group_id,canonical_itr,version_no DESC);
CREATE INDEX IF NOT EXISTS idx_source_material_type_itr_version ON source_material(material_type,canonical_itr,version_no DESC);
CREATE TABLE IF NOT EXISTS association_rule(
 rule_id TEXT PRIMARY KEY,rule_name TEXT NOT NULL,source_type TEXT NOT NULL,target_type TEXT NOT NULL,
 source_field TEXT NOT NULL,target_field TEXT NOT NULL,transform TEXT NOT NULL DEFAULT 'EXACT',
 priority INTEGER NOT NULL DEFAULT 100,status TEXT NOT NULL DEFAULT 'ACTIVE',version_no INTEGER NOT NULL DEFAULT 1,
 created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS issue_material_link(
 link_id TEXT PRIMARY KEY,knowledge_id TEXT NOT NULL,material_id TEXT NOT NULL REFERENCES source_material(material_id),
 rule_id TEXT REFERENCES association_rule(rule_id),link_status TEXT NOT NULL,match_value TEXT,
 created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(knowledge_id,material_id));
CREATE INDEX IF NOT EXISTS idx_issue_material_link_issue ON issue_material_link(knowledge_id,material_id);
CREATE INDEX IF NOT EXISTS idx_issue_material_link_material ON issue_material_link(material_id,knowledge_id);
CREATE TABLE IF NOT EXISTS material_review(
 material_id TEXT PRIMARY KEY REFERENCES source_material(material_id),review_status TEXT NOT NULL DEFAULT 'DRAFT',
 analysis_summary TEXT,root_cause TEXT,improvement_action TEXT,reviewer TEXT,
 updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS source_material_reporting_year(
 material_id TEXT PRIMARY KEY REFERENCES source_material(material_id) ON DELETE CASCADE,
 reporting_year TEXT NOT NULL,year_source TEXT NOT NULL,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE INDEX IF NOT EXISTS idx_source_material_reporting_year ON source_material_reporting_year(reporting_year,material_id);
"""


def _clean(value: Any) -> str:
    return str(value or "").strip()


def normalize_itr(value: Any) -> str:
    text = re.sub(r"\s+", "", _clean(value)).upper()
    return text[:-2] if re.fullmatch(r"ITR[0-9A-Z_-]+CS", text) else text


def year_from_itr(value: Any) -> str:
    match = re.match(r"^ITR(20\d{2})", normalize_itr(value))
    return match.group(1) if match else ""


def validate_reporting_year(value: Any) -> str:
    year = _clean(value)
    if not re.fullmatch(r"20\d{2}", year) or not 2000 <= int(year) <= 2099:
        raise ValueError("考核年份必须为2000—2099之间的四位年份")
    return year


def effective_reporting_year(raw: dict[str, Any], business_key: Any, reporting_year: Any = "", year_source: Any = "") -> str:
    """Software KPI year: manual override > KPI date > file year > ITR fallback."""
    stored=_clean(reporting_year);source=_clean(year_source)
    if stored and source in {"BATCH_MANUAL", "IMPORT_MANUAL"}:
        return stored
    kpi_month=_first(raw, "数据运营_KPI计入月份", "KPI计入月份")
    dated=re.search(r"(20\d{2})\s*(?:年|[-/.])", kpi_month)
    if dated:
        return dated.group(1)
    file_year=_first(raw, "数据运营_KPI计入年份", "数据运营_KPI计入年度", "KPI计入年份", "考核年份")
    return file_year or stored or year_from_itr(business_key)


def combine_headers(parent_row, child_row) -> list[str]:
    result, parent = [], ""
    width = max(len(parent_row or ()), len(child_row or ()))
    for index in range(width):
        top = _clean(parent_row[index] if index < len(parent_row) else "")
        child = _clean(child_row[index] if index < len(child_row) else "")
        if top:
            parent = top
        if child and parent and child != parent:
            result.append(f"{parent}_{child}")
        else:
            result.append(child or parent or f"未命名字段{index + 1}")
    return result


def _first(raw: dict[str, Any], *names: str) -> str:
    for name in names:
        value = _clean(raw.get(name))
        if value:
            return value
    return ""


def _material_view(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    raw = json.loads(item.pop("raw_json", "{}") or "{}")
    item["raw"] = raw
    item["title"] = _first(raw, "问题信息_问题主题", "问题信息_问题描述", "问题主题", "问题描述") or "未提供问题描述"
    item["month"] = _first(raw, "数据运营_KPI计入月份", "问题信息_创建月份", "创建月份") or "-"
    itr_year = year_from_itr(item.get("business_key"))
    item["year"] = effective_reporting_year(raw,item.get("business_key"),item.get("reporting_year"),item.get("year_source")) or "-"
    source_labels={"BATCH_MANUAL":"批量人工设置","IMPORT_MANUAL":"导入时人工设置","IMPORT_FILE":"原始文件","AUTO_ITR":"ITR编号默认"}
    kpi_dated=re.search(r"(20\d{2})\s*(?:年|[-/.])",item["month"])
    file_year=_first(raw,"数据运营_KPI计入年份","数据运营_KPI计入年度","KPI计入年份","考核年份")
    stored_source=_clean(item.get("year_source"))
    item["year_source"] = source_labels.get(stored_source,stored_source) if stored_source in {"BATCH_MANUAL","IMPORT_MANUAL"} else "KPI计入月份" if kpi_dated else "原始文件" if file_year else source_labels.get(stored_source,stored_source) or ("ITR编号默认" if itr_year else "未设置")
    item["product"] = _first(raw, "问题信息_产品型号", "问题信息_产品类型", "产品型号", "产品类型") or "-"
    item["domain"] = _first(raw, "问题信息_问题领域", "问题领域", "问题信息_产品类型", "产品类型") or "未分类"
    item["severity"] = _first(raw, "问题信息_问题等级", "问题等级") or "-"
    item["process_status"] = _first(raw, "过程信息_当前状态", "数据运营_审核状态", "当前状态", "审核状态") or "-"
    return item


class MaterialRepository:
    def __init__(self, db_path):
        self.db_path = str(db_path)
        with self.connect() as connection:
            connection.executescript(SCHEMA)
            for code, material_type, name, ai, insight, report in DEFAULT_GROUPS:
                connection.execute(
                    "INSERT OR IGNORE INTO data_group(group_id,group_code,group_name,material_type,include_ai,include_insight,include_report) VALUES(?,?,?,?,?,?,?)",
                    (code, code.removeprefix("DG-"), name, material_type, ai, insight, report),
                )
            defaults = (
                ("RULE-ESCAPE-ITR", "漏测分析关联ITR", "ESCAPE_ANALYSIS", "ITR_SOURCE", "ITR单号", "问题信息_ITR单号", "NORMALIZE_ITR", 10),
                ("RULE-CS-ITR", "彻底解决单关联ITR", "ITR_CS", "ITR_SOURCE", "问题信息_彻底解决单号", "问题信息_ITR单号", "STRIP_CS", 20),
                ("RULE-SWOPS-ITR", "软件运营数据关联ITR", "SOFTWARE_OPERATION", "ITR_SOURCE", "问题信息_彻底解决单号", "问题信息_ITR单号", "STRIP_CS", 30),
            )
            for row in defaults:
                connection.execute("INSERT OR IGNORE INTO association_rule(rule_id,rule_name,source_type,target_type,source_field,target_field,transform,priority) VALUES(?,?,?,?,?,?,?,?)", row)
        self.refresh_links()

    def connect(self):
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row
        return configure_connection(connection)

    def groups(self, enabled_only=True):
        where = " WHERE enabled=1" if enabled_only else ""
        with self.connect() as c:
            return [dict(x) for x in c.execute("SELECT * FROM data_group" + where + " ORDER BY group_name")]

    def group(self, group_code):
        with self.connect() as c:
            row = c.execute("SELECT * FROM data_group WHERE group_code=?", (group_code,)).fetchone()
            return dict(row) if row else None

    def add_material(self, group, business_key, raw, source_file, sheet, row_number):
        canonical = normalize_itr(business_key)
        digest = hashlib.sha256(json.dumps(raw, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()
        with self.connect() as c:
            existing = c.execute("SELECT material_id FROM source_material WHERE group_id=? AND canonical_itr=? AND source_hash=? ORDER BY version_no DESC LIMIT 1", (group["group_id"], canonical, digest)).fetchone()
            if existing:
                return existing[0], "SKIPPED"
            previous=c.execute("SELECT material_id FROM source_material WHERE group_id=? AND canonical_itr=? ORDER BY version_no DESC,created_at DESC,material_id DESC LIMIT 1",(group["group_id"],canonical)).fetchone()
            version = c.execute("SELECT COALESCE(MAX(version_no),0)+1 FROM source_material WHERE group_id=? AND canonical_itr=?", (group["group_id"], canonical)).fetchone()[0]
            material_id = f"MAT-{uuid.uuid4().hex}"
            c.execute("INSERT INTO source_material(material_id,group_id,material_type,business_key,canonical_itr,version_no,source_hash,source_file,sheet_name,row_number,raw_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)", (material_id, group["group_id"], group["material_type"], business_key, canonical, version, digest, source_file, sheet, row_number, json.dumps(raw, ensure_ascii=False, default=str)))
            if previous:
                manual=c.execute("SELECT reporting_year,year_source FROM source_material_reporting_year WHERE material_id=? AND year_source IN ('BATCH_MANUAL','IMPORT_MANUAL')",(previous[0],)).fetchone()
                if manual:c.execute("INSERT INTO source_material_reporting_year(material_id,reporting_year,year_source) VALUES(?,?,?)",(material_id,manual['reporting_year'],manual['year_source']))
            return material_id, "NEW" if version == 1 else "UPDATED"

    def refresh_links(self):
        with self.connect() as c:
            c.execute("DELETE FROM issue_material_link WHERE link_status<>'MANUAL_LINKED'")
            rules=c.execute("SELECT * FROM association_rule WHERE status='ACTIVE' ORDER BY priority").fetchall()
            allowed={}
            for rule in rules:
                if rule["source_type"]=="ESCAPE_ANALYSIS":allowed[rule["target_type"]]=rule
                else:allowed[rule["source_type"]]=rule
            try:issue_rows=c.execute("SELECT knowledge_id,business_issue_id FROM quality_issue").fetchall()
            except sqlite3.OperationalError:issue_rows=[]
            issues_by_itr={}
            for issue in issue_rows:issues_by_itr.setdefault(normalize_itr(issue['business_issue_id']),[]).append(issue['knowledge_id'])
            materials = c.execute("""SELECT material_id,material_type,canonical_itr FROM (
                SELECT m.*,ROW_NUMBER() OVER(PARTITION BY group_id,canonical_itr ORDER BY version_no DESC,created_at DESC,material_id DESC) rn
                FROM source_material m WHERE canonical_itr<>'') WHERE rn=1""").fetchall()
            for material in materials:
                rule=allowed.get(material["material_type"])
                if not rule:continue
                issues=issues_by_itr.get(normalize_itr(material['canonical_itr']),[])
                status = "LINKED" if len(issues) == 1 else ("CONFLICT" if len(issues) > 1 else "ITR_NOT_FOUND")
                if len(issues) == 1:
                    c.execute("INSERT OR IGNORE INTO issue_material_link(link_id,knowledge_id,material_id,rule_id,link_status,match_value) VALUES(?,?,?,?,?,?)", (f"LNK-{uuid.uuid4().hex}", issues[0], material["material_id"], rule["rule_id"], status, material["canonical_itr"]))

    def materials_for_issue(self, knowledge_id):
        with self.connect() as c:
            rows = c.execute("SELECT m.*,g.group_name,l.link_status FROM issue_material_link l JOIN source_material m ON m.material_id=l.material_id JOIN data_group g ON g.group_id=m.group_id WHERE l.knowledge_id=? ORDER BY m.material_type,m.version_no DESC", (knowledge_id,)).fetchall()
            result=[]
            for row in rows:
                item=dict(row); item["raw"]=json.loads(item.pop("raw_json") or "{}"); result.append(item)
            return result

    def list_materials(self, group_code="", limit=200, include_history=False):
        material_source="source_material m" if include_history else """(SELECT * FROM (
            SELECT sm.*,ROW_NUMBER() OVER(PARTITION BY sm.group_id,sm.canonical_itr ORDER BY sm.version_no DESC,sm.created_at DESC,sm.material_id DESC) rn
            FROM source_material sm) WHERE rn=1) m"""
        sql=f"SELECT m.*,g.group_code,g.group_name,y.reporting_year,y.year_source FROM {material_source} JOIN data_group g ON g.group_id=m.group_id LEFT JOIN source_material_reporting_year y ON y.material_id=m.material_id"
        values=[]
        if group_code: sql+=" WHERE g.group_code=?";values.append(group_code)
        sql+=" ORDER BY m.created_at DESC LIMIT ?";values.append(limit)
        with self.connect() as c:return [dict(x) for x in c.execute(sql,values)]

    def search_materials(self, group_code, *, q="", domain="", month="", year="", industry="", customer="", ipmt="", spdt="", product_model="", product_series="", page=1, page_size=20):
        def first_json(*keys):
            parts=[f"NULLIF(TRIM(CAST(json_extract(m.raw_json, '$.\"{key}\"') AS TEXT)),'')" for key in keys]
            return "COALESCE("+",".join(parts)+")"
        title_expr=first_json("问题信息_问题主题","问题信息_问题描述","问题主题","问题描述")
        month_expr=first_json("数据运营_KPI计入月份","问题信息_创建月份","创建月份")
        file_year_expr=first_json("数据运营_KPI计入年份","数据运营_KPI计入年度","KPI计入年份","考核年份")
        product_expr=first_json("问题信息_产品型号","问题信息_产品类型","产品型号","产品类型")
        domain_expr=first_json("问题信息_问题领域","问题领域","问题信息_产品类型","产品类型")
        dimension_exprs={
            'industry':first_json('问题信息_客户行业','客户行业'),
            'customer':first_json('问题信息_客户名称','客户名称'),
            'ipmt':first_json('问题信息_IPMT','IPMT'),
            'spdt':first_json('问题信息_SPDT','SPDT'),
            'product_model':first_json('问题信息_产品型号','产品型号'),
            'product_series':first_json('问题信息_产品系列','产品系列'),
        }
        itr_year_expr="CASE WHEN UPPER(m.business_key) GLOB 'ITR20[0-9][0-9]*' THEN SUBSTR(UPPER(m.business_key),4,4) END"
        kpi_month_year_expr=f"CASE WHEN INSTR({month_expr},'20')>0 THEN SUBSTR({month_expr},INSTR({month_expr},'20'),4) END"
        year_expr=f"CASE WHEN y.year_source IN ('BATCH_MANUAL','IMPORT_MANUAL') THEN NULLIF(TRIM(y.reporting_year),'') ELSE COALESCE({kpi_month_year_expr},{file_year_expr},NULLIF(TRIM(y.reporting_year),''),{itr_year_expr}) END"
        base=" FROM source_material m JOIN data_group g ON g.group_id=m.group_id LEFT JOIN source_material_reporting_year y ON y.material_id=m.material_id"
        latest_condition="""NOT EXISTS(SELECT 1 FROM source_material newer
            WHERE newer.group_id=m.group_id AND newer.canonical_itr=m.canonical_itr AND
              (newer.version_no>m.version_no OR (newer.version_no=m.version_no AND
               (newer.created_at>m.created_at OR (newer.created_at=m.created_at AND newer.material_id>m.material_id)))))"""
        where=["g.group_code=?",latest_condition];values=[group_code]
        if q:
            where.append(f"LOWER(COALESCE(m.business_key,'')||' '||COALESCE(m.canonical_itr,'')||' '||COALESCE({title_expr},'')||' '||COALESCE({product_expr},'')) LIKE ?")
            values.append('%'+q.strip().lower()+'%')
        if domain:
            where.append(f"LOWER(COALESCE({domain_expr},'未分类')) LIKE ?");values.append('%'+domain.lower()+'%')
        if month:
            where.append(f"COALESCE({month_expr},'-')=?");values.append(month)
        if year:
            where.append(f"COALESCE({year_expr},'-')=?");values.append(year)
        dimension_filters={'industry':industry,'customer':customer,'ipmt':ipmt,'spdt':spdt,'product_model':product_model,'product_series':product_series}
        for key,value in dimension_filters.items():
            if value:
                where.append(f"COALESCE({dimension_exprs[key]},'')=?");values.append(value)
        where_sql=" WHERE "+" AND ".join(where)
        page=max(1,int(page));page_size=max(1,min(100,int(page_size)));offset=(page-1)*page_size
        select="SELECT m.*,g.group_code,g.group_name,y.reporting_year,y.year_source"
        facet_sql="SELECT m.material_id,m.canonical_itr,m.business_key,"+','.join([
            f"COALESCE({domain_expr},'未分类') domain",f"COALESCE({month_expr},'-') month",f"COALESCE({year_expr},'-') year",
            *[f"COALESCE({expr},'') {key}" for key,expr in dimension_exprs.items()]])+base+" WHERE g.group_code=? AND "+latest_condition
        with self.connect() as c:
            total=c.execute("SELECT COUNT(*)"+base+where_sql,values).fetchone()[0]
            rows=c.execute(select+base+where_sql+" ORDER BY m.created_at DESC,m.material_id LIMIT ? OFFSET ?",[*values,page_size,offset]).fetchall()
            facets=[dict(row) for row in c.execute(facet_sql,(group_code,))]
        items=[_material_view(dict(row)) for row in rows]
        all_filters={'domain':domain,'month':month,'year':year,**dimension_filters}
        def related_options(key):
            counts={}
            for row in facets:
                if any(value and other!=key and str(row.get(other) or '')!=str(value) for other,value in all_filters.items()):continue
                label=str(row.get(key) or '').strip()
                if not label or label=='-':continue
                issue=normalize_itr(row.get('canonical_itr') or row.get('business_key')) or row['material_id']
                counts.setdefault(label,set()).add(issue)
            return [{'label':label,'count':len(ids)} for label,ids in sorted(counts.items(),key=lambda item:(-len(item[1]),item[0]))]
        month_options=related_options('month');month_options.sort(key=lambda row:row['label'],reverse=True)
        year_options=related_options('year');year_options.sort(key=lambda row:row['label'],reverse=True)
        return {"items":items,"total":total,"page":page,"page_size":page_size,
                "pages":max(1,(total+page_size-1)//page_size),
                "domains":[x['label'] for x in related_options('domain')],
                "months":[x['label'] for x in month_options],
                "years":[x['label'] for x in year_options],
                "options":{key:related_options(key) for key in dimension_exprs}}

    def duplicate_summary(self,group_code):
        group=self.group(group_code)
        if not group:raise ValueError('DATA_GROUP_NOT_FOUND')
        with self.connect() as c:
            tables={row[0] for row in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            rows=[dict(row) for row in c.execute("""SELECT * FROM (
                SELECT m.material_id,m.canonical_itr,m.version_no,
                       ROW_NUMBER() OVER(PARTITION BY m.group_id,m.canonical_itr ORDER BY m.version_no DESC,m.created_at DESC,m.material_id DESC) rn,
                       COUNT(*) OVER(PARTITION BY m.group_id,m.canonical_itr) copies
                FROM source_material m WHERE m.group_id=?) WHERE copies>1""",(group['group_id'],))]
            old=[row for row in rows if row['rn']>1];protected=set()
            if old:
                marks=','.join('?' for _ in old);ids=[row['material_id'] for row in old]
                protected.update(row[0] for row in c.execute(f"SELECT material_id FROM material_review WHERE material_id IN ({marks})",ids))
                protected.update(row[0] for row in c.execute(f"SELECT material_id FROM issue_material_link WHERE link_status='MANUAL_LINKED' AND material_id IN ({marks})",ids))
                if 'quality_scenario_evidence' in tables:
                    protected.update(row[0] for row in c.execute(f"SELECT knowledge_id FROM quality_scenario_evidence WHERE knowledge_id IN ({marks})",ids))
            return {'problem_count':len({row['canonical_itr'] for row in rows}),
                    'history_count':len(old),'removable_count':sum(row['material_id'] not in protected for row in old),
                    'protected_count':sum(row['material_id'] in protected for row in old)}

    def cleanup_duplicates(self,group_code):
        group=self.group(group_code)
        if not group:raise ValueError('DATA_GROUP_NOT_FOUND')
        with self.connect() as c:
            tables={row[0] for row in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            rows=[dict(row) for row in c.execute("""SELECT * FROM (
                SELECT m.material_id,m.canonical_itr,
                       ROW_NUMBER() OVER(PARTITION BY m.group_id,m.canonical_itr ORDER BY m.version_no DESC,m.created_at DESC,m.material_id DESC) rn
                FROM source_material m WHERE m.group_id=?) WHERE rn>1""",(group['group_id'],))]
            deleted=0;protected=0
            for row in rows:
                mid=row['material_id']
                used=bool(c.execute("SELECT 1 FROM material_review WHERE material_id=?",(mid,)).fetchone())
                used=used or bool(c.execute("SELECT 1 FROM issue_material_link WHERE material_id=? AND link_status='MANUAL_LINKED'",(mid,)).fetchone())
                if 'quality_scenario_evidence' in tables:
                    used=used or bool(c.execute("SELECT 1 FROM quality_scenario_evidence WHERE knowledge_id=?",(mid,)).fetchone())
                if used:protected+=1;continue
                c.execute("DELETE FROM issue_material_link WHERE material_id=?",(mid,))
                c.execute("DELETE FROM source_material_reporting_year WHERE material_id=?",(mid,))
                c.execute("DELETE FROM source_material WHERE material_id=?",(mid,));deleted+=1
        self.refresh_links()
        return {'deleted':deleted,'protected':protected}

    def raw_field_catalog(self,group_code):
        """List raw Excel fields in one workbench, ordered by record coverage."""
        group=self.group(group_code)
        if not group:raise ValueError('DATA_GROUP_NOT_FOUND')
        with self.connect() as c:
            rows=c.execute("""SELECT fields.key field_name,COUNT(*) record_count
                FROM source_material m JOIN json_each(m.raw_json) fields
                WHERE m.group_id=? GROUP BY fields.key
                ORDER BY record_count DESC,fields.key""",(group['group_id'],)).fetchall()
            return [dict(row) for row in rows]

    def _material_ids_protected_from_delete(self,connection,material_ids):
        ids=list(dict.fromkeys(material_ids or []));protected=set()
        if not ids:return protected
        marks=','.join('?' for _ in ids)
        tables={row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        protected.update(row[0] for row in connection.execute(f"SELECT material_id FROM material_review WHERE material_id IN ({marks})",ids))
        protected.update(row[0] for row in connection.execute(f"SELECT material_id FROM issue_material_link WHERE link_status='MANUAL_LINKED' AND material_id IN ({marks})",ids))
        if 'quality_scenario_evidence' in tables:
            protected.update(row[0] for row in connection.execute(f"SELECT knowledge_id FROM quality_scenario_evidence WHERE knowledge_id IN ({marks})",ids))
        return protected

    def _invalid_import_matches(self,connection,group_id,field_name,operator,value=''):
        operators={'HAS_FIELD','MISSING_FIELD','VALUE_CONTAINS','VALUE_NOT_CONTAINS','VALUE_EQUALS','VALUE_EMPTY','VALUE_NOT_EMPTY'}
        if operator not in operators:raise ValueError('请选择有效的过滤条件')
        field_name=_clean(field_name);value=_clean(value)
        if not field_name:raise ValueError('请选择用于识别错误数据的字段')
        if operator in {'VALUE_CONTAINS','VALUE_NOT_CONTAINS','VALUE_EQUALS'} and not value:
            raise ValueError('当前过滤条件必须填写字段值')
        exists="EXISTS(SELECT 1 FROM json_each(m.raw_json) field WHERE field.key=?)"
        value_expr="COALESCE((SELECT CAST(field.value AS TEXT) FROM json_each(m.raw_json) field WHERE field.key=? LIMIT 1),'')"
        params=[group_id]
        if operator=='HAS_FIELD':condition=exists;params.append(field_name)
        elif operator=='MISSING_FIELD':condition='NOT '+exists;params.append(field_name)
        elif operator=='VALUE_EMPTY':condition=exists+f" AND TRIM({value_expr})=''";params.extend([field_name,field_name])
        elif operator=='VALUE_NOT_EMPTY':condition=exists+f" AND TRIM({value_expr})<>''";params.extend([field_name,field_name])
        elif operator=='VALUE_CONTAINS':condition=exists+f" AND LOWER({value_expr}) LIKE ?";params.extend([field_name,field_name,'%'+value.lower()+'%'])
        elif operator=='VALUE_NOT_CONTAINS':condition=exists+f" AND LOWER({value_expr}) NOT LIKE ?";params.extend([field_name,field_name,'%'+value.lower()+'%'])
        else:condition=exists+f" AND LOWER(TRIM({value_expr}))=?";params.extend([field_name,field_name,value.lower()])
        rows=connection.execute(f"""SELECT m.material_id,m.business_key,m.canonical_itr,m.version_no,
                m.source_file,m.sheet_name,m.row_number,m.raw_json,m.created_at
            FROM source_material m WHERE m.group_id=? AND {condition}
            ORDER BY m.created_at DESC,m.source_file,m.sheet_name,m.row_number""",params).fetchall()
        return [dict(row) for row in rows]

    def invalid_import_preview(self,group_code,*,field_name,operator,value=''):
        """Preview physical import records matched by a raw-field rule, including history."""
        group=self.group(group_code)
        if not group:raise ValueError('DATA_GROUP_NOT_FOUND')
        with self.connect() as c:
            rows=self._invalid_import_matches(c,group['group_id'],field_name,operator,value)
            protected=self._material_ids_protected_from_delete(c,[row['material_id'] for row in rows])
        sources={}
        for row in rows:
            key=(row.get('source_file') or '未知文件',row.get('sheet_name') or '未知Sheet')
            sources[key]=sources.get(key,0)+1
        samples=[]
        for row in rows[:20]:
            raw=json.loads(row.pop('raw_json','{}') or '{}')
            row['title']=_first(raw,'问题信息_问题主题','问题信息_问题描述','问题主题','问题描述') or '未提供问题描述'
            row['protected']=row['material_id'] in protected;samples.append(row)
        return {'total':len(rows),'deletable':len(rows)-len(protected),'protected':len(protected),'samples':samples,
                'sources':[{'source_file':key[0],'sheet_name':key[1],'count':count} for key,count in sorted(sources.items(),key=lambda item:(-item[1],item[0]))]}

    def cleanup_invalid_import(self,group_code,*,field_name,operator,value=''):
        """Delete matched import records while preserving reviewed or referenced evidence."""
        group=self.group(group_code)
        if not group:raise ValueError('DATA_GROUP_NOT_FOUND')
        with self.connect() as c:
            rows=self._invalid_import_matches(c,group['group_id'],field_name,operator,value)
            ids=[row['material_id'] for row in rows]
            protected=self._material_ids_protected_from_delete(c,ids)
            deleted=0
            for material_id in ids:
                if material_id in protected:continue
                c.execute("DELETE FROM issue_material_link WHERE material_id=?",(material_id,))
                c.execute("DELETE FROM source_material_reporting_year WHERE material_id=?",(material_id,))
                c.execute("DELETE FROM source_material WHERE material_id=?",(material_id,));deleted+=1
        self.refresh_links()
        return {'matched':len(ids),'deleted':deleted,'protected':len(protected)}

    def software_operation_distribution(self, *, ipmt='', spdt='', product_model='', product_series='', year='', month=''):
        """Top customer/industry distribution under one associated software-operation scope."""
        scoped=self.search_materials('SW-OPS',ipmt=ipmt,spdt=spdt,product_model=product_model,
            product_series=product_series,year=year,month=month,page=1,page_size=1)
        return {'total':scoped['total'],'industries':scoped['options']['industry'][:10],
                'customers':scoped['options']['customer'][:10],'options':scoped['options'],
                'years':scoped['years'],'months':scoped['months']}

    def material(self, material_id):
        with self.connect() as c:
            row=c.execute("SELECT m.*,g.group_code,g.group_name,y.reporting_year,y.year_source FROM source_material m JOIN data_group g ON g.group_id=m.group_id LEFT JOIN source_material_reporting_year y ON y.material_id=m.material_id WHERE m.material_id=?",(material_id,)).fetchone()
            if not row:return None
            item=_material_view(dict(row))
            link=c.execute("SELECT knowledge_id,link_status FROM issue_material_link WHERE material_id=? ORDER BY created_at DESC LIMIT 1",(material_id,)).fetchone()
            item["linked_issue"]=dict(link) if link else None
            review=c.execute("SELECT * FROM material_review WHERE material_id=?",(material_id,)).fetchone()
            item["review"]=dict(review) if review else {"review_status":"DRAFT","analysis_summary":"","root_cause":"","improvement_action":"","reviewer":""}
            return item

    def set_reporting_year(self, material_ids, year, *, source="BATCH_MANUAL", preserve_manual=False):
        year=validate_reporting_year(year);ids=list(dict.fromkeys(material_ids or []))
        if not ids:raise ValueError("请至少选择一条软件考核问题")
        with self.connect() as c:
            rows=c.execute("SELECT m.material_id,g.group_code FROM source_material m JOIN data_group g ON g.group_id=m.group_id WHERE m.material_id IN ("+','.join('?' for _ in ids)+")",ids).fetchall()
            if len(rows)!=len(ids) or any(row['group_code']!='SW-OPS' for row in rows):raise ValueError("只能修改软件考核工作台中的问题年份")
            for material_id in ids:
                current=c.execute('SELECT year_source FROM source_material_reporting_year WHERE material_id=?',(material_id,)).fetchone()
                if preserve_manual and current and current[0] in ('BATCH_MANUAL','IMPORT_MANUAL'):continue
                c.execute("INSERT INTO source_material_reporting_year(material_id,reporting_year,year_source,updated_at) VALUES(?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(material_id) DO UPDATE SET reporting_year=excluded.reporting_year,year_source=excluded.year_source,updated_at=CURRENT_TIMESTAMP",(material_id,year,source))
        return len(ids)

    def save_review(self, material_id, *, review_status, analysis_summary, root_cause, improvement_action, reviewer):
        if review_status not in {"DRAFT","COMPLETED"}:raise ValueError("INVALID_REVIEW_STATUS")
        with self.connect() as c:
            if not c.execute("SELECT 1 FROM source_material WHERE material_id=?",(material_id,)).fetchone():raise KeyError(material_id)
            c.execute("""INSERT INTO material_review(material_id,review_status,analysis_summary,root_cause,improvement_action,reviewer,updated_at)
              VALUES(?,?,?,?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(material_id) DO UPDATE SET review_status=excluded.review_status,
              analysis_summary=excluded.analysis_summary,root_cause=excluded.root_cause,improvement_action=excluded.improvement_action,
              reviewer=excluded.reviewer,updated_at=CURRENT_TIMESTAMP""",(material_id,review_status,analysis_summary.strip(),root_cause.strip(),improvement_action.strip(),reviewer.strip()))

    def related_materials(self, canonical_itr, exclude_material_id=""):
        with self.connect() as c:
            rows=c.execute("SELECT m.*,g.group_code,g.group_name FROM source_material m JOIN data_group g ON g.group_id=m.group_id WHERE m.canonical_itr=? AND m.material_id<>? ORDER BY m.created_at DESC",(canonical_itr,exclude_material_id)).fetchall()
            return [_material_view(dict(row)) for row in rows]

    def rules(self):
        with self.connect() as c:return [dict(x) for x in c.execute("SELECT * FROM association_rule ORDER BY priority,rule_name")]

    def update_rule(self, rule_id, *, source_field, target_field, transform, status):
        if transform not in {"EXACT","NORMALIZE_ITR","STRIP_CS","APPEND_CS"}:raise ValueError("INVALID_TRANSFORM")
        if status not in {"ACTIVE","INACTIVE"}:raise ValueError("INVALID_STATUS")
        with self.connect() as c:
            if not c.execute("SELECT 1 FROM association_rule WHERE rule_id=?",(rule_id,)).fetchone():raise KeyError(rule_id)
            c.execute("UPDATE association_rule SET source_field=?,target_field=?,transform=?,status=?,version_no=version_no+1 WHERE rule_id=?",(source_field.strip(),target_field.strip(),transform,status,rule_id))

    def link_preview(self):
        with self.connect() as c:
            candidates=c.execute("SELECT COUNT(DISTINCT canonical_itr) FROM source_material WHERE canonical_itr<>''").fetchone()[0]
            material_itrs={normalize_itr(row[0]) for row in c.execute("SELECT canonical_itr FROM source_material WHERE canonical_itr<>''")}
            issues_by_itr={}
            try:issue_rows=c.execute("SELECT knowledge_id,business_issue_id FROM quality_issue").fetchall()
            except sqlite3.OperationalError:issue_rows=[]
            for row in issue_rows:issues_by_itr.setdefault(normalize_itr(row['business_issue_id']),set()).add(row['knowledge_id'])
            matched=sum(1 for value in material_itrs if issues_by_itr.get(value))
            conflicts=sum(1 for value in material_itrs if len(issues_by_itr.get(value,set()))>1)
            return {"candidates":candidates,"matched":matched,"unmatched":max(0,candidates-matched),"conflicts":conflicts}


class MaterialImportService:
    KEY_ALIASES = {
        "ITR_SOURCE": ("问题信息_ITR单号", "ITR单号"),
        "ITR_CS": ("问题信息_彻底解决单号", "彻底解决单号"),
        "SOFTWARE_OPERATION": ("问题信息_彻底解决单号", "彻底解决单号"),
    }

    def __init__(self, repository: MaterialRepository):
        self.repository = repository

    def import_file(self, path, group_code, header_rows=2, sheet_name="", reporting_year=""):
        group=self.repository.group(group_code)
        if not group: raise ValueError("DATA_GROUP_NOT_FOUND")
        if group["material_type"] not in {"ITR_SOURCE","ITR_CS","SOFTWARE_OPERATION"}: raise ValueError("MATERIAL_IMPORT_NOT_ENABLED")
        if reporting_year:validate_reporting_year(reporting_year)
        stats={"total":0,"new":0,"updated":0,"skipped":0,"failed":0,"errors":[]}
        workbook=load_workbook(path,read_only=True,data_only=True)
        try:
            sheets=[sheet_name] if sheet_name else workbook.sheetnames
            for name in sheets:
                sheet=workbook[name]
                rows=sheet.iter_rows(values_only=True)
                first=next(rows,())
                if int(header_rows)==2:
                    second=next(rows,());headers=combine_headers(first,second);start=3
                else:
                    headers=[_clean(x) or f"未命名字段{i+1}" for i,x in enumerate(first)];start=2
                for row_number,values in enumerate(rows,start=start):
                    if not any(_clean(x) for x in values):continue
                    raw={headers[i]:values[i] for i in range(min(len(headers),len(values)))};stats["total"]+=1
                    key=next((_clean(raw.get(alias)) for alias in self.KEY_ALIASES[group["material_type"]] if _clean(raw.get(alias))),"")
                    if not key:
                        stats["failed"]+=1;stats["errors"].append({"sheet":name,"row":row_number,"error":"BUSINESS_KEY_MISSING"});continue
                    material_id,action=self.repository.add_material(group,key,raw,Path(path).name,name,row_number);stats[action.lower()]+=1
                    if group["material_type"]=="SOFTWARE_OPERATION":
                        file_year=_first(raw,"数据运营_KPI计入年份","数据运营_KPI计入年度","KPI计入年份","考核年份")
                        kpi_year=effective_reporting_year(raw,key)
                        itr_year=year_from_itr(key);effective=reporting_year or kpi_year
                        source="IMPORT_MANUAL" if reporting_year else "IMPORT_FILE" if (kpi_year and kpi_year!=itr_year) or file_year else "AUTO_ITR"
                        if effective:self.repository.set_reporting_year([material_id],effective,source=source,preserve_manual=not bool(reporting_year))
        finally:
            workbook.close()
        self.repository.refresh_links()
        return stats
