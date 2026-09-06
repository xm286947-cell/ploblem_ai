from __future__ import annotations

import json
import hashlib
import re
import sqlite3
import uuid

from openpyxl import load_workbook
from quality_knowledge.quality_models import CUSTOMER_EXPERIENCE_MODEL, PRODUCT_QUALITY_MODEL, QUALITY_IN_USE_MODEL, model_payload
from quality_knowledge.scenario_semantics import DEFAULT_TERMS, SEMANTIC_PRINCIPLES, SEMANTIC_TYPES


LIFECYCLES = (
    ("ENGINEERING_CONFIGURATION", "工程配置", "把控制需求转化为可执行、可调试的PLC工程"),
    ("SOFTWARE_DEBUGGING", "软件调试", "把能编译的工程调成能正确工作的工程"),
    ("RUNTIME_EXECUTION", "运行执行", "保证PLC控制逻辑实时、正确、稳定地执行"),
    ("SYSTEM_INTEGRATION", "系统联动", "保证PLC与其他设备、控制系统共同完成完整业务过程"),
    ("LONG_TERM_OPERATION", "长稳运行", "保证系统在长时间、高频、复杂条件下持续稳定"),
    ("VERSION_MAINTENANCE", "版本维护", "保证已交付系统在持续变化中安全、兼容、可恢复"),
)

ACTIVITIES = (
    ("ENGINEERING_CONFIGURATION","PROJECT_INIT","工程创建与初始化","新建工程 → 选择PLC → 初始化工程环境 → 进入工程配置"),
    ("ENGINEERING_CONFIGURATION","HARDWARE_CONFIGURATION","硬件与设备组态","工程创建 → PLC配置 → IO配置 → 设备资源建立"),
    ("ENGINEERING_CONFIGURATION","NETWORK_CONFIGURATION","网络与通信配置","设备组态 → 网络拓扑 → 通信参数 → 通信关系"),
    ("ENGINEERING_CONFIGURATION","CONTROL_PARAMETER_CONFIGURATION","设备与控制参数配置","设备组态 → 参数设置 → 参数检查 → 配置完成"),
    ("ENGINEERING_CONFIGURATION","VARIABLE_DATA_MODELING","变量与数据建模","设备配置 → IO映射 → 变量定义 → 数据结构 → 程序使用"),
    ("ENGINEERING_CONFIGURATION","CONTROL_PROGRAMMING","控制程序开发","控制需求 → 程序架构 → 逻辑开发 → 功能块调用"),
    ("ENGINEERING_CONFIGURATION","BUILD_VALIDATION","工程编译与校验","配置编程 → 编译 → 错误检查 → 修改 → 编译通过"),
    ("SOFTWARE_DEBUGGING","PLC_DISCOVERY_CONNECTION","PLC发现与连接","工程准备 → 查找PLC → 网络连接 → 在线建立"),
    ("SOFTWARE_DEBUGGING","DOWNLOAD_START","工程下载与启动","PLC连接 → 下载配置/程序 → 初始化 → RUN"),
    ("SOFTWARE_DEBUGGING","ONLINE_MONITORING","程序在线监控与调试","下载运行 → 在线监控 → 变量观察 → 状态分析 → 调整"),
    ("SOFTWARE_DEBUGGING","VARIABLE_IO_DEBUG","变量与IO调试","在线运行 → 变量监控/写值 → IO检查 → 功能验证"),
    ("SOFTWARE_DEBUGGING","CONTROL_FUNCTION_DEBUG","控制功能调试","单点验证 → 功能动作 → 顺序验证 → 参数调整"),
    ("SOFTWARE_DEBUGGING","PROBLEM_DIAGNOSIS","运行问题分析与诊断","发现异常 → 信息采集 → Trace/诊断 → 定位 → 验证"),
    ("RUNTIME_EXECUTION","CONTROL_PROGRAM_EXECUTION","PLC控制程序执行","输入采集 → 程序执行 → 逻辑计算 → 输出更新"),
    ("RUNTIME_EXECUTION","TASK_CYCLE_EXECUTION","任务与周期执行","任务触发 → 程序调度 → 执行 → 周期完成"),
    ("RUNTIME_EXECUTION","STATE_DATA_PROCESSING","设备状态与数据处理","数据采集 → 状态判断 → 运算 → 状态更新"),
    ("RUNTIME_EXECUTION","POWER_LOSS_RETENTION_RECOVERY","掉电数据保持与上电恢复","正常运行 → 关键数据/状态产生 → 掉电 → 数据保持 → 重新上电 → 数据恢复 → 程序继续运行"),
    ("RUNTIME_EXECUTION","RUNTIME_EXCEPTION_HANDLING","运行异常处理","运行 → 异常触发 → 识别 → 安全逻辑 → 恢复/停机"),
    ("SYSTEM_INTEGRATION","FIELD_DEVICE_INTEGRATION","PLC与现场设备联动","PLC逻辑 → IO/驱动指令 → 设备动作 → 状态反馈"),
    ("SYSTEM_INTEGRATION","MULTI_DEVICE_SEQUENCE","多设备顺序联动","设备A完成 → 状态握手 → PLC判断 → 设备B启动"),
    ("SYSTEM_INTEGRATION","PLC_SYSTEM_INTEGRATION","PLC与PLC/控制系统联动","PLC A状态 → 通信交换 → PLC B判断 → 联动动作"),
    ("SYSTEM_INTEGRATION","HMI_HOST_INTEGRATION","PLC与HMI/上位系统联动","操作指令 → PLC接收 → 执行 → 状态反馈 → 界面更新"),
    ("SYSTEM_INTEGRATION","INTERLOCK_RECOVERY","异常联锁与协同恢复","设备异常 → 联锁传播 → 相关响应 → 处理 → 联动恢复"),
    ("LONG_TERM_OPERATION","CONTINUOUS_PRODUCTION","连续生产运行","启动生产 → 周期执行 → 重复业务 → 持续运行"),
    ("LONG_TERM_OPERATION","HIGH_LOAD_OPERATION","高频/高负载持续运行","高频输入/通信 → 高频执行 → 大量数据 → 持续输出"),
    ("LONG_TERM_OPERATION","LONG_TERM_COMMUNICATION","长期通信与设备交互","建链 → 数据交互 → 波动/断连 → 自动恢复"),
    ("LONG_TERM_OPERATION","RESOURCE_STATE_RETENTION","长期资源与状态保持","持续运行 → 资源使用 → 数据/状态累计 → 长周期检查"),
    ("VERSION_MAINTENANCE","INSTALL_ENVIRONMENT","iFA Evolution安装与运行环境构建","安装包 → 软件安装 → 环境配置 → 启动验证"),
    ("VERSION_MAINTENANCE","UPGRADE_DOWNGRADE","软件升级/降级","当前版本 → 升降级 → 环境更新 → 工程打开 → 验证"),
    ("VERSION_MAINTENANCE","PROJECT_MIGRATION","工程版本迁移","老工程 → 新版本转换 → 差异处理 → 编译下载"),
    ("VERSION_MAINTENANCE","FIRMWARE_MAINTENANCE","PLC/固件版本维护","固件升级 → 工程匹配 → 下载 → 功能验证"),
    ("VERSION_MAINTENANCE","DEVICE_CHANGE","硬件/外围设备变更","设备变更 → 配置修改 → 程序适配 → 调试验证"),
    ("VERSION_MAINTENANCE","BACKUP_RESTORE","工程备份与恢复","正常工程 → 备份 → 异常/换机 → 恢复验证"),
    ("VERSION_MAINTENANCE","FUNCTION_CHANGE_RELEASE","程序功能变更与再发布","新需求/问题 → 修改 → 编译 → 下载 → 回归"),
    ("VERSION_MAINTENANCE","VERSION_ROLLBACK","版本回退","新版本异常 → 回退判断 → 恢复旧版 → 验证"),
)

ACTIVITY_DESCRIPTIONS = {
    "POWER_LOSS_RETENTION_RECOVERY": "保证PLC在异常掉电或正常断电后，关键运行数据、参数、计数值和状态能够按预期保持，并在重新上电后正确恢复，避免业务状态丢失或设备行为异常。目标：掉电不丢关键数据，上电后恢复正确、及时、一致。",
}

SCENARIO_CAPABILITIES = {
    "ENGINEERING": (("REQUIREMENT_ENGINEERING","需求工程"),("ARCHITECTURE_DESIGN","架构与方案设计"),("DETAILED_DESIGN","详细设计"),("SOFTWARE_IMPLEMENTATION","软件实现"),("HW_SW_CO_DESIGN","软硬协同设计"),("SYSTEM_INTEGRATION","系统集成"),("BUILD_RELEASE","构建与发布工程")),
    "TEST": (("TEST_STRATEGY","测试策略"),("TEST_METHOD","测试方法"),("SCENARIO_COVERAGE","场景覆盖"),("TEST_ENVIRONMENT","测试环境"),("TEST_DATA","测试数据"),("AUTOMATION","测试自动化"),("OBSERVABILITY","可观测与诊断")),
    "MANAGEMENT": (("REVIEW","评审机制"),("BASELINE_MANAGEMENT","基线管理"),("CHANGE_MANAGEMENT","变更管理"),("CONFIGURATION_MANAGEMENT","配置管理"),("KNOWN_ISSUE_CONTROL","已知问题控制"),("RELEASE_GATE","发布门禁"),("RISK_APPROVAL","风险批准"),("HORIZONTAL_GOVERNANCE","横向治理")),
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS scenario_taxonomy_version(version_id TEXT PRIMARY KEY,version_no INTEGER NOT NULL UNIQUE,product_code TEXT NOT NULL DEFAULT 'PLC',status TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,activated_at TEXT);
CREATE TABLE IF NOT EXISTS scenario_lifecycle(version_id TEXT NOT NULL,lifecycle_code TEXT NOT NULL,label_zh TEXT NOT NULL,description TEXT,enabled INTEGER NOT NULL DEFAULT 1,sort_order INTEGER NOT NULL,PRIMARY KEY(version_id,lifecycle_code));
CREATE TABLE IF NOT EXISTS scenario_activity(version_id TEXT NOT NULL,activity_code TEXT NOT NULL,lifecycle_code TEXT NOT NULL,label_zh TEXT NOT NULL,chain_text TEXT,description TEXT,enabled INTEGER NOT NULL DEFAULT 1,sort_order INTEGER NOT NULL,PRIMARY KEY(version_id,activity_code));
CREATE TABLE IF NOT EXISTS quality_scenario(scenario_id TEXT PRIMARY KEY,scenario_code TEXT NOT NULL UNIQUE,name TEXT NOT NULL,product_code TEXT NOT NULL DEFAULT '',taxonomy_version_id TEXT,lifecycle_code TEXT,activity_code TEXT,scenario_chain TEXT,experience_requirement TEXT,concern_points TEXT,quality_attribute TEXT,quality_subcharacteristic TEXT,applicable_boundary TEXT,validation_direction TEXT,measurement_suggestion TEXT,failure_mode TEXT,failure_mechanism TEXT,trigger_conditions TEXT,preconditions TEXT,participating_systems TEXT,system_scale TEXT,user_type TEXT,affected_object TEXT,business_impact TEXT,recovery_method TEXT,status TEXT NOT NULL DEFAULT 'DRAFT',version_no INTEGER NOT NULL DEFAULT 1,created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_scenario_scope(scenario_id TEXT NOT NULL,scope_type TEXT NOT NULL,scope_value TEXT NOT NULL,PRIMARY KEY(scenario_id,scope_type,scope_value));
CREATE TABLE IF NOT EXISTS quality_scenario_generation(generation_id TEXT PRIMARY KEY,product_code TEXT,start_month TEXT,end_month TEXT,source_issue_count INTEGER,candidate_count INTEGER,model_name TEXT,created_by TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_scenario_evidence(scenario_id TEXT NOT NULL,knowledge_id TEXT NOT NULL,evidence_summary TEXT,PRIMARY KEY(scenario_id,knowledge_id));
CREATE TABLE IF NOT EXISTS quality_scenario_duplicate(candidate_id TEXT NOT NULL,existing_id TEXT NOT NULL,similarity REAL NOT NULL,reason TEXT,PRIMARY KEY(candidate_id,existing_id));
CREATE TABLE IF NOT EXISTS quality_scenario_generation_candidate(generation_id TEXT NOT NULL,scenario_id TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(generation_id,scenario_id));
CREATE TABLE IF NOT EXISTS quality_scenario_issue_classification(generation_id TEXT NOT NULL,knowledge_id TEXT NOT NULL,business_issue_id TEXT,status TEXT NOT NULL DEFAULT 'PENDING',scenario_id TEXT,activity_code TEXT,lifecycle_code TEXT,error_message TEXT,updated_at TEXT DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(generation_id,knowledge_id));
CREATE TABLE IF NOT EXISTS quality_scenario_analysis_cache(input_key TEXT PRIMARY KEY,canonical_itr TEXT,task_type TEXT NOT NULL,evidence_hash TEXT NOT NULL,taxonomy_version_id TEXT NOT NULL,scenario_id TEXT NOT NULL,source_knowledge_id TEXT,model_name TEXT,status TEXT NOT NULL DEFAULT 'COMPLETED',created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_scenario_analysis_claim(input_key TEXT PRIMARY KEY,generation_id TEXT NOT NULL,knowledge_id TEXT NOT NULL,claimed_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_scenario_industry_variant(variant_id TEXT PRIMARY KEY,scenario_id TEXT NOT NULL,industry TEXT NOT NULL,product_models TEXT,trigger_conditions TEXT,business_impact TEXT,recovery_method TEXT,evidence_count INTEGER NOT NULL DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(scenario_id,industry));
CREATE TABLE IF NOT EXISTS quality_model_version(model_code TEXT PRIMARY KEY,label_zh TEXT NOT NULL,status TEXT NOT NULL,standard_ref TEXT NOT NULL,activated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_model_term(model_code TEXT NOT NULL,term_code TEXT NOT NULL,parent_code TEXT,label_zh TEXT NOT NULL,sort_order INTEGER NOT NULL DEFAULT 0,enabled INTEGER NOT NULL DEFAULT 1,PRIMARY KEY(model_code,term_code));
CREATE TABLE IF NOT EXISTS customer_experience_quality_map(experience_code TEXT NOT NULL,target_model_code TEXT NOT NULL,target_term_code TEXT NOT NULL,PRIMARY KEY(experience_code,target_model_code,target_term_code));
CREATE TABLE IF NOT EXISTS quality_scenario_standardization_batch(batch_id TEXT PRIMARY KEY,status TEXT NOT NULL,total_count INTEGER NOT NULL,completed_count INTEGER NOT NULL DEFAULT 0,failed_count INTEGER NOT NULL DEFAULT 0,skipped_count INTEGER NOT NULL DEFAULT 0,created_by TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,finished_at TEXT);
CREATE TABLE IF NOT EXISTS quality_scenario_standardization_item(batch_id TEXT NOT NULL,scenario_id TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'PENDING',error_message TEXT,agent_id TEXT,model_name TEXT,updated_at TEXT DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(batch_id,scenario_id));
CREATE TABLE IF NOT EXISTS quality_scenario_confirmation(confirmation_id TEXT PRIMARY KEY,scenario_id TEXT NOT NULL,confirmed_by TEXT NOT NULL,previous_status TEXT,new_status TEXT,before_json TEXT,after_json TEXT,confirmed_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS quality_scenario_capability_gap(gap_id TEXT PRIMARY KEY,scenario_id TEXT NOT NULL,capability_axis TEXT NOT NULL,capability_code TEXT NOT NULL,gap_description TEXT NOT NULL,source_basis TEXT,improvement_action TEXT,verification_metric TEXT,priority TEXT NOT NULL DEFAULT 'P1',status TEXT NOT NULL DEFAULT 'OPEN',created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS scenario_semantic_term(term_id TEXT PRIMARY KEY,term_type TEXT NOT NULL,term_code TEXT NOT NULL,label_zh TEXT NOT NULL,definition TEXT NOT NULL,inclusion_criteria TEXT,exclusion_criteria TEXT,aliases_json TEXT,product_code TEXT NOT NULL DEFAULT '',status TEXT NOT NULL DEFAULT 'ACTIVE',source_scenario_id TEXT,nearest_term_code TEXT,difference_note TEXT,merged_into_code TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,updated_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(term_type,term_code,product_code));
"""


class ScenarioRepository:
    def __init__(self, db_path):
        self.db_path=str(db_path)
        with self.connect() as c:
            c.executescript(SCHEMA)
            taxonomy_columns={row['name'] for row in c.execute("PRAGMA table_info(scenario_taxonomy_version)")}
            if 'product_code' not in taxonomy_columns:c.execute("ALTER TABLE scenario_taxonomy_version ADD COLUMN product_code TEXT NOT NULL DEFAULT 'PLC'")
            lifecycle_columns={row['name'] for row in c.execute("PRAGMA table_info(scenario_lifecycle)")}
            for name in ('value_statement','objective'):
                if name not in lifecycle_columns:c.execute(f"ALTER TABLE scenario_lifecycle ADD COLUMN {name} TEXT")
            activity_columns={row['name'] for row in c.execute("PRAGMA table_info(scenario_activity)")}
            for name in ('participating_systems','objective'):
                if name not in activity_columns:c.execute(f"ALTER TABLE scenario_activity ADD COLUMN {name} TEXT")
            scenario_columns={row['name'] for row in c.execute("PRAGMA table_info(quality_scenario)")}
            for name in ('scenario_chain','quality_subcharacteristic','measurement_suggestion','failure_mode','failure_mechanism','trigger_conditions','preconditions','participating_systems','system_scale','user_type','affected_object','business_impact','recovery_method','product_code','taxonomy_version_id','customer_perception','primary_experience_code','secondary_experience_codes','quality_in_use_codes','primary_quality_characteristic_code','secondary_quality_characteristic_codes','quality_subcharacteristic_codes','quality_classification_status','quality_model_version','quality_classification_error','quality_classification_agent','quality_classification_model','quality_classification_updated_at','primary_typical_problem_code','secondary_typical_problem_codes','primary_quality_concern_code','secondary_quality_concern_codes','primary_customer_experience_statement','secondary_customer_experience_statements','operating_environment','operating_condition','duration_frequency','disturbances','extreme_conditions','environment_condition_codes'):
                if name not in scenario_columns:c.execute(f"ALTER TABLE quality_scenario ADD COLUMN {name} TEXT")
            self._seed_quality_models(c)
            self._seed_semantic_terms(c)
            columns={row['name'] for row in c.execute("PRAGMA table_info(quality_scenario_generation)")}
            for name,definition in {'status':"TEXT NOT NULL DEFAULT 'COMPLETED'",'progress_text':"TEXT",'error_message':"TEXT",'finished_at':"TEXT",'processed_count':"INTEGER NOT NULL DEFAULT 0",'classified_count':"INTEGER NOT NULL DEFAULT 0",'reused_count':"INTEGER NOT NULL DEFAULT 0",'updated_count':"INTEGER NOT NULL DEFAULT 0",'review_required_count':"INTEGER NOT NULL DEFAULT 0",'failed_count':"INTEGER NOT NULL DEFAULT 0",'unprocessed_count':"INTEGER NOT NULL DEFAULT 0",'taxonomy_version_id':"TEXT"}.items():
                if name not in columns:c.execute(f"ALTER TABLE quality_scenario_generation ADD COLUMN {name} {definition}")
            ledger_columns={row['name'] for row in c.execute("PRAGMA table_info(quality_scenario_issue_classification)")}
            for name,definition in {'agent_id':"TEXT",'model_name':"TEXT",'started_at':"TEXT",'finished_at':"TEXT",'attempt_count':"INTEGER NOT NULL DEFAULT 0"}.items():
                if name not in ledger_columns:c.execute(f"ALTER TABLE quality_scenario_issue_classification ADD COLUMN {name} {definition}")
            for row in c.execute("SELECT scenario_id,evidence_summary FROM quality_scenario_evidence").fetchall():
                try:generation_id=json.loads(row['evidence_summary'] or '{}').get('generation_id')
                except (TypeError,json.JSONDecodeError):generation_id=''
                if generation_id:c.execute("INSERT OR IGNORE INTO quality_scenario_generation_candidate(generation_id,scenario_id) VALUES(?,?)",(generation_id,row['scenario_id']))
            c.execute("""UPDATE quality_scenario_generation SET status='FAILED',progress_text='历史任务未生成有效候选',error_message='本次AI输出没有形成可用场景，请使用新版本重新生成',finished_at=COALESCE(finished_at,CURRENT_TIMESTAMP)
                       WHERE status='COMPLETED' AND COALESCE(candidate_count,0)=0 AND NOT EXISTS(SELECT 1 FROM quality_scenario_generation_candidate x WHERE x.generation_id=quality_scenario_generation.generation_id)""")
            if not c.execute("SELECT 1 FROM scenario_taxonomy_version").fetchone():
                version_id="STV-1";c.execute("INSERT INTO scenario_taxonomy_version(version_id,version_no,product_code,status,activated_at) VALUES(?,1,'PLC','ACTIVE',CURRENT_TIMESTAMP)",(version_id,))
                for order,(code,label,description) in enumerate(LIFECYCLES,1):c.execute("INSERT INTO scenario_lifecycle(version_id,lifecycle_code,label_zh,description,enabled,sort_order) VALUES(?,?,?,?,1,?)",(version_id,code,label,description,order))
                for order,(lifecycle,code,label,chain) in enumerate(ACTIVITIES,1):c.execute("INSERT INTO scenario_activity(version_id,activity_code,lifecycle_code,label_zh,chain_text,description,enabled,sort_order) VALUES(?,?,?,?,?,?,1,?)",(version_id,code,lifecycle,label,chain,ACTIVITY_DESCRIPTIONS.get(code,""),order))
            self._ensure_power_loss_activity(c)
            c.execute("UPDATE quality_scenario SET product_code='PLC' WHERE COALESCE(product_code,'')='' ")
            c.execute("""UPDATE quality_scenario SET taxonomy_version_id=(SELECT version_id FROM scenario_taxonomy_version WHERE product_code=quality_scenario.product_code AND status='ACTIVE' ORDER BY version_no DESC LIMIT 1) WHERE COALESCE(taxonomy_version_id,'')=''""")
            c.execute("""UPDATE quality_scenario SET scenario_chain=(SELECT chain_text FROM scenario_activity WHERE version_id=quality_scenario.taxonomy_version_id AND activity_code=quality_scenario.activity_code) WHERE COALESCE(scenario_chain,'')='' AND EXISTS(SELECT 1 FROM scenario_activity WHERE version_id=quality_scenario.taxonomy_version_id AND activity_code=quality_scenario.activity_code)""")
            self._backfill_context_scopes(c)

    @staticmethod
    def _seed_quality_models(c):
        payload=model_payload();versions=payload['versions']
        for code,label,ref in ((PRODUCT_QUALITY_MODEL,'产品质量模型','ISO/IEC 25010:2023'),(QUALITY_IN_USE_MODEL,'使用质量模型','ISO/IEC 25010:2011'),(CUSTOMER_EXPERIENCE_MODEL,'客户质量体验词典','公司词典 V1')):
            c.execute("INSERT OR REPLACE INTO quality_model_version(model_code,label_zh,status,standard_ref,activated_at) VALUES(?,?,'ACTIVE',?,COALESCE((SELECT activated_at FROM quality_model_version WHERE model_code=?),CURRENT_TIMESTAMP))",(code,label,ref,code))
        groups=((PRODUCT_QUALITY_MODEL,payload['product_characteristics']+payload['product_subcharacteristics']),(QUALITY_IN_USE_MODEL,payload['quality_in_use']),(CUSTOMER_EXPERIENCE_MODEL,payload['customer_experiences']))
        for model_code,items in groups:
            for order,item in enumerate(items,1):c.execute("INSERT OR REPLACE INTO quality_model_term(model_code,term_code,parent_code,label_zh,sort_order,enabled) VALUES(?,?,?,?,?,1)",(model_code,item['code'],item.get('parent_code'),item['label_zh'],order))
        c.execute("DELETE FROM customer_experience_quality_map")
        for item in payload['customer_experiences']:
            for code in item['quality_in_use_codes']:c.execute("INSERT INTO customer_experience_quality_map VALUES(?,?,?)",(item['code'],QUALITY_IN_USE_MODEL,code))
            for code in item['product_characteristic_codes']:c.execute("INSERT INTO customer_experience_quality_map VALUES(?,?,?)",(item['code'],PRODUCT_QUALITY_MODEL,code))

    @staticmethod
    def _seed_semantic_terms(c):
        for term_type,code,label,definition,inclusion,exclusion in DEFAULT_TERMS:
            c.execute("""INSERT OR IGNORE INTO scenario_semantic_term(term_id,term_type,term_code,label_zh,definition,inclusion_criteria,exclusion_criteria,aliases_json,product_code,status) VALUES(?,?,?,?,?,?,?,'[]','','ACTIVE')""",
                      (f'SST-{term_type}-{code}',term_type,code,label,definition,inclusion,exclusion))

    def semantic_dictionary(self, product_code='', include_candidates=True):
        statuses="('ACTIVE','CANDIDATE')" if include_candidates else "('ACTIVE')"
        with self.connect() as c:
            rows=[dict(x) for x in c.execute(f"SELECT * FROM scenario_semantic_term WHERE status IN {statuses} AND (product_code='' OR product_code=?) ORDER BY term_type,status,label_zh",(product_code,))]
        for row in rows:
            try:row['aliases']=json.loads(row.get('aliases_json') or '[]')
            except (TypeError,json.JSONDecodeError):row['aliases']=[]
        return {'principles':SEMANTIC_PRINCIPLES,'types':SEMANTIC_TYPES,'items':rows,**{key.lower():[x for x in rows if x['term_type']==key] for key in SEMANTIC_TYPES}}

    def save_semantic_term(self,payload):
        term_type=str(payload.get('term_type') or '')
        if term_type not in SEMANTIC_TYPES:raise ValueError('SEMANTIC_TERM_TYPE_INVALID')
        code=re.sub(r'[^A-Z0-9_]+','_',str(payload.get('term_code') or '').upper()).strip('_')
        label=str(payload.get('label_zh') or '').strip();definition=str(payload.get('definition') or '').strip()
        if not code or not label or not definition:raise ValueError('SEMANTIC_TERM_REQUIRED_FIELDS')
        status=str(payload.get('status') or 'ACTIVE')
        if status not in {'ACTIVE','CANDIDATE','REJECTED','MERGED','RETIRED'}:raise ValueError('SEMANTIC_TERM_STATUS_INVALID')
        product=str(payload.get('product_code') or '').strip();term_id=str(payload.get('term_id') or f'SST-{uuid.uuid4().hex}')
        aliases=payload.get('aliases') or []
        if isinstance(aliases,str):aliases=[x.strip() for x in re.split('[,，\n]',aliases) if x.strip()]
        with self.connect() as c:c.execute("""INSERT INTO scenario_semantic_term(term_id,term_type,term_code,label_zh,definition,inclusion_criteria,exclusion_criteria,aliases_json,product_code,status,source_scenario_id,nearest_term_code,difference_note) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(term_type,term_code,product_code) DO UPDATE SET label_zh=excluded.label_zh,definition=excluded.definition,inclusion_criteria=excluded.inclusion_criteria,exclusion_criteria=excluded.exclusion_criteria,aliases_json=excluded.aliases_json,status=excluded.status,nearest_term_code=excluded.nearest_term_code,difference_note=excluded.difference_note,updated_at=CURRENT_TIMESTAMP""",
            (term_id,term_type,code,label,definition,str(payload.get('inclusion_criteria') or ''),str(payload.get('exclusion_criteria') or ''),json.dumps(aliases,ensure_ascii=False),product,status,str(payload.get('source_scenario_id') or ''),str(payload.get('nearest_term_code') or ''),str(payload.get('difference_note') or '')))
        return code

    def review_semantic_term(self,term_id,action,target_code=''):
        if action not in {'APPROVE','MERGE','REJECT'}:raise ValueError('SEMANTIC_REVIEW_ACTION_INVALID')
        with self.connect() as c:
            row=c.execute("SELECT * FROM scenario_semantic_term WHERE term_id=? AND status='CANDIDATE'",(term_id,)).fetchone()
            if not row:raise KeyError(term_id)
            if action=='APPROVE':c.execute("UPDATE scenario_semantic_term SET status='ACTIVE',updated_at=CURRENT_TIMESTAMP WHERE term_id=?",(term_id,));return
            replacement=''
            if action=='MERGE':
                target=c.execute("SELECT 1 FROM scenario_semantic_term WHERE term_type=? AND term_code=? AND status='ACTIVE'",(row['term_type'],target_code)).fetchone()
                if not target:raise ValueError('SEMANTIC_MERGE_TARGET_INVALID')
                replacement=target_code
            primary={'TYPICAL_PROBLEM':'primary_typical_problem_code','QUALITY_CONCERN':'primary_quality_concern_code'}.get(row['term_type'])
            secondary={'TYPICAL_PROBLEM':'secondary_typical_problem_codes','QUALITY_CONCERN':'secondary_quality_concern_codes','ENVIRONMENT_CONDITION':'environment_condition_codes'}[row['term_type']]
            if primary:c.execute(f"UPDATE quality_scenario SET {primary}=? WHERE {primary}=?",(replacement,row['term_code']))
            for scenario in c.execute(f"SELECT scenario_id,{secondary} FROM quality_scenario").fetchall():
                try:codes=json.loads(scenario[secondary] or '[]')
                except (TypeError,json.JSONDecodeError):codes=[]
                if row['term_code'] not in codes:continue
                codes=[replacement if x==row['term_code'] else x for x in codes]
                codes=list(dict.fromkeys(x for x in codes if x))
                c.execute(f"UPDATE quality_scenario SET {secondary}=? WHERE scenario_id=?",(json.dumps(codes,ensure_ascii=False),scenario['scenario_id']))
            c.execute("UPDATE scenario_semantic_term SET status=?,merged_into_code=?,updated_at=CURRENT_TIMESTAMP WHERE term_id=?",('MERGED' if action=='MERGE' else 'REJECTED',replacement,term_id))

    def quality_models(self):
        with self.connect() as c:
            versions={x['model_code']:dict(x) for x in c.execute("SELECT * FROM quality_model_version WHERE status='ACTIVE'")}
            terms=[dict(x) for x in c.execute("SELECT * FROM quality_model_term WHERE enabled=1 ORDER BY model_code,sort_order")]
            mappings=[dict(x) for x in c.execute("SELECT * FROM customer_experience_quality_map")]
        by_model={}
        for row in terms:by_model.setdefault(row['model_code'],[]).append(row)
        return {'versions':versions,'product_characteristics':[x for x in by_model.get(PRODUCT_QUALITY_MODEL,[]) if not x['parent_code']], 'product_subcharacteristics':[x for x in by_model.get(PRODUCT_QUALITY_MODEL,[]) if x['parent_code']], 'quality_in_use':by_model.get(QUALITY_IN_USE_MODEL,[]), 'customer_experiences':by_model.get(CUSTOMER_EXPERIENCE_MODEL,[]), 'mappings':mappings}

    @staticmethod
    def capability_dictionary():
        axis_labels={'ENGINEERING':'研发质量工程','TEST':'测试验证','MANAGEMENT':'质量管理'}
        return {'axis_labels':axis_labels,'items':{axis:[{'code':code,'label_zh':label} for code,label in rows] for axis,rows in SCENARIO_CAPABILITIES.items()}}

    @staticmethod
    def _ensure_power_loss_activity(c):
        code="POWER_LOSS_RETENTION_RECOVERY"
        for version in c.execute("SELECT version_id FROM scenario_taxonomy_version WHERE status IN ('ACTIVE','DRAFT')").fetchall():
            version_id=version['version_id']
            if c.execute("SELECT 1 FROM scenario_activity WHERE version_id=? AND activity_code=?",(version_id,code)).fetchone():continue
            next_row=c.execute("SELECT sort_order FROM scenario_activity WHERE version_id=? AND activity_code='RUNTIME_EXCEPTION_HANDLING'",(version_id,)).fetchone()
            order=next_row[0] if next_row else c.execute("SELECT COALESCE(MAX(sort_order),0)+1 FROM scenario_activity WHERE version_id=?",(version_id,)).fetchone()[0]
            c.execute("UPDATE scenario_activity SET sort_order=sort_order+1 WHERE version_id=? AND sort_order>=?",(version_id,order))
            c.execute("INSERT INTO scenario_activity(version_id,activity_code,lifecycle_code,label_zh,chain_text,description,enabled,sort_order) VALUES(?,?,?,?,?,?,1,?)",(version_id,code,"RUNTIME_EXECUTION","掉电数据保持与上电恢复","正常运行 → 关键数据/状态产生 → 掉电 → 数据保持 → 重新上电 → 数据恢复 → 程序继续运行",ACTIVITY_DESCRIPTIONS[code],order))

    @staticmethod
    def _backfill_context_scopes(c):
        aliases={"IPMT":("问题信息_IPMT","IPMT"),"SPDT":("问题信息_SPDT","SPDT"),"PRODUCT_MODEL":("问题信息_产品型号","产品型号"),"INDUSTRY":("问题信息_客户行业","客户行业"),"CUSTOMER_NAME":("问题信息_客户名称","客户名称"),"CUSTOMER_LEVEL":("问题信息_客户分级","客户分级"),"CUSTOMER_STATUS":("问题信息_当前客户状态","问题信息_当前问题状态","问题信息_问题状态","当前客户状态","当前问题状态"),"OCCURRENCE_PHASE":("问题信息_问题发生阶段","问题发生阶段")}
        try:
            rows=c.execute("""SELECT DISTINCT e.scenario_id,m.raw_json FROM quality_scenario_evidence e JOIN issue_material_link l ON l.knowledge_id=e.knowledge_id JOIN source_material m ON m.material_id=l.material_id WHERE m.material_type='ITR_CS'""").fetchall()
        except sqlite3.OperationalError:
            return
        for row in rows:
            try:raw=json.loads(row['raw_json'] or '{}')
            except (TypeError,json.JSONDecodeError):continue
            for kind,names in aliases.items():
                value=next((str(raw.get(name) or '').strip() for name in names if str(raw.get(name) or '').strip()),'')
                if value:c.execute("INSERT OR IGNORE INTO quality_scenario_scope VALUES(?,?,?)",(row['scenario_id'],kind,value))

    def connect(self):
        c=sqlite3.connect(self.db_path);c.row_factory=sqlite3.Row;c.execute("PRAGMA foreign_keys=ON");c.execute("PRAGMA busy_timeout=5000");return c

    def versions(self, product_code=''):
        with self.connect() as c:
            if product_code:return [dict(x) for x in c.execute("SELECT * FROM scenario_taxonomy_version WHERE product_code=? ORDER BY version_no DESC",(product_code,))]
            return [dict(x) for x in c.execute("SELECT * FROM scenario_taxonomy_version ORDER BY product_code,version_no DESC")]

    def working_version(self, product_code='PLC'):
        with self.connect() as c:
            row=c.execute("SELECT * FROM scenario_taxonomy_version WHERE product_code=? ORDER BY CASE status WHEN 'DRAFT' THEN 0 ELSE 1 END,version_no DESC LIMIT 1",(product_code,)).fetchone();return dict(row) if row else None

    def taxonomy(self, version_id="", product_code='PLC'):
        version=self.working_version(product_code) if not version_id else next((x for x in self.versions() if x['version_id']==version_id),None)
        if not version:return None
        with self.connect() as c:
            version['lifecycles']=[dict(x) for x in c.execute("SELECT * FROM scenario_lifecycle WHERE version_id=? ORDER BY sort_order",(version['version_id'],))]
            version['activities']=[dict(x) for x in c.execute("SELECT * FROM scenario_activity WHERE version_id=? ORDER BY sort_order",(version['version_id'],))]
        return version

    def taxonomy_active(self, product_code='PLC'):
        active=next((x for x in self.versions(product_code) if x['status']=='ACTIVE'),None)
        return self.taxonomy(active['version_id'],product_code) if active else None

    def create_draft(self, product_code='PLC', source_product_code=''):
        with self.connect() as c:
            draft=c.execute("SELECT version_id FROM scenario_taxonomy_version WHERE product_code=? AND status='DRAFT' ORDER BY version_no DESC LIMIT 1",(product_code,)).fetchone()
            if draft:return draft[0]
            active=c.execute("SELECT version_id FROM scenario_taxonomy_version WHERE product_code=? AND status='ACTIVE' ORDER BY version_no DESC LIMIT 1",(product_code,)).fetchone()
            source=active or (c.execute("SELECT version_id FROM scenario_taxonomy_version WHERE product_code=? AND status='ACTIVE' ORDER BY version_no DESC LIMIT 1",(source_product_code,)).fetchone() if source_product_code else None)
            version_id=f"STV-{uuid.uuid4().hex}";version_no=c.execute("SELECT COALESCE(MAX(version_no),0)+1 FROM scenario_taxonomy_version").fetchone()[0]
            c.execute("INSERT INTO scenario_taxonomy_version(version_id,version_no,product_code,status) VALUES(?,?,?,'DRAFT')",(version_id,version_no,product_code))
            if source:
                c.execute("INSERT INTO scenario_lifecycle(version_id,lifecycle_code,label_zh,description,enabled,sort_order,value_statement,objective) SELECT ?,lifecycle_code,label_zh,description,enabled,sort_order,value_statement,objective FROM scenario_lifecycle WHERE version_id=?",(version_id,source['version_id']))
                c.execute("INSERT INTO scenario_activity(version_id,activity_code,lifecycle_code,label_zh,chain_text,description,enabled,sort_order,participating_systems,objective) SELECT ?,activity_code,lifecycle_code,label_zh,chain_text,description,enabled,sort_order,participating_systems,objective FROM scenario_activity WHERE version_id=?",(version_id,source['version_id']))
            return version_id

    def save_lifecycle(self, version_id, code, label, description, enabled=True, value_statement='', objective=''):
        with self.connect() as c:
            version=c.execute("SELECT status FROM scenario_taxonomy_version WHERE version_id=?",(version_id,)).fetchone()
            if not version or version['status']!='DRAFT':raise ValueError("TAXONOMY_NOT_EDITABLE")
            order=c.execute("SELECT COALESCE(MAX(sort_order),0)+1 FROM scenario_lifecycle WHERE version_id=?",(version_id,)).fetchone()[0]
            c.execute("""INSERT INTO scenario_lifecycle(version_id,lifecycle_code,label_zh,description,enabled,sort_order,value_statement,objective) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(version_id,lifecycle_code) DO UPDATE SET label_zh=excluded.label_zh,description=excluded.description,enabled=excluded.enabled,value_statement=excluded.value_statement,objective=excluded.objective""",(version_id,code.strip().upper(),label.strip(),description.strip(),int(enabled),order,value_statement.strip(),objective.strip()))

    def save_activity(self, version_id, lifecycle_code, code, label, chain, description, enabled=True, participating_systems='', objective=''):
        with self.connect() as c:
            version=c.execute("SELECT status FROM scenario_taxonomy_version WHERE version_id=?",(version_id,)).fetchone()
            if not version or version['status']!='DRAFT':raise ValueError("TAXONOMY_NOT_EDITABLE")
            if not c.execute("SELECT 1 FROM scenario_lifecycle WHERE version_id=? AND lifecycle_code=?",(version_id,lifecycle_code)).fetchone():raise ValueError("LIFECYCLE_NOT_FOUND")
            order=c.execute("SELECT COALESCE(MAX(sort_order),0)+1 FROM scenario_activity WHERE version_id=?",(version_id,)).fetchone()[0]
            c.execute("""INSERT INTO scenario_activity(version_id,activity_code,lifecycle_code,label_zh,chain_text,description,enabled,sort_order,participating_systems,objective) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(version_id,activity_code) DO UPDATE SET lifecycle_code=excluded.lifecycle_code,label_zh=excluded.label_zh,chain_text=excluded.chain_text,description=excluded.description,enabled=excluded.enabled,participating_systems=excluded.participating_systems,objective=excluded.objective""",(version_id,code.strip().upper(),lifecycle_code,label.strip(),chain.strip(),description.strip(),int(enabled),order,participating_systems.strip(),objective.strip()))

    def save_activities(self,version_id,items):
        for item in items:
            self.save_activity(version_id,item['lifecycle_code'],item['activity_code'],item['label_zh'],item.get('chain_text',''),item.get('description',''),item.get('enabled',True),item.get('participating_systems',''),item.get('objective',''))

    @staticmethod
    def _import_code(prefix,label):
        text=re.sub(r'[^A-Z0-9]+','_',str(label or '').upper()).strip('_')
        digest=hashlib.sha1(str(label).encode('utf-8')).hexdigest()[:10].upper()
        return f"{prefix}_{text[:32]}_{digest}" if text else f"{prefix}_{digest}"

    def import_taxonomy_workbook(self,version_id,source):
        taxonomy=self.taxonomy(version_id)
        if not taxonomy or taxonomy['status']!='DRAFT':raise ValueError('TAXONOMY_NOT_EDITABLE')
        workbook=load_workbook(source,read_only=True,data_only=True)
        try:
            lifecycle_sheet=workbook['使用生命周期'];activity_sheet=workbook['业务活动场景_重构']
            life_by_label={x['label_zh']:x['lifecycle_code'] for x in taxonomy['lifecycles']}
            lifecycle_count=0
            for row in lifecycle_sheet.iter_rows(min_row=4,values_only=True):
                if len(row)<4:continue
                label=str(row[0] or '').strip()
                if not label:continue
                code=life_by_label.get(label) or self._import_code('LIFE',label)
                self.save_lifecycle(version_id,code,label,str(row[2] or ''),True,str(row[1] or ''),str(row[3] or ''))
                life_by_label[label]=code;lifecycle_count+=1
            existing={(x['lifecycle_code'],x['label_zh']):x['activity_code'] for x in self.taxonomy(version_id)['activities']}
            activity_rows=[]
            for row in activity_sheet.iter_rows(min_row=4,values_only=True):
                if len(row)<6:continue
                lifecycle_label=str(row[0] or '').strip();label=str(row[1] or '').strip()
                if not lifecycle_label or not label:continue
                lifecycle_code=life_by_label.get(lifecycle_label)
                if not lifecycle_code:raise ValueError(f'IMPORT_LIFECYCLE_NOT_FOUND:{lifecycle_label}')
                activity_rows.append((row,lifecycle_code,label))
            with self.connect() as c:c.execute("DELETE FROM scenario_activity WHERE version_id=?",(version_id,))
            activity_count=0
            for row,lifecycle_code,label in activity_rows:
                code=existing.get((lifecycle_code,label)) or self._import_code('ACT',f'{lifecycle_code}_{label}')
                self.save_activity(version_id,lifecycle_code,code,label,str(row[2] or ''),str(row[4] or ''),True,str(row[3] or ''),str(row[5] or ''))
                activity_count+=1
            return {'lifecycle_count':lifecycle_count,'activity_count':activity_count}
        finally:workbook.close()

    def activate(self, version_id):
        with self.connect() as c:
            version=c.execute("SELECT product_code FROM scenario_taxonomy_version WHERE version_id=? AND status='DRAFT'",(version_id,)).fetchone()
            if not version:raise ValueError("DRAFT_NOT_FOUND")
            c.execute("UPDATE scenario_taxonomy_version SET status='RETIRED' WHERE product_code=? AND status='ACTIVE'",(version['product_code'],))
            c.execute("UPDATE scenario_taxonomy_version SET status='ACTIVE',activated_at=CURRENT_TIMESTAMP WHERE version_id=?",(version_id,))

    def scope_options(self):
        values={key:set() for key in ("IPMT","SPDT","PRODUCT_MODEL","INDUSTRY","CUSTOMER_NAME","CUSTOMER_LEVEL","CUSTOMER_STATUS","OCCURRENCE_PHASE")}
        aliases={"IPMT":("问题信息_IPMT","IPMT"),"SPDT":("问题信息_SPDT","SPDT"),"PRODUCT_MODEL":("问题信息_产品型号","产品型号"),"INDUSTRY":("问题信息_客户行业","客户行业"),"CUSTOMER_NAME":("问题信息_客户名称","客户名称"),"CUSTOMER_LEVEL":("问题信息_客户分级","客户分级"),"CUSTOMER_STATUS":("问题信息_当前客户状态","问题信息_当前问题状态","问题信息_问题状态","当前客户状态","当前问题状态"),"OCCURRENCE_PHASE":("问题信息_问题发生阶段","问题发生阶段")}
        with self.connect() as c:
            try:rows=c.execute("SELECT raw_json FROM source_material WHERE material_type='ITR_CS'").fetchall()
            except sqlite3.OperationalError:rows=[]
            for row in rows:
                raw=json.loads(row[0] or "{}")
                for kind,names in aliases.items():
                    for name in names:
                        value=str(raw.get(name) or "").strip()
                        if value:values[kind].add(value);break
            for row in c.execute("SELECT scope_type,scope_value FROM quality_scenario_scope"):
                values.setdefault(row['scope_type'],set()).add(row['scope_value'])
        return {key:sorted(items) for key,items in values.items()}

    def scenarios(self, *, ipmt="", spdt="", product_model="", industry="", customer_name="", q="", status="", generation_id="", activity_code="", experience_code="", qiu_code="", quality_code="", typical_problem_code="", environment_code=""):
        with self.connect() as c:
            rows=[dict(x) for x in c.execute("SELECT * FROM quality_scenario ORDER BY updated_at DESC")]
            scopes=c.execute("SELECT * FROM quality_scenario_scope").fetchall()
            capability_gaps=c.execute("SELECT * FROM quality_scenario_capability_gap ORDER BY priority,updated_at DESC").fetchall()
            generated={x[0] for x in c.execute("SELECT scenario_id FROM quality_scenario_generation_candidate WHERE generation_id=?",(generation_id,))} if generation_id else set()
        by_id={}
        for row in scopes:by_id.setdefault(row['scenario_id'],{}).setdefault(row['scope_type'],[]).append(row['scope_value'])
        gaps_by_id={}
        for row in capability_gaps:gaps_by_id.setdefault(row['scenario_id'],[]).append(dict(row))
        for item in rows:item['scopes']=by_id.get(item['scenario_id'],{});item['capability_gaps']=gaps_by_id.get(item['scenario_id'],[])
        def matches(item):
            s=item['scopes']
            return (not generation_id or item['scenario_id'] in generated) and (not q or q.lower() in (item['name']+' '+item['scenario_code']).lower()) and (not status or item['status']==status) and (not ipmt or ipmt in s.get('IPMT',[])) and (not spdt or spdt in s.get('SPDT',[])) and (not product_model or product_model in s.get('PRODUCT_MODEL',[])) and (not industry or industry in s.get('INDUSTRY',[])) and (not customer_name or customer_name in s.get('CUSTOMER_NAME',[]))
        items=[item for item in rows if matches(item)]
        def codes(item,key,primary=''):
            values=[item.get(primary)] if primary and item.get(primary) else []
            try:values.extend(json.loads(item.get(key) or '[]'))
            except (TypeError,json.JSONDecodeError):pass
            return {str(x) for x in values if x}
        if activity_code:items=[x for x in items if x.get('activity_code')==activity_code]
        if experience_code:items=[x for x in items if experience_code in codes(x,'secondary_experience_codes','primary_experience_code')]
        if qiu_code:items=[x for x in items if qiu_code in codes(x,'quality_in_use_codes')]
        if quality_code:items=[x for x in items if quality_code in codes(x,'secondary_quality_characteristic_codes','primary_quality_characteristic_code')]
        if typical_problem_code:items=[x for x in items if typical_problem_code in codes(x,'secondary_typical_problem_codes','primary_typical_problem_code')]
        if environment_code:items=[x for x in items if environment_code in codes(x,'environment_condition_codes')]
        return items

    def standardization_items(self):
        items=self.scenarios()
        for item in items:item['quality_classification_status']=item.get('quality_classification_status') or 'NOT_ANALYZED'
        return items

    def create_standardization_batch(self,batch_id,scenario_ids,created_by='WEB_USER'):
        ids=list(dict.fromkeys(scenario_ids))
        with self.connect() as c:
            c.execute("INSERT INTO quality_scenario_standardization_batch(batch_id,status,total_count,created_by) VALUES(?,'QUEUED',?,?)",(batch_id,len(ids),created_by))
            for sid in ids:c.execute("INSERT INTO quality_scenario_standardization_item(batch_id,scenario_id,status) VALUES(?,?,'PENDING')",(batch_id,sid))

    def mark_standardization_item(self,batch_id,scenario_id,status,*,error='',agent='',model=''):
        with self.connect() as c:
            c.execute("UPDATE quality_scenario_standardization_item SET status=?,error_message=?,agent_id=COALESCE(NULLIF(?,''),agent_id),model_name=COALESCE(NULLIF(?,''),model_name),updated_at=CURRENT_TIMESTAMP WHERE batch_id=? AND scenario_id=?",(status,error,agent,model,batch_id,scenario_id))
        self.refresh_standardization_batch(batch_id)

    def refresh_standardization_batch(self,batch_id):
        with self.connect() as c:
            counts={x['status']:x['n'] for x in c.execute("SELECT status,COUNT(*) n FROM quality_scenario_standardization_item WHERE batch_id=? GROUP BY status",(batch_id,))}
            total=sum(counts.values());completed=counts.get('COMPLETED',0);failed=counts.get('FAILED',0);skipped=counts.get('SKIPPED',0);pending=counts.get('PENDING',0)+counts.get('RUNNING',0)
            status='RUNNING' if pending else ('COMPLETED' if not failed else 'PARTIAL')
            c.execute("UPDATE quality_scenario_standardization_batch SET status=?,completed_count=?,failed_count=?,skipped_count=?,finished_at=CASE WHEN ?=0 THEN CURRENT_TIMESTAMP ELSE finished_at END WHERE batch_id=?",(status,completed,failed,skipped,pending,batch_id))
        return {'status':status,'total_count':total,'completed_count':completed,'failed_count':failed,'skipped_count':skipped,'pending_count':pending}

    def standardization_batch(self,batch_id):
        with self.connect() as c:
            row=c.execute("SELECT * FROM quality_scenario_standardization_batch WHERE batch_id=?",(batch_id,)).fetchone()
            if not row:return None
            result=dict(row);result['items']=[dict(x) for x in c.execute("SELECT i.*,s.name,s.scenario_code FROM quality_scenario_standardization_item i JOIN quality_scenario s ON s.scenario_id=i.scenario_id WHERE i.batch_id=? ORDER BY i.updated_at DESC",(batch_id,))]
            return result

    def standardization_batches(self):
        with self.connect() as c:return [dict(x) for x in c.execute("SELECT * FROM quality_scenario_standardization_batch ORDER BY created_at DESC LIMIT 20")]

    def mark_standardization(self,scenario_id,status,*,error='',agent='',model=''):
        if status not in {'NOT_ANALYZED','RUNNING','PENDING_CONFIRMATION','CONFIRMED','FAILED'}:raise ValueError('INVALID_QUALITY_CLASSIFICATION_STATUS')
        with self.connect() as c:
            if not c.execute("SELECT 1 FROM quality_scenario WHERE scenario_id=?",(scenario_id,)).fetchone():raise KeyError(scenario_id)
            c.execute("UPDATE quality_scenario SET quality_classification_status=?,quality_classification_error=?,quality_classification_agent=COALESCE(NULLIF(?,''),quality_classification_agent),quality_classification_model=COALESCE(NULLIF(?,''),quality_classification_model),quality_classification_updated_at=CURRENT_TIMESTAMP WHERE scenario_id=?",(status,error,agent,model,scenario_id))

    def save_standardization(self,scenario_id,result,*,agent='',model=''):
        current=self.scenario(scenario_id)
        if not current:raise KeyError(scenario_id)
        if current.get('quality_classification_status')=='CONFIRMED':raise ValueError('QUALITY_CLASSIFICATION_ALREADY_CONFIRMED')
        self.save_scenario(scenario_id,{**current,**result,'quality_classification_status':'PENDING_CONFIRMATION'},current.get('scopes') or {})
        self.mark_standardization(scenario_id,'PENDING_CONFIRMATION',agent=agent,model=model)

    def scenario(self, scenario_id):
        item=next(iter(self.scenarios()),None) if not scenario_id else next((x for x in self.scenarios() if x['scenario_id']==scenario_id),None)
        if item:
            for key in ('secondary_experience_codes','quality_in_use_codes','secondary_quality_characteristic_codes','quality_subcharacteristic_codes','secondary_typical_problem_codes','secondary_quality_concern_codes','secondary_customer_experience_statements','environment_condition_codes'):
                try:item[key]=json.loads(item.get(key) or '[]')
                except (TypeError,json.JSONDecodeError):item[key]=[]
            models=self.quality_models();labels={x['term_code']:x['label_zh'] for group in ('product_characteristics','product_subcharacteristics','quality_in_use','customer_experiences') for x in models[group]}
            item['quality_labels']={key:labels.get(key,key) for key in [item.get('primary_experience_code'),*item['secondary_experience_codes'],*item['quality_in_use_codes'],item.get('primary_quality_characteristic_code'),*item['secondary_quality_characteristic_codes'],*item['quality_subcharacteristic_codes']] if key}
            with self.connect() as c:item['evidence']=[dict(x) for x in c.execute("SELECT * FROM quality_scenario_evidence WHERE scenario_id=?",(item['scenario_id'],))]
            for evidence in item['evidence']:
                try:evidence['meta']=json.loads(evidence.get('evidence_summary') or '{}')
                except (TypeError,json.JSONDecodeError):evidence['meta']={'summary':evidence.get('evidence_summary')}
            with self.connect() as c:item['duplicates']=[dict(x) for x in c.execute("SELECT d.*,s.name,s.status FROM quality_scenario_duplicate d JOIN quality_scenario s ON s.scenario_id=d.existing_id WHERE d.candidate_id=? ORDER BY d.similarity DESC",(item['scenario_id'],))]
            with self.connect() as c:item['industry_variants']=[dict(x) for x in c.execute("SELECT * FROM quality_scenario_industry_variant WHERE scenario_id=? ORDER BY evidence_count DESC,industry",(item['scenario_id'],))]
            for variant in item['industry_variants']:
                try:variant['product_model_values']=json.loads(variant.get('product_models') or '[]')
                except (TypeError,json.JSONDecodeError):variant['product_model_values']=[]
            with self.connect() as c:item['confirmations']=[dict(x) for x in c.execute("SELECT * FROM quality_scenario_confirmation WHERE scenario_id=? ORDER BY confirmed_at DESC",(item['scenario_id'],))]
            with self.connect() as c:item['capability_gaps']=[dict(x) for x in c.execute("SELECT * FROM quality_scenario_capability_gap WHERE scenario_id=? ORDER BY capability_axis,priority,updated_at DESC",(item['scenario_id'],))]
            semantic=self.semantic_dictionary(item.get('product_code') or '')
            semantic_labels={x['term_code']:x['label_zh'] for x in semantic['items']}
            semantic_status={x['term_code']:x['status'] for x in semantic['items']}
            used=[item.get('primary_typical_problem_code'),*item['secondary_typical_problem_codes'],item.get('primary_quality_concern_code'),*item['secondary_quality_concern_codes'],*item['environment_condition_codes']]
            item['semantic_labels']={code:semantic_labels.get(code,code) for code in used if code}
            item['semantic_pending_codes']=[code for code in used if code and semantic_status.get(code)!='ACTIVE']
        return item

    def save_scenario_capability_gap(self,scenario_id,payload):
        if not self.scenario(scenario_id):raise KeyError(scenario_id)
        axis=str(payload.get('capability_axis') or '')
        valid={code for code,_ in SCENARIO_CAPABILITIES.get(axis,())};code=str(payload.get('capability_code') or '')
        if code not in valid:raise ValueError('SCENARIO_CAPABILITY_CODE_INVALID')
        gap_id=payload.get('gap_id') or f"QCG-{uuid.uuid4().hex}"
        priority=payload.get('priority') if payload.get('priority') in {'P0','P1','P2'} else 'P1'
        status=payload.get('status') if payload.get('status') in {'OPEN','IN_PROGRESS','VERIFIED','CLOSED'} else 'OPEN'
        with self.connect() as c:c.execute("""INSERT INTO quality_scenario_capability_gap(gap_id,scenario_id,capability_axis,capability_code,gap_description,source_basis,improvement_action,verification_metric,priority,status) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(gap_id) DO UPDATE SET capability_axis=excluded.capability_axis,capability_code=excluded.capability_code,gap_description=excluded.gap_description,source_basis=excluded.source_basis,improvement_action=excluded.improvement_action,verification_metric=excluded.verification_metric,priority=excluded.priority,status=excluded.status,updated_at=CURRENT_TIMESTAMP""",(gap_id,scenario_id,axis,code,str(payload.get('gap_description') or '').strip(),str(payload.get('source_basis') or '').strip(),str(payload.get('improvement_action') or '').strip(),str(payload.get('verification_metric') or '').strip(),priority,status))
        return gap_id

    def delete_scenario_capability_gap(self,scenario_id,gap_id):
        with self.connect() as c:
            result=c.execute("DELETE FROM quality_scenario_capability_gap WHERE scenario_id=? AND gap_id=?",(scenario_id,gap_id))
            if not result.rowcount:raise KeyError(gap_id)

    def create_generation(self,generation_id,product,start,end,source_count,created_by):
        with self.connect() as c:c.execute("INSERT INTO quality_scenario_generation(generation_id,product_code,start_month,end_month,source_issue_count,candidate_count,model_name,created_by,status,progress_text,unprocessed_count) VALUES(?,?,?,?,?,0,'',?,'QUEUED','等待开始',?)",(generation_id,product,start,end,source_count,created_by,source_count))

    def initialize_issue_classifications(self,generation_id,records):
        with self.connect() as c:
            for row in records:
                c.execute("INSERT OR IGNORE INTO quality_scenario_issue_classification(generation_id,knowledge_id,business_issue_id,status) VALUES(?,?,?,'PENDING')",(generation_id,row['knowledge_id'],row.get('business_issue_id')))
        self.refresh_generation_coverage(generation_id)

    def mark_issue_classification(self,generation_id,knowledge_id,status,*,scenario_id='',activity_code='',lifecycle_code='',error_message='',agent_id='',model_name='',started=False,finished=False):
        with self.connect() as c:
            c.execute("""UPDATE quality_scenario_issue_classification SET status=?,scenario_id=COALESCE(?,scenario_id),activity_code=COALESCE(?,activity_code),lifecycle_code=COALESCE(?,lifecycle_code),error_message=?,agent_id=COALESCE(?,agent_id),model_name=COALESCE(?,model_name),started_at=CASE WHEN ? THEN COALESCE(started_at,CURRENT_TIMESTAMP) ELSE started_at END,finished_at=CASE WHEN ? THEN CURRENT_TIMESTAMP ELSE finished_at END,attempt_count=attempt_count+CASE WHEN ? THEN 1 ELSE 0 END,updated_at=CURRENT_TIMESTAMP WHERE generation_id=? AND knowledge_id=?""",(status,scenario_id or None,activity_code or None,lifecycle_code or None,error_message or None,agent_id or None,model_name or None,int(started),int(finished),int(started),generation_id,knowledge_id))

    def cached_scenario_analysis(self,input_key):
        with self.connect() as c:
            row=c.execute("""SELECT a.*,s.lifecycle_code,s.activity_code FROM quality_scenario_analysis_cache a JOIN quality_scenario s ON s.scenario_id=a.scenario_id WHERE a.input_key=? AND a.status='COMPLETED'""",(input_key,)).fetchone()
            return dict(row) if row else None

    def latest_scenario_analysis(self,canonical_itr,task_type,taxonomy_version_id):
        with self.connect() as c:
            row=c.execute("""SELECT a.*,s.status scenario_status,s.quality_classification_status,s.scenario_code
                FROM quality_scenario_analysis_cache a JOIN quality_scenario s ON s.scenario_id=a.scenario_id
                WHERE a.canonical_itr=? AND a.task_type=? AND a.taxonomy_version_id=? AND a.status='COMPLETED'
                ORDER BY a.updated_at DESC,a.created_at DESC LIMIT 1""",(canonical_itr,task_type,taxonomy_version_id)).fetchone()
            return dict(row) if row else None

    def claim_scenario_analysis(self,input_key,generation_id,knowledge_id):
        with self.connect() as c:
            c.execute("DELETE FROM quality_scenario_analysis_claim WHERE claimed_at<datetime('now','-1 hour')")
            result=c.execute("INSERT OR IGNORE INTO quality_scenario_analysis_claim(input_key,generation_id,knowledge_id) VALUES(?,?,?)",(input_key,generation_id,knowledge_id))
            return bool(result.rowcount)

    def finish_scenario_analysis(self,input_key,*,canonical_itr,task_type,evidence_hash,taxonomy_version_id,scenario_id,source_knowledge_id,model_name):
        with self.connect() as c:
            c.execute("""INSERT INTO quality_scenario_analysis_cache(input_key,canonical_itr,task_type,evidence_hash,taxonomy_version_id,scenario_id,source_knowledge_id,model_name,status) VALUES(?,?,?,?,?,?,?,?, 'COMPLETED') ON CONFLICT(input_key) DO UPDATE SET scenario_id=excluded.scenario_id,source_knowledge_id=excluded.source_knowledge_id,model_name=excluded.model_name,status='COMPLETED',updated_at=CURRENT_TIMESTAMP""",(input_key,canonical_itr,task_type,evidence_hash,taxonomy_version_id,scenario_id,source_knowledge_id,model_name))
            c.execute("DELETE FROM quality_scenario_analysis_claim WHERE input_key=?",(input_key,))

    def release_scenario_analysis(self,input_key):
        with self.connect() as c:c.execute("DELETE FROM quality_scenario_analysis_claim WHERE input_key=?",(input_key,))

    def link_generation_candidate(self,generation_id,scenario_id):
        with self.connect() as c:c.execute("INSERT OR IGNORE INTO quality_scenario_generation_candidate(generation_id,scenario_id) VALUES(?,?)",(generation_id,scenario_id))

    def refresh_generation_coverage(self,generation_id):
        with self.connect() as c:
            counts={row['status']:row['n'] for row in c.execute("SELECT status,COUNT(*) n FROM quality_scenario_issue_classification WHERE generation_id=? GROUP BY status",(generation_id,))}
            total=sum(counts.values());classified=counts.get('CLASSIFIED',0);reused=counts.get('REUSED',0);updated=counts.get('UPDATED',0);review=counts.get('REVIEW_REQUIRED',0);failed=counts.get('FAILED',0);unprocessed=counts.get('PENDING',0)+counts.get('RUNNING',0)
            processed=classified+reused+updated+review+failed
            c.execute("UPDATE quality_scenario_generation SET processed_count=?,classified_count=?,reused_count=?,updated_count=?,review_required_count=?,failed_count=?,unprocessed_count=? WHERE generation_id=?",(processed,classified,reused,updated,review,failed,unprocessed,generation_id))
        return {'total':total,'processed_count':processed,'classified_count':classified,'reused_count':reused,'updated_count':updated,'review_required_count':review,'failed_count':failed,'unprocessed_count':unprocessed}

    def issue_classifications(self,generation_id):
        with self.connect() as c:return [dict(x) for x in c.execute("SELECT * FROM quality_scenario_issue_classification WHERE generation_id=? ORDER BY knowledge_id",(generation_id,))]

    def update_generation(self,generation_id,**values):
        allowed={'status','progress_text','error_message','candidate_count','model_name','processed_count','classified_count','reused_count','updated_count','review_required_count','failed_count','unprocessed_count','taxonomy_version_id'};data={k:v for k,v in values.items() if k in allowed}
        if not data:return
        assignments=','.join(f'{key}=?' for key in data)
        finished=",finished_at=CURRENT_TIMESTAMP" if data.get('status') in {'COMPLETED','PARTIAL','FAILED'} else ''
        with self.connect() as c:c.execute(f"UPDATE quality_scenario_generation SET {assignments}{finished} WHERE generation_id=?",(*data.values(),generation_id))

    def generation(self,generation_id):
        with self.connect() as c:
            row=c.execute("SELECT * FROM quality_scenario_generation WHERE generation_id=?",(generation_id,)).fetchone();return dict(row) if row else None

    def generations(self):
        with self.connect() as c:return [dict(x) for x in c.execute("SELECT g.*,(SELECT COUNT(*) FROM quality_scenario_generation_candidate x WHERE x.generation_id=g.generation_id) linked_candidate_count FROM quality_scenario_generation g ORDER BY g.created_at DESC LIMIT 30")]

    def delete_generation(self,generation_id):
        """Delete generation history without deleting scenarios produced by the job."""
        with self.connect() as c:
            row=c.execute("SELECT status FROM quality_scenario_generation WHERE generation_id=?",(generation_id,)).fetchone()
            if not row:raise KeyError(generation_id)
            if row['status'] in {'QUEUED','RUNNING'}:raise ValueError('SCENARIO_GENERATION_IS_RUNNING')
            self._delete_generation_rows(c,generation_id)

    def delete_finished_generations(self):
        """Clear all terminal generation records while leaving scenario assets intact."""
        with self.connect() as c:
            ids=[row[0] for row in c.execute("SELECT generation_id FROM quality_scenario_generation WHERE status NOT IN ('QUEUED','RUNNING')")]
            for generation_id in ids:self._delete_generation_rows(c,generation_id)
            return len(ids)

    @staticmethod
    def _delete_generation_rows(c,generation_id):
        tables={row[0] for row in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'quality_scenario_generation_candidate' in tables and 'quality_scenario_evidence' in tables:
            scenario_ids=[row[0] for row in c.execute("SELECT scenario_id FROM quality_scenario_generation_candidate WHERE generation_id=?",(generation_id,))]
            for scenario_id in scenario_ids:
                evidence=c.execute("SELECT knowledge_id,evidence_summary FROM quality_scenario_evidence WHERE scenario_id=?",(scenario_id,)).fetchall()
                for row in evidence:
                    try:summary=json.loads(row['evidence_summary'] or '{}')
                    except (TypeError,json.JSONDecodeError):continue
                    if isinstance(summary,dict) and summary.get('generation_id')==generation_id:
                        summary.pop('generation_id',None)
                        c.execute("UPDATE quality_scenario_evidence SET evidence_summary=? WHERE scenario_id=? AND knowledge_id=?",(json.dumps(summary,ensure_ascii=False),scenario_id,row['knowledge_id']))
        for table in ('quality_scenario_generation_candidate','quality_scenario_issue_classification','quality_scenario_analysis_claim','scenario_generation_source'):
            if table in tables:c.execute(f"DELETE FROM {table} WHERE generation_id=?",(generation_id,))
        c.execute("DELETE FROM quality_scenario_generation WHERE generation_id=?",(generation_id,))

    def insights(self, *, status=''):
        items=self.scenarios(status=status)
        product_codes={x.get('product_code') or 'PLC' for x in items} or {'PLC'}
        taxonomies=[self.taxonomy_active(code) for code in product_codes]
        taxonomies=[x for x in taxonomies if x] or [{'activities':[],'lifecycles':[]}]
        activity_terms={x['activity_code']:x for taxonomy in taxonomies for x in taxonomy['activities'] if x['enabled']}
        lifecycle_terms={x['lifecycle_code']:x for taxonomy in taxonomies for x in taxonomy['lifecycles'] if x['enabled']}
        activity_labels={code:x['label_zh'] for code,x in activity_terms.items()}
        lifecycle_labels={code:x['label_zh'] for code,x in lifecycle_terms.items()}
        activity_rows={};industry_rows={}
        for item in items:
            industries=item.get('scopes',{}).get('INDUSTRY') or ['未提供行业']
            evidence_count=len((self.scenario(item['scenario_id']) or {}).get('evidence',[]))
            activity=item.get('activity_code') or 'UNCLASSIFIED'
            a=activity_rows.setdefault(activity,{'activity_code':activity,'activity_label':activity_labels.get(activity,activity),'lifecycle_label':lifecycle_labels.get(item.get('lifecycle_code'),item.get('lifecycle_code') or '未分类'),'scenario_ids':set(),'issue_count':0,'industries':{}})
            a['scenario_ids'].add(item['scenario_id']);a['issue_count']+=evidence_count
            for industry in industries:
                a['industries'][industry]=a['industries'].get(industry,0)+evidence_count
                i=industry_rows.setdefault(industry,{'industry':industry,'scenario_ids':set(),'issue_count':0,'activities':{}})
                i['scenario_ids'].add(item['scenario_id']);i['issue_count']+=evidence_count;i['activities'][activity_labels.get(activity,activity)]=i['activities'].get(activity_labels.get(activity,activity),0)+evidence_count
        def finish(rows):
            result=[]
            for row in rows.values():
                row['scenario_count']=len(row.pop('scenario_ids'));result.append(row)
            return sorted(result,key=lambda x:(-x['issue_count'],str(x.get('activity_label') or x.get('industry'))))
        models=self.quality_models();labels={x['term_code']:x['label_zh'] for group in ('customer_experiences','quality_in_use','product_characteristics') for x in models[group]}
        detailed=[]
        for row in items:
            item=self.scenario(row['scenario_id']) or row;item['_issue_count']=len(item.get('evidence',[]));detailed.append(item)
        activities=[{'code':x['activity_code'],'label':x['label_zh']} for x in activity_terms.values() if any(y.get('activity_code')==x['activity_code'] for y in detailed)]
        def matrix(row_terms,column_terms,row_codes,column_codes,row_filter,column_filter):
            rows=[]
            for row_term in row_terms:
                cells={}
                matching=[x for x in detailed if row_term['code'] in row_codes(x)]
                for column in column_terms:
                    selected=[x for x in matching if column['term_code'] in column_codes(x)]
                    query=f"{row_filter}={row_term['code']}&{column_filter}={column['term_code']}"
                    cells[column['term_code']]={'scenario_count':len(selected),'issue_count':sum(x['_issue_count'] for x in selected),'drilldown_url':'/quality-scenarios?'+query}
                rows.append({'code':row_term['code'],'label':row_term['label'],'cells':cells})
            return {'columns':[{'code':x['term_code'],'label':x['label_zh']} for x in column_terms],'rows':rows}
        experiences=models['customer_experiences'];qiu=models['quality_in_use'];qualities=models['product_characteristics']
        activity_codes=lambda x:[x.get('activity_code')] if x.get('activity_code') else []
        experience_codes=lambda x:[x.get('primary_experience_code'),*(x.get('secondary_experience_codes') or [])]
        qiu_codes=lambda x:x.get('quality_in_use_codes') or []
        quality_codes=lambda x:[x.get('primary_quality_characteristic_code'),*(x.get('secondary_quality_characteristic_codes') or [])]
        semantic_items={}
        for product_code in product_codes:
            for term in self.semantic_dictionary(product_code,False)['items']:semantic_items[(term['term_type'],term['term_code'])]=term
        typical=[x for (kind,_),x in semantic_items.items() if kind=='TYPICAL_PROBLEM']
        conditions=[x for (kind,_),x in semantic_items.items() if kind=='ENVIRONMENT_CONDITION']
        typical_codes=lambda x:[x.get('primary_typical_problem_code'),*(x.get('secondary_typical_problem_codes') or [])]
        condition_codes=lambda x:x.get('environment_condition_codes') or []
        standardized=sum(1 for x in detailed if x.get('primary_experience_code') and x.get('primary_quality_characteristic_code'))
        capability_dict=self.capability_dictionary();capability_counts={}
        for item in detailed:
            for gap in item.get('capability_gaps') or []:
                key=(gap['capability_axis'],gap['capability_code']);entry=capability_counts.setdefault(key,{'capability_axis':gap['capability_axis'],'capability_code':gap['capability_code'],'scenario_ids':set(),'issue_count':0,'p0_count':0})
                entry['scenario_ids'].add(item['scenario_id']);entry['issue_count']+=item['_issue_count'];entry['p0_count']+=int(gap.get('priority')=='P0')
        capability_rows=[]
        for (axis,code),entry in capability_counts.items():
            labels={x['code']:x['label_zh'] for x in capability_dict['items'].get(axis,[])};entry['scenario_count']=len(entry.pop('scenario_ids'));entry['axis_label']=capability_dict['axis_labels'].get(axis,axis);entry['capability_label']=labels.get(code,code);capability_rows.append(entry)
        capability_rows.sort(key=lambda x:(-x['p0_count'],-x['issue_count'],-x['scenario_count']))
        return {'activity_rows':finish(activity_rows),'industry_rows':finish(industry_rows),'scenario_count':len(items),'issue_count':sum(x['_issue_count'] for x in detailed),'standardized_count':standardized,'standardized_rate':round(standardized*100/len(items),1) if items else 0,
                'activity_typical_problem_matrix':matrix(activities,typical,activity_codes,typical_codes,'activity_code','typical_problem_code'),'environment_typical_problem_matrix':matrix([{'code':x['term_code'],'label':x['label_zh']} for x in conditions],typical,condition_codes,typical_codes,'environment_code','typical_problem_code'),'activity_experience_matrix':matrix(activities,experiences,activity_codes,experience_codes,'activity_code','experience_code'),'activity_qiu_matrix':matrix(activities,qiu,activity_codes,qiu_codes,'activity_code','qiu_code'),'qiu_quality_matrix':matrix([{'code':x['term_code'],'label':x['label_zh']} for x in qiu],qualities,qiu_codes,quality_codes,'qiu_code','quality_code'),'activity_quality_matrix':matrix(activities,qualities,activity_codes,quality_codes,'activity_code','quality_code'),'capability_rows':capability_rows}

    def save_generated_candidate(self,code,item,scopes,generation_id,product,start,end,model,scenario_id=''):
        taxonomy=self.taxonomy_active(product)
        existing=self.scenario(scenario_id) if scenario_id else None
        payload={**item,'scenario_code':existing['scenario_code'] if existing else code,
                 'status':existing['status'] if existing else 'IN_REVIEW','product_code':product,
                 'taxonomy_version_id':taxonomy['version_id'] if taxonomy else '',
                 'applicable_boundary':item.get('applicable_boundary') or f'{product}；{start}—{end}'}
        for proposed in item.get('proposed_semantic_terms',[]):
            if not isinstance(proposed,dict):continue
            self.save_semantic_term({**proposed,'product_code':product,'status':'CANDIDATE','source_scenario_id':code})
        scenario_id=self.save_scenario(scenario_id,payload,scopes)
        summary=json.dumps({'generation_id':generation_id,'product':product,'period':f'{start}—{end}','model':model,'summary':item.get('evidence_summary'),'confidence':item.get('confidence'),'questions':item.get('confirmation_questions',[]),
                            **{k:item.get(k) for k in ('source_status','source_label','field_sources','field_evidence','source_warnings','source_material_id','source_material_ids','source_workbench','cs_material_id','itr_material_id','linked_knowledge_id','year','month','missing_leakage','lifecycle_assessment','lifecycle_reason','operating_conditions')}},ensure_ascii=False)
        with self.connect() as c:
            for knowledge_id in item.get('evidence_issue_ids',[]):c.execute("INSERT OR IGNORE INTO quality_scenario_evidence VALUES(?,?,?)",(scenario_id,knowledge_id,summary))
            c.execute("INSERT OR IGNORE INTO quality_scenario_generation_candidate(generation_id,scenario_id) VALUES(?,?)",(generation_id,scenario_id))
        self._detect_duplicates(scenario_id)
        return scenario_id

    def save_industry_variants(self,scenario_id,variants):
        with self.connect() as c:
            c.execute("DELETE FROM quality_scenario_industry_variant WHERE scenario_id=?",(scenario_id,))
            for row in variants:
                c.execute("""INSERT INTO quality_scenario_industry_variant(variant_id,scenario_id,industry,product_models,trigger_conditions,business_impact,recovery_method,evidence_count) VALUES(?,?,?,?,?,?,?,?)""",(f"QSV-{uuid.uuid4().hex}",scenario_id,row['industry'],json.dumps(row.get('product_models',[]),ensure_ascii=False),row.get('trigger_conditions',''),row.get('business_impact',''),row.get('recovery_method',''),int(row.get('evidence_count') or 0)))

    @staticmethod
    def _grams(value):
        text=''.join(ch.lower() for ch in str(value or '') if ch.isalnum())
        return {text[i:i+2] for i in range(max(1,len(text)-1))} if text else set()

    def _detect_duplicates(self,candidate_id):
        candidate=self.scenario(candidate_id)
        if not candidate:return
        a=self._grams(candidate['name']+' '+candidate.get('concern_points','')+' '+candidate.get('quality_attribute',''))
        with self.connect() as c:
            c.execute("DELETE FROM quality_scenario_duplicate WHERE candidate_id=?",(candidate_id,))
            rows=c.execute("SELECT * FROM quality_scenario WHERE scenario_id<>? AND status IN ('IN_REVIEW','PUBLISHED')",(candidate_id,)).fetchall()
            for row in rows:
                other=dict(row);b=self._grams(other['name']+' '+other.get('concern_points','')+' '+other.get('quality_attribute',''))
                lexical=len(a&b)/len(a|b) if a|b else 0
                same_activity=bool(candidate.get('activity_code') and candidate.get('activity_code')==other.get('activity_code'))
                score=min(1,lexical+(0.25 if same_activity else 0))
                if score>=0.55:c.execute("INSERT INTO quality_scenario_duplicate VALUES(?,?,?,?)",(candidate_id,other['scenario_id'],round(score,3),'业务活动一致且场景语义相近' if same_activity else '场景语义相近'))

    def save_scenario(self, scenario_id, payload, scopes):
        scenario_id=scenario_id or f"QSC-{uuid.uuid4().hex}"
        status=payload.get('status','DRAFT')
        if status not in {'DRAFT','IN_REVIEW','PUBLISHED','RETIRED'}:raise ValueError('INVALID_SCENARIO_STATUS')
        with self.connect() as c:
            existing=c.execute("SELECT * FROM quality_scenario WHERE scenario_id=?",(scenario_id,)).fetchone();version=(existing['version_no']+1 if existing else 1)
            if existing:payload={**dict(existing),**payload}
            product_code=payload.get('product_code') or (existing['product_code'] if existing else '') or 'PLC'
            taxonomy_version_id=payload.get('taxonomy_version_id') or (existing['taxonomy_version_id'] if existing else '')
            if not taxonomy_version_id:
                active=self.taxonomy_active(product_code);taxonomy_version_id=active['version_id'] if active else ''
            chain=c.execute("SELECT chain_text FROM scenario_activity WHERE version_id=? AND activity_code=?",(taxonomy_version_id,payload.get('activity_code',''))).fetchone()
            scenario_chain=(chain[0] if chain else '') or ''
            c.execute("""INSERT INTO quality_scenario(scenario_id,scenario_code,name,product_code,taxonomy_version_id,lifecycle_code,activity_code,scenario_chain,experience_requirement,concern_points,quality_attribute,quality_subcharacteristic,applicable_boundary,validation_direction,measurement_suggestion,failure_mode,failure_mechanism,trigger_conditions,preconditions,participating_systems,system_scale,user_type,affected_object,business_impact,recovery_method,status,version_no,created_at,updated_at)
             VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)
             ON CONFLICT(scenario_id) DO UPDATE SET name=excluded.name,product_code=excluded.product_code,taxonomy_version_id=excluded.taxonomy_version_id,lifecycle_code=excluded.lifecycle_code,activity_code=excluded.activity_code,scenario_chain=excluded.scenario_chain,experience_requirement=excluded.experience_requirement,concern_points=excluded.concern_points,quality_attribute=excluded.quality_attribute,quality_subcharacteristic=excluded.quality_subcharacteristic,applicable_boundary=excluded.applicable_boundary,validation_direction=excluded.validation_direction,measurement_suggestion=excluded.measurement_suggestion,failure_mode=excluded.failure_mode,failure_mechanism=excluded.failure_mechanism,trigger_conditions=excluded.trigger_conditions,preconditions=excluded.preconditions,participating_systems=excluded.participating_systems,system_scale=excluded.system_scale,user_type=excluded.user_type,affected_object=excluded.affected_object,business_impact=excluded.business_impact,recovery_method=excluded.recovery_method,status=excluded.status,version_no=excluded.version_no,updated_at=CURRENT_TIMESTAMP""",
             (scenario_id,payload['scenario_code'].strip(),payload['name'].strip(),product_code,taxonomy_version_id,payload.get('lifecycle_code',''),payload.get('activity_code',''),scenario_chain,payload.get('experience_requirement',''),payload.get('concern_points',''),payload.get('quality_attribute',''),payload.get('quality_subcharacteristic',''),payload.get('applicable_boundary',''),payload.get('validation_direction',''),payload.get('measurement_suggestion',''),payload.get('failure_mode',''),payload.get('failure_mechanism',''),payload.get('trigger_conditions',''),payload.get('preconditions',''),payload.get('participating_systems',''),payload.get('system_scale',''),payload.get('user_type',''),payload.get('affected_object',''),payload.get('business_impact',''),payload.get('recovery_method',''),status,version))
            def codes(name):
                value=payload.get(name,[])
                if isinstance(value,str):
                    try:value=json.loads(value) if value.strip().startswith('[') else [x.strip() for x in value.split(',') if x.strip()]
                    except json.JSONDecodeError:value=[]
                return list(dict.fromkeys(str(x) for x in value if str(x).strip()))
            experience=payload.get('primary_experience_code','');secondary_experiences=codes('secondary_experience_codes');qiu=codes('quality_in_use_codes')
            primary_quality=payload.get('primary_quality_characteristic_code','');secondary_quality=codes('secondary_quality_characteristic_codes');subqualities=codes('quality_subcharacteristic_codes')
            allowed={row[0]:(row[1],row[2]) for row in c.execute("SELECT term_code,label_zh,parent_code FROM quality_model_term WHERE enabled=1")}
            for code in [experience,*secondary_experiences,*qiu,primary_quality,*secondary_quality,*subqualities]:
                if code and code not in allowed:raise ValueError(f'QUALITY_MODEL_TERM_INVALID:{code}')
            if any(allowed[code][1] not in {primary_quality,*secondary_quality} for code in subqualities):raise ValueError('QUALITY_SUBCHARACTERISTIC_PARENT_MISMATCH')
            quality_labels=[allowed[x][0] for x in [primary_quality,*secondary_quality] if x in allowed]
            subquality_labels=[allowed[x][0] for x in subqualities if x in allowed]
            c.execute("""UPDATE quality_scenario SET customer_perception=?,primary_experience_code=?,secondary_experience_codes=?,quality_in_use_codes=?,primary_quality_characteristic_code=?,secondary_quality_characteristic_codes=?,quality_subcharacteristic_codes=?,quality_classification_status=?,quality_model_version=?,quality_attribute=CASE WHEN ?<>'' THEN ? ELSE quality_attribute END,quality_subcharacteristic=CASE WHEN ?<>'' THEN ? ELSE quality_subcharacteristic END WHERE scenario_id=?""",
                      (payload.get('customer_perception',''),experience,json.dumps(secondary_experiences,ensure_ascii=False),json.dumps(qiu,ensure_ascii=False),primary_quality,json.dumps(secondary_quality,ensure_ascii=False),json.dumps(subqualities,ensure_ascii=False),payload.get('quality_classification_status') or ('CONFIRMED' if status=='PUBLISHED' else 'PENDING_CONFIRMATION'),PRODUCT_QUALITY_MODEL,'、'.join(quality_labels),'、'.join(quality_labels),'、'.join(subquality_labels),'、'.join(subquality_labels),scenario_id))
            primary_problem=str(payload.get('primary_typical_problem_code') or '');secondary_problems=codes('secondary_typical_problem_codes')
            primary_concern=str(payload.get('primary_quality_concern_code') or '');secondary_concerns=codes('secondary_quality_concern_codes');environment_codes=codes('environment_condition_codes')
            semantic_rows={x['term_code']:dict(x) for x in c.execute("SELECT * FROM scenario_semantic_term WHERE status IN ('ACTIVE','CANDIDATE') AND (product_code='' OR product_code=?)",(product_code,))}
            expected=[('TYPICAL_PROBLEM',primary_problem),*[('TYPICAL_PROBLEM',x) for x in secondary_problems],('QUALITY_CONCERN',primary_concern),*[('QUALITY_CONCERN',x) for x in secondary_concerns],*[('ENVIRONMENT_CONDITION',x) for x in environment_codes]]
            for kind,code in expected:
                if code and (code not in semantic_rows or semantic_rows[code]['term_type']!=kind):raise ValueError(f'SCENARIO_SEMANTIC_TERM_INVALID:{code}')
            secondary_statements=payload.get('secondary_customer_experience_statements') or []
            if isinstance(secondary_statements,str):secondary_statements=[x.strip() for x in secondary_statements.splitlines() if x.strip()]
            c.execute("""UPDATE quality_scenario SET primary_typical_problem_code=?,secondary_typical_problem_codes=?,primary_quality_concern_code=?,secondary_quality_concern_codes=?,primary_customer_experience_statement=?,secondary_customer_experience_statements=?,operating_environment=?,operating_condition=?,duration_frequency=?,disturbances=?,extreme_conditions=?,environment_condition_codes=? WHERE scenario_id=?""",
                      (primary_problem,json.dumps(secondary_problems,ensure_ascii=False),primary_concern,json.dumps(secondary_concerns,ensure_ascii=False),str(payload.get('primary_customer_experience_statement') or ''),json.dumps(secondary_statements,ensure_ascii=False),str(payload.get('operating_environment') or ''),str(payload.get('operating_condition') or ''),str(payload.get('duration_frequency') or ''),str(payload.get('disturbances') or ''),str(payload.get('extreme_conditions') or ''),json.dumps(environment_codes,ensure_ascii=False),scenario_id))
            c.execute("DELETE FROM quality_scenario_scope WHERE scenario_id=?",(scenario_id,))
            for kind,items in scopes.items():
                for value in items:
                    if value.strip():c.execute("INSERT OR IGNORE INTO quality_scenario_scope VALUES(?,?,?)",(scenario_id,kind,value.strip()))
            new_classification=payload.get('quality_classification_status') or ('CONFIRMED' if status=='PUBLISHED' else 'PENDING_CONFIRMATION')
            old_classification=(existing['quality_classification_status'] if existing else '') or ''
            if new_classification=='CONFIRMED' and old_classification!='CONFIRMED':
                keys=('customer_perception','primary_experience_code','secondary_experience_codes','quality_in_use_codes','primary_quality_characteristic_code','secondary_quality_characteristic_codes','quality_subcharacteristic_codes')
                before={key:(existing[key] if existing else None) for key in keys};after={key:payload.get(key) for key in keys}
                c.execute("INSERT INTO quality_scenario_confirmation(confirmation_id,scenario_id,confirmed_by,previous_status,new_status,before_json,after_json) VALUES(?,?,?,?,?,?,?)",(f"QCF-{uuid.uuid4().hex}",scenario_id,payload.get('confirmed_by') or 'WEB_USER',old_classification,'CONFIRMED',json.dumps(before,ensure_ascii=False),json.dumps(after,ensure_ascii=False)))
        return scenario_id

    def delete_generated_scenarios(self,**filters):
        """Delete unconfirmed AI candidates in the selected UI scope."""
        selected=self.scenarios(**filters)
        selected_ids=[row['scenario_id'] for row in selected]
        if not selected_ids:return {'matched':0,'deleted':0,'protected':0}
        marks=','.join('?' for _ in selected_ids)
        with self.connect() as c:
            generated={row[0] for row in c.execute(f"""SELECT scenario_id FROM quality_scenario
                WHERE scenario_id IN ({marks}) AND (scenario_code LIKE 'AI-%' OR EXISTS(
                    SELECT 1 FROM quality_scenario_generation_candidate g WHERE g.scenario_id=quality_scenario.scenario_id))""",selected_ids)}
            eligible=[]
            for row in selected:
                if row['scenario_id'] in generated and row.get('status') in {'DRAFT','IN_REVIEW'} and row.get('quality_classification_status')!='CONFIRMED':
                    eligible.append(row['scenario_id'])
            for scenario_id in eligible:self._delete_scenario_rows(c,scenario_id)
        return {'matched':len(selected_ids),'deleted':len(eligible),'protected':len(selected_ids)-len(eligible)}

    def delete_scenario(self, scenario_id):
        with self.connect() as c:
            if not c.execute("SELECT 1 FROM quality_scenario WHERE scenario_id=?",(scenario_id,)).fetchone():raise KeyError(scenario_id)
            self._delete_scenario_rows(c,scenario_id)

    @staticmethod
    def _delete_scenario_rows(c,scenario_id):
        tables={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'scenario_asset_member' in tables:
            c.execute("DELETE FROM scenario_asset_member WHERE member_id=? OR asset_id=?",(scenario_id,scenario_id))
        for table in ('scenario_asset_context','scenario_asset_metric','quality_scenario_evidence','quality_scenario_industry_variant','quality_scenario_scope','quality_scenario_analysis_cache','quality_scenario_confirmation','quality_scenario_capability_gap','quality_scenario_standardization_item'):
            if table in tables:c.execute(f"DELETE FROM {table} WHERE scenario_id=?",(scenario_id,))
        generation_ids=[row[0] for row in c.execute("SELECT generation_id FROM quality_scenario_generation_candidate WHERE scenario_id=?",(scenario_id,))]
        c.execute("DELETE FROM quality_scenario_duplicate WHERE candidate_id=? OR existing_id=?",(scenario_id,scenario_id))
        c.execute("DELETE FROM quality_scenario_generation_candidate WHERE scenario_id=?",(scenario_id,))
        c.execute("UPDATE quality_scenario_issue_classification SET scenario_id=NULL,status='REVIEW_REQUIRED',error_message='关联场景已删除',updated_at=CURRENT_TIMESTAMP WHERE scenario_id=?",(scenario_id,))
        if 'scenario_semantic_term' in tables:c.execute("UPDATE scenario_semantic_term SET source_scenario_id=NULL WHERE source_scenario_id=?",(scenario_id,))
        c.execute("DELETE FROM quality_scenario WHERE scenario_id=?",(scenario_id,))
        for generation_id in generation_ids:
            c.execute("UPDATE quality_scenario_generation SET candidate_count=(SELECT COUNT(*) FROM quality_scenario_generation_candidate WHERE generation_id=?) WHERE generation_id=?",(generation_id,generation_id))
