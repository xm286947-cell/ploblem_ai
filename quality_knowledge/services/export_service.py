from __future__ import annotations
import csv
import json
from pathlib import Path
from typing import Any


class QualityKnowledgeExportService:
    """Business-readable CSV/XLSX export for M3.

    SQLite remains persistence only; callers depend on this service/repository contract.
    """

    ISSUE_COLUMNS = [
        "knowledge_id","business_type","issue_id","title","description","impact","severity","issue_type","is_defect",
        "industry","customer","month","product","platform","department","business_group","module","feature_l1","feature_l2",
        "occurrence_l1","occurrence_l2","occurrence_l3","occurrence_l4","is_escape","escape_type","escape_l1","escape_l2","escape_l3","escape_l4",
        "recurrence_risk_level",
    ]
    GAP_COLUMNS = [
        "knowledge_id","business_type","issue_id","product","platform","gap_dimension","gap_category","gap_description",
        "recommended_control","scope","confidence","analysis_run_id","model_version","prompt_version","created_at",
    ]

    def __init__(self, repository, query_service=None):
        self.repository = repository
        if query_service is None:
            from .query_service import IssueQueryService
            query_service = IssueQueryService(repository)
        self.query_service = query_service

    @staticmethod
    def _cell(v: Any):
        if isinstance(v, (dict, list, tuple)):
            return json.dumps(v, ensure_ascii=False, default=str)
        return "" if v is None else v

    def export_csv(self, output_path: str|Path, *, filters=None, dataset="issues"):
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        filters = filters or {}
        if dataset == "issues":
            rows = self.repository.query(filters, limit=1_000_000)
            columns = self.ISSUE_COLUMNS
        elif dataset == "capability_gaps":
            rows = self.repository.query_capability_gaps(filters, limit=1_000_000)
            columns = self.GAP_COLUMNS
        else:
            raise ValueError("dataset must be 'issues' or 'capability_gaps'")
        with output_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({k:self._cell(row.get(k)) for k in columns})
        return {"format":"CSV","dataset":dataset,"path":str(output_path),"rows":len(rows)}

    def export_xlsx(self, output_path: str|Path, *, filters=None, include_statistics=True):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError("XLSX export requires openpyxl (already declared in requirements.txt)") from exc

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        filters = filters or {}
        issues = self.repository.query(filters, limit=1_000_000)
        gap_filters = {k:v for k,v in filters.items() if k in {"business_type","product","platform","gap_dimension","gap_category"}}
        gaps = self.repository.query_capability_gaps(gap_filters, limit=1_000_000)
        runs = self.repository.query_analysis_runs({k:v for k,v in filters.items() if k in {"business_type","knowledge_id","status","analysis_type"}}, limit=1_000_000)

        wb = Workbook()
        ws = wb.active
        ws.title = "Issue_Knowledge"

        def write_sheet(sheet, rows, columns):
            sheet.append(columns)
            for row in rows:
                sheet.append([self._cell(row.get(k)) for k in columns])
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(color="FFFFFF", bold=True)
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            # bounded readable widths, avoiding extreme autofit on descriptions/json
            for i, col in enumerate(columns, 1):
                max_len = len(str(col))
                for row in sheet.iter_rows(min_row=2, min_col=i, max_col=i, max_row=min(sheet.max_row, 200)):
                    max_len = max(max_len, len(str(row[0].value or "")))
                sheet.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 40)
            return sheet

        write_sheet(ws, issues, self.ISSUE_COLUMNS)
        ws2 = wb.create_sheet("Capability_Gaps")
        write_sheet(ws2, gaps, self.GAP_COLUMNS)

        run_cols = ["analysis_run_id","knowledge_id","business_type","issue_id","analysis_type","model_provider","model_name","prompt_name","prompt_version","schema_version","engine_version","started_at","completed_at","status","input_hash","error_message"]
        ws3 = wb.create_sheet("Analysis_Runs")
        write_sheet(ws3, runs, run_cols)

        if include_statistics:
            stats = self.repository.statistics(business_type=filters.get("business_type"), limit=50)
            ws4 = wb.create_sheet("Statistics")
            stat_rows=[]
            for section, rows in stats.items():
                if not isinstance(rows, list):
                    continue
                for r in rows:
                    category = r.get("category") or r.get("value") or r.get("gap_category") or r.get("risk_level") or r.get("business_type") or ""
                    count = r.get("count", 0)
                    extra = {k:v for k,v in r.items() if k not in {"category","value","gap_category","risk_level","business_type","count"}}
                    stat_rows.append({"Metric":section,"Category":category,"Count":count,"Extra":self._cell(extra) if extra else ""})
            write_sheet(ws4, stat_rows, ["Metric","Category","Count","Extra"])

        wb.save(output_path)
        return {"format":"XLSX","path":str(output_path),"issues":len(issues),"capability_gaps":len(gaps),"analysis_runs":len(runs)}
