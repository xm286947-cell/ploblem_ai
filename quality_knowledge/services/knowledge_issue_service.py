from __future__ import annotations
import hashlib,json,uuid
from pathlib import Path
from openpyxl import load_workbook
from quality_knowledge.adapters import ADAPTERS
from quality_knowledge.adapters.base import clean
from quality_knowledge.config_loader import normalize_header, import_config
from quality_knowledge.mapping import MappingConfigurationService
from quality_knowledge.mapping.repository import MappingConfigurationRepository
from quality_knowledge.mapping.runtime import MappingNotInitializedError, runtime_detection_score, runtime_mapping_diagnostics
from quality_knowledge.excel_utils import repair_read_only_dimensions
from quality_knowledge.product_config import ProductConfigRepository

class NoRecordsFoundError(ValueError): pass

class KnowledgeIssueService:
    def __init__(self,repository, mapping_service=None):
        self.repository=repository
        self.mapping_service=mapping_service or MappingConfigurationService(MappingConfigurationRepository(repository.db_path))
        self.product_config=ProductConfigRepository(repository.db_path)
    def _effective(self,business_type):
        cfg=self.mapping_service.get_effective_config(business_type)
        if not cfg: raise MappingNotInitializedError(f"MAPPING_NOT_INITIALIZED: {business_type}; run knowledge-mapping-migrate --business {business_type} --apply")
        return cfg
    def _header_score(self,values,adapter_cls):
        cfg=self._effective(adapter_cls.business_type)
        return runtime_detection_score(cfg,values)
    def _detect_header(self,ws,business_type=None,max_scan=None):
        cfg=import_config()
        if max_scan is None:
            max_scan=int((cfg.get('header_scan') or {}).get('max_rows',50))
        min_score=int((cfg.get('header_scan') or {}).get('min_detection_score',20))
        best=None
        candidates=[ADAPTERS[business_type]] if business_type else [cls for cls in ADAPTERS.values() if self.mapping_service.get_effective_config(cls.business_type)]
        rows=list(ws.iter_rows(min_row=1,max_row=min(ws.max_row or max_scan,max_scan),values_only=True))
        for idx,row in enumerate(rows,1):
            for cls in candidates:
                score=self._header_score(row,cls)
                if best is None or score>best[0]:best=(score,idx,cls)
        if best and best[0]>=min_score:
            return {'header_row':best[1],'business_type':best[2].business_type,'score':best[0],'recognition_mode':'MAPPING_SCORE'}
        if business_type:
            structural=[]
            for idx,row in enumerate(rows,1):
                cells=[clean(x) for x in row if clean(x)]
                if not cells: continue
                text_like=sum(not str(x).replace('.','',1).isdigit() for x in cells)
                # Prefer the widest text-like row; earliest row wins a tie.
                structural.append(((len(cells),text_like,-idx),idx))
            if structural:
                _,idx=max(structural)
                return {'header_row':idx,'business_type':business_type,'score':best[0] if best else 0,'recognition_mode':'EXPLICIT_STRUCTURAL'}
        return None
    def import_file(self,path,business_type=None,*,sheet=None,batch_id=None,issue_domain='AUTO',issue_domain_source='IMPORT_PRESET'):
        path=Path(path); batch_id=batch_id or f'IMP-{uuid.uuid4().hex[:12]}'; bt=business_type.upper() if business_type else None
        if bt and bt not in ADAPTERS:raise ValueError(f'Unsupported business_type: {bt}')
        wb=load_workbook(path,read_only=True,data_only=True)
        wb_raw=load_workbook(path,read_only=True,data_only=False)
        sheet_names=[sheet] if sheet else wb.sheetnames
        diagnostics={'source_file':str(path),'detected_sheets':wb.sheetnames,'processed_sheets':[]}
        self.repository.begin_import(batch_id,path.name,bt or 'AUTO',diagnostics)
        stats={'batch_id':batch_id,'business_type':bt or 'AUTO','total':0,'new':0,'updated':0,'skipped':0,'failed':0,'errors':[]}
        try:
            for sn in sheet_names:
                ws=wb[sn]; raw_ws=wb_raw[sn]; repair_read_only_dimensions(ws); repair_read_only_dimensions(raw_ws); det=self._detect_header(raw_ws,bt); diagnostics['processed_sheets'].append({'sheet':sn,**(det or {'error':'HEADER_NOT_DETECTED'})})
                if not det:continue
                actual_bt=det['business_type']; effective=self._effective(actual_bt); adapter=ADAPTERS[actual_bt](effective); hr=det['header_row']; headers=[clean(x) for x in next(raw_ws.iter_rows(min_row=hr,max_row=hr,values_only=True))]; diagnostics['processed_sheets'][-1]['mapping_coverage']=runtime_mapping_diagnostics(effective,headers); diagnostics['processed_sheets'][-1]['mapping_config_id']=effective['config_id']; diagnostics['processed_sheets'][-1]['mapping_config_version']=effective['version']
                cached_rows=ws.iter_rows(min_row=hr+1,values_only=True); raw_rows=raw_ws.iter_rows(min_row=hr+1,values_only=True)
                for row_no,(cached,raw_values) in enumerate(zip(cached_rows,raw_rows),start=hr+1):
                    values=tuple(c if clean(c) else r for c,r in zip(cached,raw_values))
                    if not any(clean(v) for v in values):continue
                    raw={headers[i]:values[i] for i in range(min(len(headers),len(values))) if headers[i]}; stats['total']+=1
                    try:
                        q=adapter.adapt(raw,source_file=str(path),source_sheet=sn,source_row=row_no,batch_id=batch_id)
                        row_domain = next((clean(raw.get(normalize_header(k))) for k in ('问题领域','问题属性','issue_domain','issue domain') if clean(raw.get(normalize_header(k)))), '')
                        if row_domain:
                            q.issue_fact.issue_domain=str(row_domain).upper(); q.issue_fact.issue_domain_source='EXCEL_ROW'
                        elif issue_domain and str(issue_domain).upper() != 'AUTO':
                            q.issue_fact.issue_domain=str(issue_domain).upper(); q.issue_fact.issue_domain_source=issue_domain_source
                        elif str(q.issue_fact.issue_domain or 'AUTO').upper() == 'AUTO':
                            configured=self.product_config.get(q.issue_fact.product)
                            if configured and configured.get('default_issue_domain') not in {'', 'AUTO', None}:
                                q.issue_fact.issue_domain=configured['default_issue_domain']; q.issue_fact.issue_domain_source='PRODUCT_DEFAULT'
                            else:
                                q.issue_fact.issue_domain='AUTO'; q.issue_fact.issue_domain_source='AI'
                        # Frozen design requires a real business key; synthetic ROW ids are not accepted.
                        if q.identity.issue_id.startswith('ROW-'):raise ValueError('BUSINESS_ISSUE_ID_MISSING')
                        normalized={'fact':q.issue_fact.model_dump(),'context':q.product_context.model_dump(),'occurrence':q.occurrence.model_dump(),'escape':q.escape.model_dump(),'solution':q.solution.model_dump(),'verification':q.verification.model_dump(),'extension':q.product_extension}
                        nh=hashlib.sha256(json.dumps(normalized,ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()
                        result=self.repository.upsert_source_record(q,nh,mapping_config_id=effective['config_id'],mapping_config_version=effective['version']); stats[result['action'].lower()]+=1
                    except Exception as exc:
                        stats['failed']+=1; err={'sheet':sn,'row':row_no,'error':str(exc)};stats['errors'].append(err);self.repository.add_import_error(batch_id,source_file=path.name,source_sheet=sn,source_row=row_no,business_type=actual_bt,error_code='ROW_IMPORT_FAILED',error_message=str(exc),raw=raw)
            if stats['total']==0:
                diagnostics['error_code']='NO_RECORDS_FOUND'; self.repository.finish_import(batch_id,stats,'FAILED',diagnostics)
                raise NoRecordsFoundError(f"NO_RECORDS_FOUND: sheets={wb.sheetnames}; diagnostics={diagnostics['processed_sheets']}")
            status='COMPLETED' if stats['failed']==0 else 'PARTIAL';self.repository.finish_import(batch_id,stats,status,diagnostics);stats['status']=status;stats['diagnostics']=diagnostics;return stats
        except NoRecordsFoundError:raise
        except Exception:
            self.repository.finish_import(batch_id,stats,'FAILED',diagnostics);raise
    def get_import_batch(self,batch_id):return self.repository.get_import_batch(batch_id)
    def list_import_batches(self,limit=50):return self.repository.list_import_batches(limit)
    def get_issue(self,knowledge_id):return self.repository.get_current_issue(knowledge_id)
    def get_issue_history(self,knowledge_id):return self.repository.get_issue_history(knowledge_id)
    def get_issue_detail(self,knowledge_id):return self.repository.get_issue_detail(knowledge_id)
    def query_issues(self,filters=None,limit=100,offset=0):return self.repository.query_current_issues(filters,limit,offset)
    def count_issues(self,filters=None):return self.repository.count_current_issues(filters)

    def run_issue_analysis(self,knowledge_id,root,client=None,only_missing=None,force=False,analysis_profile=None):
        from .v1_analysis_service import KnowledgeIssueAnalysisService
        return KnowledgeIssueAnalysisService(self.repository,root,client).run_issue_analysis(knowledge_id,only_missing=only_missing,force=force,analysis_profile=analysis_profile)
    def run_batch_analysis(self,knowledge_ids,root,client=None,only_missing=None,force=False,analysis_profile=None):
        from .v1_analysis_service import KnowledgeIssueAnalysisService
        return KnowledgeIssueAnalysisService(self.repository,root,client).run_batch_analysis(knowledge_ids,only_missing=only_missing,force=force,analysis_profile=analysis_profile)
    def get_analysis_status(self,run_id):return self.repository.get_analysis_run(run_id)
    def get_latest_analysis(self,knowledge_id,analysis_type):return self.repository.get_latest_analysis(knowledge_id,analysis_type)
    def get_analysis_history(self,knowledge_id):return self.repository.get_analysis_history(knowledge_id)
    def query_capability_gaps(self,knowledge_id=None,*,filters=None,limit=1000):
        if filters:
            return self.repository.query_current_capability_gaps(filters,limit)
        return self.repository.list_capability_gaps(knowledge_id)

    def get_statistics(self,business_type=None,limit=20):
        data=self.repository.statistics(business_type,limit)
        data['common_capability_analysis']=self.repository.aggregate_common_capability_gaps(business_type=business_type,min_issues=2,limit=limit)
        return data
    def get_common_capability_gaps(self,*,business_type=None,dimension=None,min_issues=2,limit=50):return self.repository.aggregate_common_capability_gaps(business_type=business_type,dimension=dimension,min_issues=min_issues,limit=limit)
    def get_workspace_metrics(self,business_type=None):return self.repository.workspace_metrics(business_type)
    def export_issues(self,path,*,format='xlsx',filters=None,dataset='issues'):
        from .v1_export_service import KnowledgeIssueExportService
        svc=KnowledgeIssueExportService(self.repository)
        return svc.export_csv(path,filters,dataset) if format.lower()=='csv' else svc.export_xlsx(path,filters)
