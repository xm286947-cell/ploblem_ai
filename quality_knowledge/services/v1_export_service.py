from __future__ import annotations
import csv,json
from pathlib import Path
from openpyxl import Workbook

class KnowledgeIssueExportService:
    def __init__(self,repository):
        self.repository=repository
        from quality_knowledge.human_analysis import HumanAnalysisRepository,HumanAnalysisService
        self.human_service=HumanAnalysisService(HumanAnalysisRepository(repository.db_path))

    @staticmethod
    def _flat(rows):
        return [{k:(json.dumps(v,ensure_ascii=False,default=str) if isinstance(v,(dict,list)) else v) for k,v in r.items()} for r in rows]

    def _ai_rows(self, filters=None):
        issues=self.repository.export_current_issues(filters or {})
        out=[]
        for x in issues:
            kid=x['knowledge_id']
            row={
                'knowledge_id':kid,'business_type':x.get('business_type'),'business_issue_id':x.get('business_issue_id'),
                'title':x.get('title'),'product':x.get('product'),'platform':x.get('platform'),'severity':x.get('severity'),
                'version_no':x.get('version_no')
            }
            for typ in ('occurrence','escape','recurrence'):
                a=self.repository.get_latest_analysis(kid,typ)
                row[typ]=a.get('result') if a else None
                row[typ+'_analysis_run_id']=a.get('analysis_run_id') if a else None
            out.append(row)
        return out

    def export_csv(self, output_path: str|Path, filters=None, dataset='issues'):
        p=Path(output_path);p.parent.mkdir(parents=True,exist_ok=True)
        if dataset=='issues': rows=self.repository.export_current_issues(filters or {})
        elif dataset=='capability_gaps': rows=self.repository.query_current_capability_gaps(filters or {})
        elif dataset=='ai_analysis': rows=self._ai_rows(filters or {})
        elif dataset=='human_analysis':
            issues=self.repository.export_current_issues(filters or {}); rows=self.human_service.export_rows([x['knowledge_id'] for x in issues])
        else: raise ValueError('unsupported dataset: '+str(dataset))
        flat=self._flat(rows)
        fields=sorted({k for r in flat for k in r}) if flat else ['knowledge_id']
        with p.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(flat)
        return {'format':'CSV','dataset':dataset,'path':str(p),'rows':len(rows)}

    def export_xlsx(self, output_path: str|Path, filters=None):
        p=Path(output_path);p.parent.mkdir(parents=True,exist_ok=True);filters=filters or {}
        issues=self.repository.export_current_issues(filters)
        ai=self._ai_rows(filters)
        gaps=self.repository.query_current_capability_gaps({k:v for k,v in filters.items() if k in {'business_type','dimension','category'}})
        stats=self.repository.statistics(filters.get('business_type'),50)
        wb=Workbook();ws=wb.active;ws.title='Issue_Knowledge'
        def sheet(ws,rows):
            cooked=self._flat(rows)
            fields=sorted({k for r in cooked for k in r}) if cooked else ['knowledge_id'];ws.append(fields)
            for r in cooked:ws.append([r.get(k,'') for k in fields])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        sheet(ws,issues);sheet(wb.create_sheet('AI_Analysis'),ai);sheet(wb.create_sheet('Human_Analysis'),self.human_service.export_rows([x['knowledge_id'] for x in issues]));sheet(wb.create_sheet('Capability_Gaps'),gaps)
        sr=[]
        for metric,rows in stats.items():
            for r in rows:sr.append({'metric':metric,**r})
        sheet(wb.create_sheet('Statistics'),sr);wb.save(p)
        return {'format':'XLSX','path':str(p),'issues':len(issues),'ai_analysis':len(ai),'capability_gaps':len(gaps),'statistics_rows':len(sr)}
