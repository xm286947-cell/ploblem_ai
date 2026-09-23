from __future__ import annotations
import uuid
from pathlib import Path
from openpyxl import load_workbook
from quality_knowledge.adapters import ADAPTERS
from quality_knowledge.mapping.repository import MappingConfigurationRepository
from quality_knowledge.mapping.runtime import MappingNotInitializedError

class IssueImportService:
    def __init__(self, repository): self.repository=repository
    def import_excel(self, path, business_type, *, sheet=None, batch_id=None):
        business_type=business_type.upper()
        if business_type not in ADAPTERS: raise ValueError(f'Unsupported business_type: {business_type}')
        path=Path(path); batch_id=batch_id or f'IMP-{uuid.uuid4().hex[:12]}'; mr=MappingConfigurationRepository(self.repository.db_path); effective=mr.get_effective_config(business_type)
        if not effective: raise MappingNotInitializedError(f'MAPPING_NOT_INITIALIZED: {business_type}; run knowledge-mapping-migrate --business {business_type} --apply')
        adapter=ADAPTERS[business_type](effective)
        self.repository.begin_batch(batch_id,business_type,path.name)
        stats={'batch_id':batch_id,'business_type':business_type,'total':0,'imported':0,'skipped':0,'failed':0,'errors':[]}
        try:
            wb=load_workbook(path,read_only=True,data_only=True)
            sheets=[sheet] if sheet else wb.sheetnames
            for sn in sheets:
                ws=wb[sn]; it=ws.iter_rows(values_only=True)
                try: headers=[str(x).strip() if x is not None else '' for x in next(it)]
                except StopIteration: continue
                for row_no,values in enumerate(it,start=2):
                    if not any(v is not None and str(v).strip() for v in values): continue
                    stats['total']+=1; raw={headers[i]:values[i] for i in range(min(len(headers),len(values))) if headers[i]}
                    try:
                        q=adapter.adapt(raw,source_file=str(path),source_sheet=sn,source_row=row_no,batch_id=batch_id)
                        if self.repository.exists_source(business_type,q.identity.issue_id,q.source.source_hash): stats['skipped']+=1; continue
                        self.repository.save(q); stats['imported']+=1
                    except Exception as exc:
                        stats['failed']+=1; stats['errors'].append({'sheet':sn,'row':row_no,'error':str(exc)})
            self.repository.finish_batch(batch_id,stats,'COMPLETED' if stats['failed']==0 else 'PARTIAL')
            return stats
        except Exception:
            self.repository.finish_batch(batch_id,stats,'FAILED'); raise
