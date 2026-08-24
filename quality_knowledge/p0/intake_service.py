"""Clean P0 Excel Preview-before-Commit intake service."""

from __future__ import annotations

import hashlib
import json
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from quality_knowledge.p0.repository import P0RepositoryError
from quality_knowledge.config_loader import normalize_header
from quality_knowledge.standard_fields.repository import StandardFieldCatalogError, StandardFieldRepository
from quality_knowledge.standard_fields.service import StandardFieldService


class P0IntakeError(RuntimeError):
    """Stable intake error for the clean P0 Preview-before-Commit workflow."""


class P0IntakeService:
    """Preview source Excel data and commit only the immutable approved snapshot."""

    REQUIRED_PATH = "ISSUE_FACT.business_issue_id"

    def __init__(self, repository: Any):
        self.repository = repository
        project_root = Path(__file__).resolve().parents[2]
        self._field_repository = StandardFieldRepository(repository.db_path)
        self._field_service = StandardFieldService(
            self._field_repository,
            project_root / "quality_knowledge" / "config" / "plc_fields.yaml",
            technical_field_exclusions={
                "PRODUCT_EXTENSION.source_id": {"handling": "RAW_ONLY"},
            },
        )

    def preview(
        self,
        file_path: str | Path,
        business_type: str = "PLC",
        sheet: str | None = None,
    ) -> dict[str, Any]:
        source = Path(file_path)
        if not source.exists() or not source.is_file():
            raise P0IntakeError("P0_INTAKE_FILE_NOT_FOUND")
        file_sha256 = self._sha256_file(source)
        mapping, catalog, product = self._preview_contract(business_type)
        sheets = self._inspect_workbook(source, mapping["config_id"])
        selected = self._select_sheet(sheets, sheet)
        snapshot = self._build_snapshot(
            source,
            file_sha256,
            business_type,
            mapping,
            catalog,
            product,
            selected,
            sheets,
        )
        session_id = f"INTAKE-{uuid.uuid4().hex}"
        token = f"PREVIEW-{uuid.uuid4().hex}"
        with self.repository.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            connection.execute(
                """INSERT INTO intake_session(
                        intake_session_id, product_id, file_name, file_sha256, mapping_config_id,
                        catalog_version_id, status
                    ) VALUES (?, ?, ?, ?, ?, ?, 'PREVIEW_READY')""",
                (
                    session_id, product["product_id"], source.name, file_sha256,
                    mapping["config_id"], catalog["catalog_version_id"],
                ),
            )
            connection.execute(
                """INSERT INTO intake_preview_snapshot(
                        preview_token, intake_session_id, sheet_name, header_row, data_start_row,
                        mapping_content_hash, catalog_content_hash, snapshot_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    token, session_id, selected["sheet_name"], selected["header_row"],
                    selected["data_start_row"], mapping["content_hash"], catalog["content_hash"],
                    self._dump(snapshot),
                ),
            )
            connection.commit()
        return {"preview_token": token, "intake_session_id": session_id, **snapshot}

    def confirm(self, preview_token: str) -> dict[str, Any]:
        snapshot_record = self._load_snapshot(preview_token)
        if snapshot_record is None:
            raise P0IntakeError("P0_PREVIEW_NOT_FOUND")
        snapshot = snapshot_record["snapshot"]
        if snapshot.get("mapping_status") != "ACTIVE":
            raise P0IntakeError("P0_DRAFT_PREVIEW_REQUIRES_ACTIVATION")
        existing = self._existing_batch(preview_token)
        if existing is not None:
            if existing["status"] in {"COMPLETED", "PARTIAL_FAILED"}:
                return {
                    "outcome": "ALREADY_APPLIED", "import_batch_id": existing["import_batch_id"],
                    **self._batch_summary(existing, snapshot),
                }
            raise P0IntakeError("P0_IMPORT_IN_PROGRESS")
        mapping, catalog, product = self._active_contract(snapshot["business_type"])
        if (
            mapping["content_hash"] != snapshot_record["mapping_content_hash"]
            or catalog["content_hash"] != snapshot_record["catalog_content_hash"]
        ):
            raise P0IntakeError("P0_MAPPING_CHANGED_AFTER_PREVIEW")
        if snapshot["required_missing"] or snapshot["has_conflict"]:
            raise P0IntakeError("P0_INTAKE_CONFIRM_BLOCKED")
        if product["product_id"] != snapshot["product_id"]:
            raise P0IntakeError("P0_PRODUCT_CHANGED_AFTER_PREVIEW")

        batch_id = f"IMPORT-{uuid.uuid4().hex}"
        with self.repository.connect() as connection:
            connection.execute(
                "INSERT INTO import_batch(import_batch_id, preview_token, status) VALUES (?, ?, 'RUNNING')",
                (batch_id, preview_token),
            )
            connection.commit()

        success = 0
        already_applied = 0
        failed = 0
        seen_business_issue_ids: set[str] = set()
        duplicate = 0
        for row in snapshot["rows"]:
            try:
                business_issue_id = self._required_value(row["normalized_candidate"])
                if not business_issue_id:
                    raise P0IntakeError("ISSUE_FACT_BUSINESS_ISSUE_ID_REQUIRED")
                if business_issue_id in seen_business_issue_ids:
                    duplicate += 1
                else:
                    seen_business_issue_ids.add(business_issue_id)
                saved = self.repository.save_issue(
                    knowledge_id=self._knowledge_id(snapshot["business_type"], business_issue_id),
                    business_issue_id=business_issue_id,
                    raw_json=row["raw_json"],
                    normalized_snapshot=row["normalized_candidate"],
                    mapping_config_id=mapping["config_id"],
                    mapping_config_version=mapping["version"],
                    standard_catalog_version_id=catalog["catalog_version_id"],
                    source_file_sha256=snapshot["file_sha256"],
                    sheet_name=snapshot["sheet_name"],
                    row_number=row["row_number"],
                    import_batch_id=batch_id,
                    product_id=snapshot["product_id"],
                )
                if saved["outcome"] == "ALREADY_APPLIED":
                    already_applied += 1
                else:
                    success += 1
            except (P0RepositoryError, P0IntakeError, ValueError) as error:
                failed += 1
                self._record_error(batch_id, row["row_number"], str(error), "无法提交该行", {"row": row})
        status = "COMPLETED" if failed == 0 else "PARTIAL_FAILED"
        with self.repository.connect() as connection:
            connection.execute(
                "UPDATE import_batch SET status = ?, completed_at = CURRENT_TIMESTAMP WHERE import_batch_id = ?",
                (status, batch_id),
            )
            connection.execute(
                "UPDATE intake_session SET status = 'COMMITTED' WHERE intake_session_id = ?",
                (snapshot_record["intake_session_id"],),
            )
            connection.commit()
        return {
            "outcome": "APPLIED", "import_batch_id": batch_id, "status": status,
            "total": len(snapshot["rows"]), "success": success,
            "already_applied": already_applied, "duplicate": duplicate,
            "unique_business_issues": len(seen_business_issue_ids), "failed": failed,
        }

    def create_mapping_draft(
        self,
        preview_token: str,
        decisions: list[dict[str, Any]],
        actor: str,
    ) -> dict[str, Any]:
        """Apply explicit Preview differences to an editable mapping draft.

        The source file remains a Preview.  A selected source header is removed
        from every other target in the draft before being assigned, preventing
        the conflict that originally made difference handling unusable.
        """

        record = self._load_snapshot(preview_token)
        if record is None:
            raise P0IntakeError("P0_PREVIEW_NOT_FOUND")
        if not actor.strip():
            raise P0IntakeError("MAPPING_ACTOR_REQUIRED")
        snapshot = record["snapshot"]
        known_headers = {
            normalize_header(item["source_header"]): item["source_header"]
            for item in snapshot.get("field_decisions", [])
        }
        try:
            draft = self._field_repository.persist_product_starter_draft(
                snapshot["business_type"], actor.strip()
            )
        except StandardFieldCatalogError as error:
            raise P0IntakeError(str(error)) from error

        items = draft.get("items", [])
        by_path = {f"{item['target_domain']}.{item['target_field']}": item for item in items}
        normalized_aliases = {
            f"{item['target_domain']}.{item['target_field']}": {
                normalize_header(alias): alias for alias in item.get("aliases", [])
            }
            for item in items
        }
        applied: list[dict[str, str]] = []
        for decision in decisions:
            source_header = str(decision.get("source_header") or "").strip()
            normalized = normalize_header(source_header)
            action = str(decision.get("action") or "").strip().upper()
            if normalized not in known_headers:
                raise P0IntakeError("P0_MAPPING_SOURCE_HEADER_NOT_IN_PREVIEW")
            if action not in {"MAP_EXISTING", "RAW_ONLY"}:
                raise P0IntakeError("P0_MAPPING_DECISION_INVALID")
            for aliases in normalized_aliases.values():
                aliases.pop(normalized, None)
            target_path = ""
            if action == "MAP_EXISTING":
                target_path = str(decision.get("target_path") or "").strip()
                if target_path not in by_path:
                    raise P0IntakeError("P0_MAPPING_TARGET_NOT_IN_CATALOG")
                normalized_aliases[target_path][normalized] = known_headers[normalized]
            applied.append({
                "source_header": known_headers[normalized],
                "action": action,
                "target_path": target_path,
            })

        bindings = []
        for path, item in by_path.items():
            bindings.append({
                "target_domain": item["target_domain"],
                "target_field": item["target_field"],
                "source_headers": list(normalized_aliases[path].values()),
                "enabled": bool(item["enabled"]),
            })
        try:
            updated = self._field_repository.update_mapping_draft(
                draft["config_id"], bindings=bindings, updated_by=actor.strip()
            )
        except StandardFieldCatalogError as error:
            raise P0IntakeError(str(error)) from error
        return {
            "outcome": "DRAFT_UPDATED",
            "preview_token": preview_token,
            "applied_decisions": applied,
            **updated,
        }

    def _active_contract(
        self, business_type: str
    ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
        with self.repository.connect() as connection:
            mapping = connection.execute(
                "SELECT * FROM mapping_config WHERE business_type = ? AND status = 'ACTIVE'",
                (business_type,),
            ).fetchone()
            catalog = connection.execute(
                "SELECT * FROM standard_field_catalog_version WHERE status = 'ACTIVE'"
            ).fetchone()
            product = connection.execute(
                "SELECT * FROM product_config WHERE product_code = ?", (business_type,)
            ).fetchone()
        if mapping is None:
            raise P0IntakeError(f"P0_ACTIVE_MAPPING_NOT_FOUND:{business_type}")
        if catalog is None:
            raise P0IntakeError("P0_ACTIVE_CATALOG_NOT_FOUND")
        if product is None:
            raise P0IntakeError("P0_PRODUCT_NOT_FOUND")
        if not product["enabled"]:
            raise P0IntakeError("P0_PRODUCT_DISABLED")
        return dict(mapping), dict(catalog), dict(product)

    def _preview_contract(
        self, business_type: str
    ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
        """Allow a product's first upload to compare against its starter Draft."""

        code = business_type.strip().upper()
        with self.repository.connect() as connection:
            mapping = connection.execute(
                """SELECT * FROM mapping_config
                     WHERE business_type = ? AND status IN ('ACTIVE', 'DRAFT')
                     ORDER BY CASE status WHEN 'ACTIVE' THEN 0 ELSE 1 END, version DESC LIMIT 1""",
                (code,),
            ).fetchone()
            catalog = connection.execute(
                "SELECT * FROM standard_field_catalog_version WHERE status = 'ACTIVE'"
            ).fetchone()
            product = connection.execute(
                "SELECT * FROM product_config WHERE product_code = ?", (code,)
            ).fetchone()
        if mapping is None:
            raise P0IntakeError(f"P0_MAPPING_NOT_FOUND:{code}")
        if catalog is None:
            raise P0IntakeError("P0_ACTIVE_CATALOG_NOT_FOUND")
        if product is None:
            raise P0IntakeError("P0_PRODUCT_NOT_FOUND")
        if not product["enabled"]:
            raise P0IntakeError("P0_PRODUCT_DISABLED")
        return dict(mapping), dict(catalog), dict(product)

    def _inspect_workbook(self, source: Path, mapping_config_id: str) -> list[dict[str, Any]]:
        try:
            # Preserve formulas instead of requesting cached values: Excel and
            # Feishu exports often have no cached formula result, including in
            # header cells. Formula text still carries structural information.
            workbook = load_workbook(source, read_only=True, data_only=False)
        except Exception as error:
            raise P0IntakeError("P0_INTAKE_WORKBOOK_OPEN_FAILED") from error
        candidates: list[dict[str, Any]] = []
        try:
            with self.repository.connect() as connection:
                for worksheet in workbook.worksheets:
                    # Excel's stored dimension can be stale (e.g. A1:A1). Resetting
                    # makes openpyxl stream actual cells without changing the file.
                    if hasattr(worksheet, "reset_dimensions"):
                        worksheet.reset_dimensions()
                    rows = list(worksheet.iter_rows(values_only=True))
                    candidate = self._sheet_candidate(worksheet.title, rows, connection, mapping_config_id)
                    candidates.append(candidate)
        finally:
            workbook.close()
        if not candidates:
            raise P0IntakeError("P0_INTAKE_NO_SHEET")
        return candidates

    def _sheet_candidate(
        self, sheet_name: str, values: list[tuple[Any, ...]], connection: Any, mapping_config_id: str
    ) -> dict[str, Any]:
        best: dict[str, Any] | None = None
        for row_index, row in enumerate(values[:40], start=1):
            headers = [self._cell_text(value) for value in row]
            nonempty = [header for header in headers if header]
            if not nonempty:
                continue
            decisions = [self._field_decision(header, connection, mapping_config_id) for header in headers if header]
            matched = sum(item["status"] == "MATCHED" for item in decisions)
            conflicts = sum(item["status"] == "CONFLICT" for item in decisions)
            raw_only = sum(item["status"] == "RAW_ONLY" for item in decisions)
            score = matched * 10 + conflicts * 8 + raw_only * 3 + min(len(nonempty), 10) / 100
            candidate = {
                "sheet_name": sheet_name,
                "header_row": row_index,
                "data_start_row": row_index + 1,
                "headers": headers,
                "decisions": decisions,
                "score": score,
                "data_values": values[row_index:],
            }
            if best is None or candidate["score"] > best["score"]:
                best = candidate
        if best is None:
            return {
                "sheet_name": sheet_name, "header_row": 1, "data_start_row": 2,
                "headers": [], "decisions": [], "score": 0, "data_values": [],
            }
        return best

    @staticmethod
    def _select_sheet(candidates: list[dict[str, Any]], requested: str | None) -> dict[str, Any]:
        if requested:
            selected = next((item for item in candidates if item["sheet_name"] == requested), None)
            if selected is None:
                raise P0IntakeError("P0_INTAKE_SHEET_NOT_FOUND")
            return selected
        return max(candidates, key=lambda item: (item["score"], item["sheet_name"]))

    def _build_snapshot(
        self,
        source: Path,
        file_sha256: str,
        business_type: str,
        mapping: dict[str, Any],
        catalog: dict[str, Any],
        product: dict[str, Any],
        selected: dict[str, Any],
        candidates: list[dict[str, Any]],
    ) -> dict[str, Any]:
        headers = selected["headers"]
        decisions_by_header: list[dict[str, Any] | None] = []
        with self.repository.connect() as connection:
            for header in headers:
                decisions_by_header.append(
                    self._field_decision(header, connection, mapping["config_id"]) if header else None
                )
        rows: list[dict[str, Any]] = []
        missing_row_numbers: list[int] = []
        for offset, values in enumerate(selected["data_values"], start=selected["data_start_row"]):
            raw_json = {
                header: value for header, value in zip(headers, values)
                if header and value is not None and self._cell_text(value) != ""
            }
            if not raw_json:
                continue
            normalized: dict[str, dict[str, Any]] = {}
            for header, value, decision in zip(headers, values, decisions_by_header):
                if decision is None or value is None or self._cell_text(value) == "":
                    continue
                if decision["status"] != "MATCHED":
                    continue
                domain, field = decision["target"].split(".", 1)
                normalized.setdefault(domain, {})[field] = value
            if not self._required_value(normalized):
                missing_row_numbers.append(offset)
            rows.append({
                "row_number": offset,
                "raw_json": raw_json,
                "normalized_candidate": normalized,
            })
        header_missing = not any(
            decision and decision["status"] == "MATCHED" and decision["target"] == self.REQUIRED_PATH
            for decision in decisions_by_header
        )
        required_missing = [self.REQUIRED_PATH] if header_missing or missing_row_numbers else []
        required_missing_details = []
        if required_missing:
            required_missing_details.append({
                "path": self.REQUIRED_PATH,
                "label_zh": "问题编号",
                "reason": "SOURCE_HEADER_NOT_FOUND" if header_missing else "ROW_VALUE_MISSING",
                "row_numbers": missing_row_numbers,
            })
        conflicts = [item for item in decisions_by_header if item and item["status"] == "CONFLICT"]
        return {
            "file_path": str(source.resolve()),
            "file_name": source.name,
            "file_sha256": file_sha256,
            "business_type": business_type,
            "product_id": product["product_id"],
            "product_code": product["product_code"],
            "sheet_name": selected["sheet_name"],
            "header_row": selected["header_row"],
            "data_start_row": selected["data_start_row"],
            "sheet_candidates": [
                {key: candidate[key] for key in ("sheet_name", "header_row", "data_start_row", "score")}
                for candidate in candidates
            ],
            "mapping_config_id": mapping["config_id"],
            "mapping_config_version": mapping["version"],
            "mapping_status": mapping["status"],
            "mapping_content_hash": mapping["content_hash"],
            "catalog_version_id": catalog["catalog_version_id"],
            "catalog_content_hash": catalog["content_hash"],
            "field_decisions": [item for item in decisions_by_header if item],
            "has_conflict": bool(conflicts),
            "required_missing": required_missing,
            "required_missing_details": required_missing_details,
            "rows": rows,
            "total_rows": len(rows),
            "can_confirm": mapping["status"] == "ACTIVE" and not required_missing and not conflicts,
        }

    def _field_decision(self, header: str, connection: Any, mapping_config_id: str) -> dict[str, Any]:
        resolution = self._field_service.resolve_source_header(header, connection, mapping_config_id)
        status = resolution["status"]
        if status == "IGNORED":
            status = "RAW_ONLY"
        candidates = resolution.get("candidates", [])
        target = candidates[0]["path"] if status == "MATCHED" and candidates else None
        return {
            "source_header": header,
            "status": status,
            "target": target,
            "candidates": candidates,
            "required": bool(resolution.get("required", False)) if status == "MATCHED" else False,
            "handling_requirement": resolution.get("handling_requirement"),
        }

    def _load_snapshot(self, token: str) -> dict[str, Any] | None:
        with self.repository.connect() as connection:
            row = connection.execute(
                "SELECT * FROM intake_preview_snapshot WHERE preview_token = ?", (token,)
            ).fetchone()
        if row is None:
            return None
        result = dict(row)
        result["snapshot"] = json.loads(result.pop("snapshot_json"))
        return result

    def _existing_batch(self, token: str) -> dict[str, Any] | None:
        with self.repository.connect() as connection:
            row = connection.execute("SELECT * FROM import_batch WHERE preview_token = ?", (token,)).fetchone()
        return dict(row) if row else None

    def _record_error(self, batch_id: str, row_number: int, code: str, message: str, details: dict[str, Any]) -> None:
        with self.repository.connect() as connection:
            connection.execute(
                """INSERT INTO import_error(import_error_id, import_batch_id, source_row_number, code, message, details_json)
                    VALUES (?, ?, ?, ?, ?, ?)""",
                (f"IE-{uuid.uuid4().hex}", batch_id, row_number, code, message, self._dump(details)),
            )
            connection.commit()

    def _batch_summary(self, batch: dict[str, Any], snapshot: dict[str, Any]) -> dict[str, Any]:
        with self.repository.connect() as connection:
            failed = connection.execute(
                "SELECT COUNT(*) FROM import_error WHERE import_batch_id = ?",
                (batch["import_batch_id"],),
            ).fetchone()[0]
            applied = connection.execute(
                "SELECT COUNT(*) FROM issue_source_raw WHERE import_batch_id = ?",
                (batch["import_batch_id"],),
            ).fetchone()[0]
        total = len(snapshot["rows"])
        duplicate, unique_business_issues = self._duplicate_summary(snapshot)
        return {
            "status": batch["status"], "total": total, "success": applied,
            "already_applied": total - failed - applied, "duplicate": duplicate,
            "unique_business_issues": unique_business_issues, "failed": failed,
        }

    def _duplicate_summary(self, snapshot: dict[str, Any]) -> tuple[int, int]:
        """Return duplicate source-row and unique business-key counts for a preview."""

        seen: set[str] = set()
        duplicate = 0
        for row in snapshot["rows"]:
            business_issue_id = self._required_value(row["normalized_candidate"])
            if not business_issue_id:
                continue
            if business_issue_id in seen:
                duplicate += 1
            else:
                seen.add(business_issue_id)
        return duplicate, len(seen)

    @staticmethod
    def _required_value(normalized: dict[str, Any]) -> str:
        value = normalized.get("ISSUE_FACT", {}).get("business_issue_id")
        return str(value).strip() if value is not None else ""

    @staticmethod
    def _knowledge_id(business_type: str, business_issue_id: str) -> str:
        digest = hashlib.sha256(f"{business_type}:{business_issue_id}".encode("utf-8")).hexdigest()[:20]
        return f"QK-{business_type}-{digest}"

    @staticmethod
    def _cell_text(value: Any) -> str:
        return str(value).strip() if value is not None else ""

    @staticmethod
    def _sha256_file(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    @staticmethod
    def _dump(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
