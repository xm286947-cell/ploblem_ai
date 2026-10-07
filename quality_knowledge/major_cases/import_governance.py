"""R2 W2 governance metadata for the existing Major Excel intake."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TEMPLATE_CONTRACT = "major-excel-template/v1"
TEMPLATE_VERSION = "1.0"
TEMPLATE_META_SHEET = "_MAJOR_IMPORT_META"
MAPPING_CONTRACT = "major-excel-field-mapping/v1"
IMPORT_GOVERNANCE_CONTRACT = "major-excel-import-governance/v1"


def mapping_version(field_mapping: dict[str, Any]) -> str:
    payload = json.dumps(
        field_mapping,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return "sha256:" + hashlib.sha256(payload).hexdigest()


def inspect_template(path: str | Path) -> dict[str, str]:
    """Inspect version metadata only; the existing ExcelParser still parses rows."""
    source = Path(path)
    if source.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("EXCEL_XLS_READER_NOT_INSTALLED") from exc
        workbook = xlrd.open_workbook(filename=str(source), on_demand=False)
        try:
            try:
                sheet = workbook.sheet_by_name(TEMPLATE_META_SHEET)
            except xlrd.biffh.XLRDError:
                return {
                    "template_contract": TEMPLATE_CONTRACT,
                    "template_version": "LEGACY_UNVERSIONED",
                    "template_status": "LEGACY_COMPATIBLE",
                }
            values = {
                str(sheet.cell_value(row, 0) or "").strip(): str(sheet.cell_value(row, 1) or "").strip()
                for row in range(sheet.nrows)
                if sheet.ncols and str(sheet.cell_value(row, 0) or "").strip()
            }
        finally:
            workbook.release_resources()
    else:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if TEMPLATE_META_SHEET not in workbook.sheetnames:
                return {
                    "template_contract": TEMPLATE_CONTRACT,
                    "template_version": "LEGACY_UNVERSIONED",
                    "template_status": "LEGACY_COMPATIBLE",
                }
            sheet = workbook[TEMPLATE_META_SHEET]
            values: dict[str, str] = {}
            for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
                key = str(row[0] or "").strip()
                if key:
                    values[key] = str(row[1] or "").strip()
        finally:
            workbook.close()

    if values.get("template_contract") != TEMPLATE_CONTRACT:
        raise ValueError("MAJOR_EXCEL_TEMPLATE_CONTRACT_UNSUPPORTED")
    if values.get("template_version") != TEMPLATE_VERSION:
        raise ValueError("MAJOR_EXCEL_TEMPLATE_VERSION_UNSUPPORTED")
    return {
        "template_contract": TEMPLATE_CONTRACT,
        "template_version": TEMPLATE_VERSION,
        "template_status": "OFFICIAL",
    }


def add_template_metadata(workbook: Any, *, current_mapping_version: str) -> None:
    if TEMPLATE_META_SHEET in workbook.sheetnames:
        del workbook[TEMPLATE_META_SHEET]
    sheet = workbook.create_sheet(TEMPLATE_META_SHEET)
    sheet.append(["template_contract", TEMPLATE_CONTRACT])
    sheet.append(["template_version", TEMPLATE_VERSION])
    sheet.append(["mapping_contract", MAPPING_CONTRACT])
    sheet.append(["mapping_version", current_mapping_version])
    sheet.sheet_state = "hidden"
