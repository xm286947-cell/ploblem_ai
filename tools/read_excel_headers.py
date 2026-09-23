#!/usr/bin/env python3
import argparse, json, re
from pathlib import Path
from openpyxl import load_workbook

def clean(value):
    return re.sub(r'\s+', ' ', str(value)).strip() if value is not None else ''

def inspect(ws, limit):
    rows=[]
    for n in range(1, min(ws.max_row or 0, limit)+1):
        vals=[clean(v) for v in next(ws.iter_rows(min_row=n,max_row=n,values_only=True))]
        nonempty=[v for v in vals if v]
        if nonempty: rows.append({'row':n,'count':len(nonempty),'values':nonempty})
    candidates=sorted(rows,key=lambda x:(-x['count'],x['row']))
    header=candidates[0] if candidates else None
    data=next((r['row'] for r in rows if header and r['row']>header['row']),None)
    return {'sheet':ws.title,'dimension':ws.calculate_dimension(),'max_row':ws.max_row,'max_column':ws.max_column,'merged':[str(x) for x in ws.merged_cells.ranges],'hidden_rows':[n for n in range(1,(ws.max_row or 0)+1) if ws.row_dimensions[n].hidden],'header_candidate':header,'data_start_row':data,'rows':rows}

def main():
    p=argparse.ArgumentParser(description='诊断 Excel 表头和数据起始行')
    p.add_argument('excel',type=Path); p.add_argument('--rows',type=int,default=80); p.add_argument('--json',dest='out',type=Path)
    a=p.parse_args()
    if not a.excel.exists(): p.error(f'文件不存在: {a.excel}')
    if a.excel.suffix.lower() not in {'.xlsx','.xlsm'}: p.error('仅支持 .xlsx/.xlsm')
    wb=load_workbook(a.excel,data_only=False,read_only=False)
    report={'file':str(a.excel.resolve()),'sheet_count':len(wb.sheetnames),'sheets':[inspect(wb[n],a.rows) for n in wb.sheetnames]}
    print(json.dumps(report,ensure_ascii=False,indent=2))
    for s in report['sheets']:
        h=s['header_candidate']; print(f"\n[{s['sheet']}] {s['dimension']}")
        print(f"表头候选: 第{h['row']}行，共{h['count']}列: {' | '.join(h['values'])}" if h else '未找到非空行')
        print(f"数据起始: 第{s['data_start_row']}行" if s['data_start_row'] else '数据起始: 未找到')
        print(f"合并单元格: {len(s['merged'])}，隐藏行: {len(s['hidden_rows'])}")
    if a.out: a.out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8'); print(f'报告已写入: {a.out}')

if __name__=='__main__': main()
