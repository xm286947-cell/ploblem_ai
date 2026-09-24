"""Company-environment validation harness for Hardware Case M4.

This tool is intentionally data-minimizing:
- real Word/Excel files stay local;
- local absolute paths are never persisted in the report;
- report contains aggregate metrics only;
- Provider access is injected through an approved Unified Runtime structurer.

It is safe to commit because the repository contains only the harness and a
placeholder configuration, never internal source data.
"""
from __future__ import annotations

import argparse
import importlib
import json
import math
import statistics
import time
from collections import Counter
from hashlib import sha256
from pathlib import Path
from typing import Any, Callable

from repositories.hardware_case_repository import HardwareCaseRepository
from services.hardware_case_ai_adapter import FACT_FIELDS, HardwareCaseAIAdapter
from services.hardware_case_backend import HardwareCaseBackendService
from services.hardware_case_source_store import HardwareCaseSourceStore


CORE_FACTS = ("symptom", "root_cause", "actions")
VALID_TREE_TYPES = {"CIRCUIT_FEATURE", "MATERIAL_DEVICE"}


class RealValidationConfigError(RuntimeError):
    pass


def _ratio(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator / denominator, 4)


def _load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except OSError as exc:
        raise RealValidationConfigError("CONFIG_UNAVAILABLE") from exc
    except json.JSONDecodeError as exc:
        raise RealValidationConfigError("CONFIG_JSON_INVALID") from exc
    if not isinstance(value, dict):
        raise RealValidationConfigError("CONFIG_OBJECT_REQUIRED")
    return value


def load_config(path: str | Path) -> dict[str, Any]:
    config_path = Path(path).expanduser().resolve()
    config = _load_json(config_path)
    config["_config_dir"] = str(config_path.parent)
    return config


def _resolve_local_path(config: dict[str, Any], value: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = Path(config["_config_dir"]) / path
    return path.resolve()


def resolve_structurer(spec: str) -> Callable[[dict[str, Any]], dict[str, Any]]:
    """Resolve company adapter in module:function form.

    The target can be either the document->dict callable itself or a zero-arg
    factory returning that callable.
    """
    if ":" not in spec:
        raise RealValidationConfigError("STRUCTURER_SPEC_INVALID")
    module_name, attribute_name = spec.split(":", 1)
    module = importlib.import_module(module_name)
    target = getattr(module, attribute_name, None)
    if target is None or not callable(target):
        raise RealValidationConfigError("STRUCTURER_NOT_CALLABLE")

    # A company adapter may expose a factory to initialize existing Runtime
    # configuration.  Mark such factories explicitly to avoid probing a real
    # provider during resolution.
    if getattr(target, "__hardware_case_structurer_factory__", False):
        target = target()
        if not callable(target):
            raise RealValidationConfigError("STRUCTURER_FACTORY_INVALID")
    return target


def _cell_value(sheet, coordinate: str, row: int) -> Any:
    if coordinate.isalpha():
        return sheet[f"{coordinate}{row}"].value
    return sheet[coordinate].value


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_tree_xlsx(
    backend: HardwareCaseBackendService,
    *,
    path: Path,
    tree_type: str,
    sheet_name: str,
    path_columns: list[str],
    start_row: int = 2,
    id_column: str | None = None,
    active_column: str | None = None,
) -> int:
    """Load a real tree without assuming business semantics.

    The company-only config explicitly selects which columns form the path.
    This preserves the product contract while avoiding any hard-coded D/F or
    fixed-depth interpretation in public code.
    """
    if tree_type not in VALID_TREE_TYPES:
        raise RealValidationConfigError("TREE_TYPE_INVALID")
    if not path_columns:
        raise RealValidationConfigError("TREE_PATH_COLUMNS_REQUIRED")
    try:
        import openpyxl
    except ImportError as exc:
        raise RealValidationConfigError("OPENPYXL_REQUIRED") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise RealValidationConfigError("TREE_SHEET_NOT_FOUND")
    sheet = workbook[sheet_name]
    seen: set[str] = set()
    count = 0

    for row in range(int(start_row), sheet.max_row + 1):
        path_values = [_clean(_cell_value(sheet, col, row)) for col in path_columns]
        path_values = [item for item in path_values if item]
        if not path_values:
            continue
        explicit_id = (
            _clean(_cell_value(sheet, id_column, row)) if id_column else ""
        )
        identity_basis = "|".join(
            [tree_type, sheet_name, str(row), *path_values]
        )
        node_id = explicit_id or (
            ("CF-" if tree_type == "CIRCUIT_FEATURE" else "MD-")
            + sha256(identity_basis.encode("utf-8")).hexdigest()[:16].upper()
        )
        if node_id in seen:
            continue
        seen.add(node_id)

        active = True
        if active_column:
            raw_active = _clean(_cell_value(sheet, active_column, row)).lower()
            if raw_active in {"0", "false", "no", "n", "否", "停用"}:
                active = False

        backend.save_tree_node(
            {
                "node_id": node_id,
                "tree_type": tree_type,
                "name": path_values[-1],
                "parent_id": None,
                "path": path_values,
                "description": None,
                "source_ref": f"excel:{path.name}#{sheet_name}!row:{row}",
                "active": active,
                "source_metadata": {
                    "sheet": sheet_name,
                    "row": row,
                    "path_columns": list(path_columns),
                },
            }
        )
        count += 1
    workbook.close()
    return count


def load_trees_from_config(
    backend: HardwareCaseBackendService,
    config: dict[str, Any],
) -> dict[str, int]:
    result = {"CIRCUIT_FEATURE": 0, "MATERIAL_DEVICE": 0}
    for item in config.get("trees") or []:
        if not isinstance(item, dict):
            raise RealValidationConfigError("TREE_CONFIG_INVALID")
        path = _resolve_local_path(config, str(item.get("path") or ""))
        tree_type = str(item.get("tree_type") or "")
        count = load_tree_xlsx(
            backend,
            path=path,
            tree_type=tree_type,
            sheet_name=str(item.get("sheet") or ""),
            path_columns=[str(x) for x in item.get("path_columns") or []],
            start_row=int(item.get("start_row", 2)),
            id_column=(
                str(item["id_column"]) if item.get("id_column") else None
            ),
            active_column=(
                str(item["active_column"])
                if item.get("active_column")
                else None
            ),
        )
        result[tree_type] = result.get(tree_type, 0) + count
    return result


def discover_word_files(config: dict[str, Any]) -> list[Path]:
    source_dir = _resolve_local_path(
        config, str(config.get("source_dir") or "")
    )
    if not source_dir.exists() or not source_dir.is_dir():
        raise RealValidationConfigError("SOURCE_DIR_INVALID")
    recursive = bool(config.get("recursive", False))
    iterator = source_dir.rglob("*.docx") if recursive else source_dir.glob("*.docx")
    files = sorted(path for path in iterator if path.is_file())
    limit = int(config.get("max_cases", 30))
    if limit <= 0:
        raise RealValidationConfigError("MAX_CASES_INVALID")
    return files[:limit]


def _case_metrics(
    backend: HardwareCaseBackendService,
    case_id: str,
) -> dict[str, Any]:
    case = backend.get_case(case_id, role="MAINTAINER")
    facts = case.get("facts") or {}
    populated = [
        name
        for name in FACT_FIELDS
        if isinstance(facts.get(name), dict)
        and facts[name].get("candidate_value") not in (None, "")
    ]
    grounded = [
        name
        for name in populated
        if facts[name].get("evidence_refs")
    ]
    core_populated = [name for name in CORE_FACTS if name in populated]
    core_grounded = [name for name in CORE_FACTS if name in grounded]
    mappings = backend.repository.list_mappings(case_id=case_id)
    return {
        "populated_fact_count": len(populated),
        "grounded_fact_count": len(grounded),
        "core_populated_count": len(core_populated),
        "core_grounded_count": len(core_grounded),
        "has_circuit_link": any(
            item.get("tree_type") == "CIRCUIT_FEATURE"
            for item in mappings
        ),
        "has_material_link": any(
            item.get("tree_type") == "MATERIAL_DEVICE"
            for item in mappings
        ),
    }


def run_validation(
    config: dict[str, Any],
    *,
    structurer: Callable[[dict[str, Any]], dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Run a company-local batch and return sanitized aggregate metrics."""
    output_db = _resolve_local_path(
        config, str(config.get("output_db") or "./hardware_case_validation.sqlite3")
    )
    output_db.parent.mkdir(parents=True, exist_ok=True)
    backend = HardwareCaseBackendService(HardwareCaseRepository(output_db))
    source_root = _resolve_local_path(
        config,
        str(config.get("source_root") or "../data/evidence_sources"),
    )
    source_store = HardwareCaseSourceStore(output_db, source_root)

    tree_counts = load_trees_from_config(backend, config)
    if structurer is None:
        structurer = resolve_structurer(str(config.get("structurer") or ""))

    files = discover_word_files(config)
    adapter = HardwareCaseAIAdapter(
        backend,
        structurer,
        source_store=source_store,
    )

    status_counts: Counter[str] = Counter()
    failure_categories: Counter[str] = Counter()
    durations: list[float] = []
    parse_success = 0
    populated_facts = 0
    grounded_facts = 0
    core_populated = 0
    core_grounded = 0
    circuit_linked = 0
    material_linked = 0

    for path in files:
        started = time.perf_counter()
        try:
            result = adapter.ingest_docx(path)
            status = str(result.get("status") or "UNKNOWN")
            status_counts[status] += 1
            if result.get("processing_status") == "READY":
                parse_success += 1
            for warning in result.get("warnings") or []:
                failure_categories[str(warning).split(":", 1)[0]] += 1
            case_id = str(result.get("case_id") or "")
            if case_id:
                metrics = _case_metrics(backend, case_id)
                populated_facts += metrics["populated_fact_count"]
                grounded_facts += metrics["grounded_fact_count"]
                core_populated += metrics["core_populated_count"]
                core_grounded += metrics["core_grounded_count"]
                circuit_linked += int(metrics["has_circuit_link"])
                material_linked += int(metrics["has_material_link"])
        except Exception as exc:
            status_counts["FAILED"] += 1
            failure_categories[type(exc).__name__] += 1
        finally:
            durations.append(time.perf_counter() - started)

    processed = len(files)
    needs_review = status_counts.get("NEEDS_REVIEW", 0)
    failed = status_counts.get("FAILED", 0)

    report = {
        "report_version": "hardware-case-real-validation/v1",
        "input_case_count": processed,
        "tree_node_count": tree_counts,
        "parse_success_rate": _ratio(parse_success, processed),
        "field_extraction_coverage": _ratio(
            populated_facts, processed * len(FACT_FIELDS)
        ),
        "core_field_extraction_coverage": _ratio(
            core_populated, processed * len(CORE_FACTS)
        ),
        "evidence_grounding_rate": _ratio(grounded_facts, populated_facts),
        "core_evidence_grounding_rate": _ratio(
            core_grounded, core_populated
        ),
        "circuit_link_rate": _ratio(circuit_linked, processed),
        "material_link_rate": _ratio(material_linked, processed),
        "needs_review_rate": _ratio(needs_review, processed),
        "failed_rate": _ratio(failed, processed),
        "status_counts": dict(sorted(status_counts.items())),
        "failure_categories": dict(sorted(failure_categories.items())),
        "average_processing_time_seconds": (
            round(statistics.fmean(durations), 3) if durations else 0.0
        ),
        "p95_processing_time_seconds": (
            round(sorted(durations)[max(math.ceil(len(durations) * 0.95) - 1, 0)], 3)
            if durations
            else 0.0
        ),
        "security": {
            "contains_source_paths": False,
            "contains_case_text": False,
            "contains_case_ids": False,
            "contains_secrets": False,
        },
    }

    output_report = _resolve_local_path(
        config,
        str(config.get("output_report") or "./hardware_case_validation_metrics.json"),
    )
    output_report.parent.mkdir(parents=True, exist_ok=True)
    output_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return report


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run Hardware Case M4 company-local real validation."
    )
    parser.add_argument("--config", required=True)
    args = parser.parse_args()

    config = load_config(args.config)
    report = run_validation(config)
    # Print metrics only; never print source paths, case ids, or source text.
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
