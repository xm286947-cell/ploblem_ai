"""Excel parse + current-tree Diff engine for HC-TREE-IMPORT-001 M2.

The engine supports the MVP horizontal hierarchy model:
    Level1 | Level2 | ... | LevelN
Column names and depth are selected by the maintainer Mapping Profile. Real
files remain local; only filename/hash and structured import facts are persisted.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from typing import Any

import openpyxl

from repositories.hardware_tree_import_repository import HardwareTreeImportRepository
from services.hardware_tree_import_contract import (
    HardwareTreeImportContractError,
    validate_mapping_profile,
    validate_tree_type,
)


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _path_key(path: list[str]) -> tuple[str, ...]:
    return tuple(_clean(item).casefold() for item in path)


def _node_id_for_path(tree_type: str, path: list[str]) -> str:
    prefix = "CF-" if tree_type == "CIRCUIT_FEATURE" else "MD-"
    basis = tree_type + "|PATH|" + "|".join(item.casefold() for item in path)
    return prefix + sha256(basis.encode("utf-8")).hexdigest()[:16].upper()


def _change_id(job_id: str, kind: str, identity: str) -> str:
    value = f"{job_id}|{kind}|{identity}"
    return "CH-" + sha256(value.encode("utf-8")).hexdigest()[:20].upper()


def _issue_id(job_id: str, issue_type: str, identity: str) -> str:
    value = f"{job_id}|{issue_type}|{identity}"
    return "ISS-" + sha256(value.encode("utf-8")).hexdigest()[:20].upper()


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class HardwareTreeExcelParser:
    def inspect_workbook(self, path: str | Path) -> dict[str, Any]:
        source = Path(path)
        self._validate_file(source)
        try:
            workbook = openpyxl.load_workbook(
                source, read_only=True, data_only=True
            )
        except Exception as exc:
            raise HardwareTreeImportContractError("EXCEL_PARSE_FAILED") from exc
        try:
            sheets = []
            for sheet in workbook.worksheets:
                sheets.append(
                    {
                        "sheet_name": sheet.title,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                    }
                )
            return {"filename": source.name, "sheets": sheets}
        finally:
            workbook.close()

    @staticmethod
    def _validate_file(path: Path) -> None:
        if not path.is_file():
            raise HardwareTreeImportContractError("EXCEL_FILE_NOT_FOUND")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise HardwareTreeImportContractError("EXCEL_FORMAT_UNSUPPORTED")

    def preview_rows(
        self,
        path: str | Path,
        *,
        sheet_name: str,
        header_row: int,
        max_rows: int = 20,
    ) -> dict[str, Any]:
        """Return a bounded raw workbook preview for the P07 Mapping UI.

        This is intentionally read-only and data-bounded. It does not infer
        hierarchy roles or persist any import decision.
        """
        source = Path(path)
        self._validate_file(source)
        if not isinstance(header_row, int) or header_row < 1:
            raise HardwareTreeImportContractError("HEADER_ROW_INVALID")
        max_rows = max(1, min(int(max_rows), 50))
        try:
            workbook = openpyxl.load_workbook(
                source, read_only=True, data_only=True
            )
        except Exception as exc:
            raise HardwareTreeImportContractError("EXCEL_PARSE_FAILED") from exc
        try:
            if sheet_name not in workbook.sheetnames:
                raise HardwareTreeImportContractError("TREE_SHEET_NOT_FOUND")
            sheet = workbook[sheet_name]
            if header_row > sheet.max_row:
                raise HardwareTreeImportContractError("HEADER_ROW_OUT_OF_RANGE")
            header_cells = next(
                sheet.iter_rows(min_row=header_row, max_row=header_row)
            )
            headers = [_clean(cell.value) for cell in header_cells]
            columns = []
            for index, name in enumerate(headers, start=1):
                columns.append(
                    {
                        "column_index": index,
                        "column_name": name,
                        "column_key": openpyxl.utils.get_column_letter(index),
                    }
                )
            rows = []
            for excel_row, row in enumerate(
                sheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=min(sheet.max_row, header_row + max_rows),
                ),
                start=header_row + 1,
            ):
                values = [
                    _clean(cell.value if hasattr(cell, "value") else cell)
                    for cell in row
                ]
                rows.append(
                    {
                        "row_number": excel_row,
                        "values": values[: len(headers)],
                    }
                )
            return {
                "filename": source.name,
                "sheet_name": sheet_name,
                "header_row": header_row,
                "columns": columns,
                "rows": rows,
                "total_rows": max(sheet.max_row - header_row, 0),
            }
        finally:
            workbook.close()

    def parse(
        self,
        path: str | Path,
        *,
        tree_type: str,
        profile: dict[str, Any],
        job_id: str,
    ) -> dict[str, Any]:
        source = Path(path)
        self._validate_file(source)
        tree_type = validate_tree_type(tree_type)
        profile = validate_mapping_profile(profile)

        try:
            workbook = openpyxl.load_workbook(
                source, read_only=True, data_only=True
            )
        except Exception as exc:
            raise HardwareTreeImportContractError("EXCEL_PARSE_FAILED") from exc

        try:
            if profile["sheet_name"] not in workbook.sheetnames:
                raise HardwareTreeImportContractError("TREE_SHEET_NOT_FOUND")
            sheet = workbook[profile["sheet_name"]]
            if profile["header_row"] > sheet.max_row:
                raise HardwareTreeImportContractError("HEADER_ROW_OUT_OF_RANGE")

            header_values = [
                _clean(cell.value)
                for cell in next(
                    sheet.iter_rows(
                        min_row=profile["header_row"],
                        max_row=profile["header_row"],
                    )
                )
            ]
            header_positions: dict[str, list[int]] = defaultdict(list)
            for index, value in enumerate(header_values):
                if value:
                    header_positions[value].append(index)

            referenced = (
                list(profile["path_columns"])
                + list(profile["metadata_columns"])
                + (
                    [profile["business_key_column"]]
                    if profile.get("business_key_column")
                    else []
                )
            )
            for column in referenced:
                positions = header_positions.get(column) or []
                if not positions:
                    raise HardwareTreeImportContractError("MAPPING_COLUMN_NOT_FOUND")
                if len(positions) > 1:
                    raise HardwareTreeImportContractError("MAPPING_COLUMN_AMBIGUOUS")

            def value(row: tuple[Any, ...], column: str) -> str:
                idx = header_positions[column][0]
                if idx >= len(row):
                    return ""
                cell = row[idx]
                return _clean(cell.value if hasattr(cell, "value") else cell)

            nodes_by_path: dict[tuple[str, ...], dict[str, Any]] = {}
            issues: list[dict[str, Any]] = []
            seen_rows: set[str] = set()
            business_key_paths: dict[str, set[tuple[str, ...]]] = defaultdict(set)
            row_count = 0
            data_row_count = 0
            duplicate_rows = 0
            max_depth = 0

            for excel_row, row in enumerate(
                sheet.iter_rows(min_row=profile["header_row"] + 1),
                start=profile["header_row"] + 1,
            ):
                row_count += 1
                levels = [value(row, col) for col in profile["path_columns"]]
                if not any(levels):
                    continue
                data_row_count += 1

                last = max(index for index, item in enumerate(levels) if item)
                missing_parent_indexes = [
                    index for index in range(last) if not levels[index]
                ]
                if missing_parent_indexes:
                    first_missing = missing_parent_indexes[0]
                    issues.append(
                        {
                            "issue_id": _issue_id(
                                job_id,
                                "PARENT_LEVEL_MISSING",
                                f"{profile['sheet_name']}:{excel_row}:{first_missing}",
                            ),
                            "sheet_name": profile["sheet_name"],
                            "row_number": excel_row,
                            "column_name": profile["path_columns"][first_missing],
                            "original_value": levels[last],
                            "issue_type": "PARENT_LEVEL_MISSING",
                            "suggested_action": "补齐缺失的上级分类后重新解析",
                            "resolved": False,
                        }
                    )
                    continue

                path = levels[: last + 1]
                business_key = (
                    value(row, profile["business_key_column"])
                    if profile.get("business_key_column")
                    else ""
                )
                business_key = business_key or None
                business_metadata = {
                    column: value(row, column)
                    for column in profile["metadata_columns"]
                    if value(row, column)
                }
                row_signature = json.dumps(
                    {
                        "path": path,
                        "business_key": business_key,
                        "metadata": business_metadata,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                )
                if row_signature in seen_rows:
                    duplicate_rows += 1
                    issues.append(
                        {
                            "issue_id": _issue_id(
                                job_id,
                                "EXACT_DUPLICATE",
                                f"{profile['sheet_name']}:{excel_row}:{row_signature}",
                            ),
                            "sheet_name": profile["sheet_name"],
                            "row_number": excel_row,
                            "column_name": None,
                            "original_value": " / ".join(path),
                            "issue_type": "EXACT_DUPLICATE",
                            "suggested_action": "系统已折叠重复记录，无需人工处理",
                            "resolved": True,
                        }
                    )
                    continue
                seen_rows.add(row_signature)
                max_depth = max(max_depth, len(path))

                if business_key:
                    business_key_paths[business_key].add(_path_key(path))

                for depth in range(1, len(path) + 1):
                    node_path = path[:depth]
                    key = _path_key(node_path)
                    node_id = _node_id_for_path(tree_type, node_path)
                    parent_id = (
                        _node_id_for_path(tree_type, node_path[:-1])
                        if depth > 1
                        else None
                    )
                    is_leaf = depth == len(path)
                    candidate = {
                        "node_id": node_id,
                        "tree_type": tree_type,
                        "business_key": business_key if is_leaf else None,
                        "name": node_path[-1],
                        "parent_id": parent_id,
                        "path": node_path,
                        "description": None,
                        "source_ref": (
                            f"excel:{source.name}#{profile['sheet_name']}!row:{excel_row}"
                            if is_leaf
                            else f"excel:{source.name}#{profile['sheet_name']}"
                        ),
                        "active": True,
                        "source_metadata": {
                            "sheet": profile["sheet_name"],
                            "header_row": profile["header_row"],
                            "path_columns": list(profile["path_columns"]),
                            "business_metadata": business_metadata if is_leaf else {},
                            "source_row": excel_row if is_leaf else None,
                        },
                    }
                    existing = nodes_by_path.get(key)
                    if existing is None:
                        nodes_by_path[key] = candidate
                        continue
                    if is_leaf:
                        old_key = existing.get("business_key")
                        if old_key and business_key and old_key != business_key:
                            issues.append(
                                {
                                    "issue_id": _issue_id(
                                        job_id,
                                        "PATH_BUSINESS_KEY_CONFLICT",
                                        f"{profile['sheet_name']}:{excel_row}:{'|'.join(key)}",
                                    ),
                                    "sheet_name": profile["sheet_name"],
                                    "row_number": excel_row,
                                    "column_name": profile.get("business_key_column"),
                                    "original_value": business_key,
                                    "issue_type": "PATH_BUSINESS_KEY_CONFLICT",
                                    "suggested_action": "同一路径存在不同 Business Key，请人工确认",
                                    "resolved": False,
                                }
                            )
                        elif business_key and not old_key:
                            existing["business_key"] = business_key
                        if business_metadata:
                            existing["source_metadata"]["business_metadata"].update(
                                business_metadata
                            )
                            existing["source_metadata"]["source_row"] = excel_row

            for business_key, paths in business_key_paths.items():
                if len(paths) <= 1:
                    continue
                issues.append(
                    {
                        "issue_id": _issue_id(
                            job_id,
                            "BUSINESS_KEY_MULTIPLE_PATHS",
                            business_key,
                        ),
                        "sheet_name": profile["sheet_name"],
                        "row_number": None,
                        "column_name": profile.get("business_key_column"),
                        "original_value": business_key,
                        "issue_type": "BUSINESS_KEY_MULTIPLE_PATHS",
                        "suggested_action": "同一 Business Key 对应多个路径，请人工确认节点身份",
                        "resolved": False,
                    }
                )

            nodes = sorted(
                nodes_by_path.values(),
                key=lambda item: (len(item["path"]), _path_key(item["path"])),
            )
            if not nodes:
                raise HardwareTreeImportContractError("TREE_DATA_EMPTY")
            return {
                "filename": source.name,
                "tree_type": tree_type,
                "mapping_profile": profile,
                "nodes": nodes,
                "issues": issues,
                "stats": {
                    "scanned_row_count": row_count,
                    "data_row_count": data_row_count,
                    "candidate_node_count": len(nodes),
                    "duplicate_row_count": duplicate_rows,
                    "issue_count": len(issues),
                    "max_depth": max_depth,
                },
            }
        finally:
            workbook.close()


class HardwareTreeDiffEngine:
    def build(
        self,
        *,
        job_id: str,
        tree_type: str,
        candidate_nodes: list[dict[str, Any]],
        current_nodes: list[dict[str, Any]],
    ) -> dict[str, Any]:
        tree_type = validate_tree_type(tree_type)
        current_by_path: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
        current_by_business: dict[str, list[dict[str, Any]]] = defaultdict(list)
        current_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)

        for node in current_nodes:
            current_by_path[_path_key(node.get("path") or [])].append(node)
            key = _clean(
                node.get("business_key")
                or (node.get("source_metadata") or {}).get("business_key")
            )
            if key:
                current_by_business[key].append(node)
            current_by_name[_clean(node.get("name")).casefold()].append(node)

        candidate_business_counts = Counter(
            _clean(node.get("business_key"))
            for node in candidate_nodes
            if _clean(node.get("business_key"))
        )

        match_by_candidate_id: dict[str, dict[str, Any] | None] = {}
        conflicts_by_candidate_id: dict[str, str] = {}

        for node in candidate_nodes:
            temp_id = str(node["node_id"])
            business_key = _clean(node.get("business_key"))
            path_matches = current_by_path.get(_path_key(node["path"])) or []

            if business_key:
                if candidate_business_counts[business_key] > 1:
                    conflicts_by_candidate_id[temp_id] = "CANDIDATE_BUSINESS_KEY_AMBIGUOUS"
                    match_by_candidate_id[temp_id] = None
                    continue
                key_matches = current_by_business.get(business_key) or []
                if len(key_matches) > 1:
                    conflicts_by_candidate_id[temp_id] = "CURRENT_BUSINESS_KEY_AMBIGUOUS"
                    match_by_candidate_id[temp_id] = None
                    continue
                if len(key_matches) == 1:
                    match_by_candidate_id[temp_id] = key_matches[0]
                    continue

            if len(path_matches) == 1:
                match_by_candidate_id[temp_id] = path_matches[0]
                continue
            if len(path_matches) > 1:
                conflicts_by_candidate_id[temp_id] = "CURRENT_PATH_AMBIGUOUS"
                match_by_candidate_id[temp_id] = None
                continue

            if not business_key:
                same_name = current_by_name.get(_clean(node["name"]).casefold()) or []
                if len(same_name) == 1:
                    conflicts_by_candidate_id[temp_id] = "PATH_CHANGED_WITHOUT_BUSINESS_KEY"
                    match_by_candidate_id[temp_id] = same_name[0]
                    continue
            match_by_candidate_id[temp_id] = None

        canonical_id: dict[str, str] = {}
        for node in candidate_nodes:
            temp_id = str(node["node_id"])
            match = match_by_candidate_id.get(temp_id)
            canonical_id[temp_id] = str(match["node_id"]) if match else temp_id

        changes: list[dict[str, Any]] = []
        issues: list[dict[str, Any]] = []

        for node in candidate_nodes:
            temp_id = str(node["node_id"])
            after = deepcopy(node)
            after["node_id"] = canonical_id[temp_id]
            if after.get("parent_id"):
                after["parent_id"] = canonical_id.get(
                    str(after["parent_id"]), str(after["parent_id"])
                )
            conflict_code = conflicts_by_candidate_id.get(temp_id)
            match = match_by_candidate_id.get(temp_id)

            if conflict_code:
                source = node.get("source_metadata") or {}
                changes.append(
                    {
                        "change_id": _change_id(job_id, "CONFLICT", temp_id),
                        "change_type": "CONFLICT",
                        "node_id": str(match["node_id"]) if match else after["node_id"],
                        "business_key": node.get("business_key"),
                        "before": deepcopy(match) if match else None,
                        "after": after,
                        "decision": "PENDING",
                        "issue_code": conflict_code,
                    }
                )
                issues.append(
                    {
                        "issue_id": _issue_id(job_id, conflict_code, temp_id),
                        "sheet_name": source.get("sheet"),
                        "row_number": source.get("source_row"),
                        "column_name": None,
                        "original_value": " / ".join(node.get("path") or []),
                        "issue_type": conflict_code,
                        "suggested_action": "人工确认节点身份后再 Apply",
                        "resolved": False,
                    }
                )
                continue

            if match is None:
                changes.append(
                    {
                        "change_id": _change_id(job_id, "ADD", after["node_id"]),
                        "change_type": "ADD",
                        "node_id": after["node_id"],
                        "business_key": node.get("business_key"),
                        "before": None,
                        "after": after,
                        "decision": "PENDING",
                        "issue_code": None,
                    }
                )
                continue

            before = deepcopy(match)
            same_name = _clean(before.get("name")) == _clean(after.get("name"))
            same_parent = _clean(before.get("parent_id")) == _clean(after.get("parent_id"))
            same_path = _path_key(before.get("path") or []) == _path_key(after.get("path") or [])
            same_key = _clean(before.get("business_key")) == _clean(after.get("business_key"))
            before_meta = (before.get("source_metadata") or {}).get("business_metadata") or {}
            after_meta = (after.get("source_metadata") or {}).get("business_metadata") or {}
            same_metadata = before_meta == after_meta
            same_active = bool(before.get("active", True)) == bool(after.get("active", True))

            if same_name and same_parent and same_path and same_key and same_metadata and same_active:
                change_type = "NO_CHANGE"
                decision = "CONFIRMED"
                issue_code = None
            elif not same_parent and same_name:
                change_type = "MOVE"
                decision = "PENDING"
                issue_code = None
            elif same_parent and not same_name:
                change_type = "RENAME"
                decision = "PENDING"
                issue_code = None
            elif not same_parent and not same_name:
                change_type = "UPDATE"
                decision = "PENDING"
                issue_code = "MOVE_AND_RENAME"
            else:
                change_type = "UPDATE"
                decision = "PENDING"
                issue_code = None

            changes.append(
                {
                    "change_id": _change_id(
                        job_id, change_type, str(before["node_id"])
                    ),
                    "change_type": change_type,
                    "node_id": str(before["node_id"]),
                    "business_key": node.get("business_key"),
                    "before": before,
                    "after": after,
                    "decision": decision,
                    "issue_code": issue_code,
                }
            )

        summary = dict(Counter(item["change_type"] for item in changes))
        for name in ["ADD", "UPDATE", "RENAME", "MOVE", "NO_CHANGE", "CONFLICT"]:
            summary.setdefault(name, 0)
        return {"changes": changes, "issues": issues, "summary": summary}


class HardwareTreeImportAnalyzer:
    """Orchestrates local Excel parse -> diff -> persisted review package."""

    def __init__(self, repository: HardwareTreeImportRepository):
        self.repository = repository
        self.parser = HardwareTreeExcelParser()
        self.diff_engine = HardwareTreeDiffEngine()

    def inspect(self, path: str | Path) -> dict[str, Any]:
        return self.parser.inspect_workbook(path)

    def preview_rows(
        self,
        path: str | Path,
        *,
        sheet_name: str,
        header_row: int,
        max_rows: int = 20,
    ) -> dict[str, Any]:
        return self.parser.preview_rows(
            path,
            sheet_name=sheet_name,
            header_row=header_row,
            max_rows=max_rows,
        )

    def analyze(
        self,
        job_id: str,
        path: str | Path,
        profile: dict[str, Any],
    ) -> dict[str, Any]:
        job = self.repository.get_job(job_id)
        if job["status"] != "UPLOADED":
            raise HardwareTreeImportContractError("IMPORT_NOT_UPLOADED")
        source = Path(path)
        HardwareTreeExcelParser._validate_file(source)
        if source.name != job["source_filename"]:
            raise HardwareTreeImportContractError("SOURCE_FILENAME_MISMATCH")
        if _sha256_file(source) != job["source_sha256"]:
            raise HardwareTreeImportContractError("SOURCE_HASH_MISMATCH")

        self.repository.clear_analysis(job_id)
        self.repository.set_mapping_profile(job_id, profile)

        try:
            parsed = self.parser.parse(
                source,
                tree_type=job["tree_type"],
                profile=profile,
                job_id=job_id,
            )
        except HardwareTreeImportContractError:
            self.repository.advance_status(job_id, "PARSE_FAILED")
            raise

        self.repository.advance_status(job_id, "PARSED")
        self.repository.advance_status(job_id, "VALIDATING")

        current = self.repository.case_repository.list_tree_nodes(job["tree_type"])
        diff = self.diff_engine.build(
            job_id=job_id,
            tree_type=job["tree_type"],
            candidate_nodes=parsed["nodes"],
            current_nodes=current,
        )
        all_issues = [*parsed["issues"], *diff["issues"]]
        unique_issues: dict[str, dict[str, Any]] = {
            str(item["issue_id"]): item for item in all_issues
        }
        for issue in unique_issues.values():
            self.repository.add_issue(job_id, issue)
        for change in diff["changes"]:
            self.repository.save_change(job_id, change)

        counts = {
            **parsed["stats"],
            "change_summary": diff["summary"],
            "validation_issue_count": len(unique_issues),
        }
        self.repository.set_counts(job_id, counts)
        self.repository.advance_status(job_id, "REVIEW_REQUIRED")
        return {
            "job": self.repository.get_job(job_id),
            "preview": {
                "tree_type": job["tree_type"],
                "nodes": parsed["nodes"],
                "stats": parsed["stats"],
            },
            "changes": self.repository.list_changes(job_id),
            "issues": self.repository.list_issues(job_id),
            "change_summary": diff["summary"],
        }
