"""Synthetic S1-S4 harness over the existing Hardware Case backend and parser.

Mock skill outputs are derived from source blocks, never from the expected answers.
The real provider path remains the existing HardwareCaseRuntimeStructurer.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from repositories.hardware_case_repository import HardwareCaseRepository
from services.hardware_case_ai_adapter import HardwareCaseAIAdapter
from services.hardware_case_backend import HardwareCaseBackendService


FIELD_HEADINGS = {
    "背景": "background", "问题现象": "symptom", "现象记录": "symptom",
    "分析过程": "analysis_process", "排查记录": "analysis_process",
    "根因": "root_cause", "原因定位": "root_cause",
    "解决措施": "actions", "处置方案": "actions", "结论": "conclusion",
}


def load_synthetic_tree(path: Path, tree_type: str, service: HardwareCaseBackendService) -> list[dict[str, Any]]:
    """Deterministic, read-only Excel loader; each row is an independent path."""
    if tree_type not in {"CIRCUIT_FEATURE", "MATERIAL_DEVICE"}:
        raise ValueError("TREE_TYPE_INVALID")
    source_hash = sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    nodes = []
    try:
        for sheet in workbook:
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not row[0]:
                    continue
                node_id = str(row[0]).strip()
                path_parts = [str(value).strip() for value in row[1:] if value is not None and str(value).strip()]
                if not node_id or not path_parts:
                    raise ValueError("TREE_ROW_INVALID")
                node = service.save_tree_node({
                    "node_id": node_id, "tree_type": tree_type, "name": path_parts[-1],
                    "path": path_parts, "source_ref": f"synthetic:{path.name}",
                    "source_sheet": sheet.title, "source_row": row_number,
                    "source_cells": list(row), "source_file_hash": source_hash,
                    "active": True,
                })
                nodes.append(node)
    finally:
        workbook.close()
    return nodes


@dataclass
class SyntheticSkills:
    """SK-HC-01..06 minimal local implementations for mock validation."""
    circuits: list[dict[str, Any]]
    materials: list[dict[str, Any]]

    def structure(self, document: dict[str, Any]) -> dict[str, Any]:
        # SK-HC-01: heading aliases identify blocks; never fill a missing field.
        facts: dict[str, dict[str, Any]] = {}
        for block in document["blocks"]:
            if block["block_type"] not in {"PARAGRAPH", "TABLE"} or not block.get("text"):
                continue
            heading = (block.get("section_path") or [""])[-1]
            field = FIELD_HEADINGS.get(heading)
            if field:
                previous = facts.get(field)
                if previous:
                    previous["value"] += "\n" + block["text"]
                    previous["evidence_block_ids"].append(block["block_id"])
                else:
                    facts[field] = {"value": block["text"], "evidence_block_ids": [block["block_id"]]}
        # SK-HC-02: direct block references; the adapter also checks support.
        return {"facts": facts, "product_context": {},
                "circuit_feature_links": self._link(document, self.circuits),
                "material_links": self._link(document, self.materials)}

    @staticmethod
    def _link(document: dict[str, Any], nodes: list[dict[str, Any]]) -> list[dict[str, Any]]:
        # SK-HC-03/04: independent exact leaf-term candidate matching.
        links = []
        for node in nodes:
            term = node["name"]
            pattern = re.compile(r"(?<![A-Za-z])" + re.escape(term) + r"(?![A-Za-z])" if term.isascii() else re.escape(term))
            matches = []
            for block in document["blocks"]:
                if block["block_type"] not in {"PARAGRAPH", "TABLE"}:
                    continue
                text = str(block.get("text") or "")
                spans = [match.span() for match in pattern.finditer(text)]
                # Prefer an explicitly named longer device over its generic suffix.
                longer = [other["name"] for other in nodes if other["name"] != term and term in other["name"]]
                covered = [match.span() for name in longer for match in re.finditer(re.escape(name), text)]
                if any(not any(start >= left and end <= right for left, right in covered) for start, end in spans):
                    matches.append(block["block_id"])
            if matches:
                links.append({"node_id": node["node_id"], "confidence": 0.8,
                              "evidence_block_ids": matches})
        return links

    @staticmethod
    def quality(case: dict[str, Any]) -> dict[str, Any]:
        # SK-HC-05: interface kept small for the four scenarios.
        missing = [field for field in ("symptom", "root_cause", "actions")
                   if not case.get("facts", {}).get(field, {}).get("candidate_value")]
        return {"quality_status": "NEEDS_REVIEW" if missing else "READY", "missing_fields": missing}

    @staticmethod
    def interpret_query(kind: str, value: str) -> dict[str, str]:
        # SK-HC-06: explicit intents only; no cross-tree inference.
        if kind not in {"circuit", "material", "symptom", "case_id"} or not value.strip():
            raise ValueError("QUERY_INVALID")
        return {"intent": kind, "query": value.strip()}


class ScenarioPoC:
    def __init__(self, database: Path, circuit_file: Path, material_file: Path):
        self.backend = HardwareCaseBackendService(HardwareCaseRepository(database))
        circuits = load_synthetic_tree(circuit_file, "CIRCUIT_FEATURE", self.backend)
        materials = load_synthetic_tree(material_file, "MATERIAL_DEVICE", self.backend)
        self.skills = SyntheticSkills(circuits, materials)
        self.adapter = HardwareCaseAIAdapter(self.backend, self.skills.structure)

    def ingest(self, word: Path) -> dict[str, Any]:
        result = self.adapter.ingest_docx(word)
        case = self.backend.get_case(result["case_id"], role="MAINTAINER")
        quality = self.skills.quality(case)
        if quality["missing_fields"] and result["status"] == "SUCCESS":
            result["status"] = "NEEDS_REVIEW"
        return {**result, "quality": quality}

    def query(self, kind: str, value: str) -> list[dict[str, Any]]:
        interpreted = self.skills.interpret_query(kind, value)
        value = interpreted["query"]
        if kind == "case_id":
            try:
                return [self.detail(value)]
            except KeyError:
                return []
        if kind in {"circuit", "material"}:
            tree_type = "CIRCUIT_FEATURE" if kind == "circuit" else "MATERIAL_DEVICE"
            node = self.backend.repository.get_tree_node(value)
            if not node or node["tree_type"] != tree_type:
                return []
            ids = {mapping["case_id"] for mapping in self.backend.repository.list_mappings()
                   if mapping["node_id"] == value and mapping["mapping_status"] == "SUGGESTED"}
        else:
            ids = {case["case_id"] for case in self.backend.repository.list_cases()
                   if value.lower() in str(case.get("facts", {}).get("symptom", {}).get("candidate_value") or "").lower()}
        return [self.detail(case_id) for case_id in sorted(ids)]

    def detail(self, case_id: str) -> dict[str, Any]:
        case = self.backend.get_case(case_id, role="MAINTAINER")
        if not case:
            raise KeyError(case_id)
        evidence = {item["evidence_id"]: item for item in self.backend.repository.list_evidence(case_id)}
        fields = ("background", "symptom", "analysis_process", "root_cause", "actions", "conclusion")
        return {"case_id": case_id, "title": case["title"],
                **{field: case.get("facts", {}).get(field, {}).get("candidate_value") for field in fields},
                "evidence_refs": {field: case.get("facts", {}).get(field, {}).get("evidence_refs", []) for field in fields},
                "evidence": evidence, "source_refs": case["source_refs"],
                "circuit_feature_links": [m for m in self.backend.repository.list_mappings(case_id=case_id) if m["tree_type"] == "CIRCUIT_FEATURE"],
                "material_links": [m for m in self.backend.repository.list_mappings(case_id=case_id) if m["tree_type"] == "MATERIAL_DEVICE"]}
