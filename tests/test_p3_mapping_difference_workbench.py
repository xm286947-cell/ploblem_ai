from fastapi.testclient import TestClient
from openpyxl import Workbook

from quality_knowledge.web.app import create_app


def _preview(client,path):
    with path.open('rb') as stream:
        return client.post('/import/preview',data={'business_type':'HMI'},files={'file':(path.name,stream)})


def test_preview_exposes_difference_workbench(tmp_path):
    x=tmp_path/'new.xlsx'; wb=Workbook(); ws=wb.active; ws.append(['TRC单号','故障现象新列']); ws.append(['H-1','黑屏']); wb.save(x)
    response=_preview(TestClient(create_app(tmp_path/'a.db')),x)
    assert response.status_code==200
    for text in ['字段差异处理','故障现象新列','作为已有字段 Alias','新增产品扩展字段','仅保留 Raw','生成 Mapping Draft']:
        assert text in response.text


def test_difference_actions_create_single_editable_draft(tmp_path):
    x=tmp_path/'new.xlsx'; wb=Workbook(); ws=wb.active; ws.append(['TRC单号','故障现象新列']); ws.append(['H-2','花屏']); wb.save(x)
    app=create_app(tmp_path/'b.db'); client=TestClient(app); preview=_preview(client,x)
    import re
    token=re.search(r'name="intake_session_id" value="([^"]+)"',preview.text).group(1)
    effective=app.state.mapping_configuration_service.get_effective_config('HMI')
    target=next(x for x in effective['mappings'] if x['canonical_field']=='description')
    response=client.post('/import/mapping-differences/apply',data={'intake_session_id':token,'difference_count':'1','source_0':'故障现象新列','action_0':'ALIAS','target_0':target['mapping_id'],'extension_key_0':''},follow_redirects=False)
    assert response.status_code==303
    drafts=[x for x in app.state.mapping_configuration_service.list_configs('HMI') if x['status']=='DRAFT']
    assert len(drafts)==1
    description=next(x for x in drafts[0]['mappings'] if x['canonical_field']=='description')
    assert '故障现象新列' in description['aliases']


def test_required_mapping_cannot_be_disabled(tmp_path):
    app=create_app(tmp_path/'c.db'); service=app.state.mapping_configuration_service
    active=service.get_effective_config('HMI'); required=next(x for x in active['mappings'] if x['required'])
    draft=service.apply_difference_actions(active['config_id'],[],[required['mapping_id']])
    assert next(x for x in draft['mappings'] if x['canonical_field']==required['canonical_field'])['enabled'] is True

def test_mapping_settings_has_visible_excel_comparison_entry(tmp_path):
    client=TestClient(create_app(tmp_path/'d.db')); page=client.get('/settings/mapping?business_type=HMI')
    assert page.status_code==200
    assert '打开 Excel 对比表头' in page.text and '读取表头并比较差异' in page.text
    assert 'action="/settings/mapping/preview"' in page.text

def test_mapping_settings_upload_opens_difference_workbench(tmp_path):
    x=tmp_path/'mapping.xlsx'; wb=Workbook(); ws=wb.active; ws.append(['TRC单号','现场新字段']); ws.append(['H-3','值']); wb.save(x)
    client=TestClient(create_app(tmp_path/'e.db'))
    with x.open('rb') as stream: page=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(x.name,stream)})
    assert page.status_code==200
    assert '字段差异处理' in page.text and '现场新字段' in page.text

def test_multi_sheet_preview_prefers_real_data_sheet_over_summary_sheet(tmp_path):
    x=tmp_path/'multi.xlsx'; wb=Workbook(); summary=wb.active; summary.title='汇总'; summary.append(['说明页'])
    data=wb.create_sheet('HMI数据'); data.append(['TRC单号','问题描述','产品系列']); data.append(['H-4','异常','产品']); wb.save(x)
    client=TestClient(create_app(tmp_path/'f.db'))
    with x.open('rb') as stream: page=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(x.name,stream)})
    assert page.status_code==200
    assert '选中 Sheet' in page.text and 'HMI数据' in page.text
    assert '数据行' in page.text and '>1<' in page.text
