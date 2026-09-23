from fastapi.testclient import TestClient

from quality_knowledge.materials import MaterialImportService, MaterialRepository, year_from_itr
from quality_knowledge.scenario_sources import operation_records
from quality_knowledge.web.app import create_app
from test_itr_cs_material_mvp import operations_workbook


def test_kpi_year_precedes_itr_and_import_override_wins(tmp_path):
    assert year_from_itr('ITR20250605084CS') == '2025'
    source=tmp_path/'ops.xlsx';operations_workbook(source,'ITR20250605084CS')
    repo=MaterialRepository(tmp_path/'default.db')
    MaterialImportService(repo).import_file(source,'SW-OPS',2)
    item=repo.search_materials('SW-OPS')['items'][0]
    assert item['year']=='2026' and item['year_source']=='KPI计入月份'
    repo.add_material(repo.group('SW-OPS'),'ITR20241200001CS',{'数据运营_KPI计入年份':'2026'},'synthetic.xlsx','Sheet',4)
    parsed=next(x for x in repo.search_materials('SW-OPS')['items'] if x['business_key'].startswith('ITR2024'))
    assert parsed['year']=='2026' and parsed['year_source']=='原始文件'

    other=MaterialRepository(tmp_path/'manual.db')
    MaterialImportService(other).import_file(source,'SW-OPS',2,reporting_year='2027')
    item=other.search_materials('SW-OPS')['items'][0]
    assert item['year']=='2027' and item['year_source']=='导入时人工设置'


def test_batch_year_updates_workbench_and_scenario_period_without_changing_raw(tmp_path):
    app=create_app(tmp_path/'app.db');client=TestClient(app);repo=app.state.material_repository
    source=tmp_path/'ops.xlsx';operations_workbook(source,'ITR20250605084CS')
    with source.open('rb') as stream:
        response=client.post('/materials/import',data={'workbench':'software-operations','group_code':'SW-OPS','header_rows':'2'},files={'file':('ops.xlsx',stream,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==200 and 'KPI计入月份' in response.text and '2026' in response.text
    item=repo.search_materials('SW-OPS')['items'][0];original=dict(item['raw'])
    changed=client.post('/materials/software-operations/batch-year',data={'material_ids':item['material_id'],'reporting_year':'2028'},follow_redirects=False)
    assert changed.status_code==303 and changed.headers['location'].endswith('year=2028')
    item=repo.material(item['material_id'])
    assert item['year']=='2028' and item['year_source']=='批量人工设置' and item['raw']==original
    rows=operation_records(app.state.scenario_generation_service,{'year':'2028'})
    assert len(rows)==1 and rows[0]['year']=='2028'
    assert not operation_records(app.state.scenario_generation_service,{'year':'2025'})
    page=client.get('/materials/software-operations?year=2028')
    assert page.status_code==200 and '批量设置考核年份' in page.text and '批量人工设置' in page.text

    # Reimporting without an explicit year must not erase a later manual correction.
    MaterialImportService(repo).import_file(source,'SW-OPS',2)
    assert repo.material(item['material_id'])['year']=='2028'

    # A changed export creates a new current version and inherits the manual year.
    from openpyxl import load_workbook
    book=load_workbook(source);book.active.cell(row=3,column=4,value='软件异常（已更新）');book.save(source);book.close()
    changed_stats=MaterialImportService(repo).import_file(source,'SW-OPS',2)
    current=repo.search_materials('SW-OPS')['items'][0]
    assert changed_stats['updated']==1 and current['material_id']!=item['material_id']
    assert current['year']=='2028' and current['year_source']=='批量人工设置'


def test_legacy_auto_itr_year_does_not_hide_explicit_kpi_year(tmp_path):
    app=create_app(tmp_path/'legacy.db');repo=app.state.material_repository
    key='ITR20250202001CS'
    material_id,_=repo.add_material(repo.group('SW-OPS'),key,{
        '问题信息_彻底解决单号':key,'问题信息_SPDT':'中小型PLC SPDT',
        '数据运营_KPI计入月份':'2026-02','问题信息_问题描述':'历史考核问题'},'legacy.xlsx','Sheet1',1)
    repo.set_reporting_year([material_id],'2025',source='AUTO_ITR')
    assert repo.material(material_id)['year']=='2026'
    assert repo.search_materials('SW-OPS',year='2026')['total']==1
    rows=operation_records(app.state.scenario_generation_service,{'spdt':'中小型PLC SPDT','year':'2026'})
    assert len(rows)==1 and rows[0]['year']=='2026'


def test_year_validation_and_group_isolation(tmp_path):
    repo=MaterialRepository(tmp_path/'app.db')
    source=tmp_path/'ops.xlsx';operations_workbook(source,'ITR20250605084CS')
    material_id=MaterialImportService(repo).import_file(source,'SW-OPS',2) and repo.list_materials('SW-OPS')[0]['material_id']
    for bad in ('26','1999','2100','abcd'):
        try:repo.set_reporting_year([material_id],bad)
        except ValueError:pass
        else:raise AssertionError('invalid year accepted')
    try:repo.set_reporting_year([], '2026')
    except ValueError as error:assert '至少选择' in str(error)
    else:raise AssertionError('empty selection accepted')
