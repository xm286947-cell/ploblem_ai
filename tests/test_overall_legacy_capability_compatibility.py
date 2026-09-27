from __future__ import annotations

import sqlite3
import time
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.services import KnowledgeIssueService
from quality_knowledge.services.v1_analysis_service import KnowledgeIssueAnalysisService
from quality_knowledge.web.app import create_app as create_legacy_app
from quality_knowledge.web.legacy_database_binding import validate_legacy_database
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]


def _p0_db(path: Path) -> Path:
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(path)
    return path


def _legacy_db(path: Path) -> Path:
    # The legacy factory is the existing explicit local/test bootstrap. The
    # unified P0 host uses the strict router factory and never initializes it.
    create_legacy_app(path)
    return path


def _host(p0_path: Path, legacy_path: Path | None) -> TestClient:
    app = create_p0_app(
        p0_path,
        stage_runner=object(),
        project_root=ROOT,
        legacy_quality_issue_db_path=legacy_path,
    )
    return TestClient(app)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITR单号", "问题描述", "产品", "月份"])
    sheet.append(["ITR-LEGACY-IMPORT-01", "合成问题：统一入口导入确认", "PLC", "2026-09"])
    workbook.save(path)


def test_legacy_binding_rejects_p0_path_collision(tmp_path):
    p0_path = _p0_db(tmp_path / "p0.db")
    with pytest.raises(ValueError, match="LEGACY_P0_DATABASE_PATH_COLLISION"):
        create_p0_app(
            p0_path,
            stage_runner=object(),
            project_root=ROOT,
            legacy_quality_issue_db_path=p0_path.parent / "." / p0_path.name,
        )


def test_legacy_schema_validation_is_read_only(tmp_path):
    legacy_path = _legacy_db(tmp_path / "legacy.db")
    before = legacy_path.read_bytes()
    path, error = validate_legacy_database(tmp_path / "p0.db", legacy_path)
    assert error is None
    assert path == legacy_path.resolve()
    assert legacy_path.read_bytes() == before


@pytest.mark.parametrize(
    ("legacy_name", "expected_code", "create_file"),
    [
        ("missing.db", "LEGACY_DB_MISSING", False),
        ("incompatible.db", "LEGACY_DB_SCHEMA_INCOMPATIBLE", True),
    ],
)
def test_legacy_binding_fails_closed_without_creating_or_migrating_db(
    tmp_path, legacy_name, expected_code, create_file
):
    p0_path = _p0_db(tmp_path / "p0.db")
    legacy_path = tmp_path / legacy_name
    if create_file:
        with sqlite3.connect(legacy_path) as connection:
            connection.execute("CREATE TABLE unrelated(value TEXT)")
    before = legacy_path.read_bytes() if create_file else None

    client = _host(p0_path, legacy_path)
    for path in ("/analysis", "/import", "/statistics"):
        response = client.get(path)
        assert response.status_code == 503
        assert response.json() == {"detail": expected_code}
    assert legacy_path.exists() is create_file
    if create_file:
        assert legacy_path.read_bytes() == before


def test_legacy_import_analysis_statistics_share_one_host_and_isolated_database(
    tmp_path, monkeypatch
):
    p0_path = _p0_db(tmp_path / "p0.db")
    legacy_path = _legacy_db(tmp_path / "legacy.db")
    original_legacy_bytes = legacy_path.read_bytes()
    workbook_path = tmp_path / "issues.xlsx"
    _workbook(workbook_path)

    analyzed: list[str] = []

    def deterministic_analysis(self, knowledge_id, *_args, **_kwargs):
        analyzed.append(knowledge_id)
        return {"knowledge_id": knowledge_id, "status": "COMPLETED"}

    def deterministic_batch(self, knowledge_ids, *_args, progress_callback=None, **_kwargs):
        items = []
        for knowledge_id in knowledge_ids:
            item = {"knowledge_id": knowledge_id, "status": "COMPLETED", "duration_ms": 1}
            items.append(item)
            if progress_callback:
                progress_callback(item)
        return {"completed": len(items), "failed": 0, "items": items}

    monkeypatch.setattr(KnowledgeIssueService, "run_issue_analysis", deterministic_analysis)
    monkeypatch.setattr(KnowledgeIssueService, "run_batch_analysis", deterministic_batch)
    monkeypatch.setattr(KnowledgeIssueAnalysisService, "run_issue_analysis", deterministic_analysis)
    client = _host(p0_path, legacy_path)

    assert isinstance(client.app, FastAPI)
    assert client.app.state.legacy_quality_issue_status["ready"] is True
    assert client.get("/").status_code == 200
    assert client.get("/p0/issues").status_code == 200

    analysis_page = client.get("/analysis")
    import_page = client.get("/import")
    statistics_page = client.get("/statistics?business_type=PLC&month=2026-09")
    assert analysis_page.status_code == import_page.status_code == statistics_page.status_code == 200
    assert "/analysis-batch" in analysis_page.text
    assert 'action="/import/preview"' in import_page.text
    assert 'name="month"' in statistics_page.text
    assert client.get("/static/app.css").status_code == 200

    with workbook_path.open("rb") as stream:
        preview = client.post(
            "/import/preview",
            data={"business_type": "PLC"},
            files={"file": (workbook_path.name, stream)},
        )
    assert preview.status_code == 200
    assert "确认正式导入" in preview.text
    import re

    intake_id = re.search(r'name="intake_session_id" value="([^"]+)"', preview.text)
    assert intake_id
    confirmed = client.post(
        "/import/confirm",
        data={"intake_session_id": intake_id.group(1)},
        follow_redirects=False,
    )
    assert confirmed.status_code == 303
    assert confirmed.headers["location"].startswith("/imports/")
    batch_page = client.get(confirmed.headers["location"])
    assert batch_page.status_code == 200
    batch_id = confirmed.headers["location"].rsplit("/", 1)[-1]
    batch_api = client.get(f"/api/imports/{batch_id}")
    assert batch_api.status_code == 200
    assert batch_api.json()["batch"]["new_rows"] == 1

    api_workbook_path = tmp_path / "api-issues.xlsx"
    api_workbook = Workbook()
    api_sheet = api_workbook.active
    api_sheet.append(["ITR单号", "问题描述", "产品", "月份"])
    api_sheet.append(["ITR-LEGACY-API-01", "合成问题：API 导入确认", "PLC", "2026-09"])
    api_workbook.save(api_workbook_path)
    with api_workbook_path.open("rb") as stream:
        api_preview = client.post(
            "/api/import/preview",
            data={"business_type": "PLC"},
            files={"file": (api_workbook_path.name, stream)},
        )
    assert api_preview.status_code == 200
    api_confirm = client.post(
        "/api/import/confirm",
        data={"intake_session_id": api_preview.json()["intake_session_id"]},
    )
    assert api_confirm.status_code == 200
    assert api_confirm.json()["batch_id"]

    direct_api_path = tmp_path / "direct-api.xlsx"
    direct_api_workbook = Workbook()
    direct_api_sheet = direct_api_workbook.active
    direct_api_sheet.append(["ITR单号", "问题描述", "产品", "月份"])
    direct_api_sheet.append(["ITR-LEGACY-DIRECT-API-01", "合成问题：直达导入", "PLC", "2026-09"])
    direct_api_workbook.save(direct_api_path)
    with direct_api_path.open("rb") as stream:
        direct_api_import = client.post(
            "/api/issues/import",
            data={"business_type": "PLC"},
            files={"file": (direct_api_path.name, stream)},
        )
    assert direct_api_import.status_code == 200
    assert direct_api_import.json()["batch_id"]

    with sqlite3.connect(legacy_path) as connection:
        issue = connection.execute(
            "SELECT knowledge_id FROM quality_issue WHERE business_issue_id=?",
            ("ITR-LEGACY-IMPORT-01",),
        ).fetchone()
    assert issue
    knowledge_id = issue[0]
    assert client.post(f"/analysis/{knowledge_id}", follow_redirects=False).status_code == 303
    assert analyzed == [knowledge_id]
    assert client.post(
        f"/api/issues/{knowledge_id}/analyze", json={"force": False}
    ).status_code == 200
    assert analyzed == [knowledge_id, knowledge_id]

    batch_started = client.post(
        "/analysis-batch", data={"business_type": "PLC"}, follow_redirects=False
    )
    assert batch_started.status_code == 303
    batch_job_id = parse_qs(urlparse(batch_started.headers["location"]).query)["job_id"][0]
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        batch_job = client.get(f"/api/analysis-batch-jobs/{batch_job_id}")
        assert batch_job.status_code == 200
        if batch_job.json()["status"] not in {"QUEUED", "RUNNING"}:
            break
        time.sleep(0.02)
    assert batch_job.json()["status"] == "COMPLETED"
    assert batch_job.json()["completed"] == 3

    assert client.get("/api/statistics?business_type=PLC").status_code == 200
    assert client.get("/api/common-capability-gaps?business_type=PLC").status_code == 200
    assert client.get(f"/api/capability-gaps?knowledge_id={knowledge_id}").status_code == 200
    filtered_statistics = client.get("/statistics?business_type=PLC&month=2026-09")
    assert filtered_statistics.status_code == 200
    assert '<option value="2026-09" selected>' in filtered_statistics.text

    with sqlite3.connect(p0_path) as connection:
        p0_issue = connection.execute(
            "SELECT COUNT(*) FROM quality_issue WHERE knowledge_id=?", (knowledge_id,)
        ).fetchone()[0]
        p0_tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            )
        }
    assert p0_issue == 0
    assert "import_batch_v1" not in p0_tables
    assert client.app.state.p0_repository.db_path.resolve() == p0_path.resolve()
    assert Path(client.app.state.legacy_quality_issue_status["database_path"]).resolve() == legacy_path.resolve()
    # Router/service composition is read-only at startup; only the exercised
    # import, analysis and statistics operations may mutate the Legacy store.
    assert legacy_path.read_bytes() != original_legacy_bytes
