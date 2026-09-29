from __future__ import annotations

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web.p0_app import create_p0_app
from retriever.case_retriever import QueryInput


ROOT = Path(__file__).resolve().parents[1]
META_SHEET = "_MAJOR_IMPORT_META"


def _provider(_provider_input, pending_specs, _context):
    contents = {
        "ISSUE_FACT": "控制器在掉电恢复后启动失败",
        "ROOT_CAUSE": "掉电窗口内配置写入未完成",
        "ACTION": "增加原子写入和启动恢复校验",
        "VERIFICATION": "完成 100 次掉电恢复验证",
    }
    return [
        {"object_id": item["object_id"], "data": {"content": contents[item["unit_id"]]}}
        for item in pending_specs
    ]


def _client(tmp_path: Path) -> TestClient:
    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    app = create_p0_app(
        p0_db,
        stage_runner=object(),
        project_root=ROOT,
        major_case_db_path=tmp_path / "major.db",
        major_attachment_root=tmp_path / "major-attachments",
        major_artifact_root=tmp_path / "artifacts",
        major_provider=_provider,
    )
    return TestClient(app)


def _xlsx(*, blocked_second_row: bool = False) -> bytes:
    output = BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "重大问题"
    sheet.append([
        "IGR编号",
        "ITR单号",
        "问题描述",
        "TRC发生",
        "MRC发生",
        "产品",
        "模块",
        "Failure Mechanism",
        "Trigger Condition",
    ])
    sheet.append([
        "IGR-R2-W2-001",
        "ITR20269911",
        "控制器掉电恢复启动失败",
        "边界保护不足",
        "评审检查项缺失",
        "PLC-X",
        "Motion",
        "非原子状态更新",
        "掉电窗口",
    ])
    if blocked_second_row:
        sheet.append(["", "", "", "", "", "PLC-X", "", "", ""])
    book.save(output)
    return output.getvalue()


def _preview(client: TestClient, content: bytes):
    return client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY", "actor": "w2-tester"},
        files={
            "file": (
                "major.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_official_template_carries_version_and_mapping_governance(tmp_path: Path):
    client = _client(tmp_path)
    response = client.get("/api/v2/major-production/excel/template")
    assert response.status_code == 200
    assert response.headers["x-major-template-version"] == "1.0"
    assert response.headers["x-major-mapping-version"].startswith("sha256:")

    book = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert META_SHEET in book.sheetnames
    metadata = {
        str(row[0] or ""): str(row[1] or "")
        for row in book[META_SHEET].iter_rows(min_row=1, max_col=2, values_only=True)
        if row[0]
    }
    book.close()
    assert metadata["template_contract"] == "major-excel-template/v1"
    assert metadata["template_version"] == "1.0"
    assert metadata["mapping_contract"] == "major-excel-field-mapping/v1"
    assert metadata["mapping_version"].startswith("sha256:")

    preview = _preview(client, response.content)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["template_version"] == "1.0"
    assert body["template_status"] == "OFFICIAL"
    assert body["mapping"]["version"] == metadata["mapping_version"]
    batch = client.get(
        f"/api/v2/major-production/excel/batches/{body['batch_id']}"
    ).json()
    assert batch["governance"]["contract_version"] == "major-excel-import-governance/v1"
    assert batch["governance"]["preview_actor"] == "w2-tester"


def test_mapping_change_after_preview_fails_closed_without_import(tmp_path: Path):
    client = _client(tmp_path)
    preview = _preview(client, _xlsx()).json()
    restore = client.app.state.major_case_restore_service
    restore.excel_parser.field_mapping["itr_id"] = "CHANGED_AFTER_PREVIEW"

    response = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": preview["batch_id"], "actor": "w2-confirmer"},
    )
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_MAPPING_CHANGED_AFTER_PREVIEW"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0
    batch = client.get(
        f"/api/v2/major-production/excel/batches/{preview['batch_id']}"
    ).json()
    assert batch["status"] == "FAILED"
    assert batch["runs"][-1]["status"] == "FAILED"


def test_batch_with_rejected_row_is_atomic_and_creates_no_partial_case(tmp_path: Path):
    client = _client(tmp_path)
    preview_response = _preview(client, _xlsx(blocked_second_row=True))
    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert preview["total"] == 2
    assert preview["importable"] == 1

    response = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": preview["batch_id"], "actor": "w2-confirmer"},
    )
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_BATCH_PRECHECK_FAILED"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0

    batch = client.get(
        f"/api/v2/major-production/excel/batches/{preview['batch_id']}"
    ).json()
    assert batch["status"] == "FAILED"
    assert batch["result"]["imported"] == 0
    assert batch["result"]["rejected"] >= 1
    assert batch["runs"][-1]["rejected_count"] >= 1


def test_mid_commit_failure_rolls_back_database_and_attachments(tmp_path: Path):
    client = _client(tmp_path)
    preview = _preview(client, _xlsx()).json()
    restore = client.app.state.major_case_restore_service
    original = restore.repository.add_source_link

    def fail_source_link(*args, **kwargs):
        raise RuntimeError("SYNTHETIC_W2_COMMIT_FAILURE")

    restore.repository.add_source_link = fail_source_link
    try:
        response = client.post(
            "/api/v2/major-production/excel/confirm",
            data={"batch_id": preview["batch_id"], "actor": "w2-confirmer"},
        )
    finally:
        restore.repository.add_source_link = original

    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_ATOMIC_COMMIT_FAILED"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0
    assert list((tmp_path / "major-attachments").rglob("*")) == []

    batch = client.get(
        f"/api/v2/major-production/excel/batches/{preview['batch_id']}"
    ).json()
    assert batch["status"] == "FAILED"
    assert batch["result"]["atomic_rollback"] is True
    assert batch["result"]["imported"] == 0
    assert batch["runs"][-1]["failed_count"] == 1


def test_single_source_and_excel_batch_converge_on_source_fact_then_publish_repeat(
    tmp_path: Path,
):
    client = _client(tmp_path)

    source = ROOT / "tests/golden/hardware_case_scenarios/A9001-LDO 输出振荡.docx"
    single = client.post(
        "/api/v2/major-production/sources",
        data={
            "title": "单事件 Source Fact 收敛",
            "group_code": "SINGLE",
            "domain": "PLC",
            "standard_itr": "ITR-R2-W2-SINGLE-1",
        },
        files={
            "file": (
                source.name,
                source.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert single.status_code == 201, single.text
    assert single.json()["source_fact"]["source_type"] == "DOCUMENT"

    preview = _preview(client, _xlsx()).json()
    confirmed = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": preview["batch_id"], "actor": "w2-confirmer"},
    )
    assert confirmed.status_code == 200, confirmed.text
    result = confirmed.json()["result"]
    assert result["imported"] == 1
    assert result["failed"] == 0
    assert result["rejected"] == 0
    assert result["run_id"].startswith("MIR-")
    case_id = result["case_ids"][0]

    detail = client.get(f"/api/v2/major-production/cases/{case_id}").json()
    assert len(detail["events"]) == 1
    event_id = detail["events"][0]["event_id"]
    assert any(
        item["source_type"] == "MAJOR_EXCEL_SOURCE_FACT"
        for item in detail["source_links"]
    )

    analysis = client.post(
        f"/api/v2/major-production/cases/{case_id}/analysis",
        params={"event_id": event_id},
    )
    assert analysis.status_code == 200, analysis.text
    candidates = analysis.json()["candidates"]
    assert {item["entry_type"] for item in candidates} == {
        "ISSUE_FACT", "ROOT_CAUSE", "ACTION", "VERIFICATION",
    }

    for candidate in candidates:
        reviewed = client.post(
            f"/api/v2/major-production/entries/{candidate['entry_id']}/confirm",
            json={
                "reviewer": "w2-reviewer",
                "content": candidate["content"],
                "reason": "W2 governed human review",
            },
        )
        assert reviewed.status_code == 200, reviewed.text
        assert reviewed.json()["status"] == "CONFIRMED"

    published = client.post(f"/api/v2/major-production/events/{event_id}/publish")
    assert published.status_code == 200, published.text
    historical_case_id = published.json()["case_id"]

    historical = client.get(f"/api/v2/historical-cases/{historical_case_id}")
    assert historical.status_code == 200
    assert historical.json()["root_cause"] == "掉电窗口内配置写入未完成"

    repeat = client.app.state.historical_case_service.search_repeat_cases(
        QueryInput(text="控制器 掉电恢复"),
        top_k=5,
    )
    assert any(item["case_id"] == historical_case_id for item in repeat["candidates"])

    with client.app.state.major_case_repository.connect() as connection:
        excel_fact_count = connection.execute(
            """SELECT COUNT(*) FROM kb_source_fact_revision
               WHERE case_id=? AND source_type='EXCEL'""",
            (case_id,),
        ).fetchone()[0]
    assert excel_fact_count == 1
