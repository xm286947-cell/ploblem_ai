from fastapi.testclient import TestClient
from openpyxl import Workbook
from zipfile import ZIP_DEFLATED,ZipFile

from quality_knowledge.web.app import create_app


def _book(path,headers):
    wb=Workbook(); ws=wb.active; ws.append(headers); ws.append(['A']*len(headers)); wb.save(path)


def test_explicit_business_type_allows_new_headers_into_safe_preview(tmp_path):
    excel=tmp_path/'new_hmi.xlsx'; _book(excel,['全新编号列','全新问题列','全新分类列'])
    client=TestClient(create_app(tmp_path/'a.db'))
    with excel.open('rb') as stream:
        response=client.post('/import/preview',data={'business_type':'HMI'},files={'file':(excel.name,stream)})
    assert response.status_code==200
    assert '全新编号列' in response.text and '未匹配' in response.text
    assert '必填缺失' in response.text


def test_auto_detection_accepts_four_known_non_id_fields(tmp_path):
    excel=tmp_path/'no_id_plc.xlsx'; _book(excel,['问题描述','问题原因定位','问题解决方案','管理改进措施'])
    client=TestClient(create_app(tmp_path/'b.db'))
    with excel.open('rb') as stream:
        response=client.post('/import/preview',data={'business_type':''},files={'file':(excel.name,stream)})
    assert response.status_code==200
    assert 'PLC' in response.text and '必填缺失' in response.text

def test_preview_uses_detected_header_below_title_rows(tmp_path):
    excel=tmp_path/'title_then_header.xlsx'
    wb=Workbook(); ws=wb.active; ws.append(['质量问题清单']); ws.append(['说明']); ws.append(['TRC单号','问题描述']); ws.append(['HMI-1','显示异常']); wb.save(excel)
    client=TestClient(create_app(tmp_path/'c.db'))
    with excel.open('rb') as stream:
        response=client.post('/import/preview',data={'business_type':'HMI'},files={'file':(excel.name,stream)})
    assert response.status_code==200
    assert 'TRC单号' in response.text and '质量问题清单' not in response.text

def test_explicit_detection_allows_duplicate_new_headers(tmp_path):
    excel=tmp_path/'duplicate_headers.xlsx'
    wb=Workbook(); ws=wb.active; ws.append(['质量问题清单']); ws.append(['新字段','新字段','另一字段']); ws.append(['A','B','C']); wb.save(excel)
    client=TestClient(create_app(tmp_path/'d.db'))
    with excel.open('rb') as stream:
        response=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(excel.name,stream)})
    assert response.status_code==200
    assert '字段差异处理' in response.text and '另一字段' in response.text

def test_explicit_detection_accepts_single_column_sheet(tmp_path):
    excel=tmp_path/'single.xlsx'; wb=Workbook(); ws=wb.active; ws.append(['新问题编号']); ws.append(['A-1']); wb.save(excel)
    client=TestClient(create_app(tmp_path/'e.db'))
    with excel.open('rb') as stream:
        response=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(excel.name,stream)})
    assert response.status_code==200 and '新问题编号' in response.text

def test_wide_real_header_is_not_collapsed(tmp_path):
    excel=tmp_path/'wide.xlsx'; headers=['TRC单号','问题描述']+[f'新字段{i}' for i in range(49)]
    wb=Workbook(); ws=wb.active; ws.append(headers); ws.append(['H-51','异常']+['值']*49); wb.save(excel)
    client=TestClient(create_app(tmp_path/'f.db'))
    with excel.open('rb') as stream:
        response=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(excel.name,stream)})
    assert response.status_code==200
    assert '源字段</span><strong>51</strong>' in response.text
    assert '数据行</span><b>1</b>' in response.text
    assert '诊断 ID' in response.text and 'data_intake.log' in response.text

def test_broken_a1_dimension_metadata_is_recalculated(tmp_path):
    good=tmp_path/'good.xlsx'; broken=tmp_path/'broken.xlsx'; headers=['TRC单号','问题描述']+[f'字段{i}' for i in range(49)]
    wb=Workbook(); ws=wb.active; ws.append(headers); ws.append(['H-1','异常']+['值']*49); wb.save(good)
    with ZipFile(good,'r') as src,ZipFile(broken,'w',ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data=src.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml': data=data.replace(b'<dimension ref="A1:AY2"/>',b'<dimension ref="A1:A1"/>')
            dst.writestr(item,data)
    client=TestClient(create_app(tmp_path/'g.db'))
    with broken.open('rb') as stream:
        response=client.post('/settings/mapping/preview',data={'business_type':'HMI'},files={'file':(broken.name,stream)})
    assert response.status_code==200
    assert '源字段</span><strong>51</strong>' in response.text
    assert '数据行</span><b>1</b>' in response.text
