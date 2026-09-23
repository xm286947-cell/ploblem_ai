from pathlib import Path
from openpyxl import load_workbook
from quality_knowledge.repositories import SqliteIssueKnowledgeRepository
from quality_knowledge.models.issue import *
from quality_knowledge.services import IssueQueryService, QualityKnowledgeExportService


def seed_issue(repo, kid, business, issue_id, occ, esc, product='P', platform='X'):
    q=QualityIssueDTO(
        identity=IssueIdentity(knowledge_id=kid,case_id=kid,business_type=business,issue_id=issue_id,source_record_id=kid),
        source=IssueSource(source_hash='hash-'+kid),
        issue_fact=IssueFact(title='t-'+kid,description='d-'+kid,product=product,platform=platform,severity='S2',month='2026-08'),
        product_context=ProductContext(product=product,platform=platform,module='M'),
        occurrence=OccurrenceFact(cause_l1=occ,cause_l2='L2'),
        escape=EscapeFact(is_escape='是',escape_l1=esc,escape_l2='E2'),
        solution=SolutionFact(improvement_action='action'),raw_record={'raw':'keep'}
    )
    repo.save(q)


def seed_gap(repo, kid, business_gap='MANAGEMENT', category='CHANGE_MANAGEMENT'):
    repo.ensure_m2()
    run='run-'+kid
    repo.start_analysis_run({'analysis_run_id':run,'knowledge_id':kid,'analysis_type':'capability_gap','status':'RUNNING','model_provider':'fake','model_name':'fake','prompt_name':'p','prompt_version':'1','schema_version':'1','engine_version':'M2','input_hash':'x'})
    repo.save_capability_gaps(kid,run,[{'gap_dimension':business_gap,'gap_category':category,'gap_description':'gap','recommended_control':'control','scope':'ALL','confidence':0.8,'evidence_refs':[]}],'fake','1')
    repo.finish_analysis_run(run,'SUCCESS','fake')


def test_combined_query_and_aggregations(tmp_path):
    repo=SqliteIssueKnowledgeRepository(tmp_path/'q.db')
    seed_issue(repo,'h1','HMI','H1','代码实现','场景覆盖')
    seed_issue(repo,'p1','PLC','P1','代码实现','场景覆盖')
    seed_issue(repo,'i1','IFA','I1','设计','用例缺失')
    seed_gap(repo,'h1'); seed_gap(repo,'p1')
    svc=IssueQueryService(repo)
    rows=svc.query_issues(gap_dimension='MANAGEMENT',gap_category='CHANGE_MANAGEMENT',limit=20)
    assert {r['business_type'] for r in rows}=={'HMI','PLC'}
    assert svc.aggregate_by_cause()[0]=={'value':'代码实现','count':2}
    assert svc.aggregate_by_escape()[0]=={'value':'场景覆盖','count':2}
    gaps=svc.aggregate_by_capability_gap()
    assert gaps[0]['gap_category']=='CHANGE_MANAGEMENT' and gaps[0]['business_count']==2
    stats=svc.statistics()
    assert stats['cross_business_capability_gaps'][0]['business_count']==2


def test_csv_and_xlsx_business_exports(tmp_path):
    repo=SqliteIssueKnowledgeRepository(tmp_path/'q.db')
    seed_issue(repo,'h1','HMI','H1','代码实现','场景覆盖')
    seed_gap(repo,'h1','TECHNICAL','TEST_CAPABILITY')
    svc=QualityKnowledgeExportService(repo)
    csv_path=tmp_path/'issues.csv'
    out=svc.export_csv(csv_path,filters={'business_type':'HMI'},dataset='issues')
    assert out['rows']==1 and csv_path.exists()
    assert 'knowledge_id' in csv_path.read_text(encoding='utf-8-sig').splitlines()[0]

    xlsx_path=tmp_path/'knowledge.xlsx'
    out=svc.export_xlsx(xlsx_path,filters={'business_type':'HMI'})
    assert out['issues']==1 and out['capability_gaps']==1
    wb=load_workbook(xlsx_path,read_only=True)
    assert wb.sheetnames==['Issue_Knowledge','Capability_Gaps','Analysis_Runs','Statistics']
    assert wb['Issue_Knowledge'].max_row==2
    assert wb['Capability_Gaps'].max_row==2
