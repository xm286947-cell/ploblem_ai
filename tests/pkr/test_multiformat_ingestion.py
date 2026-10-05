from __future__ import annotations

import hashlib
import io

from fastapi.testclient import TestClient
from docx import Document

from public_knowledge_rag import app as service
from public_knowledge_rag.store import Store


def setup_client(monkeypatch, tmp_path):
    store = Store(tmp_path)
    monkeypatch.setattr(service, "store", store)
    monkeypatch.setattr(service, "retriever", service.SQLiteLexicalRetriever(store))
    return TestClient(service.app), store


def minimal_pdf(text: str = "Public datasheet endurance 3000 cycles") -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(output)


def upload(client, filename, content, media_type, title=None, classification="PUBLIC", source_uri=None):
    data = {"title": title or filename, "classification": classification}
    if source_uri:
        data["source_uri"] = source_uri
    return client.post("/sources/import-file", data=data,
                       files={"file": (filename, content, media_type)})


def test_pdf_upload_keeps_byte_exact_snapshot_and_page_citation(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    payload = minimal_pdf()
    response = upload(client, "规格书.pdf", payload, "application/pdf", source_uri="https://vendor.example/spec.pdf")
    assert response.status_code == 200, response.text
    imported = response.json()
    assert imported["source_sha256"] == hashlib.sha256(payload).hexdigest()
    assert imported["locator_status"] == "READY"
    hits = client.post("/search", json={"query": "endurance", "top_k": 5}).json()["hits"]
    assert hits and '"page":1' in hits[0]["locator"]
    citation = client.get(f"/citations/{hits[0]['hit_id']}").json()
    assert citation["source_sha256"] == imported["source_sha256"]
    snapshot = client.get(imported["snapshot_url"])
    assert snapshot.content == payload
    assert snapshot.headers["x-source-snapshot"] == "immutable"
    assert snapshot.headers["x-source-sha256"] == imported["source_sha256"]


def test_docx_heading_paragraph_and_table_row_locators(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    doc = Document()
    doc.add_heading("Endurance", level=1)
    doc.add_paragraph("Program endurance is 3000 cycles.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Condition", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "25 C", "3000 cycles"
    stream = io.BytesIO()
    doc.save(stream)
    response = upload(client, "spec.docx", stream.getvalue(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["locator_status"] == "READY"
    source = client.get("/sources/" + data["source_id"]).json()
    assert source["revisions"][0]["element_counts"]["table_row"] == 1
    hits = client.post("/search", json={"query": "Condition", "top_k": 5}).json()["hits"]
    assert hits and '"type":"docx_table_row"' in hits[0]["locator"]


def test_html_snapshot_hash_anchor_and_table_structure(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    payload = b'<html><head><title>Endurance</title></head><body><h1 id="pe">P/E</h1><p>3000 cycles</p><table><tr><th>Condition</th><th>Value</th></tr><tr><td>25 C</td><td>3000</td></tr></table></body></html>'
    response = upload(client, "spec.html", payload, "text/html")
    assert response.status_code == 200, response.text
    imported = response.json()
    hits = client.post("/search", json={"query": "3000", "top_k": 10}).json()["hits"]
    assert hits
    citation = client.get("/citations/" + hits[0]["hit_id"]).json()
    locator = __import__("json").loads(citation["locator"])
    assert locator["snapshot_sha256"] == hashlib.sha256(payload).hexdigest()
    assert locator["anchor"] == "pe"
    assert client.get(imported["snapshot_url"]).content == payload


def test_markdown_and_text_have_deterministic_line_locators_and_raw_dedup(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    payload = b"# Endurance\n\nP/E minimum: 3000 cycles.\n"
    first = upload(client, "spec.md", payload, "text/markdown").json()
    repeated = upload(client, "spec.md", payload, "text/markdown").json()
    revised = upload(client, "spec.md", payload + b"Updated.\n", "text/markdown").json()
    assert first["created"] is True and repeated["created"] is False
    assert revised["created"] is True and revised["source_revision"] != first["source_revision"]
    hit = client.post("/search", json={"query": "minimum", "top_k": 5}).json()["hits"][0]
    assert '"line":3' in hit["locator"]
    assert upload(client, "note.txt", b"retention is 10 years", "text/plain").status_code == 200


def test_classification_gate_runs_before_file_parser(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    monkeypatch.setattr(service.file_parser, "parse", lambda *args: (_ for _ in ()).throw(AssertionError("parser invoked")))
    response = upload(client, "private.pdf", minimal_pdf(), "application/pdf", classification="INTERNAL")
    assert response.status_code == 422


def test_bad_format_encrypted_or_empty_document_fails_closed(monkeypatch, tmp_path):
    client, _ = setup_client(monkeypatch, tmp_path)
    assert upload(client, "file.pdf", b"not a pdf", "application/pdf").status_code == 422
    assert upload(client, "file.docx", b"", "application/vnd.openxmlformats-officedocument.wordprocessingml.document").status_code == 422
    assert upload(client, "file.exe", b"content", "application/octet-stream").status_code == 422


def test_incomplete_locator_coverage_is_stored_as_partial_without_indexing(monkeypatch, tmp_path):
    from public_knowledge_rag.contracts import DocumentElement, ParsedDocument
    client, _ = setup_client(monkeypatch, tmp_path)
    monkeypatch.setattr(service.file_parser, "parse", lambda *args: ParsedDocument(
        "partial evidence", ("page:1",), "test-parser", "1",
        (DocumentElement("partial evidence", "page:1"),), False, {"text": 1}))
    result = upload(client, "partial.pdf", minimal_pdf(), "application/pdf").json()
    assert result["locator_status"] == "PARTIAL_NOT_EVIDENCE_READY"
    assert result["chunk_count"] == 0
    assert client.post("/search", json={"query": "partial evidence"}).json()["hits"] == []
    assert client.get(result["snapshot_url"]).status_code == 200


def test_real_public_datasheet_golden_if_present(monkeypatch, tmp_path):
    # Local Golden smoke uses the already staged, public Kioxia PDF without copying it into Git.
    from pathlib import Path
    path = Path("/Users/xiamin/dev/存储/public-knowledge-rag-391-run/deliverables/package-staging/sources/byte-snapshots/STDL-EXT-001_kioxia_endurance.pdf")
    if not path.is_file():
        return
    client, _ = setup_client(monkeypatch, tmp_path)
    content = path.read_bytes()
    response = upload(client, path.name, content, "application/pdf")
    assert response.status_code == 200, response.text
    assert response.json()["source_sha256"] == hashlib.sha256(content).hexdigest()
    assert response.json()["element_counts"].get("text", 0) > 0
    assert response.json()["element_counts"].get("table_row", 0) > 0
    assert client.get(response.json()["snapshot_url"]).content == content
