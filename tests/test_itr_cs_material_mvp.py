import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from fastapi.testclient import TestClient

from quality_knowledge.materials import MaterialImportService, MaterialRepository, combine_headers, normalize_itr
from quality_knowledge.web.app import create_app


def workbook(path, material_type, key):
    book=Workbook();sheet=book.active
    sheet.append(["问题信息", "问题信息", "问题信息"])
    sheet.append(["彻底解决单号" if material_type=="ITR_CS" else "ITR单号", "产品类型", "问题描述"])
    sheet.append([key, "PLC", "测试问题"]);book.save(path)


def operations_workbook(path, key):
    book=Workbook();sheet=book.active
    sheet.append(["数据运营","数据运营","问题信息","问题信息","技术根因分析与纠正"])
    sheet.append(["审核状态","KPI计入月份","彻底解决单号","问题描述","软件模块"])
    sheet.append(["已审核","2026-06",key,"软件异常","通信模块"]);book.save(path)


def test_two_row_headers_and_itr_normalization():
    assert combine_headers(("问题信息", None), ("ITR单号", "产品类型")) == ["问题信息_ITR单号", "问题信息_产品类型"]
    assert normalize_itr(" itr20260605084cs ") == "ITR20260605084"


def test_material_import_always_closes_excel_workbook(tmp_path, monkeypatch):
    import quality_knowledge.materials as materials_module

    source=tmp_path/"itr.xlsx";workbook(source,"ITR_SOURCE","ITR20260605084")
    real_loader=materials_module.load_workbook
    closed=[]

    def tracked_loader(*args,**kwargs):
        book=real_loader(*args,**kwargs)
        real_close=book.close
        def tracked_close():
            closed.append(True)
            real_close()
        book.close=tracked_close
        return book

    monkeypatch.setattr(materials_module,"load_workbook",tracked_loader)
    repository=MaterialRepository(tmp_path/"app.db")
    with repository.connect() as connection:
        connection.execute("CREATE TABLE quality_issue(knowledge_id TEXT PRIMARY KEY,business_issue_id TEXT)")
    MaterialImportService(repository).import_file(source,"ITR",2)
    assert closed == [True]
    Path(source).unlink()


def test_group_isolation_versioning_and_linking(tmp_path):
    db=tmp_path/"app.db";repo=MaterialRepository(db)
    with repo.connect() as c:
        c.execute("CREATE TABLE quality_issue(knowledge_id TEXT PRIMARY KEY,business_issue_id TEXT)")
        c.execute("INSERT INTO quality_issue VALUES('K-1','ITR20260605084')")
    itr=tmp_path/"itr.xlsx";cs=tmp_path/"cs.xlsx"
    workbook(itr,"ITR_SOURCE","ITR20260605084");workbook(cs,"ITR_CS","ITR20260605084CS")
    service=MaterialImportService(repo)
    assert service.import_file(itr,"ITR",2)["new"] == 1
    assert service.import_file(cs,"ITR-CS",2)["new"] == 1
    assert service.import_file(cs,"ITR-CS",2)["skipped"] == 1
    assert len(repo.list_materials("ITR")) == 1
    assert len(repo.list_materials("ITR-CS")) == 1
    linked=repo.materials_for_issue("K-1")
    assert {x["material_type"] for x in linked} == {"ITR_SOURCE","ITR_CS"}
    assert all(x["canonical_itr"]=="ITR20260605084" for x in linked)
    assert json.dumps(linked,ensure_ascii=False).find("测试问题") >= 0


def test_incremental_import_lists_latest_and_cleans_unreferenced_history(tmp_path):
    repo=MaterialRepository(tmp_path/'incremental.db');group=repo.group('ITR-CS')
    key='ITR20260605084CS'
    first,_=repo.add_material(group,key,{'问题信息_彻底解决单号':key,'问题信息_问题描述':'旧内容'},'a.xlsx','Sheet1',3)
    same,action=repo.add_material(group,key.lower(),{'问题信息_彻底解决单号':key,'问题信息_问题描述':'旧内容'},'b.xlsx','Sheet1',3)
    assert action=='SKIPPED' and same==first
    latest,_=repo.add_material(group,key,{'问题信息_彻底解决单号':key,'问题信息_问题描述':'新内容'},'c.xlsx','Sheet1',3)
    listing=repo.search_materials('ITR-CS')
    assert listing['total']==1 and listing['items'][0]['material_id']==latest and listing['items'][0]['version_no']==2
    summary=repo.duplicate_summary('ITR-CS')
    assert summary['history_count']==1 and summary['removable_count']==1
    cleaned=repo.cleanup_duplicates('ITR-CS')
    assert cleaned=={'deleted':1,'protected':0}
    assert len(repo.list_materials('ITR-CS',include_history=True))==1


def test_duplicate_cleanup_preserves_reviewed_old_version(tmp_path):
    repo=MaterialRepository(tmp_path/'protected.db');group=repo.group('ITR')
    key='ITR20260605084'
    old,_=repo.add_material(group,key,{'问题信息_ITR单号':key,'问题信息_问题描述':'旧内容'},'a.xlsx','Sheet1',3)
    repo.save_review(old,review_status='COMPLETED',analysis_summary='人工结论',root_cause='',improvement_action='',reviewer='质量组')
    repo.add_material(group,key,{'问题信息_ITR单号':key,'问题信息_问题描述':'新内容'},'b.xlsx','Sheet1',3)
    assert repo.duplicate_summary('ITR')['protected_count']==1
    assert repo.cleanup_duplicates('ITR')=={'deleted':0,'protected':1}
    assert len(repo.list_materials('ITR',include_history=True))==2


def test_cs_suffixed_analysis_issue_links_to_resolution_record(tmp_path):
    db=tmp_path/'cs-issue.db';repo=MaterialRepository(db)
    with repo.connect() as c:
        c.execute("CREATE TABLE quality_issue(knowledge_id TEXT PRIMARY KEY,business_issue_id TEXT)")
        c.execute("INSERT INTO quality_issue VALUES('K-CS','ITR20260605084CS')")
    source=tmp_path/'cs.xlsx';workbook(source,'ITR_CS','ITR20260605084CS')
    MaterialImportService(repo).import_file(source,'ITR-CS',2)
    assert [x['material_type'] for x in repo.materials_for_issue('K-CS')]==['ITR_CS']
    assert repo.link_preview()['matched']==1


def test_default_groups_do_not_enter_ai_or_insight(tmp_path):
    repo=MaterialRepository(tmp_path/"app.db")
    groups={x["material_type"]:x for x in repo.groups(False)}
    assert groups["ESCAPE_ANALYSIS"]["include_ai"] == 1
    assert groups["ITR_SOURCE"]["include_ai"] == 0
    assert groups["ITR_CS"]["include_insight"] == 0
    assert groups["SOFTWARE_OPERATION"]["include_report"] == 0


def test_material_web_import_and_navigation(tmp_path):
    db=tmp_path/"web.db";client=TestClient(create_app(db))
    assert client.get("/materials",follow_redirects=False).headers["location"]=="/materials/itr"
    page=client.get("/materials/itr")
    assert page.status_code==200 and "ITR问题工作台" in page.text and "两级表头" in page.text
    assert "软件 / 硬件 / 机械 / 跨领域" in page.text
    assert "仅软件" in client.get("/materials/software-operations").text
    itr=tmp_path/"itr.xlsx";workbook(itr,"ITR_SOURCE","ITR20260605084")
    with itr.open("rb") as stream:
        result=client.post("/materials/import",data={"workbench":"itr","group_code":"ITR","header_rows":"2"},files={"file":("itr.xlsx",stream,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert result.status_code==200 and "ITR20260605084" in result.text and "全新问题 1" in result.text
    assert "搜索编号、问题描述、产品" in result.text and "查看 / 分析" in result.text
    material_id=client.app.state.material_repository.list_materials("ITR")[0]["material_id"]
    detail=client.get(f"/materials/itr/{material_id}")
    assert detail.status_code==200 and "完整原始字段" in detail.text and "同一 ITR 的相关数据" in detail.text and "本工作台独立分析" in detail.text
    saved_review=client.post(f"/materials/itr/{material_id}/review",data={"review_status":"COMPLETED","analysis_summary":"客户停机","root_cause":"变更管理不足","improvement_action":"补充准入规则","reviewer":"质量组"},follow_redirects=False)
    assert saved_review.status_code==303
    reviewed=client.get(f"/materials/itr/{material_id}")
    assert "客户停机" in reviewed.text and "变更管理不足" in reviewed.text and "已完成" in reviewed.text
    filtered=client.get("/materials/itr?q=ITR20260605084")
    assert filtered.status_code==200 and "测试问题" in filtered.text
    assert "ITR工作台" in client.get("/issues").text
    settings=client.get("/settings/associations")
    assert settings.status_code==200 and "关联预检" in settings.text and "移除末尾CS" in settings.text
    saved=client.post("/settings/associations/RULE-CS-ITR",data={"source_field":"问题信息_彻底解决单号","target_field":"问题信息_ITR单号","transform":"STRIP_CS","status":"INACTIVE"},follow_redirects=False)
    assert saved.status_code==303
    assert next(x for x in client.app.state.material_repository.rules() if x["rule_id"]=="RULE-CS-ITR")["status"]=="INACTIVE"


def test_workbench_can_preview_and_clean_duplicate_history(tmp_path):
    app=create_app(tmp_path/'cleanup-web.db');client=TestClient(app);repo=app.state.material_repository
    group=repo.group('SW-OPS');key='ITR20260605084CS'
    repo.add_material(group,key,{'问题信息_彻底解决单号':key,'问题信息_问题描述':'旧内容'},'a.xlsx','Sheet1',3)
    repo.add_material(group,key,{'问题信息_彻底解决单号':key,'问题信息_问题描述':'新内容'},'b.xlsx','Sheet1',3)
    page=client.get('/materials/software-operations')
    assert page.status_code==200 and '历史重复问题' in page.text and '可安全清理 1 条' in page.text
    cleaned=client.post('/materials/software-operations/cleanup-duplicates',follow_redirects=False)
    assert cleaned.status_code==303 and 'cleaned=1' in cleaned.headers['location']
    result=client.get(cleaned.headers['location']).text
    assert '删除无引用的历史版本 1 条' in result and '历史重复问题' not in result


def test_wrong_sheet_records_can_be_previewed_and_safely_deleted_by_raw_field(tmp_path):
    repo=MaterialRepository(tmp_path/'wrong-sheet.db');group=repo.group('ITR')
    valid,_=repo.add_material(group,'ITR20260605001',{'问题信息_ITR单号':'ITR20260605001','问题信息_问题描述':'正常问题','过程信息_当前状态':'处理中'},'itr.xlsx','正确Sheet',3)
    wrong,_=repo.add_material(group,'ITR20260605002',{'问题信息_ITR单号':'ITR20260605002','问题信息_问题描述':'错误Sheet数据','错误表_统计口径':'仅供汇总'},'itr.xlsx','错误Sheet',3)
    protected,_=repo.add_material(group,'ITR20260605003',{'问题信息_ITR单号':'ITR20260605003','问题信息_问题描述':'受保护错误数据','错误表_统计口径':'仅供汇总'},'itr.xlsx','错误Sheet',4)
    repo.save_review(protected,review_status='COMPLETED',analysis_summary='已人工使用',root_cause='',improvement_action='',reviewer='质量组')
    catalog={row['field_name']:row['record_count'] for row in repo.raw_field_catalog('ITR')}
    assert catalog['错误表_统计口径']==2
    preview=repo.invalid_import_preview('ITR',field_name='错误表_统计口径',operator='HAS_FIELD')
    assert preview['total']==2 and preview['deletable']==1 and preview['protected']==1
    assert preview['sources']==[{'source_file':'itr.xlsx','sheet_name':'错误Sheet','count':2}]
    result=repo.cleanup_invalid_import('ITR',field_name='错误表_统计口径',operator='VALUE_CONTAINS',value='汇总')
    assert result=={'matched':2,'deleted':1,'protected':1}
    assert repo.material(valid) is not None and repo.material(wrong) is None and repo.material(protected) is not None


def test_wrong_sheet_filter_distinguishes_missing_field_from_value_not_contains(tmp_path):
    repo=MaterialRepository(tmp_path/'wrong-sheet-operators.db');group=repo.group('ITR-CS')
    repo.add_material(group,'ITR20260605011CS',{'问题信息_彻底解决单号':'ITR20260605011CS','标记':'错误数据'},'cs.xlsx','A',3)
    repo.add_material(group,'ITR20260605012CS',{'问题信息_彻底解决单号':'ITR20260605012CS','标记':'正确数据'},'cs.xlsx','A',4)
    repo.add_material(group,'ITR20260605013CS',{'问题信息_彻底解决单号':'ITR20260605013CS'},'cs.xlsx','B',3)
    missing=repo.invalid_import_preview('ITR-CS',field_name='标记',operator='MISSING_FIELD')
    not_contains=repo.invalid_import_preview('ITR-CS',field_name='标记',operator='VALUE_NOT_CONTAINS',value='正确')
    assert missing['total']==1 and missing['samples'][0]['business_key']=='ITR20260605013CS'
    assert not_contains['total']==1 and not_contains['samples'][0]['business_key']=='ITR20260605011CS'


def test_all_workbenches_expose_wrong_sheet_cleanup_page(tmp_path):
    app=create_app(tmp_path/'wrong-sheet-web.db');client=TestClient(app);repo=app.state.material_repository
    group=repo.group('SW-OPS')
    repo.add_material(group,'ITR20260605101CS',{'问题信息_彻底解决单号':'ITR20260605101CS','错误表_标识':'DELETE'},'ops.xlsx','汇总Sheet',3)
    page=client.get('/materials/software-operations')
    assert page.status_code==200 and '清理误导入数据' in page.text
    preview=client.get('/materials/software-operations/data-cleanup',params={'preview':1,'field_name':'错误表_标识','operator':'VALUE_EQUALS','value':'DELETE'})
    assert preview.status_code==200 and '汇总Sheet' in preview.text and '确认删除 1 条误导入记录' in preview.text
    deleted=client.post('/materials/software-operations/data-cleanup',data={'field_name':'错误表_标识','operator':'VALUE_EQUALS','value':'DELETE','confirmed':'yes'},follow_redirects=False)
    assert deleted.status_code==303 and 'deleted=1' in deleted.headers['location']
    assert repo.search_materials('SW-OPS')['total']==0


def test_disabled_rule_stops_automatic_link_but_keeps_other_type(tmp_path):
    db=tmp_path/"rules.db";repo=MaterialRepository(db)
    with repo.connect() as c:
        c.execute("CREATE TABLE quality_issue(knowledge_id TEXT PRIMARY KEY,business_issue_id TEXT)")
        c.execute("INSERT INTO quality_issue VALUES('K-1','ITR20260605084')")
    itr=tmp_path/"itr.xlsx";cs=tmp_path/"cs.xlsx";workbook(itr,"ITR_SOURCE","ITR20260605084");workbook(cs,"ITR_CS","ITR20260605084CS")
    service=MaterialImportService(repo);service.import_file(itr,"ITR",2);service.import_file(cs,"ITR-CS",2)
    repo.update_rule("RULE-CS-ITR",source_field="问题信息_彻底解决单号",target_field="问题信息_ITR单号",transform="STRIP_CS",status="INACTIVE");repo.refresh_links()
    assert {x["material_type"] for x in repo.materials_for_issue("K-1")}=={"ITR_SOURCE"}


def test_software_operation_is_isolated_and_links_by_cs_number(tmp_path):
    db=tmp_path/"ops.db";repo=MaterialRepository(db)
    with repo.connect() as c:
        c.execute("CREATE TABLE quality_issue(knowledge_id TEXT PRIMARY KEY,business_issue_id TEXT)")
        c.execute("INSERT INTO quality_issue VALUES('K-OPS','ITR20260605084')")
    source=tmp_path/"ops.xlsx";operations_workbook(source,"ITR20260605084CS")
    stats=MaterialImportService(repo).import_file(source,"SW-OPS",2)
    assert stats["new"]==1
    material=repo.list_materials("SW-OPS")[0]
    assert material["material_type"]=="SOFTWARE_OPERATION" and material["canonical_itr"]=="ITR20260605084"
    linked=repo.materials_for_issue("K-OPS")
    assert linked[0]["raw"]["数据运营_KPI计入月份"]=="2026-06"
