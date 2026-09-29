from pathlib import Path
import sqlite3

from fastapi.testclient import TestClient
from openpyxl import Workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web import create_app
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]


def _seed_issues(client: TestClient, tmp_path: Path) -> dict[str, str]:
    source = tmp_path / "association-issues.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([
        "ITR 单号",
        "问题标题",
        "问题描述",
        "是否漏测",
        "产品",
        "平台",
        "严重程度",
    ])
    sheet.append([
        "ITR-R2-ASSOC-1",
        "关联闭环问题",
        "用于验证四工作台共享同一问题身份",
        "是",
        "PLC",
        "IDE",
        "A",
    ])
    sheet.append([
        "ITR-R2-ASSOC-2",
        "无扩展关联问题",
        "只存在 Canonical Problem / ITR",
        "否",
        "PLC",
        "IDE",
        "B",
    ])
    book.save(source)
    with source.open("rb") as stream:
        response = client.post(
            "/api/issues/import",
            files={
                "file": (
                    source.name,
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"business_type": "PLC"},
        )
    assert response.status_code == 200
    items = client.get("/api/issues", params={"limit": 20}).json()["items"]
    return {item["business_issue_id"]: item["knowledge_id"] for item in items}


def _import_material(
    client: TestClient,
    path: Path,
    *,
    workbench: str,
    group_code: str,
) -> None:
    with path.open("rb") as stream:
        response = client.post(
            "/materials/import",
            data={
                "workbench": workbench,
                "group_code": group_code,
                "header_rows": "2",
            },
            files={
                "file": (
                    path.name,
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 200


def _resolution_source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["问题信息", "问题信息", "原因分析", "解决措施"])
    sheet.append(["彻底解决单号", "问题描述", "根因", "技术措施"])
    sheet.append([
        "ITR-R2-ASSOC-1CS",
        "彻底解决历史记录",
        "历史根因",
        "历史措施",
    ])
    book.save(path)


def _assessment_source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["问题信息", "问题信息", "考核信息", "责任信息"])
    sheet.append(["彻底解决单号", "问题描述", "考核结果", "责任部门"])
    sheet.append([
        "ITR-R2-ASSOC-1CS",
        "软件考核历史记录",
        "已确认",
        "软件研发部",
    ])
    book.save(path)


def test_current_problem_four_workbench_association_closure_is_read_only(tmp_path: Path):
    legacy_db = tmp_path / "association-legacy.db"
    legacy_client = TestClient(create_app(legacy_db))
    ids = _seed_issues(legacy_client, tmp_path)

    resolution = tmp_path / "resolution.xlsx"
    _resolution_source(resolution)
    _import_material(
        legacy_client,
        resolution,
        workbench="cs",
        group_code="ITR-CS",
    )

    assessment = tmp_path / "assessment.xlsx"
    _assessment_source(assessment)
    _import_material(
        legacy_client,
        assessment,
        workbench="software-operations",
        group_code="SW-OPS",
    )

    # Historical source facts can survive while the projection link table is
    # stale/missing. R2 must recover only an exact unique canonical relation,
    # and must not mutate the Legacy DB merely by reading it.
    with sqlite3.connect(legacy_db) as connection:
        connection.execute("DELETE FROM issue_material_link")
        assert connection.execute(
            "SELECT COUNT(*) FROM issue_material_link"
        ).fetchone()[0] == 0

    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    client = TestClient(
        create_p0_app(
            p0_db,
            stage_runner=object(),
            project_root=ROOT,
            legacy_quality_issue_db_path=legacy_db,
        )
    )

    resolution_page = client.get(
        "/p0/itr-resolution", params={"q": "ITR-R2-ASSOC-1"}
    )
    assert resolution_page.status_code == 200
    assert "ITR-R2-ASSOC-1CS" in resolution_page.text
    assert f"/issues/{ids['ITR-R2-ASSOC-1']}" in resolution_page.text

    assessment_page = client.get(
        "/p0/software-assessment", params={"q": "ITR-R2-ASSOC-1"}
    )
    assert assessment_page.status_code == 200
    assert "ITR-R2-ASSOC-1CS" in assessment_page.text
    assert f"/issues/{ids['ITR-R2-ASSOC-1']}" in assessment_page.text

    missed_page = client.get(
        "/p0/missed-test-analysis", params={"q": "ITR-R2-ASSOC-1"}
    )
    assert missed_page.status_code == 200
    assert "ITR-R2-ASSOC-1" in missed_page.text

    detail = client.get(
        f"/issues/{ids['ITR-R2-ASSOC-1']}",
        params={
            "return_to": "/p0/missed-test-analysis?q=ITR-R2-ASSOC-1"
        },
    )
    assert detail.status_code == 200
    assert "4/4 已关联" in detail.text
    assert "/p0/itr-recovery?q=ITR-R2-ASSOC-1" in detail.text
    assert "/p0/itr-resolution?q=ITR-R2-ASSOC-1" in detail.text
    assert "/p0/software-assessment?q=ITR-R2-ASSOC-1" in detail.text
    assert "/p0/missed-test-analysis?q=ITR-R2-ASSOC-1" in detail.text
    assert "返回漏测分析" in detail.text

    contract = client.get(
        f"/api/v2/issues/{ids['ITR-R2-ASSOC-1']}/workbench-relations"
    )
    assert contract.status_code == 200
    payload = contract.json()
    assert payload["contract_version"] == "canonical-problem/v1"
    assert payload["identity"]["master_object_ref"] == {
        "domain": "EXISTING_PROBLEM",
        "object_type": "quality_issue",
        "object_id": ids["ITR-R2-ASSOC-1"],
    }
    assert payload["workbench_count"] == 4
    assert all(
        relation["relation_contract_version"] == "canonical-problem-relation/v1"
        for relation in payload["relations"]
    )
    assert all(
        relation["relation_policy"] == "EXPLICIT_OR_EXACT_UNIQUE_ONLY"
        for relation in payload["relations"]
    )
    assert all(
        relation["return_context_contract"] == "overall-return-context/v1"
        for relation in payload["relations"]
    )
    for key in ("RESOLUTION", "SOFTWARE_ASSESSMENT"):
        relation = next(item for item in payload["relations"] if item["key"] == key)
        assert relation["source_refs"]
        assert relation["source_refs"][0]["version_no"] == 1
        assert relation["source_refs"][0]["source_hash"]
        assert relation["source_refs"][0]["object_ref"]["object_type"] == "source_material"

    no_relation = client.get(f"/p0/issues/{ids['ITR-R2-ASSOC-2']}")
    assert no_relation.status_code == 200
    assert "1/4 已关联" in no_relation.text
    assert no_relation.text.count("未发现已确认关联") >= 3

    with sqlite3.connect(legacy_db) as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM issue_material_link"
        ).fetchone()[0] == 0



def test_current_problem_ambiguous_canonical_relation_fails_closed(tmp_path: Path):
    legacy_db = tmp_path / "association-ambiguous.db"
    legacy_client = TestClient(create_app(legacy_db))
    ids = _seed_issues(legacy_client, tmp_path)

    resolution = tmp_path / "resolution-ambiguous.xlsx"
    _resolution_source(resolution)
    _import_material(
        legacy_client,
        resolution,
        workbench="cs",
        group_code="ITR-CS",
    )

    assessment = tmp_path / "assessment-ambiguous.xlsx"
    _assessment_source(assessment)
    _import_material(
        legacy_client,
        assessment,
        workbench="software-operations",
        group_code="SW-OPS",
    )

    # Same normalized ITR across two business-type uniqueness domains is
    # intentionally ambiguous to the cross-workbench resolver. Reading the
    # Common Problem View must never choose one by recency/title/product.
    with sqlite3.connect(legacy_db) as connection:
        connection.execute("DELETE FROM issue_material_link")
        connection.execute(
            """
            INSERT INTO quality_issue(
                knowledge_id,business_type,business_issue_id,status
            ) VALUES(?,?,?,?)
            """,
            (
                "QK-HMI-R2-AMBIGUOUS",
                "HMI",
                "ITR-R2-ASSOC-1",
                "ACTIVE",
            ),
        )

    p0_db = tmp_path / "p0-ambiguous.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    client = TestClient(
        create_p0_app(
            p0_db,
            stage_runner=object(),
            project_root=ROOT,
            legacy_quality_issue_db_path=legacy_db,
        )
    )

    response = client.get(
        f"/api/v2/issues/{ids['ITR-R2-ASSOC-1']}/workbench-relations"
    )
    assert response.status_code == 200
    payload = response.json()
    by_key = {relation["key"]: relation for relation in payload["relations"]}

    assert by_key["ITR"]["present"] is True
    assert by_key["MISSED_TEST"]["present"] is True
    for key in ("RESOLUTION", "SOFTWARE_ASSESSMENT"):
        relation = by_key[key]
        assert relation["present"] is False
        assert relation["relation_status"] == "NO_RELATION"
        assert relation["href"] == ""
        assert relation["source_refs"] == []
        assert relation["no_relation_reason"] == "RELATION_NOT_FOUND_OR_AMBIGUOUS"

    page = client.get(f"/p0/issues/{ids['ITR-R2-ASSOC-1']}")
    assert page.status_code == 200
    assert "2/4 已关联" in page.text
    assert page.text.count("未发现已确认关联") >= 2
