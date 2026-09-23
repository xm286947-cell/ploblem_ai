from pathlib import Path
from openpyxl import Workbook
from quality_knowledge.repositories import IssueKnowledgeRepository
from quality_knowledge.services import KnowledgeIssueService, NoRecordsFoundError

def make_xlsx(path,desc='D1',with_preamble=True):
    wb=Workbook();ws=wb.active;ws.title='问题清单'
    if with_preamble: ws.append(['PLC历史问题清单']);ws.append(['说明：以下为数据'])
    ws.append(['ITR单号','问题描述','问题原因定位','问题解决方案','是否漏测','原因分类一级','新二级分类'])
    ws.append(['ITR-001',desc,'RC','SOL','是','测试覆盖','功能'])
    wb.save(path)

def test_auto_detect_incremental_version_and_current(tmp_path):
    x=tmp_path/'plc.xlsx';db=tmp_path/'q.db';make_xlsx(x,'D1')
    svc=KnowledgeIssueService(IssueKnowledgeRepository(db))
    r1=svc.import_file(x); assert r1['new']==1 and r1['business_type']=='AUTO'
    r2=svc.import_file(x); assert r2['skipped']==1
    make_xlsx(x,'D2');r3=svc.import_file(x); assert r3['updated']==1
    rows=svc.query_issues({'business_type':'PLC'});assert len(rows)==1 and rows[0]['description']=='D2' and rows[0]['version_no']==2
    hist=svc.get_issue_history(rows[0]['knowledge_id']);assert [x['version_no'] for x in hist]==[2,1]

def test_import_error_isolated(tmp_path):
    x=tmp_path/'plc.xlsx';db=tmp_path/'q.db';make_xlsx(x)
    wb=load_workbook(x);ws=wb.active;ws.append(['','bad','','','','','']);wb.save(x)
    svc=KnowledgeIssueService(IssueKnowledgeRepository(db));r=svc.import_file(x)
    assert r['new']==1 and r['failed']==1 and r['status']=='PARTIAL'
    batch=svc.get_import_batch(r['batch_id']);assert len(batch['errors'])==1

def test_no_records_found_is_explicit(tmp_path):
    x=tmp_path/'empty.xlsx';db=tmp_path/'q.db';wb=Workbook();wb.active.append(['hello']);wb.save(x)
    svc=KnowledgeIssueService(IssueKnowledgeRepository(db))
    try:svc.import_file(x)
    except NoRecordsFoundError as e:assert 'NO_RECORDS_FOUND' in str(e)
    else:assert False

from openpyxl import load_workbook
