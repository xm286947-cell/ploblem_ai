from quality_knowledge.human_analysis import HumanAnalysisRepository,HumanAnalysisService
from quality_knowledge.repositories import IssueKnowledgeRepository
from quality_knowledge.services.v1_export_service import KnowledgeIssueExportService
from openpyxl import load_workbook
def test_human_query_and_export_rows(tmp_path):
 db=tmp_path/'k.db'; h=HumanAnalysisService(HumanAnalysisRepository(db))
 f=h.create_field(field_key='review',field_name='复盘结论',field_type='TEXT')
 h.save_analysis('K1','V1',{f['field_id']:'需要补强测试'})
 h.save_analysis('K2','V1',{f['field_id']:'设计整改'})
 assert h.query_issue_ids(f['field_id'],'测试')==['K1']
 rows=h.export_rows(); assert rows[0]['人工分析｜复盘结论']
def test_xlsx_has_human_analysis_sheet(tmp_path):
 db=tmp_path/'k.db'; repo=IssueKnowledgeRepository(db); h=HumanAnalysisService(HumanAnalysisRepository(db))
 f=h.create_field(field_key='review',field_name='人工结论',field_type='TEXT');h.save_analysis('K1','V1',{f['field_id']:'人工值'})
 out=tmp_path/'x.xlsx';KnowledgeIssueExportService(repo).export_xlsx(out,{})
 wb=load_workbook(out,read_only=True);assert 'Human_Analysis' in wb.sheetnames
 ws=wb['Human_Analysis'];headers=[x.value for x in next(ws.iter_rows())];assert '人工分析｜人工结论' in headers
def test_csv_human_analysis_dataset(tmp_path):
 db=tmp_path/'k.db';repo=IssueKnowledgeRepository(db);h=HumanAnalysisService(HumanAnalysisRepository(db))
 f=h.create_field(field_key='x',field_name='人工字段',field_type='TEXT');h.save_analysis('K1','V1',{f['field_id']:'abc'})
 out=tmp_path/'h.csv';r=KnowledgeIssueExportService(repo).export_csv(out,{},'human_analysis')
 assert r['rows']==1 and '人工分析｜人工字段' in out.read_text(encoding='utf-8-sig')
