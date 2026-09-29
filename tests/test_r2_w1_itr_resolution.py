# R2 W1 focused regression: ITR resolution vertical slice.
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web import create_app
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]


def _seed_issue(client: TestClient, tmp_path: Path) -> str:
    source = tmp_path / "issue.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([
        "ITR 单号",
        "问题描述",
        "问题原因定位（×开发填写×）",
        "问题解决方案（×开发填写×）",
        "是否漏测",
        "产品",
        "平台",
        "严重程度",
    ])
    sheet.append([
        "ITR-R2-W1-1",
        "现场通信异常",
        "边界处理遗漏",
        "补充边界处理",
        "是",
        "PLC",
        "IDE",
        "A",
    ])
    book.save(source)
    with source.open("rb") as stream:
        response = client.post(
            "/api/issues/import",
            files={
                "file": (
                    "issue.xlsx",
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"business_type": "PLC"},
        )
    assert response.status_code == 200
    return client.get("/api/issues").json()["items"][0]["knowledge_id"]


def _resolution_source(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append([
        "问题信息",
        "问题信息",
        "原因分析",
        "解决措施",
        "验证信息",
        "流程信息",
        "流程信息",
    ])
    sheet.append([
        "彻底解决单号",
        "问题描述",
        "根因",
        "技术措施",
        "验证结果",
        "业务状态",
        "责任人",
    ])
    sheet.append([
        "ITR-R2-W1-1CS",
        "现场通信异常彻底解决记录",
        "通信边界处理遗漏",
        "增加通信边界保护与回归测试",
        "验证通过",
        "待源系统关闭",
        "研发A",
    ])
    book.save(path)


def _import_resolution(client: TestClient, source: Path) -> None:
    with source.open("rb") as stream:
        imported = client.post(
            "/materials/import",
            data={
                "workbench": "cs",
                "group_code": "ITR-CS",
                "header_rows": "2",
            },
            files={
                "file": (
                    "resolution.xlsx",
                    stream,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert imported.status_code == 200


def test_r2_w1_resolution_projects_source_facts_and_round_trips(tmp_path: Path):
    db = tmp_path / "r2-w1.db"
    app = create_app(db)
    client = TestClient(app)
    knowledge_id = _seed_issue(client, tmp_path)

    issue_page = client.get(f"/issues/{knowledge_id}")
    assert issue_page.status_code == 200
    assert 'href="/itr/resolution-workbench"' in issue_page.text

    list_page = client.get("/issues")
    assert list_page.status_code == 200
    assert 'href="/itr/resolution-workbench"' in list_page.text

    source = tmp_path / "resolution.xlsx"
    _resolution_source(source)
    _import_resolution(client, source)

    workbench = client.get("/itr/resolution-workbench")
    assert workbench.status_code == 200
    assert "ITR 彻底解决工作台" in workbench.text
    assert "ITR-R2-W1-1CS" in workbench.text
    assert "ITR-R2-W1-1" in workbench.text
    assert "现场通信异常彻底解决记录" in workbench.text
    assert "通信边界处理遗漏" in workbench.text
    assert "增加通信边界保护与回归测试" in workbench.text
    assert "验证通过" in workbench.text
    assert "待源系统关闭" in workbench.text
    assert "研发A" in workbench.text
    assert "SOURCE_ACTION_BINDING_PENDING" in workbench.text
    assert "不得由 Overall 自建状态机" in workbench.text

    filtered = client.get(
        "/itr/resolution-workbench",
        params={"q": "验证通过"},
    )
    assert filtered.status_code == 200
    assert "ITR-R2-W1-1CS" in filtered.text
    assert "return_to=%2Fitr%2Fresolution-workbench%3Fq%3D" in filtered.text

    linked_issue = client.get(
        f"/issues/{knowledge_id}",
        params={"return_to": "/itr/resolution-workbench?q=验证通过"},
    )
    assert linked_issue.status_code == 200
    assert "返回彻底解决工作台" in linked_issue.text
    assert 'href="/itr/resolution-workbench?q=验证通过"' in linked_issue.text

    rejected = client.get(
        f"/issues/{knowledge_id}",
        params={"return_to": "https://example.invalid/escape"},
    )
    assert rejected.status_code == 400

    # No Overall-owned save / submit / transition endpoint is created.
    assert client.post("/itr/resolution-workbench").status_code == 405


def test_r2_w1_resolution_p0_shell_reuses_same_source_facts(tmp_path: Path):
    legacy_db = tmp_path / "legacy-resolution.db"
    legacy_app = create_app(legacy_db)
    legacy_client = TestClient(legacy_app)
    knowledge_id = _seed_issue(legacy_client, tmp_path)

    source = tmp_path / "resolution.xlsx"
    _resolution_source(source)
    _import_resolution(legacy_client, source)

    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)

    app = create_p0_app(
        p0_db,
        stage_runner=object(),
        project_root=ROOT,
        legacy_quality_issue_db_path=legacy_db,
    )
    client = TestClient(app)

    page = client.get("/p0/itr-resolution", params={"q": "通信边界"})
    assert page.status_code == 200
    assert "SOURCE-ALIGNED RESOLUTION VIEW" in page.text
    assert "ITR-R2-W1-1CS" in page.text
    assert "通信边界处理遗漏" in page.text
    assert "验证通过" in page.text
    assert "待源系统关闭" in page.text
    assert "SOURCE_ACTION_BINDING_PENDING" in page.text
    assert "return_to=%2Fp0%2Fitr-resolution%3Fq%3D" in page.text

    detail = client.get(
        f"/issues/{knowledge_id}",
        params={"return_to": "/p0/itr-resolution?q=通信边界"},
    )
    assert detail.status_code == 200
    assert "返回彻底解决单" in detail.text
    assert 'href="/p0/itr-resolution?q=通信边界"' in detail.text

    rejected = client.get(
        f"/issues/{knowledge_id}",
        params={"return_to": "https://example.invalid/escape"},
    )
    assert rejected.status_code == 400

    issues = client.get("/p0/issues")
    assert issues.status_code == 200
    assert 'href="/p0/itr-resolution"' in issues.text

    assert client.post("/p0/itr-resolution").status_code == 405


def test_r2_w1_resolution_workbench_does_not_guess_unlinked_itr(tmp_path: Path):
    app = create_app(tmp_path / "r2-w1-unlinked.db")
    client = TestClient(app)

    source = tmp_path / "unlinked.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["问题信息", "问题信息"])
    sheet.append(["彻底解决单号", "问题描述"])
    sheet.append(["ITR-NOT-FOUND-CS", "无匹配 ITR"])
    book.save(source)

    _import_resolution(client, source)

    workbench = client.get("/itr/resolution-workbench")
    assert workbench.status_code == 200
    assert "未匹配 ITR" in workbench.text
    assert "?return_to=%2Fitr%2Fresolution-workbench" not in workbench.text
