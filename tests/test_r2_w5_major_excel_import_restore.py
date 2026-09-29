from __future__ import annotations

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]


def _client(tmp_path: Path) -> TestClient:
    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    app = create_p0_app(
        p0_db,
        stage_runner=object(),
        project_root=ROOT,
        major_case_db_path=tmp_path / "major.db",
        major_attachment_root=tmp_path / "major-attachments",
        major_artifact_root=tmp_path / "artifacts",
    )
    return TestClient(app)


def _xlsx() -> bytes:
    output = BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "重大问题"
    sheet.append([
        "IGR编号",
        "ITR单号",
        "问题描述",
        "TRC发生",
        "MRC发生",
        "产品",
        "模块",
        "Failure Mechanism",
        "Trigger Condition",
    ])
    sheet.append([
        "IGR-R2-EXCEL-001",
        "ITR20269901",
        "Excel 批量导入恢复验证",
        "边界保护不足",
        "评审检查项缺失",
        "PLC-X",
        "Motion",
        "非原子状态更新",
        "掉电窗口",
    ])
    book.save(output)
    return output.getvalue()


def test_major_product_restores_excel_entry_and_formal_template(tmp_path: Path):
    client = _client(tmp_path)

    page = client.get("/p0/major-production")
    assert page.status_code == 200
    assert "Excel 批量导入" in page.text
    assert "下载正式 Excel 模板" in page.text
    assert "PDF / DOCX / DOC" in page.text

    template = client.get("/api/v2/major-production/excel/template")
    assert template.status_code == 200
    assert "MAJOR_CASE_IMPORT_TEMPLATE_V1.0.xlsx" in template.headers["content-disposition"]
    book = load_workbook(BytesIO(template.content), read_only=True)
    headers = [cell.value for cell in next(book.active.iter_rows(max_row=1))]
    book.close()
    assert "ITR单号" in headers
    assert "问题描述" in headers
    assert "TRC发生" in headers
    assert "Failure Mechanism" in headers


def test_major_excel_preview_mapping_confirm_import_reuses_existing_store(tmp_path: Path):
    client = _client(tmp_path)

    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={
            "file": (
                "major.xlsx",
                _xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200
    body = preview.json()
    assert body["total"] == 1
    assert body["importable"] == 1
    assert body["mapping"]["contract"] == "major-excel-field-mapping/v1"
    assert body["mapping"]["fields"]["itr_id"] == "ITR单号"
    assert body["rows"][0]["itrs"] == ["ITR20269901"]
    assert body["rows"][0]["completeness"]["retrieval_ready"] is True

    confirm = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": body["batch_id"]},
    )
    assert confirm.status_code == 200
    result = confirm.json()["result"]
    assert result["failed"] == 0
    assert result["created_cases"] == 1
    assert result["source_fact_revisions"] == 1
    assert result["events"] == 1
    case_id = result["case_ids"][0]

    batch = client.get(
        f"/api/v2/major-production/excel/batches/{body['batch_id']}"
    )
    assert batch.status_code == 200
    assert batch.json()["status"] == "COMPLETED"

    detail = client.get(f"/api/v2/major-production/cases/{case_id}")
    assert detail.status_code == 200
    case = detail.json()
    assert case["legacy_case_id"].startswith("IGR:")
    assert len(case["events"]) == 1
    assert case["events"][0]["standard_itr"] == "ITR20269901"
    assert any(
        link["source_type"] == "MAJOR_EXCEL_SOURCE_FACT"
        for link in case["source_links"]
    )

    # Excel import feeds the same canonical Major production service; there is
    # no second Web/app/port and no second problem/case master.
    assert client.get("/p0/cases").status_code == 200
    assert client.get("/p0/major-production").status_code == 200


def test_major_excel_invalid_file_fails_closed_and_document_intake_remains(tmp_path: Path):
    client = _client(tmp_path)

    invalid = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.csv", b"a,b\n1,2", "text/csv")},
    )
    assert invalid.status_code == 400
    assert invalid.json()["detail"] == "MAJOR_EXCEL_TYPE_UNSUPPORTED"

    # Existing document intake contract is still mounted and rejects Excel,
    # proving the two inputs share Major production without repurposing routes.
    document = client.post(
        "/api/v2/major-production/sources",
        data={
            "title": "document path",
            "group_code": "MAJOR",
            "standard_itr": "ITR20269902",
            "domain": "QUALITY",
        },
        files={
            "file": (
                "not-document.xlsx",
                _xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert document.status_code == 400
    assert document.json()["detail"] == "MAJOR_SOURCE_TYPE_UNSUPPORTED"
