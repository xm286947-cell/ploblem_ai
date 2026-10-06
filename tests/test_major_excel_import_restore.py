from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from quality_knowledge.major_cases.document_parser import ParseResult, ParsedFragment
import quality_knowledge.major_cases.restore as restore_module
from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]


def _client(tmp_path: Path) -> TestClient:
    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    app = create_p0_app(
        p0_db,
        stage_runner=object(),
        project_root=ROOT,
        major_case_db_path=tmp_path / "major.db",
        major_attachment_root=tmp_path / "major-attachments",
        major_artifact_root=tmp_path / "artifacts",
    )
    return TestClient(app)


def _xlsx() -> bytes:
    output = BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "重大问题"
    sheet.append([
        "IGR编号",
        "ITR单号",
        "问题描述",
        "TRC发生",
        "MRC发生",
        "产品",
        "模块",
        "Failure Mechanism",
        "Trigger Condition",
    ])
    sheet.append([
        "IGR-R2-EXCEL-001",
        "ITR20269901",
        "Excel 批量导入恢复验证",
        "边界保护不足",
        "评审检查项缺失",
        "PLC-X",
        "Motion",
        "非原子状态更新",
        "掉电窗口",
    ])
    book.save(output)
    return output.getvalue()


def _xlsx_with_rows(rows: list[list[str]], headers: list[str] | None = None) -> bytes:
    output = BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "重大问题"
    sheet.append(headers or [
        "IGR编号", "ITR单号", "问题描述", "TRC发生", "MRC发生", "产品", "模块",
        "Failure Mechanism", "Trigger Condition", "报告文件名",
    ])
    for row in rows:
        sheet.append(row)
    book.save(output)
    return output.getvalue()


def test_major_product_restores_excel_entry_and_formal_template(tmp_path: Path):
    client = _client(tmp_path)

    page = client.get("/p0/major-production")
    assert page.status_code == 200
    assert "Excel 批量导入" in page.text
    assert "下载正式 Excel 模板" in page.text
    assert "PDF / DOCX / DOC" in page.text

    template = client.get("/api/v2/major-production/excel/template")
    assert template.status_code == 200
    assert "MAJOR_CASE_IMPORT_TEMPLATE_V1.0.xlsx" in template.headers["content-disposition"]
    book = load_workbook(BytesIO(template.content), read_only=True)
    headers = [cell.value for cell in next(book.active.iter_rows(max_row=1))]
    book.close()
    assert "ITR单号" in headers
    assert "问题描述" in headers
    assert "TRC发生" in headers
    assert "Failure Mechanism" in headers
    metadata = load_workbook(BytesIO(template.content), read_only=True, data_only=True)
    assert metadata["_MAJOR_IMPORT_META"]["B1"].value == "major-excel-template/v1"
    assert metadata["_MAJOR_IMPORT_META"]["B2"].value == "1.0"
    metadata.close()


def test_major_excel_preview_mapping_confirm_import_reuses_existing_store(tmp_path: Path):
    client = _client(tmp_path)
    assert client.app.state.major_case_repository.schema_version() == 3

    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={
            "file": (
                "major.xlsx",
                _xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200
    body = preview.json()
    assert body["total"] == 1
    assert body["importable"] == 1
    assert body["mapping"]["contract"] == "major-excel-field-mapping/v1"
    assert body["mapping"]["fields"]["itr_id"] == "ITR单号"
    assert body["rows"][0]["itrs"] == ["ITR20269901"]
    assert body["rows"][0]["completeness"]["retrieval_ready"] is True
    assert body["blocked"] == 0
    assert body["rows"][0]["blocking_reasons"] == []
    # Preview is read-only against the formal Major store.
    assert client.app.state.major_case_repository.list_cases()["total"] == 0

    confirm = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": body["batch_id"]},
    )
    assert confirm.status_code == 200
    result = confirm.json()["result"]
    assert result["failed"] == 0
    assert result["created_cases"] == 1
    assert result["source_fact_revisions"] == 1
    assert result["events"] == 1
    case_id = result["case_ids"][0]

    batch = client.get(
        f"/api/v2/major-production/excel/batches/{body['batch_id']}"
    )
    assert batch.status_code == 200
    assert batch.json()["status"] == "COMPLETED"

    detail = client.get(f"/api/v2/major-production/cases/{case_id}")
    assert detail.status_code == 200
    case = detail.json()
    assert case["legacy_case_id"].startswith("IGR:")
    assert len(case["events"]) == 1
    assert case["events"][0]["standard_itr"] == "ITR20269901"
    assert any(
        link["source_type"] == "MAJOR_EXCEL_SOURCE_FACT"
        for link in case["source_links"]
    )
    assert len(case["source_fact_revisions"]) == 1
    assert case["source_fact_revisions"][0]["source_type"] == "EXCEL"

    # Excel import feeds the same canonical Major production service; there is
    # no second Web/app/port and no second problem/case master.
    assert client.get("/p0/cases").status_code == 200
    assert client.get("/p0/major-production").status_code == 200


def test_major_excel_invalid_file_fails_closed_and_document_intake_remains(tmp_path: Path):
    client = _client(tmp_path)

    invalid = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.csv", b"a,b\n1,2", "text/csv")},
    )
    assert invalid.status_code == 400
    assert invalid.json()["detail"] == "MAJOR_EXCEL_TYPE_UNSUPPORTED"

    # Existing document intake contract is still mounted and rejects Excel,
    # proving the two inputs share Major production without repurposing routes.
    document = client.post(
        "/api/v2/major-production/sources",
        data={
            "title": "document path",
            "group_code": "MAJOR",
            "standard_itr": "ITR20269902",
            "domain": "QUALITY",
        },
        files={
            "file": (
                "not-document.xlsx",
                _xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert document.status_code == 400
    assert document.json()["detail"] == "MAJOR_SOURCE_TYPE_UNSUPPORTED"


def test_major_excel_mapping_change_after_preview_fails_closed(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    client.app.state.major_case_restore_service.excel_parser.field_mapping["itr_id"] = "变更后的ITR列"
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_MAPPING_CHANGED_AFTER_PREVIEW"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_major_excel_preview_hash_change_fails_closed(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    repository = client.app.state.major_case_repository
    with repository.transaction() as connection:
        connection.execute(
            "UPDATE kb_major_import_batch SET preview_json=? WHERE batch_id=?",
            ('{"rows":[]}', preview["batch_id"]),
        )
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_PREVIEW_CHANGED_AFTER_PREVIEW"
    assert repository.list_cases()["total"] == 0


def test_major_excel_staged_file_change_fails_closed(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    batch = client.app.state.major_case_restore_service.batch(preview["batch_id"])
    staged_excel = Path(batch["staging_path"]) / batch["source_file"]
    staged_excel.write_bytes(b"changed after preview")
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_SOURCE_CHANGED_AFTER_PREVIEW"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_major_excel_commit_serializes_against_existing_major_writes(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    guard = client.app.state.major_mutation_guard
    assert guard._lock.acquire(blocking=False)
    try:
        response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    finally:
        guard._lock.release()
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_MUTATION_BUSY"
    assert client.app.state.major_case_restore_service.batch(preview["batch_id"])["status"] == "PREVIEW"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_major_excel_atomic_rollback_restores_case_and_attachment_state(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    repository = client.app.state.major_case_repository
    add_source_link = repository.add_source_link

    def fail_after_first_write(*args, **kwargs):
        add_source_link(*args, **kwargs)
        raise RuntimeError("injected mid-commit failure")

    repository.add_source_link = fail_after_first_write
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    repository.add_source_link = add_source_link
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_ATOMIC_COMMIT_FAILED"
    assert repository.list_cases()["total"] == 0
    batch = client.app.state.major_case_restore_service.batch(preview["batch_id"])
    assert batch["status"] == "FAILED"
    assert batch["result"]["atomic_rollback"] is True


def test_source_fact_revision_is_idempotent_and_changed_content_versions(tmp_path: Path):
    client = _client(tmp_path)
    repository = client.app.state.major_case_repository
    case = repository.create_case("revision test", "MAJOR")
    first = repository.add_source_fact_revision(
        case["case_id"], source_type="EXCEL", source_ref="book.xlsx#sheet:2",
        raw={"问题描述": "A"}, normalized={"original_description": "A"}, actor="tester",
    )
    duplicate = repository.add_source_fact_revision(
        case["case_id"], source_type="EXCEL", source_ref="book.xlsx#sheet:2",
        raw={"问题描述": "A"}, normalized={"original_description": "A"}, actor="tester",
    )
    changed = repository.add_source_fact_revision(
        case["case_id"], source_type="EXCEL", source_ref="book.xlsx#sheet:2",
        raw={"问题描述": "B"}, normalized={"original_description": "B"}, actor="tester",
    )
    assert first["created"] is True
    assert duplicate["created"] is False
    assert duplicate["source_fact_revision_id"] == first["source_fact_revision_id"]
    assert changed["created"] is True
    assert changed["revision_no"] == 2


def test_non_importable_row_blocks_the_whole_batch(tmp_path: Path):
    client = _client(tmp_path)
    content = _xlsx_with_rows([[
        "", "", "", "", "", "PLC-X", "", "", "", "",
    ]])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("blocked.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert preview["blocked"] == 1
    assert "ROW_NOT_IMPORTABLE" in preview["rows"][0]["blocking_reasons"]
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_BATCH_PRECHECK_FAILED"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_ambiguous_report_match_blocks_the_whole_batch(tmp_path: Path):
    client = _client(tmp_path)
    content = _xlsx_with_rows([[
        "IGR-AMB-001", "ITR20269903", "Ambiguous report", "TRC", "MRC", "PLC-X", "Motion",
        "Mechanism", "Trigger", "missing.pdf",
    ]])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files=[
            ("file", ("ambiguous.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ("materials", ("ITR20269903-a.pdf", b"a", "application/pdf")),
            ("materials", ("ITR20269903-b.pdf", b"b", "application/pdf")),
        ],
    ).json()
    assert preview["rows"][0]["report_match"]["match_type"] == "AMBIGUOUS"
    assert "AMBIGUOUS_REPORT_MATCH" in preview["rows"][0]["blocking_reasons"]
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    assert response.json()["detail"] == "MAJOR_EXCEL_BATCH_PRECHECK_FAILED"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_case_identity_conflict_added_after_preview_fails_closed(tmp_path: Path):
    client = _client(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("major.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    repository = client.app.state.major_case_repository
    restore = client.app.state.major_case_restore_service
    case_by_igr = repository.create_case("IGR owner", "MAJOR")
    case_by_source = repository.create_case("Source owner", "MAJOR")
    row = preview["rows"][0]
    restore._set_identity(case_by_igr["case_id"], "MAJOR", "IGR", row["igr"])
    restore._set_identity(case_by_source["case_id"], "MAJOR", "SOURCE_KEY", row["source_key"])
    response = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]})
    assert response.status_code == 409
    batch = restore.batch(preview["batch_id"])
    assert batch["status"] == "FAILED"
    assert any(item["error"] == "CASE_IDENTITY_CONFLICT" for item in batch["result"]["errors"])


def test_import_result_lists_every_case_and_case_detail_shows_events_and_facts(tmp_path: Path):
    client = _client(tmp_path)
    content = _xlsx_with_rows([
        ["IGR-MULTI-001", "ITR20269911", "First case", "TRC", "MRC", "PLC-X", "Motion", "Mechanism", "Trigger", ""],
        ["IGR-MULTI-002", "ITR20269912", "Second case", "TRC", "MRC", "PLC-X", "Motion", "Mechanism", "Trigger", ""],
    ])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("multi.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert preview["blocked"] == 0
    result = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]}).json()["result"]
    assert len(result["case_ids"]) == 2
    for case_id in result["case_ids"]:
        detail = client.get(f"/api/v2/major-production/cases/{case_id}").json()
        assert detail["events"]
        assert detail["source_fact_revisions"]
        assert detail["source_fact_revisions"][0]["source_type"] == "EXCEL"
    page = client.get("/p0/major-production")
    assert page.status_code == 200
    script = client.get("/p0/static/major_production.js?v=major-excel-restore-v1")
    assert "查看 Case / Event / Source Fact" in script.text


def test_multi_itr_excel_source_fact_is_case_scoped_and_analyzable_for_each_event(tmp_path: Path):
    client = _client(tmp_path)
    content = _xlsx_with_rows([[
        "IGR-MULTI-ITR-001", "ITR20269921", "One Excel row spans two ITRs",
        "Shared TRC occurrence", "Shared TRC escape", "Shared MRC occurrence",
        "Shared MRC escape", "", "ITR20269922",
    ]], headers=[
        "IGR编号", "ITR单号", "问题描述", "TRC发生", "TRC流出", "MRC发生", "MRC流出",
        "报告文件名", "关联ITR",
    ])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("multi-itr.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert preview["blocked"] == 0
    imported = client.post(
        "/api/v2/major-production/excel/confirm", data={"batch_id": preview["batch_id"]},
    )
    assert imported.status_code == 200, imported.text
    case_id = imported.json()["result"]["case_ids"][0]
    detail = client.get(f"/api/v2/major-production/cases/{case_id}").json()
    assert [item["standard_itr"] for item in detail["events"]] == ["ITR20269921", "ITR20269922"]
    assert detail["source_fact_revisions"][0]["normalized"]["itrs"] == ["ITR20269921", "ITR20269922"]

    fact_links = [
        item for item in detail["source_links"]
        if item["source_type"] == "MAJOR_EXCEL_SOURCE_FACT"
    ]
    assert len(fact_links) == 1
    shared_link = fact_links[0]
    assert shared_link["event_id"] is None
    assert shared_link["standard_itr"] == ""
    assert json.loads(shared_link["snapshot_json"])["binding_scope"] == "CASE_SHARED"

    for event in detail["events"]:
        response = client.post(
            f"/api/v2/major-production/cases/{case_id}/analysis",
            json={"event_id": event["event_id"]},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["event_id"] == event["event_id"]
        assert all(item["event_id"] == event["event_id"] for item in body["candidates"])
        occurrence = next(item for item in body["candidates"] if item["entry_type"] == "TRC_OCCURRENCE")
        assert occurrence["content"] == "Shared TRC occurrence"
        assert {item["source_link_id"] for item in occurrence["evidence"]} == {shared_link["source_link_id"]}



def test_multi_itr_report_without_unique_event_reference_fails_closed(tmp_path: Path):
    client = _client(tmp_path)
    content = _xlsx_with_rows([[
        "IGR-MULTI-REPORT-001", "ITR20269931", "Multi-event report requires binding",
        "TRC", "MRC", "PLC-X", "Motion", "review.pdf", "ITR20269932",
    ]], headers=[
        "IGR编号", "ITR单号", "问题描述", "TRC发生", "MRC发生", "产品", "模块",
        "报告文件名", "关联ITR",
    ])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files=[
            ("file", (
                "multi-report.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )),
            ("materials", ("review.pdf", b"%PDF-1.4\n", "application/pdf")),
        ],
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["blocked"] == 1
    assert body["rows"][0]["report_match"]["event_binding_status"] == "REVIEW_REQUIRED"
    assert body["rows"][0]["report_match"]["event_binding_itr"] == ""
    assert "MULTI_EVENT_REPORT_BINDING_AMBIGUOUS" in body["rows"][0]["blocking_reasons"]

    confirm = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": body["batch_id"]},
    )
    assert confirm.status_code == 409
    assert confirm.json()["detail"] == "MAJOR_EXCEL_BATCH_PRECHECK_FAILED"
    assert client.app.state.major_case_repository.list_cases()["total"] == 0


def test_multi_itr_report_filename_uniquely_binds_second_event(
    tmp_path: Path,
    monkeypatch,
):
    client = _client(tmp_path)
    monkeypatch.setattr(
        restore_module,
        "parse_document",
        lambda _path: ParseResult(
            "PDF",
            [
                ParsedFragment(
                    ordinal=1,
                    section_path="",
                    location_type="PAGE",
                    location_ref="page:1",
                    fragment_type="TEXT",
                    text="review evidence",
                )
            ],
        ),
    )
    report_name = "ITR20269942-review.pdf"
    content = _xlsx_with_rows([[
        "IGR-MULTI-REPORT-002", "ITR20269941", "Report belongs to second event",
        "TRC", "MRC", "PLC-X", "Motion", report_name, "ITR20269942",
    ]], headers=[
        "IGR编号", "ITR单号", "问题描述", "TRC发生", "MRC发生", "产品", "模块",
        "报告文件名", "关联ITR",
    ])
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files=[
            ("file", (
                "multi-report-bound.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )),
            ("materials", (report_name, b"%PDF-1.4\n", "application/pdf")),
        ],
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["blocked"] == 0
    match = body["rows"][0]["report_match"]
    assert match["event_binding_status"] == "BOUND"
    assert match["event_binding_itr"] == "ITR20269942"

    confirm = client.post(
        "/api/v2/major-production/excel/confirm",
        data={"batch_id": body["batch_id"]},
    )
    assert confirm.status_code == 200, confirm.text
    case_id = confirm.json()["result"]["case_ids"][0]
    detail = client.get(f"/api/v2/major-production/cases/{case_id}").json()
    events = {event["standard_itr"]: event for event in detail["events"]}
    assert set(events) == {"ITR20269941", "ITR20269942"}

    fact_links = [
        item for item in detail["source_links"]
        if item["source_type"] == "MAJOR_EXCEL_SOURCE_FACT"
    ]
    assert len(fact_links) == 1
    assert fact_links[0]["event_id"] is None
    assert json.loads(fact_links[0]["snapshot_json"])["binding_scope"] == "CASE_SHARED"

    document_links = [
        item for item in detail["source_links"]
        if item["source_type"] == "MAJOR_SOURCE_DOCUMENT"
    ]
    assert len(document_links) == 1
    link = document_links[0]
    assert link["event_id"] == events["ITR20269942"]["event_id"]
    assert link["event_id"] != events["ITR20269941"]["event_id"]
    snapshot = json.loads(link["snapshot_json"])
    assert snapshot["event_binding_basis"] == "REPORT_FILENAME_ITR"
    assert snapshot["event_binding_itr"] == "ITR20269942"



def test_repeated_excel_source_is_idempotent_then_changed_content_adds_revision(tmp_path: Path):
    client = _client(tmp_path)
    endpoint = "/api/v2/major-production/excel/preview"
    confirm_endpoint = "/api/v2/major-production/excel/confirm"
    def import_xlsx(description: str) -> dict:
        content = _xlsx_with_rows([[
            "IGR-REV-001", "ITR20269921", description, "TRC", "MRC", "PLC-X", "Motion",
            "Mechanism", "Trigger", "",
        ]])
        preview = client.post(
            endpoint,
            data={"group_code": "MAJOR", "domain": "QUALITY"},
            files={"file": ("revision.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        return client.post(confirm_endpoint, data={"batch_id": preview["batch_id"]}).json()["result"]

    first = import_xlsx("same source fact")
    repeated = import_xlsx("same source fact")
    changed = import_xlsx("updated source fact")
    assert first["created_cases"] == 1
    assert repeated["reused_cases"] == 1
    assert repeated["source_fact_revisions"] == 0
    assert changed["reused_cases"] == 1
    assert changed["source_fact_revisions"] == 1
    case_id = first["case_ids"][0]
    facts = client.get(f"/api/v2/major-production/cases/{case_id}").json()["source_fact_revisions"]
    assert [fact["revision_no"] for fact in facts] == [2, 1]
