"""Regenerate the entirely synthetic Word, Excel, and expected answers."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
CASES = [
    ("A9001", "LDO 输出振荡", "LDO 样机验证。", "LDO 输出振荡。", "测量电容 ESR。", "电容 ESR 偏低。", "更换电容。", "LDO 恢复稳定。", ["CF-LDO"], ["MD-CAP"]),
    ("A9002", "GPIO 控制异常", "GPIO 输入测试。", "GPIO 控制异常。", "检查输入电平。", "IC 输入悬空。", "增加上拉。", "GPIO 恢复。", ["CF-GPIO"], ["MD-IC"]),
    ("A9003", "CAN 通信中断", "CAN 通信测试。", "CAN 通信中断。", "测量收发器供电。", "通信 IC 供电跌落。", "调整通信 IC 供电。", "CAN 通信恢复。", ["CF-CAN"], ["MD-COMM"]),
    ("A9004", "功率驱动过热", "功率驱动测试。", "功率驱动过热。", "测量 MOS 温升。", "MOS 导通损耗过高。", "更换 MOS。", "温升达标。", ["CF-DRIVE"], ["MD-MOS"]),
    ("A9005", "接口信号丢失", "接口振动测试。", "接口信号间歇丢失。", "检查连接器。", "连接器接触不良。", "更换连接器。", "接口信号稳定。", ["CF-PORT"], ["MD-CONN"]),
    ("A9006", "多器件复位", "电源测试。", "电源复位。", "检查电容及 MOS。", "电容与 MOS 参数不匹配。", "更换电容和 MOS。", "电源复位消失。", ["CF-POWER"], ["MD-CAP", "MD-MOS"]),
    ("A9007", "多电路异常", "GPIO 与 CAN 联测。", "GPIO 与 CAN 同时异常。", "检查 GPIO 和 CAN 供电。", "通信 IC 供电不足。", "调整通信 IC 供电。", "GPIO 和 CAN 均恢复。", ["CF-GPIO", "CF-CAN"], ["MD-COMM"]),
    ("A9008", "LDO 输出偏低", "LDO 负载测试。", "LDO 输出偏低。", "检查 LDO 负载。", None, "减小负载。", "LDO 输出恢复。", ["CF-LDO"], []),
    ("A9009", "接口接触异常", "接口测试。", "接口接触异常。", "检查连接器触点。", "连接器触点氧化。", None, "待措施验证。", ["CF-PORT"], ["MD-CONN"]),
    ("A9010", "电源纹波", "电源测试。", "电源纹波超标。", "观察电源波形。", "电容老化。", "更换电容。", "电源纹波恢复。", ["CF-POWER"], ["MD-CAP"]),
    ("A9011", "非标准章节", "功率驱动测试。", "功率驱动启动失败。", "测量 MOS 栅压。", "MOS 栅压不足。", "调整栅极电阻。", "功率驱动启动正常。", ["CF-DRIVE"], ["MD-MOS"]),
    ("A9012", "未分类现象", "通用试验台。", "偶发指示灯闪烁。", "记录环境温度。", "环境温度突变。", "增加温控。", "闪烁消失。", [], []),
]
CIRCUITS = [("CF-LDO", "电源", "LDO"), ("CF-GPIO", "控制", "GPIO"), ("CF-CAN", "通信", "CAN"), ("CF-DRIVE", "驱动", "功率驱动"), ("CF-PORT", "接口", "接口"), ("CF-POWER", "电源", "电源")]
MATERIALS = [("MD-CAP", "电子器件", "电容"), ("MD-IC", "电子器件", "IC"), ("MD-COMM", "电子器件", "通信 IC"), ("MD-MOS", "电子器件", "MOS"), ("MD-CONN", "结构件", "连接器")]


def main() -> None:
    HERE.mkdir(parents=True, exist_ok=True)
    for filename, nodes in (("circuit_feature.xlsx", CIRCUITS), ("material.xlsx", MATERIALS)):
        book = Workbook()
        sheet = book.active
        sheet.title = "Synthetic Tree"
        sheet.append(["node_id", "level_1", "level_2"])
        for row in nodes:
            sheet.append(row)
        book.save(HERE / filename)
    expected = {}
    for case_id, title, background, symptom, analysis, cause, actions, conclusion, circuit, material in CASES:
        sections = [("背景", background), ("问题现象", symptom), ("分析过程", analysis), ("根因", cause), ("解决措施", actions), ("结论", conclusion)]
        if case_id == "A9011":
            sections = [("项目背景", background), ("现象记录", symptom), ("排查记录", analysis), ("原因定位", cause), ("处置方案", actions), ("结论", conclusion)]
            sections[0] = ("背景", background)
        doc = Document()
        for heading, value in sections:
            if value is not None:
                doc.add_heading(heading, level=1)
                doc.add_paragraph(value)
        if case_id == "A9010":
            doc.add_heading("附图", level=1)
            # Tiny synthetic image kept as an attachment, with no image semantics.
            import base64
            from io import BytesIO
            from docx.shared import Inches
            pixel = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL/nwAAAABJRU5ErkJggg==")
            doc.add_picture(BytesIO(pixel), width=Inches(0.1))
        doc.save(HERE / f"{case_id}-{title}.docx")
        expected[case_id] = {"title": title, "background": background, "symptom": symptom,
                             "analysis_process": analysis, "root_cause": cause, "actions": actions,
                             "conclusion": conclusion, "circuit_feature_links": circuit,
                             "material_links": material,
                             "expected_evidence_fields": [field for field, value in zip(
                                 ("background", "symptom", "analysis_process", "root_cause", "actions", "conclusion"),
                                 (background, symptom, analysis, cause, actions, conclusion)) if value is not None]}
    (HERE / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
