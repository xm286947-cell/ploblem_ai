from pathlib import Path
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from quality_knowledge.web import create_app

def _seed(c,tmp_path):
    p=tmp_path/'p.xlsx'; w=Workbook(); s=w.active
    s.append(['ITR 单号','问题描述','产品','平台','严重程度','问题原因定位（×开发填写×）'])
    s.append(['ITR-WEB-EXP-1','导出验证问题','PLC','IDE','A','设计原因'])
    w.save(p)
    with p.open('rb') as f:
        r=c.post('/api/issues/import',files={'file':('p.xlsx',f,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},data={'business_type':'PLC'})
    assert r.status_code==200

def test_web_export_page_and_filtered_link(tmp_path):
    c=TestClient(create_app(tmp_path/'q.db'));_seed(c,tmp_path)
    r=c.get('/issues?business_type=PLC&severity=A')
    assert r.status_code==200 and '导出当前结果' in r.text and '/export?' in r.text
    e=c.get('/export?business_type=PLC&severity=A')
    assert e.status_code==200 and '下载导出文件' in e.text and 'business_type = PLC' in e.text

def test_web_export_download_xlsx_and_csv(tmp_path):
    c=TestClient(create_app(tmp_path/'q.db'));_seed(c,tmp_path)
    x=c.post('/export/download',data={'format':'xlsx','dataset':'issues','scope':'filtered','business_type':'PLC'})
    assert x.status_code==200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in x.headers['content-type']
    fp=tmp_path/'out.xlsx';fp.write_bytes(x.content); wb=load_workbook(fp,read_only=True)
    assert {'Issue_Knowledge','AI_Analysis','Capability_Gaps','Statistics'} <= set(wb.sheetnames)
    csv=c.post('/export/download',data={'format':'csv','dataset':'issues','scope':'filtered','business_type':'PLC'})
    assert csv.status_code==200 and 'text/csv' in csv.headers['content-type'] and 'ITR-WEB-EXP-1' in csv.content.decode('utf-8-sig')
