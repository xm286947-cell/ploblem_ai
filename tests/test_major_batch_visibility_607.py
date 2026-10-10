"""Major Windows P0: full batch visibility and explainable preflight.

Synthetic workbook and isolated repositories only; no product DB or live Provider.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfWriter

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.web.p0_app import create_p0_app


ROOT = Path(__file__).resolve().parents[1]
XLS_FIXTURE = ROOT / "tests/fixtures/major_v5/cases.xls"


def client_for(tmp_path: Path) -> TestClient:
    p0_db = tmp_path / "p0.db"
    P0Initializer(
        manifest_path=ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=ROOT / "quality_knowledge/config/plc_fields.yaml",
    ).initialize(p0_db)
    return TestClient(create_p0_app(
        p0_db, stage_runner=object(), project_root=ROOT,
        major_case_db_path=tmp_path / "major.db",
        major_attachment_root=tmp_path / "attachments",
        major_artifact_root=tmp_path / "artifacts",
    ))


def _pdf_bytes() -> bytes:
    out = BytesIO()
    pdf = PdfWriter()
    pdf.add_blank_page(width=72, height=72)
    pdf.write(out)
    return out.getvalue()


def _matched_preview(client: TestClient) -> dict:
    resp = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files=[
            ("file", ("cases.xls", XLS_FIXTURE.read_bytes(), "application/vnd.ms-excel")),
            ("materials", ("ITR20269951-review.pdf", _pdf_bytes(), "application/pdf")),
        ],
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_preflight_explains_confirmable_batch_without_mutation(tmp_path: Path) -> None:
    client = client_for(tmp_path)
    preview = _matched_preview(client)
    batch_id = preview["batch_id"]
    before = client.app.state.major_case_restore_service.batch(batch_id)
    response = client.get(f"/api/v2/major-production/excel/batches/{batch_id}/preflight")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["batch_id"] == batch_id
    assert body["status"] == "PREVIEW"
    assert body["row_count"] == len(preview["rows"]) == 1
    assert body["confirmable"] is True
    assert body["errors"] == []
    after = client.app.state.major_case_restore_service.batch(batch_id)
    assert after["status"] == before["status"] == "PREVIEW"
    assert after["governance"]["preview_sha256"] == before["governance"]["preview_sha256"]

    result = client.post("/api/v2/major-production/excel/confirm", data={"batch_id": batch_id})
    assert result.status_code == 200, result.text
    completed = client.get(f"/api/v2/major-production/excel/batches/{batch_id}/preflight").json()
    assert completed["confirmable"] is False
    assert completed["errors"][0]["error"] == "MAJOR_EXCEL_BATCH_NOT_CONFIRMABLE"


def test_preflight_returns_all_row_blockers_without_mutating_failed_batch(tmp_path: Path) -> None:
    client = client_for(tmp_path)
    preview = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files=[
            ("file", ("cases.xls", XLS_FIXTURE.read_bytes(), "application/vnd.ms-excel")),
            ("materials", ("ITR20269951-review.pdf", _pdf_bytes(), "application/pdf")),
            ("materials", ("itr20269951-REVIEW.PDF", _pdf_bytes(), "application/pdf")),
        ],
    )
    assert preview.status_code == 200, preview.text
    batch_id = preview.json()["batch_id"]
    preflight = client.get(f"/api/v2/major-production/excel/batches/{batch_id}/preflight")
    assert preflight.status_code == 200
    assert preflight.json()["confirmable"] is False
    assert any(item["error"] == "AMBIGUOUS_REPORT_MATCH" for item in preflight.json()["errors"])
    assert client.app.state.major_case_restore_service.batch(batch_id)["status"] == "PREVIEW"


def test_real_69_row_excel_preview_is_not_server_truncated(tmp_path: Path) -> None:
    client = client_for(tmp_path)
    template = client.get("/api/v2/major-production/excel/template")
    assert template.status_code == 200, template.text
    wb = load_workbook(BytesIO(template.content))
    ws = wb["重大问题"]
    header = [str(c.value or "") for c in ws[1]]
    original = [c.value for c in ws[2]]
    assert "IGR编号" in header and "ITR单号" in header
    for row_num in range(2, 71):
        content = dict(zip(header, original))
        content["IGR编号"] = f"IGR-2026-{row_num:04d}"
        content["ITR单号"] = f"ITR2027{row_num:04d}"
        content["问题描述"] = f"独立人工确认测试源问题 {row_num}：用于 69 行分页边界"
        for col, field in enumerate(header, 1):
            ws.cell(row_num, col).value = content.get(field)
    payload = BytesIO()
    wb.save(payload)
    response = client.post(
        "/api/v2/major-production/excel/preview",
        data={"group_code": "MAJOR", "domain": "QUALITY"},
        files={"file": ("sixty-nine.xlsx", payload.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["total"] == 69
    assert len(preview["rows"]) == 69
    assert preview["rows"][0]["excel_row"] < preview["rows"][-1]["excel_row"]
    batch_id = preview["batch_id"]
    restored = client.get(f"/api/v2/major-production/excel/batches/{batch_id}")
    assert restored.status_code == 200
    assert len(restored.json()["preview"]["rows"]) == 69
    preflight = client.get(f"/api/v2/major-production/excel/batches/{batch_id}/preflight")
    assert preflight.status_code == 200
    assert preflight.json()["row_count"] == 69


def test_frontend_uses_shared_paginated_view_and_never_silently_disables_confirm() -> None:
    js = (ROOT / "quality_knowledge/web/static/major_production.js").read_text(encoding="utf-8")
    html = (ROOT / "quality_knowledge/web/templates/major_production.html").read_text(encoding="utf-8")
    assert "IMPORT_PAGE_SIZE = 20" in js
    assert "renderPagedRows(host, items" in js
    assert "renderBatchTable(" in js
    assert "renderCommitResults(" in js
    assert "data-major-recovered-confirm" in js
    assert "data-major-blocking-reasons" in js
    assert "data-open-import-case" in js
    assert "slice(0, 20).map(row" not in js
    assert "slice(0, 100).map(row" not in js
    assert "await restoreBatch(data.batch_id)" in js
    assert "if (!state.caseId)" in js
    assert "if (!state.eventId)" in js
    assert "服务器拒绝执行（HTTP " in js
    assert "major-batch607-v1" in html
