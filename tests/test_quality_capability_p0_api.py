from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from quality_knowledge.p0.initializer import P0Initializer
from quality_knowledge.p0.repository import P0Repository
from quality_knowledge.web.p0_app import create_p0_app


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def initializer() -> P0Initializer:
    return P0Initializer(
        manifest_path=PROJECT_ROOT / "quality_knowledge/config/p0_seed_manifest.json",
        plc_seed_path=PROJECT_ROOT / "quality_knowledge/config/plc_fields.yaml",
    )


def test_uninitialized_p0_web_exposes_only_blocking_diagnostic(tmp_path):
    client = TestClient(create_p0_app(tmp_path / "missing.db"))

    status = client.get("/api/v2/initialization/status")
    assert status.status_code == 503
    assert status.json()["error"] == "P0_DATABASE_NOT_INITIALIZED"
    assert client.get("/api/v2/products").status_code == 404


def test_ready_p0_web_exposes_dynamic_products_taxonomy_and_59_fields(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path))

    assert client.get("/api/v2/initialization/status").json()["initialization_state"] == "READY"
    products = client.get("/api/v2/products").json()
    assert products["total"] == 1
    assert {item["product_code"] for item in products["items"]} == {"PLC"}
    fields = client.get("/api/v2/standard-fields?limit=1000").json()
    assert fields["total"] == 59
    assert all(item["target_field"] != "source_id" for item in fields["items"])
    assert client.get("/api/v2/taxonomies").json()["total"] > 0
    assert client.get("/p0/insights").status_code == 200
    assert client.get("/p0/static/p0_insights.css").status_code == 200


def test_dynamic_product_mapping_draft_can_bind_validate_activate_and_import(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path))

    product = client.post("/api/v2/products", json={
        "product_code": "ROBOT", "product_name": "机器人", "product_type": "MECHATRONIC",
    })
    assert product.status_code == 201, product.text
    assert product.json()["product_code"] == "ROBOT"
    draft = client.post("/api/v2/products/ROBOT/mapping-drafts", json={"created_by": "admin"})
    assert draft.status_code == 201, draft.text
    assert len(draft.json()["items"]) == 59
    assert all(not item["aliases"] for item in draft.json()["items"])
    config_id = draft.json()["config_id"]

    invalid = client.post(f"/api/v2/mappings/{config_id}/validate")
    assert invalid.json()["valid"] is False
    updated = client.put(f"/api/v2/mappings/{config_id}", json={
        "updated_by": "admin",
        "bindings": [{
            "target_domain": "ISSUE_FACT", "target_field": "business_issue_id",
            "source_headers": ["机器人问题编号"], "enabled": True,
        }],
    })
    assert updated.status_code == 200, updated.text
    assert client.post(f"/api/v2/mappings/{config_id}/validate").json()["valid"] is True
    activated = client.post(f"/api/v2/mappings/{config_id}/activate", json={"actor": "admin"})
    assert activated.status_code == 200, activated.text
    assert activated.json()["status"] == "ACTIVE"

    workbook_path = tmp_path / "robot.xlsx"
    workbook = Workbook()
    workbook.active.append(["机器人问题编号", "未配置备注"])
    workbook.active.append(["R-001", "保留原始信息"])
    workbook.save(workbook_path)
    workbook.close()
    with workbook_path.open("rb") as stream:
        preview = client.post(
            "/api/v2/intake/preview",
            files={"file": ("robot.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"business_type": "ROBOT"},
        )
    assert preview.status_code == 200, preview.text
    assert preview.json()["product_code"] == "ROBOT"
    assert preview.json()["required_missing"] == []
    confirmed = client.post("/api/v2/intake/confirm", json={"preview_token": preview.json()["preview_token"]})
    assert confirmed.status_code == 200, confirmed.text
    issues = client.get("/api/v2/issues", params={"business_type": "ROBOT"}).json()
    assert issues["total"] == 1 and issues["items"][0]["product_name"] == "机器人"


def test_new_product_can_preview_against_draft_resolve_differences_then_activate_and_import(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path))
    assert client.post("/api/v2/products", json={
        "product_code": "VISION", "product_name": "视觉", "product_type": "SOFTWARE",
    }).status_code == 201
    draft = client.post(
        "/api/v2/products/VISION/mapping-drafts", json={"created_by": "quality-owner"}
    ).json()

    workbook_path = tmp_path / "vision.xlsx"
    workbook = Workbook()
    workbook.active.append(["视觉问题编号", "现象说明", "现场备注"])
    workbook.active.append(["V-001", "识别结果漂移", "保留原始信息"])
    workbook.save(workbook_path)
    workbook.close()
    with workbook_path.open("rb") as stream:
        preview = client.post(
            "/api/v2/intake/preview",
            files={"file": ("vision.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"business_type": "VISION"},
        )
    assert preview.status_code == 200, preview.text
    preview_payload = preview.json()
    assert preview_payload["mapping_status"] == "DRAFT"
    assert preview_payload["can_confirm"] is False
    assert client.post(
        "/api/v2/intake/confirm", json={"preview_token": preview_payload["preview_token"]}
    ).status_code == 409

    updated = client.post("/api/v2/intake/mapping-draft", json={
        "preview_token": preview_payload["preview_token"],
        "actor": "quality-owner",
        "decisions": [
            {"source_header": "视觉问题编号", "action": "MAP_EXISTING",
             "target_path": "ISSUE_FACT.business_issue_id"},
            {"source_header": "现象说明", "action": "MAP_EXISTING",
             "target_path": "ISSUE_FACT.description"},
            {"source_header": "现场备注", "action": "RAW_ONLY", "target_path": ""},
        ],
    })
    assert updated.status_code == 201, updated.text
    assert updated.json()["config_id"] == draft["config_id"]
    assert client.post(f"/api/v2/mappings/{draft['config_id']}/validate").json()["valid"] is True
    assert client.post(
        f"/api/v2/mappings/{draft['config_id']}/activate", json={"actor": "quality-owner"}
    ).status_code == 200

    with workbook_path.open("rb") as stream:
        final_preview = client.post(
            "/api/v2/intake/preview",
            files={"file": ("vision.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"business_type": "VISION"},
        )
    assert final_preview.status_code == 200
    final_payload = final_preview.json()
    assert final_payload["mapping_status"] == "ACTIVE"
    assert final_payload["can_confirm"] is True
    decisions = {item["source_header"]: item for item in final_payload["field_decisions"]}
    assert decisions["视觉问题编号"]["target"] == "ISSUE_FACT.business_issue_id"
    assert decisions["现象说明"]["target"] == "ISSUE_FACT.description"
    assert decisions["现场备注"]["status"] == "UNMATCHED"
    confirmed = client.post(
        "/api/v2/intake/confirm", json={"preview_token": final_payload["preview_token"]}
    )
    assert confirmed.status_code == 200 and confirmed.json()["success"] == 1


def test_mapping_validation_rejects_one_source_header_bound_to_multiple_targets(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path))
    draft = client.post(
        "/api/v2/products/PLC/mapping-drafts", json={"created_by": "quality-owner"}
    ).json()
    items = draft["items"]
    bindings = [{
        "target_domain": item["target_domain"],
        "target_field": item["target_field"],
        "source_headers": (["同一个表头"] if item["target_field"] in {"title", "description"} else item["aliases"]),
        "enabled": bool(item["enabled"]),
    } for item in items]
    assert client.put(f"/api/v2/mappings/{draft['config_id']}", json={
        "updated_by": "quality-owner", "bindings": bindings,
    }).status_code == 200
    validation = client.post(f"/api/v2/mappings/{draft['config_id']}/validate").json()
    assert validation["valid"] is False
    assert "SOURCE_HEADER_TARGET_CONFLICT" in validation["errors"]


def test_active_plc_mapping_can_be_cloned_edited_validated_and_activated(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path))

    active = client.get("/api/v2/mappings", params={"business_type": "PLC"}).json()["items"][0]
    assert active["status"] == "ACTIVE"
    active_detail = client.get(f"/api/v2/mappings/{active['config_id']}").json()
    created = client.post("/api/v2/products/PLC/mapping-drafts", json={"created_by": "quality-owner"})
    assert created.status_code == 201, created.text
    draft = created.json()
    assert draft["status"] == "DRAFT" and draft["version"] == active["version"] + 1
    assert draft["source_type"] == "ACTIVE_CLONE"
    expected_aliases = {
        item["target_field"]: item["aliases"] for item in active_detail["items"]
        if item["target_field"] != "source_id"
    }
    assert {item["target_field"]: item["aliases"] for item in draft["items"]} == expected_aliases

    bindings = [{
        "target_domain": item["target_domain"], "target_field": item["target_field"],
        "source_headers": ([*item["aliases"], "问题标题补充"] if item["target_field"] == "title" else item["aliases"]),
        "enabled": bool(item["enabled"]),
    } for item in draft["items"]]
    saved = client.put(f"/api/v2/mappings/{draft['config_id']}", json={
        "updated_by": "quality-owner", "bindings": bindings,
    })
    assert saved.status_code == 200, saved.text
    title = next(item for item in saved.json()["items"] if item["target_field"] == "title")
    assert "问题标题补充" in title["aliases"]
    assert client.post(f"/api/v2/mappings/{draft['config_id']}/validate").json()["valid"] is True
    activated = client.post(f"/api/v2/mappings/{draft['config_id']}/activate", json={"actor": "quality-owner"})
    assert activated.status_code == 200 and activated.json()["status"] == "ACTIVE"


def test_analysis_endpoint_does_not_fall_back_for_an_unknown_native_p0_issue(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    client = TestClient(create_p0_app(db_path, stage_runner=None))

    response = client.post("/api/v2/issues/UNKNOWN/analysis", json={})
    assert response.status_code == 404
    assert response.json()["detail"] == "ISSUE_NOT_FOUND"


def test_p0_intake_api_previews_and_confirms_into_the_clean_repository(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    workbook_path = tmp_path / "plc.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ITR单号", "问题描述", "SourceID"])
    sheet.append(["PLC-API-001", "API 接入问题", "feishu-row-1"])
    workbook.save(workbook_path)
    workbook.close()
    client = TestClient(create_p0_app(db_path))

    with workbook_path.open("rb") as stream:
        preview = client.post(
            "/api/v2/intake/preview",
            files={"file": ("plc.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"business_type": "PLC"},
        )
    assert preview.status_code == 200, preview.text
    assert preview.json()["required_missing"] == []
    source_id = next(item for item in preview.json()["field_decisions"] if item["source_header"] == "SourceID")
    assert source_id["status"] == "RAW_ONLY"

    confirmed = client.post("/api/v2/intake/confirm", json={"preview_token": preview.json()["preview_token"]})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["success"] == 1
    issue_list = client.get("/api/v2/issues", params={"business_type": "PLC"}).json()
    assert issue_list["total"] == 1
    assert issue_list["items"][0]["product_code"] == "PLC"


def add_insight_issue(repository: P0Repository, number: int) -> None:
    issue = repository.save_issue(
        knowledge_id=f"K-API-{number}",
        business_issue_id=f"B-API-{number}",
        raw_json={"问题编号": f"B-API-{number}"},
        normalized_snapshot={"ISSUE_FACT": {"business_issue_id": f"B-API-{number}", "month": "2026-08", "severity": "H"}},
        mapping_config_id="MAP-PLC-V1",
        mapping_config_version=1,
        standard_catalog_version_id="SFC-PLC-V1",
        source_file_sha256="api-source",
        sheet_name="issues",
        row_number=number,
        product_id=None,
    )
    repository.save_analysis_set({
        "analysis_set_id": f"AS-API-{number}",
        "knowledge_id": f"K-API-{number}",
        "issue_version_id": issue["issue_version_id"],
        "taxonomy_version_id": "QUALITY_TAXONOMY_P0_V1",
        "input_hash": f"input-api-{number}",
        "status": "COMPLETED",
        "tags": [
            {"stage": "occurrence", "axis": "DOMAIN", "tag_code": "SOFTWARE", "confidence": 0.6},
            {"stage": "occurrence", "axis": "LIFECYCLE", "tag_code": "IMPLEMENTATION", "confidence": 0.6},
        ],
        "mrc": [{"side": "OCCURRENCE", "mrc_code": "CHANGE_IMPACT_NOT_ASSESSED", "confidence": 0.6}],
        "capability_gaps": [{
            "capability_axis": "QUALITY_ENGINEERING",
            "capability_code": "TEST_VERIFICATION",
            "governance_scope": "PRODUCT",
            "control_status": "NOT_DEFINED",
            "details": {"priority": "P0"},
            "confidence": 0.6,
        }],
    })


def test_insight_api_uses_semantic_scope_and_returns_409_when_scope_changes(tmp_path):
    db_path = tmp_path / "p0.db"
    initializer().initialize(db_path)
    repository = P0Repository(db_path)
    add_insight_issue(repository, 1)
    client = TestClient(create_p0_app(db_path))

    issue_list = client.get("/api/v2/issues", params={"business_type": "PLC", "month": "2026-08"})
    assert issue_list.status_code == 200
    assert issue_list.json()["total"] == 1
    assert issue_list.json()["items"][0]["analysis_status"] == "COMPLETED"
    assert issue_list.json()["items"][0]["product_name"] is None
    assert client.get("/api/v2/issues/facets").json()["months"] == ["2026-08"]
    detail = client.get("/api/v2/issues/K-API-1")
    assert detail.status_code == 200
    assert detail.json()["issue"]["business_type"] == "PLC"
    assert detail.json()["effective_analysis"]["analysis_set_id"] == "AS-API-1"

    overview = client.get(
        "/api/v2/insights/business-contradictions",
        params={"business_type": "PLC", "issue_domain": "SOFTWARE", "lifecycle_phase": "IMPLEMENTATION"},
    )
    assert overview.status_code == 200
    payload = overview.json()
    contradiction = payload["quality_engineering_top3"][0]
    path = f"/api/v2/insights/business-contradictions/{contradiction['contradiction_key']}/issues"
    parameters = {
        "analysis_scope_hash": payload["analysis_scope_hash"],
        "business_type": "PLC",
        "issue_domain": "SOFTWARE",
        "lifecycle_phase": "IMPLEMENTATION",
        "page": 1,
        "page_size": 100,
    }
    drilldown = client.get(path, params=parameters)
    assert drilldown.status_code == 200
    assert drilldown.json()["total"] == 1

    add_insight_issue(repository, 2)
    navigation = client.get(
        "/api/v2/issues/K-API-1/navigation",
        params={"business_type": "PLC", "month": "2026-08"},
    )
    assert navigation.status_code == 200
    assert navigation.json()["total"] == 2
    assert {navigation.json()["previous_id"], navigation.json()["next_id"]} == {None, "K-API-2"}
    stale = client.get(path, params=parameters)
    assert stale.status_code == 409
    assert stale.json()["detail"] == "INSIGHT_SCOPE_CHANGED"
