from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from html.parser import HTMLParser

from .contracts import Chunk, DocumentElement, ParsedDocument


def _locator(**values: object) -> str:
    return json.dumps(values, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _document(elements: list[DocumentElement], parser_id: str, version: str, raw_text: str | None = None) -> ParsedDocument:
    elements = [item for item in elements if item.text.strip()]
    if not elements:
        raise ValueError("Source content is empty or contains no extractable text.")
    text = raw_text if raw_text is not None else "\n\n".join(item.text for item in elements)
    counts = Counter(item.element_type for item in elements)
    return ParsedDocument(text, tuple(item.locator for item in elements), parser_id, version,
                          tuple(elements), all(item.locator for item in elements), dict(counts))


class PlainTextParser:
    """Parser retained for the legacy JSON import contract."""

    parser_id = "plain-text"
    parser_version = "1"

    def parse(self, content: bytes, media_type: str) -> ParsedDocument:
        normalized_type = media_type.split(";", 1)[0].strip().lower()
        if normalized_type not in {"text/plain", "text/markdown"}:
            raise ValueError("This parser accepts text/plain or text/markdown only.")
        text = content.decode("utf-8-sig").strip()
        if not text:
            raise ValueError("Source content is empty.")
        return ParsedDocument(text=text, locators=("text:1",), parser_id=self.parser_id, parser_version=self.parser_version)

    @staticmethod
    def _structured_text(content: bytes, media_type: str) -> ParsedDocument:
        text = content.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        if not text.strip():
            raise ValueError("Source content is empty.")
        elements: list[DocumentElement] = []
        section = "Document"
        for line_number, line in enumerate(text.splitlines(), 1):
            if not line.strip():
                continue
            heading = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
            if heading:
                section = heading.group(1).strip()
            elements.append(DocumentElement(line, _locator(type="text_line", line=line_number, section=section),
                                            "heading" if heading else "text"))
        return _document(elements, "plain-text-file", "1", text.strip())


class FileParser:
    """Deterministic local file adapters. No model/provider is involved in parsing."""

    parser_id = "public-file-adapters"
    parser_version = "1"
    MEDIA_TYPES = {
        ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".html": "text/html", ".htm": "text/html", ".md": "text/markdown", ".markdown": "text/markdown",
        ".txt": "text/plain",
    }

    def parse(self, content: bytes, media_type: str, filename: str) -> ParsedDocument:
        from pathlib import PurePath
        suffix = PurePath(filename).suffix.lower()
        expected = self.MEDIA_TYPES.get(suffix)
        actual = media_type.split(";", 1)[0].strip().lower()
        aliases = {"application/x-pdf": "application/pdf", "application/octet-stream": expected}
        actual = aliases.get(actual, actual)
        if expected is None or actual != expected:
            raise ValueError("File extension and declared media type do not match a supported format.")
        if not content:
            raise ValueError("Source file is empty.")
        if suffix == ".pdf":
            return self._pdf(content)
        if suffix == ".docx":
            return self._docx(content)
        if suffix in {".html", ".htm"}:
            return self._html(content)
        return PlainTextParser._structured_text(content, actual)

    def _pdf(self, content: bytes) -> ParsedDocument:
        import io
        import pdfplumber

        elements: list[DocumentElement] = []
        normalized_pages: list[str] = []
        section = "Document"
        complete = True
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page_no, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text() or ""
                if not page_text.strip():
                    complete = False
                page_lines: list[str] = []
                for line_no, line in enumerate(page_text.splitlines(), 1):
                    value = line.strip()
                    if not value:
                        continue
                    if len(value) < 100 and (value.isupper() or re.match(r"^\d+(?:\.\d+)*\s+\S", value)):
                        section = value
                    page_lines.append(value)
                    elements.append(DocumentElement(value, _locator(type="pdf_text", page=page_no, line=line_no, section=section), "text"))
                normalized_pages.append("\n".join(page_lines))
                for table_no, table in enumerate(page.extract_tables() or [], 1):
                    rows = [[str(cell or "").strip() for cell in row] for row in table]
                    if not rows or not any(any(cell for cell in row) for row in rows):
                        continue
                    header = rows[0]
                    md = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
                    for row_no, row in enumerate(rows[1:] or rows[:1], 1):
                        padded = row + [""] * max(0, len(header) - len(row))
                        rendered = "| " + " | ".join(padded[:len(header)]) + " |"
                        md.append(rendered)
                        elements.append(DocumentElement("Header: " + " | ".join(header) + "\n" + rendered,
                            _locator(type="pdf_table_row", page=page_no, section=section, table=table_no, row=row_no), "table_row"))
                    normalized_pages.append("\n".join(md))
        parsed = _document(elements, "pdfplumber", getattr(pdfplumber, "__version__", "unknown"), "\n\n".join(x for x in normalized_pages if x))
        if not complete:
            return ParsedDocument(parsed.text, parsed.locators, parsed.parser_id, parsed.parser_version,
                                  parsed.elements, False, parsed.element_counts)
        return parsed

    def _docx(self, content: bytes) -> ParsedDocument:
        import io
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(io.BytesIO(content))
        elements: list[DocumentElement] = []
        normalized: list[str] = []
        section = "Document"
        p_no = t_no = 0
        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, doc)
                value = paragraph.text.strip()
                p_no += 1
                if not value:
                    continue
                style = paragraph.style.name if paragraph.style else ""
                if style.lower().startswith("heading"):
                    section = value
                    rendered = "## " + value
                    kind = "heading"
                else:
                    rendered, kind = value, "paragraph"
                normalized.append(rendered)
                elements.append(DocumentElement(value, _locator(type="docx_paragraph", heading=section, paragraph=p_no, style=style), kind))
            elif isinstance(child, CT_Tbl):
                table = Table(child, doc)
                t_no += 1
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
                if not rows:
                    continue
                header = rows[0]
                normalized.append("| " + " | ".join(header) + " |")
                normalized.append("| " + " | ".join("---" for _ in header) + " |")
                for row_no, row in enumerate(rows[1:] or rows[:1], 1):
                    cells = row + [""] * max(0, len(header) - len(row))
                    rendered = "| " + " | ".join(cells[:len(header)]) + " |"
                    normalized.append(rendered)
                    elements.append(DocumentElement("Header: " + " | ".join(header) + "\n" + rendered,
                        _locator(type="docx_table_row", heading=section, table=t_no, row=row_no,
                                 columns=[{"column": i + 1, "text": value} for i, value in enumerate(cells[:len(header)])]), "table_row"))
        return _document(elements, "python-docx", "1", "\n\n".join(normalized))

    def _html(self, content: bytes) -> ParsedDocument:
        raw = content.decode("utf-8-sig")
        snapshot_sha = hashlib.sha256(content).hexdigest()

        class Extractor(HTMLParser):
            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.items: list[tuple[str, str, str | None, str]] = []
                self.stack: list[str] = []
                self.current_tag: str | None = None
                self.current_parts: list[str] = []
                self.current_anchor: str | None = None
                self.current_dom = ""
                self.anchor: str | None = None
                self.section = "Document"
                self.title = False
                self.page_title = ""

            def handle_starttag(self, tag, attrs):
                attrs_map = dict(attrs)
                self.stack.append(tag)
                if attrs_map.get("id"):
                    self.anchor = attrs_map["id"]
                if tag == "title":
                    self.title = True
                if tag in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr"}:
                    self.current_tag, self.current_parts = tag, []
                    self.current_anchor = self.anchor
                    self.current_dom = "/" + "/".join(self.stack)

            def handle_data(self, data):
                value = data.strip()
                if not value:
                    return
                if self.title:
                    self.page_title += value
                if self.current_tag:
                    self.current_parts.append(value)

            def handle_endtag(self, tag):
                if tag in {"td", "th"} and self.current_tag == "tr":
                    self.current_parts.append(" | ")
                if tag == "title":
                    self.title = False
                if tag == self.current_tag:
                    value = " ".join(self.current_parts).strip()
                    if value:
                        if tag.startswith("h"):
                            self.section = value
                        self.items.append((tag, value, self.current_anchor, self.current_dom))
                    self.current_tag, self.current_parts = None, []
                if tag in self.stack:
                    self.stack = self.stack[:len(self.stack) - 1 - self.stack[::-1].index(tag)]

        extractor = Extractor()
        extractor.feed(raw)
        elements: list[DocumentElement] = []
        normalized: list[str] = []
        counts: dict[str, int] = {}
        for tag, value, anchor, dom_path in extractor.items:
            counts[tag] = counts.get(tag, 0) + 1
            locator = _locator(type="html_element", snapshot_sha256=snapshot_sha, title=extractor.page_title,
                               section=extractor.section, anchor=anchor, dom_path=dom_path, dom_tag=tag, element=counts[tag])
            is_heading = tag.startswith("h")
            is_row = tag == "tr"
            normalized.append(("## " if is_heading else ("| " + value.strip(" |") + " |" if is_row else value)))
            elements.append(DocumentElement(value, locator, "heading" if is_heading else "table_row" if is_row else "text"))
        return _document(elements, "html-parser", "stdlib-1", "\n\n".join(normalized))


class WindowChunker:
    """Deterministic chunker that retains source-native locators where available."""

    def __init__(self, chunk_size: int = 1200, overlap: int = 120) -> None:
        if chunk_size < 100 or overlap < 0 or overlap >= chunk_size:
            raise ValueError("Invalid chunk size or overlap.")
        self.chunk_size = chunk_size
        self.overlap = overlap

    def chunk(self, parsed: ParsedDocument) -> list[Chunk]:
        chunks: list[Chunk] = []
        legacy = not parsed.elements
        elements = parsed.elements or (DocumentElement(parsed.text, "text:1"),)
        for element in elements:
            text = element.text
            start = 0
            while start < len(text):
                end = min(len(text), start + self.chunk_size)
                if end < len(text):
                    boundary = text.rfind(" ", start + self.chunk_size // 2, end)
                    if boundary > start:
                        end = boundary
                locator = (f"char:{start}-{end}" if legacy else element.locator if start == 0 and end == len(text)
                           else _locator(source=element.locator, char_start=start, char_end=end))
                chunks.append(Chunk(len(chunks), text[start:end].strip(), locator))
                if end >= len(text):
                    break
                start = max(start + 1, end - self.overlap)
        return chunks
